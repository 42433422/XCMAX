#!/usr/bin/env python3
"""FastAPI/Uvicorn entry; packaged desktop uses the host/port supplied by Electron."""

from __future__ import annotations

import argparse
import os
import sys
import time
from pathlib import Path

BOUNDARY_ERRORS: tuple[type[Exception], ...] = (Exception,)
RECOVERABLE_ERRORS: tuple[type[Exception], ...] = (
    OSError,
    ValueError,
    TypeError,
    AttributeError,
    RuntimeError,
    ImportError,
    LookupError,
    ArithmeticError,
)

_XCAGI_DIR = Path(__file__).resolve().parent
_REPO_ROOT = _XCAGI_DIR.parent


def _runtime_port_file() -> Path:
    """Write runtime ports to desktop userData, never signed application resources."""
    data_root = (
        os.environ.get("XCAGI_DATA_DIR") or os.environ.get("XCAGI_DESKTOP_DATA_DIR") or ""
    ).strip()
    if data_root:
        return Path(data_root).expanduser().resolve() / ".runtime" / "api.port"
    return _REPO_ROOT / ".runtime" / "api.port"


def _persist_runtime_port(port: int) -> None:
    """将实际监听端口写入 .runtime/api.port，供前端 Vite 读取联动。"""
    try:
        port_file = _runtime_port_file()
        port_file.parent.mkdir(parents=True, exist_ok=True)
        port_file.write_text(str(port), encoding="utf-8")
    except OSError:
        pass


def _load_dotenv_if_present(env_path: Path) -> None:
    if not env_path.is_file():
        return
    try:
        from dotenv import dotenv_values

        for k, v in dotenv_values(str(env_path)).items():
            if v is not None and k not in os.environ:
                os.environ[k] = v
    except ImportError:
        pass


_LOCAL_MARKET_ENV_KEYS = frozenset(
    {
        "XCAGI_MARKET_BASE_URL",
        "MODSTORE_LOCAL_AUTOMATION",
        "MODSTORE_LOCAL_BASE_URL",
        "MODSTORE_DIGEST_BASE_URL",
        "MODSTORE_ALL_HANDS_BASE_URL",
        "MODSTORE_DIGEST_ADMIN_USER",
        "MODSTORE_DIGEST_ADMIN_PASSWORD",
        "XCMAX_MONOREPO_ROOT",
    }
)


def _load_dotenv_override(env_path: Path, keys: frozenset[str] | None = None) -> None:
    if not env_path.is_file():
        return
    try:
        from dotenv import dotenv_values

        for k, v in dotenv_values(str(env_path)).items():
            if v is None:
                continue
            if keys is None or k in keys:
                os.environ[k] = v
    except ImportError:
        pass


def _env_truthy(name: str) -> bool:
    return os.environ.get(name, "").strip().lower() in {"1", "true", "yes", "on"}


def _apply_desktop_local_market_env() -> None:
    """Production by default; explicit remote preference overrides explicit local market."""
    os.environ.setdefault("XCAGI_MARKET_BASE_URL", "https://xiu-ci.com")
    if _env_truthy("XCAGI_USE_REMOTE_MARKET"):
        _load_dotenv_override(
            _XCAGI_DIR / ".env.online-market", frozenset({"XCAGI_MARKET_BASE_URL"})
        )
        if not os.environ.get("XCAGI_MARKET_BASE_URL"):
            os.environ["XCAGI_MARKET_BASE_URL"] = "https://xiu-ci.com"
        return
    if _env_truthy("XCAGI_USE_LOCAL_MARKET"):
        _load_dotenv_override(_XCAGI_DIR / ".env.local-market", _LOCAL_MARKET_ENV_KEYS)
        if not os.environ.get("XCAGI_MARKET_BASE_URL"):
            os.environ["XCAGI_MARKET_BASE_URL"] = "http://127.0.0.1:8788"


def _ensure_sys_path() -> None:
    for p in (_REPO_ROOT, _XCAGI_DIR):
        s = str(p)
        if s not in sys.path:
            sys.path.insert(0, s)


def _is_frozen() -> bool:
    return bool(getattr(sys, "frozen", False)) or hasattr(sys, "_MEIPASS")


def _force_stdio_utf8(data_dir: str | None = None) -> None:
    """Keep Windows pipes UTF-8; windowed Python needs streams before Uvicorn logging."""
    if sys.platform != "win32":
        return
    fallback = None
    for name, descriptor in (("stdout", 1), ("stderr", 2)):
        if getattr(sys, name) is not None:
            continue
        try:
            stream = os.fdopen(os.dup(descriptor), "w", encoding="utf-8", buffering=1)
        except OSError:
            if fallback is None:
                root = Path(
                    data_dir
                    or os.environ.get("XCAGI_DATA_DIR")
                    or os.environ.get("XCAGI_DESKTOP_DATA_DIR")
                    or (Path(os.environ.get("APPDATA") or Path.home()) / "XCAGI")
                )
                logs = root.expanduser().resolve() / "logs"
                logs.mkdir(parents=True, exist_ok=True)
                fallback = (logs / "backend-bootstrap.log").open("a", encoding="utf-8", buffering=1)
            stream = fallback
        setattr(sys, name, stream)
    for stream in (sys.stdout, sys.stderr):
        reconfigure = getattr(stream, "reconfigure", None)
        if reconfigure is None:
            continue
        try:
            reconfigure(encoding="utf-8")
        except (ValueError, OSError):
            pass


def _bootstrap_stage(stage: str) -> None:
    if sys.platform == "win32" and _is_frozen():
        print(f"[bootstrap] {time.time():.3f} {stage}", file=sys.stderr, flush=True)


def _verify_frozen_elt_runtime() -> None:
    """Exercise safe XML parsing and the real XLSX slice path without user data or AI."""
    from hashlib import sha256
    from io import BytesIO

    from defusedxml import ElementTree as safe_xml
    from defusedxml.common import DTDForbidden
    from openpyxl import Workbook, load_workbook
    from openpyxl.worksheet.print_settings import PrintArea

    from app.application.reviewed_excel_template import (
        derive_document_workbook,
        inspect_workbook,
    )

    # Importing alone does not prove the bundled parser or derivation path works.
    try:
        safe_xml.fromstring(b"<!DOCTYPE probe><probe/>", forbid_dtd=True)
    except DTDForbidden:
        pass
    else:
        raise RuntimeError("frozen ELT runtime probe did not reject a DTD")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Synthetic"
    sheet.append(["Item", "Quantity"])
    sheet.append(["Sample", 2])
    sheet.append(["Total", "=SUM(B2:B2)"])
    sheet.auto_filter.ref = "A1:B3"
    sheet.print_area = "A1:B3"
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    source = stream.getvalue()
    source_sha = sha256(source).digest()
    if inspect_workbook(source)["sheets"] != ["Synthetic"]:
        raise RuntimeError("frozen ELT runtime probe could not inspect its workbook")
    derived = derive_document_workbook(
        source, [{"sheet": "Synthetic", "range": "A1:B3"}], print_area_mode="selected"
    )
    checked = load_workbook(BytesIO(derived), data_only=False, keep_links=True, rich_text=True)
    try:
        result = checked["Synthetic"]
        if (
            checked.sheetnames != ["Synthetic"]
            or list(result.values)
            != [("Item", "Quantity"), ("Sample", 2), ("Total", "=SUM(B2:B2)")]
            or result.auto_filter.ref != "A1:B3"
            or {str(area) for area in PrintArea.from_string(result.print_area).ranges} != {"A1:B3"}
            or sha256(source).digest() != source_sha
        ):
            raise RuntimeError("frozen ELT runtime probe did not preserve its workbook")
    finally:
        checked.close()
    print("[run_fastapi] frozen ELT runtime probe OK: defusedxml + XLSX inspect/derive/readback")


def _verify_frozen_critical_runtime() -> None:
    """Exercise critical office and voice dependencies in the frozen executable."""
    import tempfile

    print("[frozen-runtime-probe] voice and PDF imports begin", file=sys.stderr, flush=True)
    import av
    import faster_whisper
    from pypdf import PdfReader
    from reportlab.lib.pagesizes import A4  # type: ignore[import-untyped]
    from reportlab.pdfbase import pdfmetrics  # type: ignore[import-untyped]
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont  # type: ignore[import-untyped]
    from reportlab.pdfgen import canvas  # type: ignore[import-untyped]

    print("[frozen-runtime-probe] ELT derivation begin", file=sys.stderr, flush=True)
    _verify_frozen_elt_runtime()
    print("[frozen-runtime-probe] ELT derivation complete", file=sys.stderr, flush=True)

    # Validate string-loaded sync appliers that static PyInstaller analysis cannot discover.
    print("[frozen-runtime-probe] sync applier import begin", file=sys.stderr, flush=True)
    from app.domain.shared_services.xcmax_sync_service import _ENTITY_APPLIERS

    print("[frozen-runtime-probe] sync applier import complete", file=sys.stderr, flush=True)
    required_sync_appliers = {
        "personnel",
        "department",
        "attendance",
        "approval",
        "approval_flow",
        "print_job",
        "template",
        "model_config",
        "ecosystem",
        "im_message",
        "im_read_state",
        "workflow_employee",
        "account_entitlements",
        "private_mod_delivery",
    }
    missing_sync_appliers = required_sync_appliers.difference(_ENTITY_APPLIERS)
    if missing_sync_appliers:
        raise RuntimeError(
            f"frozen XCMAX sync applier probe missing: {sorted(missing_sync_appliers)!r}"
        )

    with tempfile.TemporaryDirectory(prefix="xcagi-office-probe-") as tmp:
        output = Path(tmp) / "probe.pdf"
        font_name = "STSong-Light"
        try:
            pdfmetrics.getFont(font_name)
        except KeyError:
            pdfmetrics.registerFont(UnicodeCIDFont(font_name))
        pdf = canvas.Canvas(str(output), pagesize=A4)
        pdf.setFont(font_name, 11)
        pdf.drawString(72, 770, "XCAGI PDF runtime probe")
        pdf.save()
        reader = PdfReader(str(output))
        if len(reader.pages) != 1 or output.stat().st_size < 512:
            raise RuntimeError("frozen PDF runtime probe produced an invalid document")

        if _is_frozen():
            print("[frozen-runtime-probe] desktop seed import begin", file=sys.stderr, flush=True)
            from app.desktop_runtime.desktop_runtime_paths import configure_desktop_environment

            print("[frozen-runtime-probe] desktop seed setup begin", file=sys.stderr, flush=True)
            data_root = Path(os.environ["XCAGI_DATA_DIR"])
            configure_desktop_environment(data_root)
            print("[frozen-runtime-probe] desktop seed setup complete", file=sys.stderr, flush=True)
            bundled_root = Path(sys._MEIPASS) / "mods" / "_employees"  # type: ignore[attr-defined]
            installed_root = data_root / "mods" / "_employees"
            bundled = {
                path.name
                for path in bundled_root.iterdir()
                if path.is_dir() and (path / "manifest.json").is_file()
            }
            installed = {
                path.name
                for path in installed_root.iterdir()
                if path.is_dir() and (path / "manifest.json").is_file()
            }
            required = {"pdf-generate-employee", "pdf-full-read-employee"}
            if not required.issubset(installed) or installed != bundled:
                raise RuntimeError(
                    "frozen Office employee seed mismatch: "
                    f"bundled={sorted(bundled)!r} installed={sorted(installed)!r}"
                )
    print(
        "[run_fastapi] frozen critical runtime probe OK: "
        "XCMAX sync appliers + pypdf + reportlab + "
        f"PyAV {av.__version__} + faster-whisper {faster_whisper.__version__}"
    )


def _run_frozen_critical_probe(data_dir: str | None) -> None:
    """Isolate diagnostics before importing application bootstrap packages."""
    import faulthandler
    import tempfile
    import traceback

    with tempfile.TemporaryDirectory(prefix="xcagi-critical-probe-") as temporary:
        root = Path(data_dir or temporary).resolve()
        for name in ("data", "logs", "mods"):
            (root / name).mkdir(parents=True, exist_ok=True)
        database_url = "sqlite:///" + (root / "data" / "probe.db").as_posix()
        isolated_env = {
            "PYTHON_DOTENV_DISABLED": "1",
            "XCAGI_DESKTOP_MODE": "1",
            "XCAGI_DATA_DIR": str(root),
            "XCAGI_DESKTOP_DATA_DIR": str(root),
            "XCAGI_MODS_ROOT": str(root / "mods"),
            "XCAGI_LOG_DIR": str(root / "logs"),
            "DATABASE_URL": database_url,
            "VECTOR_DB_URL": database_url,
            "XCAGI_DESKTOP_DATABASE_URL": database_url,
            "XCAGI_DESKTOP_VECTOR_DB_URL": database_url,
            "DATABASE_PATH": str(root / "data"),
            "WORKSPACE_ROOT": str(root),
            "XCAGI_WORKSPACE_ROOT": str(root),
            "XCAGI_FORCE_SYNC_TASKS": "1",
            "XCAGI_DISABLE_REDIS": "1",
            "CACHE_REDIS_URL": "",
            "CELERY_BROKER_URL": "memory://",
            "CELERY_RESULT_BACKEND": "cache+memory://",
        }
        previous_env = {name: os.environ.get(name) for name in isolated_env}
        os.environ.update(isolated_env)
        try:
            _force_stdio_utf8(str(root))
            print("[frozen-runtime-probe] isolated entry", file=sys.stderr, flush=True)
            _ensure_sys_path()
            print("[frozen-runtime-probe] source path ready", file=sys.stderr, flush=True)
            with (root / "logs" / "frozen-runtime-probe.log").open(
                "w", encoding="utf-8", buffering=1
            ) as diagnostics:
                # Dump Python stacks if an import stalls; the caller also bounds
                # this one process. No locals, arguments or environment are logged.
                faulthandler.dump_traceback_later(60, repeat=True, file=diagnostics)
                try:
                    _verify_frozen_critical_runtime()
                except Exception as exc:  # noqa: BLE001 - isolated command boundary
                    traceback.print_exc(file=diagnostics)
                    diagnostics.flush()
                    traceback.print_exc(file=sys.stderr)
                    # A windowed PyInstaller executable displays unhandled errors
                    # in a blocking dialog. An explicit failed exit keeps CI usable.
                    raise SystemExit(1) from exc
                finally:
                    faulthandler.cancel_dump_traceback_later()
        finally:
            for name, value in previous_env.items():
                if value is None:
                    os.environ.pop(name, None)
                else:
                    os.environ[name] = value


def _parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="XCAGI FastAPI server")
    parser.add_argument(
        "--desktop", action="store_true", help="桌面模式（本地 SQLite、禁用 reload）"
    )
    parser.add_argument("--headless", action="store_true", help="无控制台窗口（由 Electron 托管）")
    parser.add_argument("--host", default=None, help="监听地址")
    parser.add_argument("--port", type=int, default=None, help="监听端口")
    parser.add_argument("--data-dir", default=None, help="桌面数据目录")
    parser.add_argument("--verify-backup", default="", help="只读校验 SQLite 备份后退出")
    parser.add_argument(
        "--export-model-usage-ledger",
        default="",
        help="后端停写后将计量账本交接给旧版；成功后退出，不启动服务器",
    )
    parser.add_argument("--migrate-only", action="store_true", help="仅执行数据库迁移后退出")
    parser.add_argument(
        "--backup", action="store_true", help="迁移前备份（与 --migrate-only 合用）"
    )
    parser.add_argument(
        "--verify-frozen-critical-runtime",
        action="store_true",
        help="验证冻结包内办公与语音关键依赖后退出",
    )
    return parser.parse_args(argv)


def _apply_desktop_bootstrap(args: argparse.Namespace) -> None:
    if args.desktop:
        os.environ["XCAGI_DESKTOP_MODE"] = "1"
    if args.data_dir:
        os.environ["XCAGI_DATA_DIR"] = str(args.data_dir)
    _ensure_sys_path()
    try:
        from app.desktop_runtime import configure_desktop_environment, is_desktop_mode

        if is_desktop_mode():
            configure_desktop_environment(os.environ.get("XCAGI_DATA_DIR"))
    except RECOVERABLE_ERRORS as exc:  # noqa: BLE001 - desktop bootstrap must stay best-effort
        print(f"[run_fastapi] desktop bootstrap warning: {exc}", file=sys.stderr)


def _apply_no_console_child_defaults() -> None:
    """打包桌面后端自身无控制台：默认让子进程不开新控制台窗口（避免点功能闪黑窗）。"""
    try:
        from app.desktop_runtime.no_console_children import (
            install_no_console_child_defaults,
        )

        install_no_console_child_defaults()
    except BOUNDARY_ERRORS:
        pass


def _resolve_reload(desktop: bool) -> bool:
    if _is_frozen() or desktop:
        return False
    return _env_truthy("XCAGI_UVICORN_RELOAD") if "XCAGI_UVICORN_RELOAD" in os.environ else True


def main(argv: list[str] | None = None) -> None:
    args = _parse_args(argv)
    if args.verify_frozen_critical_runtime:
        _run_frozen_critical_probe(args.data_dir)
        return
    _force_stdio_utf8(args.data_dir)
    _bootstrap_stage("entry")
    _ensure_sys_path()
    _apply_no_console_child_defaults()
    _bootstrap_stage("child-defaults-ready")

    if args.export_model_usage_ledger:
        from app.infrastructure.billing.model_usage_export import main as export_ledger

        ledger = str(Path(args.export_model_usage_ledger).expanduser().resolve())
        result = export_ledger(["--ledger", ledger, "--output", ledger, "--handoff"])
        if result:
            raise SystemExit(result)
        print(f"XCAGI_MODEL_USAGE_HANDOFF={ledger}", flush=True)
        return

    if args.verify_backup:
        from app.desktop_runtime.db import integrity_check_ok

        backup = Path(args.verify_backup)
        if not backup.is_file() or not backup.stat().st_size or not integrity_check_ok(backup):
            raise SystemExit(1)
        print(f"XCAGI_BACKUP_VALID={backup}")
        return

    _load_dotenv_if_present(_XCAGI_DIR / ".env")
    _load_dotenv_if_present(_REPO_ROOT / ".env")

    # Clash 常注入 ALL_PROXY=socks5://…；无 socksio 时 httpx 会直接失败。
    try:
        from app.utils.security.proxy_env import sanitize_socks_all_proxy

        sanitize_socks_all_proxy()
    except BOUNDARY_ERRORS:
        pass

    if args.desktop or args.data_dir or args.migrate_only:
        _bootstrap_stage("desktop-environment-begin")
        _apply_desktop_bootstrap(args)
        _bootstrap_stage("desktop-environment-ready")

    desktop = args.desktop or _env_truthy("XCAGI_DESKTOP_MODE")
    if desktop:
        _apply_desktop_local_market_env()

    if args.migrate_only:
        from app.desktop_runtime.desktop_runtime_paths import (
            configure_desktop_environment,
            ensure_desktop_dirs,
        )
        from app.desktop_runtime.migrate import (
            backup_database,
            migration_lock,
            run_alembic_upgrade,
        )

        configure_desktop_environment(args.data_dir)
        version = os.environ.get("XCAGI_VERSION", "unknown")
        # Backup and migration share the lock with concurrent normal startup.
        with migration_lock(args.data_dir):
            if args.backup:
                dirs = ensure_desktop_dirs(args.data_dir)
                database_exists = (dirs["data"] / "xcagi.db").exists()
                backup_path = backup_database(args.data_dir, version)
                if database_exists and backup_path is None:
                    raise RuntimeError("migration backup failed; refusing to continue")
                if backup_path is not None:
                    print(f"XCAGI_MIGRATION_BACKUP={backup_path}", flush=True)
            run_alembic_upgrade(args.data_dir)
        return

    if desktop:
        from app.desktop_runtime.migrate import ensure_startup_migration

        _bootstrap_stage("startup-migration-begin")
        ensure_startup_migration(args.data_dir)
        _bootstrap_stage("startup-migration-ready")

    if args.host:
        os.environ["FASTAPI_HOST"] = args.host
        os.environ["XCAGI_API_HOST"] = args.host
    if args.port is not None:
        os.environ["FASTAPI_PORT"] = str(args.port)
        os.environ["XCAGI_API_PORT"] = str(args.port)

    from app.utils.device_system.listen_port import (
        DESKTOP_API_PORT,
        find_available_port,
        publish_listen_address,
        resolve_listen_host,
        resolve_listen_port,
    )

    host = resolve_listen_host()
    port = resolve_listen_port()
    if args.port is not None and not 1 <= args.port <= 65535:
        raise ValueError("HTTP port must be between 1 and 65535")
    if not desktop:
        free_port = find_available_port(host, port, range(DESKTOP_API_PORT, DESKTOP_API_PORT + 21))
        if free_port != port:
            print(
                f"[run_fastapi] 端口 {port} 被占用，自动切换到 {free_port}",
                file=sys.stderr,
            )
            port = free_port
    publish_listen_address(host, port)
    _persist_runtime_port(port)
    reload = _resolve_reload(desktop)

    _bootstrap_stage("server-import-begin")
    import uvicorn

    # 打包后避免 ``"module:attr"`` 字符串导入（PyInstaller 常无法解析，导致桌面端启动失败/假死）
    if _is_frozen():
        from app.fastapi_app import create_fastapi_app

        _bootstrap_stage("server-import-ready")
        target = create_fastapi_app
    else:
        target = "app.fastapi_app:create_fastapi_app"

    uvicorn.run(
        target,
        factory=True,
        host=host,
        port=port,
        reload=reload,
        reload_dirs=[str(_REPO_ROOT)] if reload else None,
        log_level="info",
        timeout_graceful_shutdown=30,
    )


if __name__ == "__main__":
    # Dispatch frozen multiprocessing workers before parsing application arguments.
    import multiprocessing

    multiprocessing.freeze_support()
    main()
