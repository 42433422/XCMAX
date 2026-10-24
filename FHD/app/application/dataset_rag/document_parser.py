"""Admitted document decoding and chunk segmentation."""

from __future__ import annotations

from collections.abc import Callable
from pathlib import Path
from typing import Any

from app.infrastructure.rag import SemanticChunker
from app.infrastructure.rag.semantic_chunker import Chunk
from app.utils.security.safe_download_path import resolve_under_allowed_dirs


class DocumentParser:
    """Resolve admitted paths, decode document formats, and segment text."""

    def __init__(
        self,
        *,
        chunker: SemanticChunker,
        allowed_roots: list[Path] | None,
        default_roots: Callable[[], list[Path]],
    ) -> None:
        self._chunker = chunker
        self._allowed_roots = allowed_roots
        self._default_roots = default_roots

    def _resolve_file_path(self, file_path: str) -> Path:
        return resolve_under_allowed_dirs(file_path, self._allowed_file_roots())

    def _allowed_file_roots(self) -> list[Path]:
        if self._allowed_roots is not None:
            return [Path(root).resolve() for root in self._allowed_roots]
        return self._default_roots()

    def _extract_file_text(self, path: Path) -> tuple[str, str, dict[str, Any]]:
        if not path.exists() or not path.is_file():
            raise ValueError(f"file not found: {path}")
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            return self._extract_pdf_text(path)
        if suffix == ".docx":
            return self._extract_docx_text(path)
        if suffix in {".xlsx", ".xls"}:
            return self._extract_excel_text(path)
        if suffix in {".txt", ".md", ".csv", ".json", ".log"}:
            return (
                path.read_text(encoding="utf-8", errors="replace"),
                "text_file",
                {"extension": suffix},
            )
        raise ValueError(f"unsupported document type: {suffix or '<none>'}")

    def _extract_pdf_text(self, path: Path) -> tuple[str, str, dict[str, Any]]:
        try:
            import pdfplumber
        except ImportError as exc:
            raise RuntimeError("pdfplumber is required to ingest PDF documents") from exc
        pages: list[str] = []
        with pdfplumber.open(str(path)) as pdf:
            for index, page in enumerate(pdf.pages, start=1):
                page_text = (page.extract_text() or "").strip()
                if page_text:
                    pages.append(f"[page {index}]\n{page_text}")
            page_count = len(pdf.pages)
        return ("\n\n".join(pages), "pdfplumber", {"extension": ".pdf", "page_count": page_count})

    def _extract_docx_text(self, path: Path) -> tuple[str, str, dict[str, Any]]:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("python-docx is required to ingest DOCX documents") from exc
        doc = Document(str(path))
        parts: list[str] = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return ("\n".join(parts), "python-docx", {"extension": ".docx"})

    def _extract_excel_text(self, path: Path) -> tuple[str, str, dict[str, Any]]:
        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            return self._extract_xlsx_text(path)
        if suffix == ".xls":
            return self._extract_xls_text(path)
        raise ValueError(f"unsupported document type: {suffix or '<none>'}")

    def _extract_xlsx_text(self, path: Path) -> tuple[str, str, dict[str, Any]]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RuntimeError("openpyxl is required to ingest XLSX documents") from exc
        max_rows_per_sheet = 5000
        max_cols = 100
        parts: list[str] = []
        sheet_count = 0
        row_count = 0
        wb = load_workbook(str(path), read_only=True, data_only=True)
        try:
            for sheet in wb.worksheets:
                sheet_count += 1
                sheet_lines: list[str] = []
                for index, row in enumerate(sheet.iter_rows(values_only=True)):
                    if index >= max_rows_per_sheet:
                        sheet_lines.append(f"... truncated after {max_rows_per_sheet} rows")
                        break
                    cells: list[str] = []
                    for col_index, cell in enumerate(row or ()):
                        if col_index >= max_cols:
                            break
                        if cell is None:
                            continue
                        text = str(cell).strip()
                        if text:
                            cells.append(text)
                    if cells:
                        sheet_lines.append("\t".join(cells))
                        row_count += 1
                if sheet_lines:
                    parts.append(f"[sheet {sheet.title}]\n" + "\n".join(sheet_lines))
        finally:
            wb.close()
        return (
            "\n\n".join(parts),
            "openpyxl",
            {"extension": ".xlsx", "sheet_count": sheet_count, "row_count": row_count},
        )

    def _extract_xls_text(self, path: Path) -> tuple[str, str, dict[str, Any]]:
        try:
            import xlrd
        except ImportError as exc:
            raise RuntimeError("xlrd is required to ingest XLS documents") from exc
        max_rows_per_sheet = 5000
        max_cols = 100
        parts: list[str] = []
        sheet_count = 0
        row_count = 0
        book = xlrd.open_workbook(str(path))
        for sheet in book.sheets():
            sheet_count += 1
            sheet_lines: list[str] = []
            nrows = min(int(sheet.nrows), max_rows_per_sheet)
            for row_index in range(nrows):
                cells: list[str] = []
                ncols = min(int(sheet.ncols), max_cols)
                for col_index in range(ncols):
                    value = sheet.cell_value(row_index, col_index)
                    if value is None or value == "":
                        continue
                    text = str(value).strip()
                    if text:
                        cells.append(text)
                if cells:
                    sheet_lines.append("\t".join(cells))
                    row_count += 1
            if int(sheet.nrows) > max_rows_per_sheet:
                sheet_lines.append(f"... truncated after {max_rows_per_sheet} rows")
            if sheet_lines:
                parts.append(f"[sheet {sheet.name}]\n" + "\n".join(sheet_lines))
        return (
            "\n\n".join(parts),
            "xlrd",
            {"extension": ".xls", "sheet_count": sheet_count, "row_count": row_count},
        )

    def _split_text(
        self, text: str, *, strategy: str, chunk_size: int, chunk_overlap: int
    ) -> list[Chunk]:
        if strategy == "semantic":
            return self._chunker.split_by_semantic(text)
        return self._chunker.split_by_fixed(
            text,
            chunk_size=max(50, min(int(chunk_size), 5000)),
            chunk_overlap=max(0, min(int(chunk_overlap), 500)),
        )
