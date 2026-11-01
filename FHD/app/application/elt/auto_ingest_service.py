"""Auto-ingest KB + business domains after reviewed-excel intake.

Product paths reused:
- Knowledge: KnowledgeAdapter / DatasetRag ``office-docking``
- Business: ``detect_etl_target`` → ETL upload → preview → execute
- Withdraw: KnowledgeAdapter.rollback_row + EltService.rollback
"""

from __future__ import annotations

import hashlib
import json
import logging
import threading
from contextvars import copy_context
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

from sqlalchemy.orm import Session

from app.application.elt.errors import EltError
from app.application.elt.target_detection import detect_etl_target
from app.application.shipment_excel_etl_security import (
    ShipmentEtlPathError,
    resolve_etl_path,
)
from app.infrastructure.tenant_scope import tenant_id_for_write
from app.request_active_mod_ctx import get_request_active_mod_id
from app.utils.operational_errors import BOUNDARY_ERRORS, RECOVERABLE_ERRORS
from app.utils.path_io.path_utils import get_app_data_dir

logger = logging.getLogger(__name__)

_LOCK = threading.RLock()


def _utcnow() -> str:
    return datetime.now(UTC).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def _ledger_path(tenant_id: int | str) -> Path:
    root = Path(get_app_data_dir()).resolve() / "etl" / "auto_ingest" / str(tenant_id)
    root.mkdir(parents=True, exist_ok=True)
    return root / "batches.json"


def _load_ledger(tenant_id: int | str) -> dict[str, Any]:
    path = _ledger_path(tenant_id)
    if not path.is_file():
        return {"batches": {}}
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(raw, dict) and isinstance(raw.get("batches"), dict):
            return raw
    except RECOVERABLE_ERRORS:
        logger.warning("auto-ingest ledger load failed: %s", path, exc_info=True)
    return {"batches": {}}


def _save_ledger(tenant_id: int | str, data: dict[str, Any]) -> None:
    path = _ledger_path(tenant_id)
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(path)


def _progress(
    stage: str, pct: int, *, message: str = "", error: str | None = None
) -> dict[str, Any]:
    pct = max(0, min(100, int(pct)))
    out: dict[str, Any] = {
        "stage": stage,
        "pct": pct,
        "message": message,
        "updated_at": _utcnow(),
    }
    if error:
        out["error"] = error
    return out


def _patch_batch(tenant_id: int | str, batch_id: str, **fields: Any) -> dict[str, Any] | None:
    with _LOCK:
        ledger = _load_ledger(tenant_id)
        batch = (ledger.get("batches") or {}).get(batch_id)
        if not isinstance(batch, dict):
            return None
        batch.update(fields)
        batch["updated_at"] = _utcnow()
        ledger.setdefault("batches", {})[batch_id] = batch
        _save_ledger(tenant_id, ledger)
        return dict(batch)


def _extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        lines: list[str] = []
        wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            for ws in wb.worksheets[:20]:
                lines.append(f"# sheet:{ws.title}")
                for index, row in enumerate(ws.iter_rows(values_only=True)):
                    cells = [str(c).strip() for c in row if c not in (None, "")]
                    if cells:
                        lines.append(" | ".join(cells))
                    if index >= 120:
                        break
        finally:
            wb.close()
        return "\n".join(lines)[:80_000]
    return path.read_text(encoding="utf-8", errors="replace")[:80_000]



def _resolve_auto_ingest_source(source_path: str | Path) -> Path:
    """Jail auto-ingest inputs under ETL trusted roots (CodeQL path-injection)."""
    try:
        return resolve_etl_path(source_path)
    except ShipmentEtlPathError as exc:
        raise EltError(
            "ETL_AUTO_INGEST_PATH_INVALID",
            f"自动入库路径无效：{exc}",
        ) from exc


class AutoIngestService:
    """Create an ingest batch: knowledge + matching business domain writes."""

    @staticmethod
    def detect_target(path: Path, labels: list[str] | None = None) -> dict[str, Any]:
        detection = detect_etl_target(path, suffix=path.suffix.lower())
        label_blob = " ".join(labels or [])
        if any(tok in label_blob for tok in ("出货", "送货", "发货", "shipment")):
            detection = {
                **detection,
                "target_type": "shipment_records",
                "reason": "classification_label_shipment",
                "confidence": max(float(detection.get("confidence") or 0), 0.9),
            }
        elif any(tok in label_blob for tok in ("考勤", "attendance")):
            detection = {
                **detection,
                "target_type": "attendance",
                "reason": "classification_label_attendance",
                "confidence": max(float(detection.get("confidence") or 0), 0.9),
            }
        elif any(tok in label_blob for tok in ("物料", "材料", "material")):
            detection = {
                **detection,
                "target_type": "products",
                "reason": "classification_label_material",
                "confidence": max(float(detection.get("confidence") or 0), 0.85),
            }
        return detection

    def knowledge_plan(
        self, *, path: Path, file_name: str, note: str, owner_user_id: int
    ) -> dict[str, Any]:
        from app.application.elt.service_support import dump_json
        from app.application.elt.targets.knowledge import KnowledgeAdapter

        text = _extract_text(path)
        body = text if not note.strip() else f"{note.strip()}\n\n{text}"
        content_hash = hashlib.sha256(body.encode("utf-8")).hexdigest()
        decision = KnowledgeAdapter().preview(
            None,
            {"content": body, "source_key": file_name},
            allowed_update_fields={"content", "document_path"},
            context={
                "owner_user_id": owner_user_id,
                "file_name": file_name,
                "file_sha256": content_hash,
            },
        )
        if decision.action not in {"new", "update", "skip"}:
            raise EltError(
                "ELT_KNOWLEDGE_PREVIEW_FAILED", "知识库预演未通过，未开始写入", status_code=409
            )
        fingerprint = hashlib.sha256(
            dump_json(
                {
                    "action": decision.action,
                    "match_ref": decision.match_ref,
                    "before": decision.before,
                    "content_hash": content_hash,
                }
            ).encode()
        ).hexdigest()
        return {"action": decision.action, "documents": 1, "fingerprint": fingerprint}

    def validate_prepared(
        self,
        db: Session,
        *,
        plan: dict[str, Any],
        path: Path,
        owner_user_id: int,
        tenant_id: int,
        preview_revision: str,
    ) -> dict[str, Any]:
        from app.application.elt.preview_confirmation import verify_preview
        from app.application.elt.service import get_elt_service
        from app.db.models.elt import EltRun
        from app.request_active_mod_ctx import get_request_active_mod_id

        run = (
            db.query(EltRun)
            .filter(
                EltRun.id == plan.get("run_id"),
                EltRun.tenant_id == tenant_id,
                EltRun.owner_user_id == owner_user_id,
            )
            .first()
        )
        if (
            run is None
            or plan.get("mod_id") != get_request_active_mod_id()
            or run.target_type != plan.get("target_type")
            or run.upload_id != plan.get("upload_id")
            or run.file_sha256 != plan.get("candidate_sha256")
            or hashlib.sha256(path.read_bytes()).hexdigest() != plan.get("candidate_sha256")
            or run.status != "preview_ready"
        ):
            raise EltError(
                "ELT_INTAKE_PREVIEW_CHANGED",
                "入库范围或业务空间已变化，请重新预演",
                status_code=409,
            )
        verify_preview(db, run, preview_revision)
        knowledge = self.knowledge_plan(
            path=path, file_name=plan["name"], note=plan["note"], owner_user_id=owner_user_id
        )
        if knowledge != plan.get("knowledge"):
            raise EltError("ELT_PREVIEW_STALE", "知识库已变化，请重新预演并确认", status_code=409)
        return dict(get_elt_service().get_run(db, run_id=run.id, owner_user_id=owner_user_id))

    def run(
        self,
        db: Session,
        *,
        owner_user_id: int,
        tenant_id: int,
        source_path: str | Path,
        file_name: str = "",
        template_id: str = "",
        review_id: str = "",
        classification_id: str = "",
        composer_note: str = "",
        labels: list[str] | None = None,
        batch_id: str | None = None,
        prepared_plan: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        path = _resolve_auto_ingest_source(source_path)
        if not path.is_file():
            raise EltError("ETL_AUTO_INGEST_SOURCE_MISSING", "自动入库找不到源文件")
        safe_name = (file_name or path.name).strip() or path.name
        batch_id = batch_id or f"aib-{uuid4().hex[:16]}"
        run_token = f"auto-ingest:{batch_id}"
        batch = {
            "batch_id": batch_id,
            "status": "running",
            "created_at": _utcnow(),
            "updated_at": _utcnow(),
            "tenant_id": int(tenant_id),
            "mod_id": get_request_active_mod_id(),
            "owner_user_id": int(owner_user_id),
            "template_id": template_id,
            "review_id": review_id,
            "classification_id": classification_id,
            "file_name": safe_name,
            "composer_note": composer_note,
            "detection": {},
            "knowledge": {},
            "business": {},
            "kb_progress": _progress("queued", 0, message="等待写入知识库"),
            "business_progress": _progress("queued", 0, message="等待写入业务域"),
            "withdrawn_at": None,
        }
        with _LOCK:
            ledger = _load_ledger(tenant_id)
            ledger.setdefault("batches", {})[batch_id] = batch
            _save_ledger(tenant_id, ledger)

        try:
            if prepared_plan:
                self.validate_prepared(
                    db,
                    plan=prepared_plan,
                    path=path,
                    owner_user_id=owner_user_id,
                    tenant_id=tenant_id,
                    preview_revision=prepared_plan["preview_revision"],
                )
            _patch_batch(
                tenant_id,
                batch_id,
                kb_progress=_progress("extracting", 10, message="提取文档正文"),
                business_progress=_progress("waiting_kb", 0, message="等待知识库完成"),
            )
            text = _extract_text(path)
            detection = (
                prepared_plan["detection"] if prepared_plan else self.detect_target(path, labels)
            )
            _patch_batch(
                tenant_id,
                batch_id,
                detection=detection,
                kb_progress=_progress("writing", 55, message="写入知识库"),
            )
            content_hash = hashlib.sha256(text.encode("utf-8")).hexdigest()
            kb_receipt = self._ingest_knowledge(
                text=text,
                source_key=safe_name,
                run_id=run_token,
                owner_user_id=owner_user_id,
                content_hash=content_hash,
                note=composer_note,
                tenant_id=tenant_id,
                expected_plan=prepared_plan.get("knowledge") if prepared_plan else None,
            )
            kb_done = _progress(
                "done" if kb_receipt.get("action") != "skip" else "skipped",
                100,
                message=str(kb_receipt.get("reason") or "知识库已写入"),
            )
            _patch_batch(
                tenant_id,
                batch_id,
                knowledge=kb_receipt,
                kb_progress=kb_done,
                business_progress=_progress("uploading", 15, message="上传业务文件"),
            )
            business = self._ingest_business(
                db,
                owner_user_id=owner_user_id,
                path=path,
                file_name=safe_name,
                target_type=str(detection.get("target_type") or "customer_products"),
                batch_id=batch_id,
                prepared_plan=prepared_plan,
                progress_cb=lambda stage, pct, message="": _patch_batch(
                    tenant_id,
                    batch_id,
                    business_progress=_progress(stage, pct, message=message),
                ),
            )
            if business.get("skipped"):
                biz_prog = _progress(
                    "skipped", 100, message=str(business.get("reason") or "跳过业务域")
                )
            elif business.get("failed") or business.get("pending_preview"):
                biz_prog = _progress(
                    "failed" if business.get("failed") else "pending",
                    70 if business.get("pending_preview") else 40,
                    message=str(
                        business.get("message") or business.get("error") or "业务写入未完成"
                    ),
                    error=str(business.get("error") or "") or None,
                )
            elif business.get("pending_execution"):
                biz_prog = _progress("executing", 80, message="业务写入处理中，请等待实际回执")
            else:
                biz_prog = _progress("done", 100, message="业务域已写入")
            final = _patch_batch(
                tenant_id,
                batch_id,
                status=(
                    "running"
                    if business.get("pending_execution")
                    else "awaiting_confirmation"
                    if business.get("pending_preview")
                    else "failed"
                    if business.get("failed")
                    else "applied"
                ),
                detection=detection,
                knowledge=kb_receipt,
                business=business,
                kb_progress=kb_done,
                business_progress=biz_prog,
            )
            return final or batch
        except (
            BOUNDARY_ERRORS
        ) as exc:  # Never erase a successful knowledge receipt after a business failure.
            with _LOCK:
                saved = (_load_ledger(tenant_id).get("batches") or {}).get(batch_id) or {}
            kb_progress = (
                saved.get("kb_progress")
                if saved.get("knowledge")
                else _progress(
                    "failed", 0, message="知识入库未完成，请核对结果", error=type(exc).__name__
                )
            )
            _patch_batch(
                tenant_id,
                batch_id,
                status="failed",
                kb_progress=kb_progress,
                business_progress=_progress(
                    "failed",
                    0,
                    message="业务写入未完成，请核对已写入部分，勿重复提交",
                    error=type(exc).__name__,
                ),
                error=type(exc).__name__,
            )
            if isinstance(exc, EltError):
                raise
            raise EltError("ETL_AUTO_INGEST_FAILED", f"自动入库失败：{type(exc).__name__}") from exc

    def start_background(
        self,
        *,
        owner_user_id: int,
        tenant_id: int,
        source_path: str | Path,
        file_name: str = "",
        template_id: str = "",
        review_id: str = "",
        classification_id: str = "",
        composer_note: str = "",
        labels: list[str] | None = None,
        prepared_plan: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        """Create ledger row and finish ingest off-request so FE can poll true progress."""
        path = _resolve_auto_ingest_source(source_path)
        if not path.is_file():
            raise EltError("ETL_AUTO_INGEST_SOURCE_MISSING", "自动入库找不到源文件")
        batch_id = f"aib-{uuid4().hex[:16]}"
        safe_name = (file_name or path.name).strip() or path.name
        batch = {
            "batch_id": batch_id,
            "status": "running",
            "created_at": _utcnow(),
            "updated_at": _utcnow(),
            "tenant_id": int(tenant_id),
            "mod_id": get_request_active_mod_id(),
            "owner_user_id": int(owner_user_id),
            "template_id": template_id,
            "review_id": review_id,
            "classification_id": classification_id,
            "file_name": safe_name,
            "composer_note": composer_note,
            "detection": {},
            "knowledge": {},
            "business": {},
            "kb_progress": _progress("queued", 0, message="排队写入知识库"),
            "business_progress": _progress("queued", 0, message="排队写入业务域"),
            "withdrawn_at": None,
            "source_path": str(path),
        }
        with _LOCK:
            ledger = _load_ledger(tenant_id)
            ledger.setdefault("batches", {})[batch_id] = batch
            _save_ledger(tenant_id, ledger)

        def _worker() -> None:
            from app.db.session import get_db
            from app.infrastructure.tenant_scope import tenant_scope

            try:
                with tenant_scope(tenant_id), get_db() as db:
                    self.run(
                        db,
                        owner_user_id=owner_user_id,
                        tenant_id=tenant_id,
                        source_path=path,
                        file_name=safe_name,
                        template_id=template_id,
                        review_id=review_id,
                        classification_id=classification_id,
                        composer_note=composer_note,
                        labels=labels,
                        batch_id=batch_id,
                        prepared_plan=prepared_plan,
                    )
                    db.commit()
            except BOUNDARY_ERRORS:
                logger.exception("background auto-ingest failed batch_id=%s", batch_id)

        context = copy_context()
        threading.Thread(
            target=lambda: context.run(_worker), name=f"auto-ingest-{batch_id}", daemon=True
        ).start()
        return dict(batch)

    def _refresh_business(self, batch: dict[str, Any]) -> dict[str, Any]:
        if batch.get("mod_id") != get_request_active_mod_id():
            return batch
        business = batch.get("business") or {}
        if batch.get("status") not in {
            "running",
            "awaiting_confirmation",
            "failed",
        } or not business.get("run_id"):
            return batch
        from app.application.elt.service import get_elt_service
        from app.db.session import get_db
        from app.infrastructure.tenant_scope import tenant_scope

        try:
            with tenant_scope(int(batch["tenant_id"])), get_db() as db:
                run = get_elt_service().get_run(
                    db, run_id=business["run_id"], owner_user_id=int(batch["owner_user_id"])
                )
            status = str(run.get("status") or "")
            if status == "completed":
                batch_status, progress = "applied", _progress("done", 100, message="业务写入已完成")
            elif status in {"failed", "interrupted", "outcome_unknown"}:
                batch_status, progress = (
                    "failed",
                    _progress(
                        "failed",
                        int(run.get("progress") or 0),
                        message="业务写入失败或结果待核对；请核对已写入部分，勿重复提交",
                    ),
                )
            elif status == "preview_ready":
                batch_status, progress = (
                    "awaiting_confirmation",
                    _progress("pending", 70, message="业务预演待确认；请到历史导入任务核对"),
                )
            else:
                batch_status, progress = (
                    "running",
                    _progress(
                        "executing",
                        int(run.get("progress") or 0),
                        message="业务处理仍在进行，请等待实际回执",
                    ),
                )
            with _LOCK:
                ledger = _load_ledger(batch["tenant_id"])
                latest = (ledger.get("batches") or {}).get(batch["batch_id"])
                if (
                    not isinstance(latest, dict)
                    or latest.get("status") == "withdrawn"
                    or latest.get("status") != batch.get("status")
                    or latest.get("business") != batch.get("business")
                ):
                    return dict(latest) if isinstance(latest, dict) else batch
                latest.update(
                    status=batch_status,
                    business={
                        **business,
                        "status": status,
                        "summary": run.get("summary"),
                        "error": run.get("error"),
                    },
                    business_progress=progress,
                    updated_at=_utcnow(),
                )
                _save_ledger(batch["tenant_id"], ledger)
                return dict(latest)
        except BOUNDARY_ERRORS:
            # Reading a receipt must never resubmit execution or erase a known partial write.
            return batch

    def list_batches(
        self, *, tenant_id: int, owner_user_id: int | None = None, limit: int = 50
    ) -> list[dict[str, Any]]:
        with _LOCK:
            ledger = _load_ledger(tenant_id)
        items = [
            b
            for b in (ledger.get("batches") or {}).values()
            if b.get("mod_id") == get_request_active_mod_id()
        ]
        if owner_user_id is not None:
            items = [b for b in items if int(b.get("owner_user_id") or 0) == int(owner_user_id)]
        items.sort(key=lambda b: str(b.get("created_at") or ""), reverse=True)
        return [
            self._refresh_business(dict(batch))
            for batch in items[: max(1, min(int(limit or 50), 200))]
        ]

    def get_batch(
        self, *, tenant_id: int, batch_id: str, owner_user_id: int | None = None
    ) -> dict[str, Any] | None:
        with _LOCK:
            ledger = _load_ledger(tenant_id)
        item = (ledger.get("batches") or {}).get(batch_id)
        if not isinstance(item, dict) or item.get("mod_id") != get_request_active_mod_id():
            return None
        if owner_user_id is not None:
            if int(item.get("owner_user_id") or 0) != owner_user_id:
                raise EltError("ETL_AUTO_INGEST_FORBIDDEN", "无权查看该批次", status_code=403)
            return self._refresh_business(dict(item))
        return dict(item)

    def get_progress(self, *, tenant_id: int, owner_user_id: int, batch_id: str) -> dict[str, Any]:
        batch = self.get_batch(tenant_id=tenant_id, batch_id=batch_id, owner_user_id=owner_user_id)
        if batch is None:
            raise EltError("ETL_AUTO_INGEST_NOT_FOUND", "入库批次不存在", status_code=404)
        if int(batch.get("owner_user_id") or 0) != int(owner_user_id):
            raise EltError("ETL_AUTO_INGEST_FORBIDDEN", "无权查看该批次", status_code=403)
        kb = (
            batch.get("kb_progress")
            if isinstance(batch.get("kb_progress"), dict)
            else _progress("unknown", 0)
        )
        biz = (
            batch.get("business_progress")
            if isinstance(batch.get("business_progress"), dict)
            else _progress("unknown", 0)
        )
        return {
            "batch_id": batch_id,
            "status": batch.get("status"),
            "kb_progress": kb,
            "business_progress": biz,
            "updated_at": batch.get("updated_at"),
        }

    def withdraw(
        self,
        db: Session,
        *,
        tenant_id: int,
        owner_user_id: int,
        batch_id: str,
    ) -> dict[str, Any]:
        with _LOCK:
            ledger = _load_ledger(tenant_id)
            batch = (ledger.get("batches") or {}).get(batch_id)
            if not isinstance(batch, dict):
                raise EltError("ETL_AUTO_INGEST_NOT_FOUND", "入库批次不存在", status_code=404)
            if batch.get("mod_id") != get_request_active_mod_id():
                raise EltError(
                    "ELT_AUTO_INGEST_FORBIDDEN",
                    "批次业务空间未核实或已变化，不能撤回此批次",
                    status_code=403,
                )
            if int(batch.get("owner_user_id") or 0) != int(owner_user_id):
                raise EltError(
                    "ETL_AUTO_INGEST_FORBIDDEN", "只能撤回自己的入库批次", status_code=403
                )
            if batch.get("status") == "withdrawn":
                return batch
            if batch.get("status") != "applied":
                raise EltError(
                    "ETL_AUTO_INGEST_BAD_STATE",
                    f"当前状态不可撤回：{batch.get('status')}",
                    status_code=400,
                )

        kb_rollback = self._withdraw_knowledge(batch)
        biz_rollback = self._withdraw_business(db, batch=batch, owner_user_id=owner_user_id)

        with _LOCK:
            ledger = _load_ledger(tenant_id)
            current = (ledger.get("batches") or {}).get(batch_id) or batch
            current["status"] = "withdrawn"
            current["withdrawn_at"] = _utcnow()
            current["updated_at"] = current["withdrawn_at"]
            current["withdraw"] = {"knowledge": kb_rollback, "business": biz_rollback}
            ledger.setdefault("batches", {})[batch_id] = current
            _save_ledger(tenant_id, ledger)
            return current

    def _ingest_knowledge(
        self,
        *,
        text: str,
        source_key: str,
        run_id: str,
        owner_user_id: int,
        content_hash: str,
        note: str,
        tenant_id: int,
        expected_plan: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        from app.application.elt.targets.knowledge import KnowledgeAdapter
        from app.infrastructure.tenant_scope import tenant_id_for_write, tenant_scope

        body = text if not note.strip() else f"{note.strip()}\n\n{text}"
        adapter = KnowledgeAdapter()
        data = {
            "content": body,
            "source_key": source_key,
        }
        context = {
            "owner_user_id": owner_user_id,
            "run_id": run_id,
            "file_name": source_key,
            "file_sha256": content_hash,
        }
        with tenant_scope(tenant_id):
            decision = adapter.preview(
                None, data, allowed_update_fields={"content", "document_path"}, context=context
            )
            if expected_plan:
                from app.application.elt.service_support import dump_json

                fingerprint = hashlib.sha256(
                    dump_json(
                        {
                            "action": decision.action,
                            "match_ref": decision.match_ref,
                            "before": decision.before,
                            "content_hash": hashlib.sha256(body.encode("utf-8")).hexdigest(),
                        }
                    ).encode()
                ).hexdigest()
                if fingerprint != expected_plan.get("fingerprint"):
                    raise EltError("ELT_PREVIEW_STALE", "知识库已变化，请重新预演", status_code=409)
            if decision.action == "skip":
                return {
                    "action": "skip",
                    "reason": decision.reason,
                    "match_ref": decision.match_ref,
                    "run_id": run_id,
                    "content_hash": content_hash,
                }
            result = adapter.execute_row(
                None,
                data,
                action=decision.action,
                match_ref=decision.match_ref or "",
                allowed_update_fields={"content", "document_path"},
                context={**context, "preview_before": decision.before},
            )
            return {
                "action": decision.action,
                "reason": decision.reason,
                "match_ref": result.get("match_ref"),
                "after": result.get("after"),
                "run_id": run_id,
                "content_hash": content_hash,
                "tenant_id": str(tenant_id_for_write()),
            }

    def _withdraw_knowledge(self, batch: dict[str, Any]) -> dict[str, Any]:
        from app.application.elt.targets.knowledge import KnowledgeAdapter
        from app.infrastructure.tenant_scope import tenant_scope

        kb = batch.get("knowledge") or {}
        after = kb.get("after") or {}
        match_ref = str(kb.get("match_ref") or after.get("document_id") or "")
        if not match_ref or kb.get("action") == "skip":
            return {"skipped": True, "reason": "no_knowledge_receipt"}
        adapter = KnowledgeAdapter()
        with tenant_scope(int(batch.get("tenant_id") or 0)):
            adapter.rollback_row(
                None,
                match_ref=match_ref,
                before=None,
                after={
                    "document_id": match_ref,
                    "version": after.get("version"),
                    "content_hash": after.get("content_hash") or kb.get("content_hash"),
                },
                context={"run_id": kb.get("run_id"), "owner_user_id": batch.get("owner_user_id")},
            )
        return {"rolled_back": True, "match_ref": match_ref}

    def _ingest_business(
        self,
        db: Session,
        *,
        owner_user_id: int,
        path: Path,
        file_name: str,
        target_type: str,
        batch_id: str,
        poll_seconds: float = 30.0,
        prepared_plan: dict[str, Any] | None = None,
        progress_cb=None,
    ) -> dict[str, Any]:
        import time

        from app.application.elt.service import get_elt_service

        def _notify(stage: str, pct: int, message: str = "") -> None:
            if progress_cb is None:
                return
            try:
                progress_cb(stage, pct, message)
            except RECOVERABLE_ERRORS:
                logger.debug("auto-ingest progress_cb failed", exc_info=True)

        # Knowledge-only targets: KB path already covered.
        if target_type == "knowledge":
            _notify("skipped", 100, "knowledge_only_target")
            return {"skipped": True, "reason": "knowledge_only_target", "target_type": target_type}

        service = get_elt_service()
        if prepared_plan:
            run_id = str(prepared_plan["run_id"])
            upload_id = str(prepared_plan["upload_id"])
            current = service.get_run(db, run_id=run_id, owner_user_id=owner_user_id)
            if (
                current.get("target_type") != target_type
                or current.get("preview_revision") != prepared_plan["preview_revision"]
            ):
                raise EltError(
                    "ELT_PREVIEW_STALE", "已确认的业务预演已变化，未执行", status_code=409
                )
        else:
            _notify("uploading", 20, "上传业务文件")
            with path.open("rb") as handle:
                upload = service.save_upload(
                    db,
                    owner_user_id=owner_user_id,
                    file_name=file_name,
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    stream=handle,
                    batch_id=batch_id,
                    relative_path=file_name,
                )
            upload_id = str(upload.get("upload_id") or "")
            _notify("preview", 45, "生成业务预演")
            run = service.create_preview(
                db,
                owner_user_id=owner_user_id,
                upload_id=upload_id,
                target_type=target_type,
            )
            run_id = str(run.get("id") or "")
            deadline = time.monotonic() + max(1.0, float(poll_seconds))
            current = run
            while time.monotonic() < deadline:
                status = str(current.get("status") or "")
                if status == "preview_ready":
                    break
                if status in {"failed", "interrupted"}:
                    return {
                        "target_type": target_type,
                        "upload_id": upload_id,
                        "run_id": run_id,
                        "status": status,
                        "error": current.get("error"),
                        "failed": True,
                    }
                time.sleep(0.2)
                current = service.get_run(db, run_id=run_id, owner_user_id=owner_user_id)
            if str(current.get("status") or "") != "preview_ready":
                return {
                    "target_type": target_type,
                    "upload_id": upload_id,
                    "run_id": run_id,
                    "status": current.get("status"),
                    "pending_preview": True,
                    "message": "业务预演尚未完成；知识库已写入，可稍后在历史任务中执行或撤回",
                }
        preview_revision = current.get("preview_revision")
        _notify("executing", 80, "执行业务写入")
        _patch_batch(
            tenant_id_for_write(),
            batch_id,
            business={
                "run_id": run_id,
                "upload_id": upload_id,
                "target_type": target_type,
                "status": "executing",
            },
        )
        executed = service.execute(
            db,
            run_id=run_id,
            owner_user_id=owner_user_id,
            confirmed=True,
            valid_rows_only=bool(prepared_plan.get("valid_rows_only")) if prepared_plan else True,
            preview_revision=str(preview_revision or ""),
        )
        return {
            "target_type": target_type,
            "upload_id": upload_id,
            "run_id": run_id,
            "preview_revision": preview_revision,
            "summary": executed.get("summary") if isinstance(executed, dict) else None,
            "status": executed.get("status") if isinstance(executed, dict) else None,
            "pending_execution": isinstance(executed, dict)
            and executed.get("status") in {"queued", "executing"},
            "failed": not isinstance(executed, dict)
            or executed.get("status") in {"failed", "interrupted", "outcome_unknown"},
        }

    def _withdraw_business(
        self, db: Session, *, batch: dict[str, Any], owner_user_id: int
    ) -> dict[str, Any]:
        from app.application.elt.service import get_elt_service

        business = batch.get("business") or {}
        run_id = str(business.get("run_id") or "")
        if not run_id or business.get("skipped"):
            return {"skipped": True, "reason": "no_business_run"}
        rolled = get_elt_service().rollback(db, run_id=run_id, owner_user_id=owner_user_id)
        return {
            "rolled_back": True,
            "run_id": run_id,
            "status": rolled.get("status") if isinstance(rolled, dict) else None,
            "rollback_status": rolled.get("rollback_status") if isinstance(rolled, dict) else None,
        }


_SERVICE: AutoIngestService | None = None


def get_auto_ingest_service() -> AutoIngestService:
    global _SERVICE
    if _SERVICE is None:
        _SERVICE = AutoIngestService()
    return _SERVICE


__all__ = ["AutoIngestService", "get_auto_ingest_service"]
