"""Auto-ingest KB + business domains after reviewed-excel intake.

Product paths reused:
- Knowledge: KnowledgeAdapter / DatasetRag ``office-docking``
- Business: ``detect_etl_target`` → ETL upload → preview → execute
- Withdraw: KnowledgeAdapter.rollback_row + EtlService.rollback
"""

from __future__ import annotations

import hashlib
import json
import logging
import threading
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

from sqlalchemy.orm import Session

from app.application.etl.errors import EtlError
from app.application.etl.target_detection import detect_etl_target
from app.utils.operational_errors import RECOVERABLE_ERRORS
from app.utils.path_io.path_utils import get_app_data_dir

logger = logging.getLogger(__name__)

_LOCK = threading.RLock()


def _utcnow() -> str:
    return datetime.now(UTC).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def _ledger_path(tenant_id: int | str) -> Path:
    root = Path(get_app_data_dir()).resolve() / "etl" / "auto_ingest" / str(tenant_id)
    root.mkdir(parents=True, exist_ok=True)
    return root / "batches.json"


def _load_ledger(tenant_id: int | str) -> dict[str, Any]:
    path = _ledger_path(tenant_id)
    if not path.is_file():
        return {"batches": {}}
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(raw, dict) and isinstance(raw.get("batches"), dict):
            return raw
    except RECOVERABLE_ERRORS:
        logger.warning("auto-ingest ledger load failed: %s", path, exc_info=True)
    return {"batches": {}}


def _save_ledger(tenant_id: int | str, data: dict[str, Any]) -> None:
    path = _ledger_path(tenant_id)
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(path)


def _extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        lines: list[str] = []
        wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            for ws in wb.worksheets[:20]:
                lines.append(f"# sheet:{ws.title}")
                for index, row in enumerate(ws.iter_rows(values_only=True)):
                    cells = [str(c).strip() for c in row if c not in (None, "")]
                    if cells:
                        lines.append(" | ".join(cells))
                    if index >= 120:
                        break
        finally:
            wb.close()
        return "\n".join(lines)[:80_000]
    return path.read_text(encoding="utf-8", errors="replace")[:80_000]


class AutoIngestService:
    """Create an ingest batch: knowledge + matching business domain writes."""

    def run(
        self,
        db: Session,
        *,
        owner_user_id: int,
        tenant_id: int,
        source_path: str | Path,
        file_name: str = "",
        template_id: str = "",
        review_id: str = "",
        classification_id: str = "",
        composer_note: str = "",
        labels: list[str] | None = None,
    ) -> dict[str, Any]:
        path = Path(source_path).expanduser().resolve()
        if not path.is_file():
            raise EtlError("ETL_AUTO_INGEST_SOURCE_MISSING", "自动入库找不到源文件")
        safe_name = (file_name or path.name).strip() or path.name
        batch_id = f"aib-{uuid4().hex[:16]}"
        run_token = f"auto-ingest:{batch_id}"
        text = _extract_text(path)
        detection = detect_etl_target(path, suffix=path.suffix.lower())
        # Prefer explicit labels for domain hint
        label_blob = " ".join(labels or [])
        if any(tok in label_blob for tok in ("出货", "送货", "发货", "shipment")):
            detection = {
                **detection,
                "target_type": "shipment_records",
                "reason": "classification_label_shipment",
                "confidence": max(float(detection.get("confidence") or 0), 0.9),
            }
        elif any(tok in label_blob for tok in ("考勤", "attendance")):
            detection = {
                **detection,
                "target_type": "attendance",
                "reason": "classification_label_attendance",
                "confidence": max(float(detection.get("confidence") or 0), 0.9),
            }
        elif any(tok in label_blob for tok in ("物料", "材料", "material")):
            detection = {
                **detection,
                "target_type": "products",
                "reason": "classification_label_material",
                "confidence": max(float(detection.get("confidence") or 0), 0.85),
            }

        content_hash = hashlib.sha256(text.encode("utf-8")).hexdigest()
        kb_receipt = self._ingest_knowledge(
            text=text,
            source_key=safe_name,
            run_id=run_token,
            owner_user_id=owner_user_id,
            content_hash=content_hash,
            note=composer_note,
            tenant_id=tenant_id,
        )
        business = self._ingest_business(
            db,
            owner_user_id=owner_user_id,
            path=path,
            file_name=safe_name,
            target_type=str(detection.get("target_type") or "customer_products"),
            batch_id=batch_id,
        )

        batch = {
            "batch_id": batch_id,
            "status": "applied",
            "created_at": _utcnow(),
            "updated_at": _utcnow(),
            "tenant_id": int(tenant_id),
            "owner_user_id": int(owner_user_id),
            "template_id": template_id,
            "review_id": review_id,
            "classification_id": classification_id,
            "file_name": safe_name,
            "composer_note": composer_note,
            "detection": detection,
            "knowledge": kb_receipt,
            "business": business,
            "withdrawn_at": None,
        }
        with _LOCK:
            ledger = _load_ledger(tenant_id)
            ledger.setdefault("batches", {})[batch_id] = batch
            _save_ledger(tenant_id, ledger)
        return batch

    def list_batches(
        self, *, tenant_id: int, owner_user_id: int | None = None, limit: int = 50
    ) -> list[dict[str, Any]]:
        with _LOCK:
            ledger = _load_ledger(tenant_id)
        items = list((ledger.get("batches") or {}).values())
        if owner_user_id is not None:
            items = [b for b in items if int(b.get("owner_user_id") or 0) == int(owner_user_id)]
        items.sort(key=lambda b: str(b.get("created_at") or ""), reverse=True)
        return items[: max(1, min(int(limit or 50), 200))]

    def get_batch(self, *, tenant_id: int, batch_id: str) -> dict[str, Any] | None:
        with _LOCK:
            ledger = _load_ledger(tenant_id)
        item = (ledger.get("batches") or {}).get(batch_id)
        return dict(item) if isinstance(item, dict) else None

    def withdraw(
        self,
        db: Session,
        *,
        tenant_id: int,
        owner_user_id: int,
        batch_id: str,
    ) -> dict[str, Any]:
        with _LOCK:
            ledger = _load_ledger(tenant_id)
            batch = (ledger.get("batches") or {}).get(batch_id)
            if not isinstance(batch, dict):
                raise EtlError("ETL_AUTO_INGEST_NOT_FOUND", "入库批次不存在", status_code=404)
            if int(batch.get("owner_user_id") or 0) != int(owner_user_id):
                raise EtlError("ETL_AUTO_INGEST_FORBIDDEN", "只能撤回自己的入库批次", status_code=403)
            if batch.get("status") == "withdrawn":
                return batch
            if batch.get("status") != "applied":
                raise EtlError(
                    "ETL_AUTO_INGEST_BAD_STATE",
                    f"当前状态不可撤回：{batch.get('status')}",
                    status_code=400,
                )

        kb_rollback = self._withdraw_knowledge(batch)
        biz_rollback = self._withdraw_business(db, batch=batch, owner_user_id=owner_user_id)

        with _LOCK:
            ledger = _load_ledger(tenant_id)
            current = (ledger.get("batches") or {}).get(batch_id) or batch
            current["status"] = "withdrawn"
            current["withdrawn_at"] = _utcnow()
            current["updated_at"] = current["withdrawn_at"]
            current["withdraw"] = {"knowledge": kb_rollback, "business": biz_rollback}
            ledger.setdefault("batches", {})[batch_id] = current
            _save_ledger(tenant_id, ledger)
            return current

    def _ingest_knowledge(
        self,
        *,
        text: str,
        source_key: str,
        run_id: str,
        owner_user_id: int,
        content_hash: str,
        note: str,
        tenant_id: int,
    ) -> dict[str, Any]:
        from app.application.etl.targets.knowledge import KnowledgeAdapter
        from app.infrastructure.tenant_scope import tenant_id_for_write, tenant_scope

        body = text if not note.strip() else f"{note.strip()}\n\n{text}"
        adapter = KnowledgeAdapter()
        data = {
            "content": body,
            "source_key": source_key,
        }
        context = {
            "owner_user_id": owner_user_id,
            "run_id": run_id,
            "file_name": source_key,
            "file_sha256": content_hash,
        }
        with tenant_scope(tenant_id):
            decision = adapter.preview(
                None, data, allowed_update_fields={"content", "document_path"}, context=context
            )
            if decision.action == "skip":
                return {
                    "action": "skip",
                    "reason": decision.reason,
                    "match_ref": decision.match_ref,
                    "run_id": run_id,
                    "content_hash": content_hash,
                }
            result = adapter.execute_row(
                None,
                data,
                action=decision.action,
                match_ref=decision.match_ref or "",
                allowed_update_fields={"content", "document_path"},
                context={**context, "preview_before": decision.before},
            )
            return {
                "action": decision.action,
                "reason": decision.reason,
                "match_ref": result.get("match_ref"),
                "after": result.get("after"),
                "run_id": run_id,
                "content_hash": content_hash,
                "tenant_id": str(tenant_id_for_write()),
            }

    def _withdraw_knowledge(self, batch: dict[str, Any]) -> dict[str, Any]:
        from app.application.etl.targets.knowledge import KnowledgeAdapter
        from app.infrastructure.tenant_scope import tenant_scope

        kb = batch.get("knowledge") or {}
        after = kb.get("after") or {}
        match_ref = str(kb.get("match_ref") or after.get("document_id") or "")
        if not match_ref or kb.get("action") == "skip":
            return {"skipped": True, "reason": "no_knowledge_receipt"}
        adapter = KnowledgeAdapter()
        with tenant_scope(int(batch.get("tenant_id") or 0)):
            adapter.rollback_row(
                None,
                match_ref=match_ref,
                before=None,
                after={
                    "document_id": match_ref,
                    "version": after.get("version"),
                    "content_hash": after.get("content_hash") or kb.get("content_hash"),
                },
                context={"run_id": kb.get("run_id"), "owner_user_id": batch.get("owner_user_id")},
            )
        return {"rolled_back": True, "match_ref": match_ref}

    def _ingest_business(
        self,
        db: Session,
        *,
        owner_user_id: int,
        path: Path,
        file_name: str,
        target_type: str,
        batch_id: str,
        poll_seconds: float = 30.0,
    ) -> dict[str, Any]:
        import time

        from app.application.etl.service import get_etl_service

        # Knowledge-only targets: KB path already covered.
        if target_type == "knowledge":
            return {"skipped": True, "reason": "knowledge_only_target", "target_type": target_type}

        service = get_etl_service()
        with path.open("rb") as handle:
            upload = service.save_upload(
                db,
                owner_user_id=owner_user_id,
                file_name=file_name,
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                stream=handle,
                batch_id=batch_id,
                relative_path=file_name,
            )
        upload_id = str(upload.get("upload_id") or "")
        run = service.create_preview(
            db,
            owner_user_id=owner_user_id,
            upload_id=upload_id,
            target_type=target_type,
        )
        run_id = str(run.get("id") or "")
        deadline = time.monotonic() + max(1.0, float(poll_seconds))
        current = run
        while time.monotonic() < deadline:
            status = str(current.get("status") or "")
            if status == "preview_ready":
                break
            if status in {"failed", "interrupted"}:
                return {
                    "target_type": target_type,
                    "upload_id": upload_id,
                    "run_id": run_id,
                    "status": status,
                    "error": current.get("error"),
                    "failed": True,
                }
            time.sleep(0.2)
            current = service.get_run(db, run_id=run_id, owner_user_id=owner_user_id)
        if str(current.get("status") or "") != "preview_ready":
            return {
                "target_type": target_type,
                "upload_id": upload_id,
                "run_id": run_id,
                "status": current.get("status"),
                "pending_preview": True,
                "message": "业务预演尚未完成；知识库已写入，可稍后在历史任务中执行或撤回",
            }
        preview_revision = current.get("preview_revision")
        executed = service.execute(
            db,
            run_id=run_id,
            owner_user_id=owner_user_id,
            confirmed=True,
            valid_rows_only=True,
            preview_revision=str(preview_revision or ""),
        )
        return {
            "target_type": target_type,
            "upload_id": upload_id,
            "run_id": run_id,
            "preview_revision": preview_revision,
            "summary": executed.get("summary") if isinstance(executed, dict) else None,
            "status": executed.get("status") if isinstance(executed, dict) else None,
        }

    def _withdraw_business(
        self, db: Session, *, batch: dict[str, Any], owner_user_id: int
    ) -> dict[str, Any]:
        from app.application.etl.service import get_etl_service

        business = batch.get("business") or {}
        run_id = str(business.get("run_id") or "")
        if not run_id or business.get("skipped"):
            return {"skipped": True, "reason": "no_business_run"}
        rolled = get_etl_service().rollback(db, run_id=run_id, owner_user_id=owner_user_id)
        return {
            "rolled_back": True,
            "run_id": run_id,
            "status": rolled.get("status") if isinstance(rolled, dict) else None,
            "rollback_status": rolled.get("rollback_status") if isinstance(rolled, dict) else None,
        }


_SERVICE: AutoIngestService | None = None


def get_auto_ingest_service() -> AutoIngestService:
    global _SERVICE
    if _SERVICE is None:
        _SERVICE = AutoIngestService()
    return _SERVICE


__all__ = ["AutoIngestService", "get_auto_ingest_service"]
