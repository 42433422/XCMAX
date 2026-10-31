"""Rule-based Excel workbook labeling, without model inference."""

from __future__ import annotations

from openpyxl.utils import get_column_letter

from app.application.reviewed_excel_evidence import bounded_range


def classify_review_workbook(review: dict) -> list[dict]:
    """Suggest non-exclusive document labels from inspected cell text only."""
    import re

    from openpyxl.utils.cell import coordinate_from_string

    from app.application.elt.target_detection import _ATTENDANCE_RE, _DELIVERY_RE, _PURCHASE_RE

    patterns = {
        "quotation": re.compile(r"报价|報價|quotation|sales quote", re.I),
        "delivery": _DELIVERY_RE,
        "purchase": _PURCHASE_RE,
        "attendance": _ATTENDANCE_RE,
    }
    results = []
    for sheet in review["sheets"]:
        cells = [cell for cell in review["cells"] if cell["sheet"] == sheet]
        evidence = []
        for label, pattern in patterns.items():
            for cell in cells:
                match = pattern.search(cell["value"])
                if match is not None:
                    evidence.append(
                        {"label": label, "cell": cell["cell"], "matched": match.group(0)}
                    )
                    break
        headings = {
            "quotation": re.compile(r"^(报价单|報價單|quotation|sales quote)$", re.I),
            "delivery": re.compile(r"^(送货单|发货单|出货单|delivery note)$", re.I),
            "purchase": re.compile(r"^(采购订单|purchase order)$", re.I),
            "attendance": re.compile(r"^(考勤表|attendance sheet)$", re.I),
        }
        anchors: dict[int, list[dict]] = {}
        for cell in cells:
            for label, pattern in headings.items():
                if pattern.fullmatch(cell["value"].strip()):
                    row = coordinate_from_string(cell["cell"])[1]
                    anchors.setdefault(row, []).append(
                        {"label": label, "cell": cell["cell"], "matched": cell["value"]}
                    )
        dimensions: dict = next(
            (item for item in review["sheet_overview"] if item["sheet"] == sheet), {}
        )
        max_row = max(1, int(dimensions.get("rows", 1)))
        max_col = max(1, int(dimensions.get("columns", 1)))
        merges = dimensions.get("merged_regions", [])
        section_headings = {
            "production": r"生产(?:任务|计划)?单|production (?:order|sheet)",
            "process": r"工艺(?:流程|说明|要求)?|工序(?:说明|要求)?|process(?: instructions?| sheet)?",
            "material": r"原(?:辅)?料(?:清单|明细)?|物料(?:清单|明细)|materials?(?: list)?|bill of materials",
            "inspection": r"检验(?:说明|要求|标准|记录|结果)?|质检(?:说明|要求|标准|记录)?|inspection(?: instructions?| standards?| report)?",
        }
        rows_of_cells: dict[int, list[dict]] = {}
        for cell in cells:
            rows_of_cells.setdefault(coordinate_from_string(cell["cell"])[1], []).append(cell)
        merged_titles = {
            f"{get_column_letter(c1)}{r1}"
            for c1, r1, c2, _r2 in map(bounded_range, merges)
            if c2 > c1
        }
        for cell in cells:
            column, row = coordinate_from_string(cell["cell"])
            merged_title = cell["cell"] in merged_titles
            sparse_title = (
                column == "A"
                and len(rows_of_cells[row]) == 1
                and any(
                    len(rows_of_cells.get(next_row, [])) >= 2 for next_row in (row + 1, row + 2)
                )
            )
            if cell.get("kind") == "formula" or not (merged_title or sparse_title):
                continue
            value = cell["value"].strip().rstrip(":：")
            for label, section_pattern in section_headings.items():
                if re.fullmatch(section_pattern, value, re.I):
                    item = {"label": label, "cell": cell["cell"], "matched": cell["value"]}
                    anchors.setdefault(row, []).append(item)
                    if not any(ref["label"] == label for ref in evidence):
                        evidence.append(item)
        rows = sorted(anchors)
        if (
            len(rows) > 1
            and {item["label"] for item in anchors[rows[0]]} == {"production"}
            and any(
                item["label"] in {"process", "material", "inspection"} for item in anchors[rows[1]]
            )
        ):
            # A workbook title remains in the first section, rather than becoming a title-only slice.
            rows = rows[1:]
        split = len(rows) > 1 and all(b - a >= 3 for a, b in zip(rows, rows[1:]))
        if split and any(
            min_row < boundary <= max_row_of_merge
            for merged in merges
            for _min_col, min_row, _max_col, max_row_of_merge in [bounded_range(merged)]
            for boundary in rows[1:]
        ):
            split = False
        possible_shared = [
            {"cell": cell["cell"], "kind": "可能的图例或说明，关联待确认"}
            for cell in cells
            if re.search(r"图例|圖例|legend|说明|說明|备注|備註", cell["value"], re.I)
        ]
        header_pattern = re.compile(
            r"^(产品名称|商品名称|品名|数量|单价|金额|合计|客户|供应商|姓名|日期|product|quantity|price|amount|customer|supplier|date)$",
            re.I,
        )
        boundaries = [1] + rows[1:] if split else [1]
        regions = []
        for index, start in enumerate(boundaries):
            end = boundaries[index + 1] - 1 if index + 1 < len(boundaries) else max_row
            region_cells = [
                cell for cell in cells if start <= coordinate_from_string(cell["cell"])[1] <= end
            ]
            region_evidence = [
                item for row, items in anchors.items() if start <= row <= end for item in items
            ]
            region_labels = list(dict.fromkeys(item["label"] for item in region_evidence)) or [
                "unknown"
            ]
            header_cells = [
                {"cell": cell["cell"], "value": cell["value"]}
                for cell in region_cells
                if header_pattern.fullmatch(cell["value"].strip())
            ]
            merged_in_region = []
            for merged in merges:
                _min_col, min_row, _max_col, max_row_of_merge = bounded_range(merged)
                if min_row <= end and max_row_of_merge >= start:
                    merged_in_region.append(merged)
            regions.append(
                {
                    "range": f"A{start}:{get_column_letter(max_col)}{end}",
                    "labels": region_labels,
                    "label_source": "rule",
                    "range_source": "rule",
                    "evidence": region_evidence,
                    "parent_heading": region_evidence[0]["matched"] if region_evidence else "",
                    "header_cells": header_cells,
                    "merged_ranges": merged_in_region,
                    "possible_shared_context": possible_shared,
                    "checked_cells": len(region_cells),
                    "needs_confirmation": True,
                    "dependency_status": "跨区图例和依赖未分析，边界需确认",
                }
            )
        results.append(
            {
                "sheet": sheet,
                "labels": [item["label"] for item in evidence] or ["unknown"],
                "evidence": evidence,
                "checked_cells": len(cells),
                "regions": regions,
            }
        )
    return results
