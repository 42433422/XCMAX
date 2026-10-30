"""Workbook classification: bounded evidence and validated semantic regions."""

from __future__ import annotations

from typing import Any

from openpyxl.utils import get_column_letter

from app.application.etl.errors import ReviewedExcelError
from app.application.reviewed_excel_evidence import bounded_range
from app.application.reviewed_excel_evidence import (
    observe_fragment_relationships as observe_fragment_relationships,
)
from app.application.reviewed_excel_evidence import (
    plan_semantic_batches as plan_semantic_batches,
)
from app.application.reviewed_excel_partition import (
    DOCUMENT_LABELS as DOCUMENT_LABELS,
)
from app.application.reviewed_excel_partition import (
    apply_classification_corrections as apply_classification_corrections,
)


def classify_review_workbook(review: dict) -> list[dict]:
    """Suggest non-exclusive document labels from inspected cell text only."""
    import re

    from openpyxl.utils.cell import coordinate_from_string

    from app.application.etl.target_detection import _ATTENDANCE_RE, _DELIVERY_RE, _PURCHASE_RE

    patterns = {
        "quotation": re.compile(r"报价|報價|quotation|sales quote", re.I),
        "delivery": _DELIVERY_RE,
        "purchase": _PURCHASE_RE,
        "attendance": _ATTENDANCE_RE,
    }
    results = []
    for sheet in review["sheets"]:
        cells = [cell for cell in review["cells"] if cell["sheet"] == sheet]
        evidence = []
        for label, pattern in patterns.items():
            for cell in cells:
                match = pattern.search(cell["value"])
                if match is not None:
                    evidence.append(
                        {"label": label, "cell": cell["cell"], "matched": match.group(0)}
                    )
                    break
        headings = {
            "quotation": re.compile(r"^(报价单|報價單|quotation|sales quote)$", re.I),
            "delivery": re.compile(r"^(送货单|发货单|出货单|delivery note)$", re.I),
            "purchase": re.compile(r"^(采购订单|purchase order)$", re.I),
            "attendance": re.compile(r"^(考勤表|attendance sheet)$", re.I),
        }
        anchors: dict[int, list[dict]] = {}
        for cell in cells:
            for label, pattern in headings.items():
                if pattern.fullmatch(cell["value"].strip()):
                    row = coordinate_from_string(cell["cell"])[1]
                    anchors.setdefault(row, []).append(
                        {"label": label, "cell": cell["cell"], "matched": cell["value"]}
                    )
        dimensions: dict = next(
            (item for item in review["sheet_overview"] if item["sheet"] == sheet), {}
        )
        max_row = max(1, int(dimensions.get("rows", 1)))
        max_col = max(1, int(dimensions.get("columns", 1)))
        merges = dimensions.get("merged_regions", [])
        section_headings = {
            "production": r"生产(?:任务|计划)?单|production (?:order|sheet)",
            "process": r"工艺(?:流程|说明|要求)?|工序(?:说明|要求)?|process(?: instructions?| sheet)?",
            "material": r"原(?:辅)?料(?:清单|明细)?|物料(?:清单|明细)|materials?(?: list)?|bill of materials",
            "inspection": r"检验(?:说明|要求|标准|记录|结果)?|质检(?:说明|要求|标准|记录)?|inspection(?: instructions?| standards?| report)?",
        }
        rows_of_cells: dict[int, list[dict]] = {}
        for cell in cells:
            rows_of_cells.setdefault(coordinate_from_string(cell["cell"])[1], []).append(cell)
        merged_titles = {
            f"{get_column_letter(c1)}{r1}"
            for c1, r1, c2, _r2 in map(bounded_range, merges)
            if c2 > c1
        }
        for cell in cells:
            column, row = coordinate_from_string(cell["cell"])
            merged_title = cell["cell"] in merged_titles
            sparse_title = (
                column == "A"
                and len(rows_of_cells[row]) == 1
                and any(
                    len(rows_of_cells.get(next_row, [])) >= 2 for next_row in (row + 1, row + 2)
                )
            )
            if cell.get("kind") == "formula" or not (merged_title or sparse_title):
                continue
            value = cell["value"].strip().rstrip(":：")
            for label, section_pattern in section_headings.items():
                if re.fullmatch(section_pattern, value, re.I):
                    item = {"label": label, "cell": cell["cell"], "matched": cell["value"]}
                    anchors.setdefault(row, []).append(item)
                    if not any(ref["label"] == label for ref in evidence):
                        evidence.append(item)
        rows = sorted(anchors)
        if (
            len(rows) > 1
            and {item["label"] for item in anchors[rows[0]]} == {"production"}
            and any(
                item["label"] in {"process", "material", "inspection"} for item in anchors[rows[1]]
            )
        ):
            # A workbook title remains in the first section, rather than becoming a title-only slice.
            rows = rows[1:]
        split = len(rows) > 1 and all(b - a >= 3 for a, b in zip(rows, rows[1:]))
        if split and any(
            min_row < boundary <= max_row_of_merge
            for merged in merges
            for _min_col, min_row, _max_col, max_row_of_merge in [bounded_range(merged)]
            for boundary in rows[1:]
        ):
            split = False
        possible_shared = [
            {"cell": cell["cell"], "kind": "可能的图例或说明，关联待确认"}
            for cell in cells
            if re.search(r"图例|圖例|legend|说明|說明|备注|備註", cell["value"], re.I)
        ]
        header_pattern = re.compile(
            r"^(产品名称|商品名称|品名|数量|单价|金额|合计|客户|供应商|姓名|日期|product|quantity|price|amount|customer|supplier|date)$",
            re.I,
        )
        boundaries = [1] + rows[1:] if split else [1]
        regions = []
        for index, start in enumerate(boundaries):
            end = boundaries[index + 1] - 1 if index + 1 < len(boundaries) else max_row
            region_cells = [
                cell for cell in cells if start <= coordinate_from_string(cell["cell"])[1] <= end
            ]
            region_evidence = [
                item for row, items in anchors.items() if start <= row <= end for item in items
            ]
            region_labels = list(dict.fromkeys(item["label"] for item in region_evidence)) or [
                "unknown"
            ]
            header_cells = [
                {"cell": cell["cell"], "value": cell["value"]}
                for cell in region_cells
                if header_pattern.fullmatch(cell["value"].strip())
            ]
            merged_in_region = []
            for merged in merges:
                _min_col, min_row, _max_col, max_row_of_merge = bounded_range(merged)
                if min_row <= end and max_row_of_merge >= start:
                    merged_in_region.append(merged)
            regions.append(
                {
                    "range": f"A{start}:{get_column_letter(max_col)}{end}",
                    "labels": region_labels,
                    "label_source": "rule",
                    "range_source": "rule",
                    "evidence": region_evidence,
                    "parent_heading": region_evidence[0]["matched"] if region_evidence else "",
                    "header_cells": header_cells,
                    "merged_ranges": merged_in_region,
                    "possible_shared_context": possible_shared,
                    "checked_cells": len(region_cells),
                    "needs_confirmation": True,
                    "dependency_status": "跨区图例和依赖未分析，边界需确认",
                }
            )
        results.append(
            {
                "sheet": sheet,
                "labels": [item["label"] for item in evidence] or ["unknown"],
                "evidence": evidence,
                "checked_cells": len(cells),
                "regions": regions,
            }
        )
    return results


def _classify_semantic_batch(
    review: dict,
    *,
    prepare_only: bool = False,
    reference_ids: bool = False,
    expected_request_sha256: str | None = None,
) -> dict:
    """Request bounded semantic regions; validate every reference before adoption."""
    import json
    import os

    from openpyxl.utils.cell import (
        column_index_from_string,
        coordinate_from_string,
        get_column_letter,
    )

    if (
        not prepare_only
        and os.environ.get("FHD_TEMPLATE_LLM_PROVIDER", "").strip().lower() != "mimo"
    ):
        raise ReviewedExcelError("模板模型供应商未配置，未进行模型分类，未调用模型")

    allowed_labels = DOCUMENT_LABELS
    rules = [] if reference_ids else classify_review_workbook(review)
    from app.application.reviewed_excel_evidence import (
        MAX_MODEL_CELLS,
        MAX_MODEL_TEXT_BYTES,
        MAX_VALUE_CHARS,
        _model_cell_cost,
        build_relationship_context,
        build_rule_hints,
        build_structure_map,
    )

    if reference_ids:
        source_cells = review.get("full_cells", review["cells"])
        structure_map = [
            {
                **sheet,
                "regions": [],
                "nonempty_cells": sum(c["sheet"] == sheet["sheet"] for c in source_cells),
            }
            for sheet in review["sheet_overview"]
        ]
    else:
        structure_map = review.get("full_structure_map") or build_structure_map(review, rules)
    chosen = list(review["cells"])
    drawings = [d for cell in chosen for d in cell.get("drawing_evidence", [])]
    if (
        len(chosen) + len(drawings) > MAX_MODEL_CELLS
        or any(len(cell["value"]) > MAX_VALUE_CHARS for cell in chosen)
        or any(len(d["text"]) > MAX_VALUE_CHARS for d in drawings)
        or sum(map(_model_cell_cost, chosen))
        + sum(len(json.dumps(d, ensure_ascii=False).encode()) + 64 for d in drawings)
        > MAX_MODEL_TEXT_BYTES
    ):
        raise ReviewedExcelError("本批模型证据超过安全预算，未调用模型")
    sent = {(c["sheet"], c["cell"]): c for c in chosen}
    primary: dict[str, list[tuple[int, int]]] = {}
    primary_keys = {(c["sheet"], c["cell"]) for c in chosen if not c.get("context_only")}
    for cell in chosen:
        if not cell.get("context_only"):
            col, row = coordinate_from_string(cell["cell"])
            primary.setdefault(cell["sheet"], []).append((column_index_from_string(col), row))
    overview = []
    merge_budget = 100
    for sheet in review["sheet_overview"]:
        merges = sheet.get("merged_regions", [])
        included = merges[:merge_budget]
        merge_budget -= len(included)
        overview.append(
            {
                "sheet": sheet["sheet"],
                "rows": sheet["rows"],
                "columns": sheet["columns"],
                "merged": included,
                "merged_omitted": len(merges) - len(included),
                "sheet_state": sheet.get("sheet_state", "visible"),
                "hidden_rows_count": len(sheet.get("hidden_rows", [])),
                "hidden_columns_count": len(sheet.get("hidden_columns", [])),
                "hidden_rows": sheet.get("hidden_rows", [])[:100],
                "hidden_rows_omitted": max(0, len(sheet.get("hidden_rows", [])) - 100),
                "hidden_column_ranges": sheet.get("hidden_column_ranges", [])[:100],
                "hidden_columns_omitted": max(0, len(sheet.get("hidden_column_ranges", [])) - 100),
            }
        )
    rule_hints = build_rule_hints(rules, sent)
    cell_schema = {
        "type": "object",
        "required": ["sheet", "cell"],
        "additionalProperties": False,
        "properties": {"sheet": {"type": "string"}, "cell": {"type": "string"}},
    }
    region_schema: dict[str, Any] = {
        "type": "object",
        "required": ["sheet", "range", "labels", "evidence", "dependencies", "reason"],
        "additionalProperties": False,
        "properties": {
            "sheet": {"type": "string"},
            "range": {"type": "string"},
            "labels": {"type": "array", "minItems": 1, "items": {"type": "string"}},
            "evidence": {"type": "array", "maxItems": 100, "items": cell_schema},
            "dependencies": {"type": "array", "maxItems": 100, "items": cell_schema},
            "reason": {"type": "string"},
        },
    }
    if drawings:
        region_schema["required"].append("drawing_evidence")
        region_schema["properties"]["drawing_evidence"] = {
            "type": "array",
            "maxItems": 100,
            "items": {
                "type": "object",
                "required": ["sheet", "part", "object_id"],
                "additionalProperties": False,
                "properties": {key: {"type": "string"} for key in ("sheet", "part", "object_id")},
            },
        }
    fragment_schema = {
        "type": "object",
        "required": ["sheet", "range"],
        "additionalProperties": False,
        "properties": {"sheet": {"type": "string"}, "range": {"type": "string"}},
    }
    document_schema = {
        "type": "object",
        "required": ["document_id", "fragments"],
        "additionalProperties": False,
        "properties": {
            "document_id": {"type": "string"},
            "fragments": {
                "type": "array",
                "minItems": 1,
                "maxItems": 100,
                "items": fragment_schema,
            },
        },
    }
    schema = {
        "type": "object",
        "required": ["physical_regions", "logical_documents", "layout_references"],
        "additionalProperties": False,
        "properties": {
            "physical_regions": {
                "type": "array",
                "minItems": 1,
                "maxItems": 100,
                "items": region_schema,
            },
            "logical_documents": {"type": "array", "maxItems": 100, "items": document_schema},
            "layout_references": {"type": "array", "maxItems": 100, "items": fragment_schema},
        },
    }
    structure_regions = []
    region_budget = 500
    for sheet in structure_map:
        included = sheet["regions"][:region_budget]
        region_budget -= len(included)
        structure_regions.append(
            {
                "sheet": sheet["sheet"],
                "sheet_state": sheet.get("sheet_state", "visible"),
                "rows": sheet["rows"],
                "columns": sheet["columns"],
                "nonempty_cells": sheet["nonempty_cells"],
                "regions": [
                    {
                        "range": region["range"],
                        "evidence_coordinates": [ref["cell"] for ref in region["evidence"][:20]],
                        "header_coordinates": [ref["cell"] for ref in region["header_cells"][:20]],
                        "shared_context_coordinates": [
                            ref["cell"] for ref in region["possible_shared_context"][:20]
                        ],
                        "evidence_omitted": max(0, len(region["evidence"]) - 20),
                        "headers_omitted": max(0, len(region["header_cells"]) - 20),
                        "shared_context_omitted": max(
                            0, len(region["possible_shared_context"]) - 20
                        ),
                    }
                    for region in included
                ],
                "regions_omitted": len(sheet["regions"]) - len(included),
            }
        )
    messages = [
        {
            "role": "system",
            "content": "",  # Populated only by the independent evidence contract below.
        },
        {
            "role": "user",
            "content": json.dumps(
                {
                    "output_schema": schema,
                    "labels": sorted(allowed_labels),
                    "sheets": overview,
                    "structure_map": structure_regions,
                    "relationship_context": build_relationship_context(structure_map, chosen),
                    "batch_coverage": [
                        {
                            "sheet": sheet["sheet"],
                            "total_nonempty": sheet["nonempty_cells"],
                            "sent_nonempty": sum(
                                cell["sheet"] == sheet["sheet"] for cell in chosen
                            ),
                            "unread_nonempty": max(
                                0,
                                sheet["nonempty_cells"]
                                - sum(cell["sheet"] == sheet["sheet"] for cell in chosen),
                            ),
                        }
                        for sheet in structure_map
                    ],
                    "cells": [
                        {
                            "sheet": c["sheet"],
                            "cell": c["cell"],
                            "value": c["value"],
                            "kind": c["kind"],
                            "context_only": bool(c.get("context_only")),
                        }
                        for c in chosen
                    ],
                    "rule_hints": rule_hints,
                },
                ensure_ascii=False,
            ),
        },
    ]
    from app.application.reviewed_excel_template import _template_mimo_complete

    if drawings:
        payload = json.loads(messages[1]["content"])
        payload["drawing_evidence"] = drawings
        messages[1]["content"] = json.dumps(payload, ensure_ascii=False)
    provider_schema = schema
    references: dict = {}
    if reference_ids:
        from app.application.reviewed_excel_reference_ids import prepare_reference_messages

        references, provider_schema = prepare_reference_messages(messages, schema)
    else:
        from app.application.reviewed_excel_reference_ids import neutralize_messages

        neutralize_messages(messages)

    if reference_ids:
        from app.application.reviewed_excel_reference_ids import verify_contract

        verify_contract(messages[0]["content"], provider_schema, "identify")
    from app.application.reviewed_excel_budget import output_budget, validate_request_budget

    max_output_tokens = output_budget(chosen)
    audit_request = {
        "messages": messages,
        "schema": provider_schema,
        "max_tokens": max_output_tokens,
    }
    validate_request_budget(audit_request)
    if prepare_only:
        return audit_request
    if expected_request_sha256 is not None:
        from app.application.reviewed_excel_audit import verify_request

        verify_request(audit_request, expected_request_sha256)
    response = _template_mimo_complete(messages, provider_schema, max_output_tokens)
    data = response.data
    raw_reference_response = None
    if reference_ids:
        from app.application.reviewed_excel_reference_ids import normalize_reference_ids

        raw_reference_response = data
        data = normalize_reference_ids(data, references, provider_schema, schema)
    if not isinstance(data, dict):
        raise ReviewedExcelError("模型区域结构无效，原分类未改变")
    # Legacy in-process replays contain only regions; the live provider receives
    # and validates the new schema above and cannot return that shape.
    legacy = set(data) == {"regions"}
    if not legacy and set(data) != {"physical_regions", "logical_documents", "layout_references"}:
        raise ReviewedExcelError("模型区域结构无效，原分类未改变")
    physical = data.get("regions" if legacy else "physical_regions")
    if not isinstance(physical, list) or not 1 <= len(physical) <= 100:
        raise ReviewedExcelError("模型区域结构无效，原分类未改变")
    by_dimensions = {s["sheet"]: s for s in review["sheet_overview"]}
    rects: dict[str, list[tuple[int, int, int, int]]] = {}
    accepted = []
    range_map = {}
    for region in physical:
        if not isinstance(region, dict) or set(region) != {
            "sheet",
            "range",
            "labels",
            "evidence",
            "dependencies",
            "reason",
        } | ({"drawing_evidence"} if drawings else set()):
            raise ReviewedExcelError("模型区域格式无效，原分类未改变")
        sheet, area, labels = region["sheet"], region["range"], region["labels"]
        original_area = area
        if (
            not isinstance(sheet, str)
            or sheet not in by_dimensions
            or not isinstance(area, str)
            or not isinstance(region["reason"], str)
        ):
            raise ReviewedExcelError("模型区域来源无效，原分类未改变")
        if (
            not isinstance(labels, list)
            or not labels
            or any(not isinstance(x, str) or x not in allowed_labels for x in labels)
            or len(set(labels)) != len(labels)
            or ("unknown" in labels and len(labels) > 1)
        ):
            raise ReviewedExcelError("模型分类无效，原分类未改变")
        try:
            c1, r1, c2, r2 = bounded_range(area)
            dimensions = by_dimensions[sheet]
            if not all(isinstance(x, int) for x in (c1, r1, c2, r2)) or not (
                1 <= r1 <= r2 <= dimensions["rows"] and 1 <= c1 <= c2 <= dimensions["columns"]
            ):
                raise ReviewedExcelError("bounds")
        except (ValueError, TypeError):
            raise ReviewedExcelError("模型区域坐标越界，原分类未改变") from None
        if not legacy and not any(
            c1 <= col <= c2 and r1 <= row <= r2 for col, row in primary.get(sheet, [])
        ):
            raise ReviewedExcelError("模型区域仅包含上下文或未读内容，原分类未改变")
        fragment_metadata = {}
        if not legacy:
            unread_in_region = []
            for cell in review.get("full_cells", review["cells"]):
                if cell["sheet"] != sheet or (sheet, cell["cell"]) in primary_keys:
                    continue
                col, row = coordinate_from_string(cell["cell"])
                if c1 <= column_index_from_string(col) <= c2 and r1 <= row <= r2:
                    unread_in_region.append(row)
            if unread_in_region:
                observed_rows = [
                    row
                    for col, row in primary.get(sheet, [])
                    if c1 <= col <= c2 and r1 <= row <= r2
                ]
                first, last = min(observed_rows), max(observed_rows)
                if any(first <= row <= last for row in unread_in_region):
                    raise ReviewedExcelError("模型区域包含本批未读内容，原分类未改变")
                fragment_metadata = {"model_original_range": area, "open_fragment": True}
                r1, r2 = first, last
                area = f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}"
                region = {**region, "range": area}
        for old_c1, old_r1, old_c2, old_r2 in rects.setdefault(sheet, []):
            if c1 <= old_c2 and old_c1 <= c2 and r1 <= old_r2 and old_r1 <= r2:
                raise ReviewedExcelError("模型区域重叠，原分类未改变")
        rects[sheet].append((c1, r1, c2, r2))
        for merge in dimensions.get("merged_regions", []):
            mc1, mr1, mc2, mr2 = bounded_range(merge)
            intersects = c1 <= mc2 and mc1 <= c2 and r1 <= mr2 and mr1 <= r2
            if intersects and not (c1 <= mc1 and mc2 <= c2 and r1 <= mr1 and mr2 <= r2):
                raise ReviewedExcelError("模型区域切断合并单元格，原分类未改变")
        for field in ("evidence", "dependencies"):
            refs = region[field]
            if (
                not isinstance(refs, list)
                or len(refs) > 100
                or any(
                    not isinstance(x, dict)
                    or set(x) != {"sheet", "cell"}
                    or not isinstance(x["sheet"], str)
                    or not isinstance(x["cell"], str)
                    or (x["sheet"], x["cell"]) not in sent
                    for x in refs
                )
            ):
                raise ReviewedExcelError("模型证据不在已发送范围，原分类未改变")
            if field == "evidence":
                if labels != ["unknown"] and not refs:
                    raise ReviewedExcelError("模型分类缺少依据，原分类未改变")
                for ref in refs:
                    col, row = coordinate_from_string(ref["cell"])
                    if not legacy and sent[(ref["sheet"], ref["cell"])].get("context_only"):
                        raise ReviewedExcelError("模型分类证据仅为重复上下文，原分类未改变")
                    if ref["sheet"] != sheet or not (
                        c1 <= column_index_from_string(col) <= c2 and r1 <= row <= r2
                    ):
                        raise ReviewedExcelError("模型证据不在区域内，原分类未改变")
        if drawings:
            from app.application.reviewed_excel_drawing import validate_region_drawings

            validate_region_drawings(review, drawings, region)
        accepted.append({**region, **fragment_metadata, "needs_confirmation": True})
        range_map[(sheet, original_area)] = (sheet, area)
    suggested_documents, layout_references = [], []
    if not legacy:
        documents, layout = data["logical_documents"], data["layout_references"]
        if (
            not isinstance(documents, list)
            or len(documents) > 100
            or not isinstance(layout, list)
            or len(layout) > 100
        ):
            raise ReviewedExcelError("模型逻辑文档或布局引用格式无效，原分类未改变")
        known = {(region["sheet"], region["range"]) for region in accepted}
        assigned, document_ids = set(), set()
        for document in documents:
            if (
                not isinstance(document, dict)
                or set(document) != {"document_id", "fragments"}
                or not isinstance(document["document_id"], str)
                or not document["document_id"].strip()
                or len(document["document_id"]) > 80
                or document["document_id"] in document_ids
                or not isinstance(document["fragments"], list)
                or not 1 <= len(document["fragments"]) <= 100
            ):
                raise ReviewedExcelError("模型逻辑文档格式或编号无效，原分类未改变")
            document_ids.add(document["document_id"])
            refs = []
            for ref in document["fragments"]:
                if (
                    not isinstance(ref, dict)
                    or set(ref) != {"sheet", "range"}
                    or not isinstance(ref["sheet"], str)
                    or not isinstance(ref["range"], str)
                ):
                    raise ReviewedExcelError("模型文档片段格式无效，原分类未改变")
                key = range_map.get((ref["sheet"], ref["range"]), (ref["sheet"], ref["range"]))
                if key not in known or key in assigned:
                    raise ReviewedExcelError("模型文档片段来源无效或重复，原分类未改变")
                assigned.add(key)
                refs.append({"sheet": key[0], "range": key[1]})
            suggested_documents.append(
                {"document_id": document["document_id"], "fragments": refs, "status": "candidate"}
            )
        allowed_merges = {
            (sheet["sheet"], merge) for sheet in overview for merge in sheet["merged"]
        }
        for ref in layout:
            if (
                not isinstance(ref, dict)
                or set(ref) != {"sheet", "range"}
                or not isinstance(ref["sheet"], str)
                or not isinstance(ref["range"], str)
                or (ref["sheet"], ref["range"]) not in allowed_merges
            ):
                raise ReviewedExcelError("模型布局引用不在已发送的合并区域，原分类未改变")
            layout_references.append(ref)
    from app.application.reviewed_excel_evidence import summarize_model_input

    input_summary = summarize_model_input(
        messages,
        chosen,
        source_sha256=review["sha256"],
        provider=getattr(response, "provider", "") or "unreported",
        model=response.model,
        max_output_tokens=max_output_tokens,
    )
    input_summary["sent_drawing_objects"] = [
        {k: d[k] for k in ("sheet", "part", "object_id")} for d in drawings
    ]
    if raw_reference_response is not None:
        from app.application.reviewed_excel_reference_ids import add_reference_trace

        add_reference_trace(input_summary, raw_reference_response, references)
    # Final coverage is computed from the source and all accepted batches by the combiner.
    coverage = [{"sheet": item["sheet"], "merged_sent": len(item["merged"])} for item in overview]
    return {
        "regions": accepted,
        "coverage": coverage,
        "input_summary": input_summary,
        "logical_documents": suggested_documents,
        "logical_document_count": None if legacy else len(suggested_documents),
        "layout_references": layout_references,
        "contract": "legacy_regions" if legacy else "physical_logical_v2",
    }


def combine_semantic_batches(
    review: dict, batches: list[list[dict]], results: list[dict], max_calls: int
) -> dict:
    """Combine only validated results; reject overlap across batches."""
    from app.application.reviewed_excel_evidence import build_structure_map

    rules = classify_review_workbook(review)
    accepted = []
    rectangles: dict[str, list[tuple[int, int, int, int, dict]]] = {}
    sent: set[tuple[str, str]] = set()
    primary_sent: set[tuple[str, str]] = set()
    input_summaries = []
    merge_sent: dict[str, int] = {}
    for batch, result in zip(batches, results, strict=True):
        input_summaries.append(result["input_summary"])
        sent.update((cell["sheet"], cell["cell"]) for cell in batch)
        primary_sent.update(
            (cell["sheet"], cell["cell"]) for cell in batch if not cell.get("context_only")
        )
        for coverage in result["coverage"]:
            merge_sent[coverage["sheet"]] = max(
                merge_sent.get(coverage["sheet"], 0), coverage["merged_sent"]
            )
        for region in result["regions"]:
            c1, r1, c2, r2 = bounded_range(region["range"])
            prior = rectangles.setdefault(region["sheet"], [])
            overlaps = [
                old
                for old in prior
                if c1 <= old[2] and old[0] <= c2 and r1 <= old[3] and old[1] <= r2
            ]
            if overlaps:
                if any(old[4] == region for old in overlaps):
                    continue
                raise ReviewedExcelError("分批模型区域冲突，原分类未改变")
            prior.append((c1, r1, c2, r2, region))
            accepted.append(region)
    from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

    overview = {sheet["sheet"]: sheet for sheet in review["sheet_overview"]}
    coverage = []
    for name in review["sheets"]:
        cells = [cell for cell in review["cells"] if cell["sheet"] == name]
        unread = [cell["cell"] for cell in cells if (name, cell["cell"]) not in primary_sent]
        unassigned = []
        for cell in cells:
            col, row = coordinate_from_string(cell["cell"])
            number = column_index_from_string(col)
            if not any(
                c1 <= number <= c2 and r1 <= row <= r2
                for c1, r1, c2, r2, _ in rectangles.get(name, [])
            ):
                unassigned.append(cell["cell"])
        merges_total = len(overview[name].get("merged_regions", []))
        from app.application.reviewed_excel_evidence import (
            MAX_MODEL_CELLS,
            MAX_MODEL_TEXT_BYTES,
            MAX_VALUE_CHARS,
            _model_cell_cost,
        )

        rows: dict[int, list[dict]] = {}
        for cell in cells:
            _col, row = coordinate_from_string(cell["cell"])
            rows.setdefault(row, []).append(cell)
        deferred = [
            cell["cell"]
            for row_cells in rows.values()
            if len(row_cells) > MAX_MODEL_CELLS
            or any(len(cell["value"]) > MAX_VALUE_CHARS for cell in row_cells)
            or sum(map(_model_cell_cost, row_cells)) > MAX_MODEL_TEXT_BYTES
            for cell in row_cells
            if cell["cell"] in unread
        ]
        from app.application.reviewed_excel_drawing import drawing_coverage

        coverage.append(
            {
                **drawing_coverage(review, batches, name, rectangles.get(name, [])),
                "sheet": name,
                "total": len(cells),
                "sent": len(cells) - len(unread),
                "unread": unread,
                "unassigned": unassigned,
                "merged_sent": merge_sent.get(name, 0),
                "merged_omitted": merges_total - merge_sent.get(name, 0),
                "value_truncated": [],
                "deferred_large_values": [
                    cell["cell"] for cell in cells if len(cell["value"]) > MAX_VALUE_CHARS
                ],
                "deferred_oversized_rows": deferred,
            }
        )
    from app.application.reviewed_excel_evidence import (
        candidate_fragment_links,
        flag_repeated_header_conflicts,
        value_extent,
    )

    accepted = flag_repeated_header_conflicts(
        accepted, [cell for cell in review["cells"] if (cell["sheet"], cell["cell"]) in sent]
    )
    batch_keys = [{(cell["sheet"], cell["cell"]) for cell in batch} for batch in batches]
    for region in accepted:
        region["value_extent"] = value_extent(region, review["cells"])
        c1, r1, c2, r2 = bounded_range(region["range"])
        region_keys = set()
        for cell in review["cells"]:
            if cell["sheet"] != region["sheet"]:
                continue
            col, row = coordinate_from_string(cell["cell"])
            if c1 <= column_index_from_string(col) <= c2 and r1 <= row <= r2:
                region_keys.add((cell["sheet"], cell["cell"]))
        unread_region = region_keys - primary_sent
        if unread_region:
            if region["labels"] != ["unknown"]:
                region.setdefault("suggested_labels", region["labels"])
                region["labels"] = ["unknown"]
            region["context_warning"] = (
                region.get("context_warning", "")
                + f"区域内仍有 {len(unread_region)} 个非空格未发送；不能确认分类。"
            )
        elif region_keys and not any(region_keys <= keys for keys in batch_keys):
            region["context_warning"] = (
                region.get("context_warning", "")
                + "区域内容分批发送，没有一次完整上下文；模型建议需核对。"
            )
    sent_cells = [cell for cell in review["cells"] if (cell["sheet"], cell["cell"]) in sent]
    links = candidate_fragment_links(accepted, sent_cells)
    by_ref = {(region["sheet"], region["range"]): region for region in accepted}
    existing = {
        (item["sheet"], frozenset((item["first"], item["second"]))) for item in links["items"]
    }
    suggested_documents = []
    for batch_index, batch_result in enumerate(results):
        for document in batch_result.get("logical_documents", []):
            suggested_documents.append(
                {
                    **document,
                    "model_document_id": document["document_id"],
                    "document_id": f"batch-{batch_index + 1}:{document['document_id']}",
                    "batch": batch_index + 1,
                }
            )
            for first, second in zip(document["fragments"], document["fragments"][1:]):
                if first["sheet"] != second["sheet"]:
                    continue
                key = (first["sheet"], frozenset((first["range"], second["range"])))
                if key in existing:
                    continue
                pair = [dict(by_ref[(ref["sheet"], ref["range"])]) for ref in (first, second)]
                item = candidate_fragment_links(
                    pair,
                    sent_cells,
                    model_pairs={(first["sheet"], first["range"], second["range"])},
                )["items"][0]
                existing.add(key)
                if len(links["items"]) < 200:
                    links["items"].append(item)
                else:
                    links["omitted"] += 1
    return {
        "source_sha256": review["sha256"],
        "model_status": "returned",
        "regions": accepted,
        "fragment_links": links,
        "suggested_documents": suggested_documents,
        "suggested_document_count": None,
        "document_count_status": "partial",
        "layout_references": [
            ref for result in results for ref in result.get("layout_references", [])
        ],
        "coverage": coverage,
        "input_summaries": input_summaries,
        "structure_map": build_structure_map(review, rules),
        "structure_map_source": "program_observation_not_sent_to_model",
        "budget": {
            "calls": len(batches),
            "max_calls": max_calls,
            "cells_per_call": 500,
            "unread": sum(len(item["unread"]) for item in coverage),
        },
    }


def classify_semantic_workbook(review: dict) -> dict:
    """Run the bounded standalone path, used by isolated replay and tests."""
    batches, max_calls = plan_semantic_batches(review, neutral=True)
    from app.application.reviewed_excel_evidence import build_structure_map

    full_structure_map = build_structure_map(review, classify_review_workbook(review))
    results = []
    for batch in batches[:max_calls]:
        results.append(
            _classify_semantic_batch(
                {
                    **review,
                    "cells": batch,
                    "full_cells": review["cells"],
                    "full_structure_map": full_structure_map,
                }
            )
        )
        combined = combine_semantic_batches(review, batches[: len(results)], results, max_calls)
    return combined
