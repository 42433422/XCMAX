"""Bounded OOXML drawing evidence extraction, without changing inspection gates."""

from __future__ import annotations

import posixpath
from zipfile import ZipFile

from defusedxml import ElementTree as ET

from app.application.etl.errors import ReviewedExcelError

S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
P = "http://schemas.openxmlformats.org/package/2006/relationships"
X = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _relations(z: ZipFile, part: str) -> dict:
    directory, filename = posixpath.split(part)
    path = f"{directory}/_rels/{filename}.rels"
    if path not in z.namelist():
        return {}
    root = ET.fromstring(z.read(path))
    if root.tag != f"{{{P}}}Relationships":
        raise ReviewedExcelError("绘图关系格式无效")
    result = {}
    for rel in root:
        key, target = rel.get("Id"), rel.get("Target", "")
        if (
            rel.tag != f"{{{P}}}Relationship"
            or rel.get("TargetMode", "Internal") != "Internal"
            or not key
            or key in result
            or not target
            or target.startswith("\\")
            or "\\" in target
            or "?" in target
            or "#" in target
        ):
            raise ReviewedExcelError("绘图关系无效或含外部目标")
        target = posixpath.normpath(
            target.lstrip("/") if target.startswith("/") else posixpath.join(directory, target)
        )
        if not target.startswith("xl/"):
            raise ReviewedExcelError("绘图关系目标越界")
        result[key] = (rel.get("Type", ""), target)
    return result


def _marker(node) -> dict:
    result = {}
    if node is None or sorted(child.tag for child in node) != sorted(
        f"{{{X}}}{key}" for key in ("col", "colOff", "row", "rowOff")
    ):
        raise ReviewedExcelError("绘图锚点字段重复或未知")
    for key in ("col", "colOff", "row", "rowOff"):
        value = node.findtext(f"{{{X}}}{key}") if node is not None else None
        if value is None or not value.isdecimal():
            raise ReviewedExcelError("绘图锚点无效")
        result[key] = int(value)
    if result["row"] >= 1048576 or result["col"] >= 16384:
        raise ReviewedExcelError("绘图锚点越界")
    return result


def extract_drawing_evidence(z: ZipFile) -> list[dict]:
    """Extract stable object identities and content; reject unrecognized relationships.

    This is only a building block: callers must implement budgeting, model
    coverage and lossless export before enabling any supported drawing input.
    """
    names = set(z.namelist())
    if len(names) != len(z.namelist()):
        raise ReviewedExcelError("工作簿存在重复部件，无法核验")
    root = ET.fromstring(z.read("xl/workbook.xml"))
    workbook_rels = _relations(z, "xl/workbook.xml")
    evidence: list[dict] = []
    seen_parts: set[str] = set()
    for sheet in root.findall(f"{{{S}}}sheets/{{{S}}}sheet"):
        name, rid = sheet.get("name"), sheet.get(f"{{{R}}}id")
        if not name or rid not in workbook_rels:
            raise ReviewedExcelError("工作表关系无效")
        typ, part = workbook_rels[rid]
        if typ != f"{R}/worksheet" or part not in names:
            raise ReviewedExcelError("工作表关系不受支持")
        ws = ET.fromstring(z.read(part))
        drawings = ws.findall(f"{{{S}}}drawing")
        if len(drawings) > 1:
            raise ReviewedExcelError("工作表绘图关系不受支持")
        if not drawings:
            continue
        rels = _relations(z, part)
        drawing_id = drawings[0].get(f"{{{R}}}id")
        if drawing_id not in rels:
            raise ReviewedExcelError("工作表绘图关系缺失")
        typ, drawing_part = rels[drawing_id]
        if (
            typ != f"{R}/drawing"
            or not drawing_part.startswith("xl/drawings/")
            or drawing_part not in names
            or drawing_part in seen_parts
        ):
            raise ReviewedExcelError("工作表绘图关系不受支持")
        seen_parts.add(drawing_part)
        if "[Content_Types].xml" not in names:
            raise ReviewedExcelError("绘图内容类型声明缺失")
        type_root = ET.fromstring(z.read("[Content_Types].xml"))
        declarations = [node for node in type_root if node.get("PartName") == "/" + drawing_part]
        if (
            len(declarations) != 1
            or declarations[0].tag
            != "{http://schemas.openxmlformats.org/package/2006/content-types}Override"
            or declarations[0].get("ContentType")
            != "application/vnd.openxmlformats-officedocument.drawing+xml"
        ):
            raise ReviewedExcelError("绘图内容类型声明不受支持或重复")
        if _relations(z, drawing_part):
            raise ReviewedExcelError("绘图存在未支持的对象关系")
        drawing = ET.fromstring(z.read(drawing_part))
        if drawing.tag != f"{{{X}}}wsDr":
            raise ReviewedExcelError("绘图格式不受支持")
        if any(
            not isinstance(el.tag, str)
            or not el.tag.startswith((f"{{{X}}}", f"{{{A}}}"))
            or el.tag.split("}")[-1] in {"extLst", "hlinkClick", "hlinkHover", "oleObj"}
            for el in drawing.iter()
        ):
            raise ReviewedExcelError("绘图包含不受支持的扩展或引用")
        allowed_tags = {
            "wsDr",
            "oneCellAnchor",
            "twoCellAnchor",
            "from",
            "to",
            "col",
            "colOff",
            "row",
            "rowOff",
            "ext",
            "clientData",
            "sp",
            "cxnSp",
            "nvSpPr",
            "nvCxnSpPr",
            "cNvPr",
            "cNvSpPr",
            "cNvCxnSpPr",
            "spPr",
            "txBody",
            "style",
            "xfrm",
            "off",
            "avLst",
            "gd",
            "prstGeom",
            "custGeom",
            "bodyPr",
            "lstStyle",
            "p",
            "pPr",
            "r",
            "rPr",
            "t",
            "br",
            "endParaRPr",
            "defPPr",
            "defRPr",
            "latin",
            "ea",
            "cs",
            "sym",
            "noAutofit",
            "normAutofit",
            "spAutoFit",
            "solidFill",
            "noFill",
            "gradFill",
            "pattFill",
            "schemeClr",
            "scrgbClr",
            "srgbClr",
            "sysClr",
            "prstClr",
            "alpha",
            "tint",
            "shade",
            "lumMod",
            "lumOff",
            "satMod",
            "satOff",
            "hueMod",
            "hueOff",
            "ln",
            "lnRef",
            "fillRef",
            "effectRef",
            "fontRef",
            "effectLst",
            "effectDag",
            "outerShdw",
            "innerShdw",
            "glow",
            "softEdge",
            "reflection",
            "blur",
            "scene3d",
            "sp3d",
            "stCxn",
            "endCxn",
            "headEnd",
            "tailEnd",
            "prstDash",
            "round",
            "bevel",
            "miter",
        }
        allowed_tags.update(f"lvl{i}pPr" for i in range(1, 10))
        x_tags = {
            "wsDr",
            "oneCellAnchor",
            "twoCellAnchor",
            "from",
            "to",
            "col",
            "colOff",
            "row",
            "rowOff",
            "ext",
            "clientData",
            "sp",
            "cxnSp",
            "nvSpPr",
            "nvCxnSpPr",
            "cNvPr",
            "cNvSpPr",
            "cNvCxnSpPr",
            "spPr",
            "txBody",
            "style",
        }
        if any(
            el.tag
            not in (
                {f"{{{X}}}{tag}" for tag in x_tags}
                | {f"{{{A}}}{tag}" for tag in allowed_tags - x_tags}
                | {f"{{{A}}}ext"}
            )
            for el in drawing.iter()
        ):
            raise ReviewedExcelError("绘图包含未知对象结构")
        ids = set()
        for anchor in drawing:
            kind = anchor.tag.removeprefix(f"{{{X}}}")
            if kind not in ("oneCellAnchor", "twoCellAnchor"):
                raise ReviewedExcelError("绘图锚点类型不受支持")
            objects = [c for c in anchor if c.tag in (f"{{{X}}}sp", f"{{{X}}}cxnSp")]
            permitted = {f"{{{X}}}{x}" for x in ("from", "to", "ext", "sp", "cxnSp", "clientData")}
            if len(objects) != 1 or any(c.tag not in permitted for c in anchor):
                raise ReviewedExcelError("绘图对象类型不受支持")
            counts = {
                tag: sum(child.tag == f"{{{X}}}{tag}" for child in anchor)
                for tag in ("from", "to", "ext", "clientData")
            }
            if counts != {
                "from": 1,
                "to": int(kind == "twoCellAnchor"),
                "ext": int(kind == "oneCellAnchor"),
                "clientData": 1,
            }:
                raise ReviewedExcelError("绘图锚点结构重复或不完整")
            obj = objects[0]
            if (
                len(list(obj.iter(f"{{{X}}}cNvPr"))) != 1
                or len(obj.findall(f"{{{X}}}spPr")) != 1
                or len(obj.findall(f"{{{X}}}txBody")) > 1
            ):
                raise ReviewedExcelError("绘图对象结构重复或不完整")
            properties = obj.find(f".//{{{X}}}cNvPr")
            oid = properties.get("id") if properties is not None else None
            if not oid or not oid.isdecimal() or oid in ids:
                raise ReviewedExcelError("绘图对象编号无效")
            ids.add(oid)
            if any(k.startswith(f"{{{R}}}") for el in obj.iter() for k in el.attrib):
                raise ReviewedExcelError("绘图对象引用不受支持")
            if any(el.tag == f"{{{A}}}xfrm" and el.get("rot", "0") != "0" for el in obj.iter()):
                raise ReviewedExcelError("旋转绘图对象的边界暂不能核验")
            start = _marker(anchor.find("{" + X + "}from"))
            end = _marker(anchor.find(f"{{{X}}}to")) if kind == "twoCellAnchor" else None
            if end and (
                (end["row"], end["rowOff"]) < (start["row"], start["rowOff"])
                or (end["col"], end["colOff"]) < (start["col"], start["colOff"])
            ):
                raise ReviewedExcelError("绘图锚点顺序无效")
            ext = anchor.find(f"{{{X}}}ext")
            if kind == "oneCellAnchor":
                if ext is None or any(
                    not (ext.get(k) or "").isdecimal() or int(ext.get(k) or "0") <= 0
                    for k in ("cx", "cy")
                ):
                    raise ReviewedExcelError("绘图尺寸无效")
                size = {k: int(ext.get(k) or "0") for k in ("cx", "cy")}
            else:
                size = None
            transforms = list(obj.iter(f"{{{A}}}xfrm"))
            if any(
                node.find(f"{{{A}}}chOff") is not None or node.find(f"{{{A}}}chExt") is not None
                for node in transforms
            ):
                raise ReviewedExcelError("分组绘图变换暂不能核验")
            text = "\n".join(
                "".join(
                    "\n" if node.tag == f"{{{A}}}br" else node.text or ""
                    for node in paragraph.iter()
                    if node.tag in {f"{{{A}}}t", f"{{{A}}}br"}
                )
                for paragraph in obj.iter(f"{{{A}}}p")
            )
            body_props = obj.find(f"{{{X}}}txBody/{{{A}}}bodyPr")
            body_nodes = list(obj.iter(f"{{{A}}}bodyPr"))
            if len(body_nodes) > 1 or (body_nodes and body_props is None):
                raise ReviewedExcelError("绘图文字布局结构重复或未知")
            # Theme stroke/effect references (lnRef/effectRef) and ordinary xfrm
            # markers do not by themselves expand unverifiable cell coverage.
            # Group child transforms already fail closed above via chOff/chExt.
            visual_overflow = any(
                node.tag.split("}")[-1]
                in {
                    "effectLst",
                    "effectDag",
                    "scene3d",
                    "sp3d",
                    "custGeom",
                    "spAutoFit",
                }
                for node in obj.iter()
            )
            if text and (
                body_props is None
                or body_props.get("rot", "0") != "0"
                or body_props.get("vertOverflow", "clip") != "clip"
                or body_props.get("horzOverflow", "clip") != "clip"
            ):
                visual_overflow = True

            if len(evidence) >= 20000:
                raise ReviewedExcelError("绘图对象超过安全预算")
            evidence.append(
                {
                    "sheet": name,
                    "part": drawing_part,
                    "object_id": oid,
                    "kind": obj.tag.split("}")[-1],
                    "anchor_type": kind,
                    "from": start,
                    "to": end,
                    "extent": size,
                    "text": text,
                    **(
                        {
                            "geometry_error": "绘图含独立变换或未约束视觉外溢，边界尚未核验，不能分批或派生"
                        }
                        if visual_overflow
                        else {}
                    ),
                }
            )
    actual = {n for n in names if n.startswith("xl/drawings/") and not n.endswith("/")}
    if actual != seen_parts:
        raise ReviewedExcelError("工作簿含未关联或不受支持的绘图部件")
    return evidence


def drawing_row_boundaries(sheet, drawings: list[dict]) -> list[int]:
    """Conservative row boundaries; objects and merges are indivisible."""
    from app.application.reviewed_excel_evidence import bounded_range

    max_row = sheet.max_row
    unsafe: set[int] = set()
    for merged in sheet.merged_cells.ranges:
        _, first, _, last = bounded_range(str(merged))
        unsafe.update(range(first + 1, last + 1))
    for obj in drawings:
        first = obj["from"]["row"] + 1
        if not 1 <= first <= max_row:
            raise ReviewedExcelError("绘图锚点超出工作表范围")
        if obj.get("geometry_error"):
            raise ReviewedExcelError("绘图边界尚未核验")
        if obj.get("last_row"):
            last = int(obj["last_row"])
        elif obj.get("to") is not None:
            last = int(obj["to"]["row"]) + 1
            obj["last_row"] = last
            obj["last_col"] = int(obj["to"]["col"]) + 1
        else:
            raise ReviewedExcelError("绘图边界尚未核验")
        if last > max_row:
            raise ReviewedExcelError("绘图超出工作表范围，无法安全确定边界")
        unsafe.update(range(first + 1, last + 1))
    return [row for row in range(1, max_row + 2) if row not in unsafe]


def plan_drawing_batches(review: dict) -> list[list[dict]]:
    """Partition at explicit legal boundaries, never splitting merges or shapes."""
    import json
    from collections import defaultdict

    from openpyxl.utils.cell import coordinate_from_string

    from app.application.reviewed_excel_evidence import (
        MAX_MODEL_CELLS,
        MAX_MODEL_TEXT_BYTES,
        MAX_VALUE_CHARS,
        _model_cell_cost,
    )

    rows = defaultdict(list)
    for cell in review["cells"]:
        _, row = coordinate_from_string(cell["cell"])
        rows[(cell["sheet"], row)].append(dict(cell))
    for drawing in review["drawing_evidence"]:
        if drawing.get("geometry_error"):
            raise ReviewedExcelError(drawing["geometry_error"])
        key = (drawing["sheet"], drawing["from"]["row"] + 1)
        if key not in rows:
            raise ReviewedExcelError("绘图锚点所在行没有单元格证据，暂不能安全分批")
        rows[key][0].setdefault("drawing_evidence", []).append(drawing)

    def cost(cell):
        return _model_cell_cost(cell) + sum(
            len(json.dumps(d, ensure_ascii=False).encode("utf-8")) + 64
            for d in cell.get("drawing_evidence", [])
        )

    def count(cell):
        return 1 + len(cell.get("drawing_evidence", []))

    result = []
    overview = {entry["sheet"]: entry for entry in review["sheet_overview"]}
    for sheet in review["sheets"]:
        boundaries = overview[sheet]["drawing_row_boundaries"]
        current: list[dict] = []
        amount, used, first = 0, 0, boundaries[0]
        for start, end in zip(boundaries, boundaries[1:]):
            group = [cell for row in range(start, end) for cell in rows.get((sheet, row), [])]
            n, size = sum(map(count, group)), sum(map(cost, group))
            if (
                n > MAX_MODEL_CELLS
                or size > MAX_MODEL_TEXT_BYTES
                or any(len(c["value"]) > MAX_VALUE_CHARS for c in group)
                or any(
                    len(d["text"]) > MAX_VALUE_CHARS
                    for c in group
                    for d in c.get("drawing_evidence", [])
                )
            ):
                raise ReviewedExcelError("不可分绘图或合并区域超过模型安全预算，未调用模型")
            if current and (amount + n > MAX_MODEL_CELLS or used + size > MAX_MODEL_TEXT_BYTES):
                current[0]["drawing_batch_bounds"] = [first, start]
                current[0]["drawing_legal_boundaries"] = [
                    x for x in boundaries if first <= x <= start
                ]
                result.append(current)
                current, amount, used, first = [], 0, 0, start
            current.extend(group)
            amount += n
            used += size
        if current:
            current[0]["drawing_batch_bounds"] = [first, boundaries[-1]]
            current[0]["drawing_legal_boundaries"] = [
                x for x in boundaries if first <= x <= boundaries[-1]
            ]
            result.append(current)
    if sum(len(c.get("drawing_evidence", [])) for batch in result for c in batch) != len(
        review["drawing_evidence"]
    ):
        raise ReviewedExcelError("绘图证据未能完整分批，未调用模型")
    if not result:
        raise ReviewedExcelError("没有可发送的证据")
    return result


def rectangles_cover(bounds: tuple, rectangles: list[tuple]) -> bool:
    """Prove full two-dimensional union coverage, including empty space and holes."""
    x1, y1, x2, y2 = bounds
    bands = sorted(
        {
            y1,
            y2 + 1,
            *(max(y1, a) for _, a, _, b in rectangles if a <= y2 and b >= y1),
            *(min(y2 + 1, b + 1) for _, a, _, b in rectangles if a <= y2 and b >= y1),
        }
    )
    for top, bottom in zip(bands, bands[1:]):
        right = x1 - 1
        for left, end in sorted(
            (max(x1, a), min(x2, c))
            for a, b, c, d in rectangles
            if b <= top and bottom - 1 <= d and a <= x2 and c >= x1
        ):
            if left > right + 1:
                break
            right = max(right, end)
        if right < x2:
            return False
    return True


def drawing_bounds(obj: dict) -> tuple:
    if obj.get("geometry_error"):
        raise ReviewedExcelError(obj["geometry_error"])
    return (obj["from"]["col"] + 1, obj["from"]["row"] + 1, obj["last_col"], obj["last_row"])


def _row_height_emu(sheet, row: int) -> float:
    """Worksheet row height in EMUs (1 pt = 12700 EMU). Uses sheet records only."""
    height = sheet.row_dimensions[row].height
    if height is None:
        height = sheet.sheet_format.defaultRowHeight or 15
    return float(height) * 12700


def _col_width_emu(sheet, col: int) -> float:
    """Worksheet column width in EMUs from dimension records (ECMA-376 pixel path).

    Uses stored width and OOXML default max-digit width 7 when theme fonts are
    not separately verified — not a glyph-advance estimate from cell fonts.
    """
    from openpyxl.utils import get_column_letter

    letter = get_column_letter(col)
    width = sheet.column_dimensions[letter].width
    if width is None:
        width = sheet.sheet_format.defaultColWidth or 8.43
    max_digit = 7
    pixels = int(((256 * float(width) + int(128 / max_digit)) / 256) * max_digit)
    return float(max(pixels, 1) * 9525)


def _extent_last_cell(sheet, start: dict, cx: int, cy: int) -> tuple[int, int]:
    """Map oneCell from+extent into inclusive 1-based last_row/last_col via sheet sizes."""
    last_col = int(start["col"]) + 1
    last_row = int(start["row"]) + 1
    remain_x = float(start.get("colOff") or 0) + float(cx)
    remain_y = float(start.get("rowOff") or 0) + float(cy)
    guard = 0
    while True:
        guard += 1
        if guard > 1024 or last_col > int(sheet.max_column) + 64:
            raise ReviewedExcelError("绘图水平范围超出可核验列")
        width = _col_width_emu(sheet, last_col)
        if remain_x <= width:
            break
        remain_x -= width
        last_col += 1
    guard = 0
    while True:
        guard += 1
        if guard > 1024 or last_row > int(sheet.max_row) + 64:
            raise ReviewedExcelError("绘图垂直范围超出可核验行")
        height = _row_height_emu(sheet, last_row)
        if remain_y <= height:
            break
        remain_y -= height
        last_row += 1
    return last_row, last_col


def set_drawing_columns(sheet, drawings: list[dict]) -> int:
    """Resolve drawing cell coverage without guessing glyph advances.

    twoCell anchors: use to-marker cells (sub-cell offsets stay inside end cells).
    oneCell anchors: map extent EMUs through worksheet row/column dimension records.
    Prior geometry_error (transform/overflow) stays fail-closed.
    """
    for obj in drawings:
        if obj.get("geometry_error"):
            continue
        end = obj["to"]
        if end is not None:
            obj["last_row"] = end["row"] + 1
            obj["last_col"] = end["col"] + 1
            continue
        extent = obj.get("extent") or {}
        cx, cy = extent.get("cx"), extent.get("cy")
        if not isinstance(cx, int) or not isinstance(cy, int) or cx < 0 or cy < 0:
            obj["geometry_error"] = "绘图边界依赖字体指标或行高，尚未核验，不能分批或派生"
            continue
        try:
            last_row, last_col = _extent_last_cell(sheet, obj["from"], cx, cy)
        except ReviewedExcelError:
            obj["geometry_error"] = "绘图边界依赖字体指标或行高，尚未核验，不能分批或派生"
            continue
        obj["last_row"] = last_row
        obj["last_col"] = last_col
    return max([int(sheet.max_column), *(int(obj.get("last_col", 0)) for obj in drawings)])


def select_drawing_objects(info: dict, sheet: str, rectangles: list[tuple]) -> list[dict]:
    """Accept arbitrary rectangles only when each touched object is wholly contained."""
    selected = []
    for obj in info.get("drawing_evidence", []):
        if obj["sheet"] != sheet:
            continue
        x1, y1, x2, y2 = bounds = drawing_bounds(obj)
        intersects = any(x1 <= c and a <= x2 and y1 <= d and b <= y2 for a, b, c, d in rectangles)
        if intersects and not rectangles_cover(bounds, rectangles):
            raise ReviewedExcelError("所选范围切断绘图对象，暂不能提取")
        if intersects:
            selected.append(obj)
    return selected


def _anchor_identity(anchor) -> str:
    node = anchor.find(f".//{{{X}}}cNvPr")
    identifier = node.get("id") if node is not None else None
    if not isinstance(identifier, str) or not identifier.isdecimal():
        raise ReviewedExcelError("绘图对象编号无效")
    return identifier


def restore_drawing_parts(source: bytes, derived: bytes, selected: list[dict]) -> bytes:
    """Restore selected validated shape XML and the sheet relationship after save.

    The regular extractor keeps original row and column coordinates; anchors
    therefore require no translation. Reject all connected-object references.
    """
    from io import BytesIO
    from xml.etree import ElementTree as XML
    from zipfile import ZIP_DEFLATED

    if not selected:
        return derived
    if len({d["part"] for d in selected}) != 1:
        raise ReviewedExcelError("跨绘图部件裁切暂不支持")
    part = selected[0]["part"]
    wanted = {d["object_id"] for d in selected}
    content_type = "http://schemas.openxmlformats.org/package/2006/content-types"
    drawing_type = "application/vnd.openxmlformats-officedocument.drawing+xml"
    with ZipFile(BytesIO(source)) as old, ZipFile(BytesIO(derived)) as result:
        if part in result.namelist():
            raise ReviewedExcelError("派生绘图部件冲突")
        if "xl/theme/theme1.xml" in old.namelist() and (
            "xl/theme/theme1.xml" not in result.namelist()
            or old.read("xl/theme/theme1.xml") != result.read("xl/theme/theme1.xml")
        ):
            raise ReviewedExcelError("绘图主题发生变化，不能保证保真")
        root = ET.fromstring(old.read(part))
        original_xml = {
            _anchor_identity(anchor): XML.tostring(anchor, encoding="utf-8") for anchor in root
        }
        for anchor in list(root):
            if any(n.tag.split("}")[-1] in {"stCxn", "endCxn"} for n in anchor.iter()):
                raise ReviewedExcelError("绘图连接对象引用暂不能安全裁切")
            node = anchor.find(f".//{{{X}}}cNvPr")
            if node is None or node.get("id") not in wanted:
                root.remove(anchor)
        drawing_data = XML.tostring(root, encoding="utf-8", xml_declaration=True)
        workbook = ET.fromstring(result.read("xl/workbook.xml"))
        sheets = workbook.findall(f"{{{S}}}sheets/{{{S}}}sheet")
        sheets = [sheet for sheet in sheets if sheet.get("name") == selected[0]["sheet"]]
        if len(sheets) != 1:
            raise ReviewedExcelError("派生工作表关系无法核验")
        links = _relations(result, "xl/workbook.xml")
        rid = sheets[0].get(f"{{{R}}}id")
        if rid not in links or links[rid][0] != f"{R}/worksheet":
            raise ReviewedExcelError("派生工作表关系无法核验")
        sheet_part = links[rid][1]
        directory, name = posixpath.split(sheet_part)
        rel_part = f"{directory}/_rels/{name}.rels"
        rel_root = (
            ET.fromstring(result.read(rel_part))
            if rel_part in result.namelist()
            else XML.Element(f"{{{P}}}Relationships")
        )
        existing_ids = {rel.get("Id") for rel in rel_root}
        new_id = next(
            f"rId{i}" for i in range(1, len(existing_ids) + 2) if f"rId{i}" not in existing_ids
        )
        XML.SubElement(
            rel_root,
            f"{{{P}}}Relationship",
            {"Id": new_id, "Type": f"{R}/drawing", "Target": posixpath.relpath(part, directory)},
        )
        sheet_root = ET.fromstring(result.read(sheet_part))
        if sheet_root.find(f"{{{S}}}drawing") is not None:
            raise ReviewedExcelError("派生工作表已含未预期绘图")
        drawing = XML.Element(f"{{{S}}}drawing", {f"{{{R}}}id": new_id})
        after = {
            "legacyDrawing",
            "legacyDrawingHF",
            "picture",
            "oleObjects",
            "controls",
            "webPublishItems",
            "tableParts",
            "extLst",
        }
        position = next(
            (i for i, child in enumerate(sheet_root) if child.tag.split("}")[-1] in after),
            len(sheet_root),
        )
        sheet_root.insert(position, drawing)
        types = ET.fromstring(result.read("[Content_Types].xml"))
        XML.SubElement(
            types,
            f"{{{content_type}}}Override",
            {"PartName": "/" + part, "ContentType": drawing_type},
        )
        replacements = {
            sheet_part: XML.tostring(sheet_root, encoding="utf-8", xml_declaration=True),
            rel_part: XML.tostring(rel_root, encoding="utf-8", xml_declaration=True),
            "[Content_Types].xml": XML.tostring(types, encoding="utf-8", xml_declaration=True),
            part: drawing_data,
        }
        out = BytesIO()
        with ZipFile(out, "w", ZIP_DEFLATED) as z:
            for entry in result.infolist():
                z.writestr(entry, replacements.pop(entry.filename, result.read(entry.filename)))
            for name, data in replacements.items():
                z.writestr(name, data)
    output = out.getvalue()
    with ZipFile(BytesIO(output)) as z:
        found = extract_drawing_evidence(z)
        written_root = ET.fromstring(z.read(part))
        written_xml = {
            _anchor_identity(anchor): XML.tostring(anchor, encoding="utf-8")
            for anchor in written_root
        }
    if set(written_xml) != wanted or any(written_xml[key] != original_xml[key] for key in wanted):
        raise ReviewedExcelError("派生绘图对象 XML 写出后核验失败")
    # Geometry, text, types and identity must agree exactly with the source.
    keys = ("sheet", "part", "object_id", "kind", "anchor_type", "from", "to", "extent", "text")
    expected = [{k: d[k] for k in keys} for d in selected]
    actual = [{k: d[k] for k in keys} for d in found]
    if actual != expected:
        raise ReviewedExcelError("派生绘图对象写出后核验失败")
    return output


def validate_region_drawings(review: dict, sent: list[dict], region: dict) -> None:
    from app.application.reviewed_excel_evidence import bounded_range

    selected = select_drawing_objects(review, region["sheet"], [bounded_range(region["range"])])
    known = {(d["sheet"], d["part"], d["object_id"]) for d in sent}
    inside = {(d["sheet"], d["part"], d["object_id"]) for d in selected}
    if not inside <= known:
        raise ReviewedExcelError("模型区域包含未发送绘图对象")
    refs = region["drawing_evidence"]
    if (
        not isinstance(refs, list)
        or len(refs) > 100
        or any(
            not isinstance(ref, dict)
            or set(ref) != {"sheet", "part", "object_id"}
            or (ref["sheet"], ref["part"], ref["object_id"]) not in inside
            for ref in refs
        )
    ):
        raise ReviewedExcelError("模型绘图证据不在区域内已发送范围")


def drawing_coverage(review: dict, batches: list, sheet: str, rectangles: list) -> dict:
    objects = [obj for obj in review.get("drawing_evidence", []) if obj["sheet"] == sheet]
    sent = {
        (d["sheet"], d["part"], d["object_id"])
        for batch in batches
        for cell in batch
        for d in cell.get("drawing_evidence", [])
    }
    return {
        "drawing_unread": [
            obj["object_id"]
            for obj in objects
            if (sheet, obj["part"], obj["object_id"]) not in sent
        ],
        "drawing_unassigned": [
            obj["object_id"]
            for obj in objects
            if not any(rectangles_cover(drawing_bounds(obj), [r[:4]]) for r in rectangles)
        ],
    }


def grouping_drawings(review: dict, fragments: list[dict], cells: list[dict]) -> list[dict]:
    import json

    from app.application.reviewed_excel_evidence import bounded_range

    chosen = {}
    for part in fragments:
        for obj in select_drawing_objects(review, part["sheet"], [bounded_range(part["range"])]):
            chosen[(obj["sheet"], obj["part"], obj["object_id"])] = obj
    if len(chosen) != len(review["drawing_evidence"]):
        raise ReviewedExcelError("绘图证据覆盖不完整，不能分组")
    drawings = list(chosen.values())
    from app.application.reviewed_excel_evidence import MAX_MODEL_CELLS, MAX_MODEL_TEXT_BYTES

    if (
        len(cells) + len(drawings) > MAX_MODEL_CELLS
        or sum(len(c["value"].encode()) + 64 for c in cells)
        + sum(len(json.dumps(obj, ensure_ascii=False).encode()) + 64 for obj in drawings)
        > MAX_MODEL_TEXT_BYTES
    ):
        raise ReviewedExcelError("分组绘图证据超过安全预算")
    return drawings


def validate_group_drawings(group_input: dict, group: dict) -> None:
    from app.application.reviewed_excel_evidence import bounded_range

    known = {(d["sheet"], d["part"], d["object_id"]): d for d in group_input["drawing_evidence"]}
    refs = group["drawing_evidence"]
    if not isinstance(refs, list) or len(refs) > 100:
        raise ReviewedExcelError("模型分组绘图证据格式无效")
    seen = set()
    for ref in refs:
        if not isinstance(ref, dict) or set(ref) != {"sheet", "part", "object_id"}:
            raise ReviewedExcelError("模型分组绘图证据格式无效")
        key = tuple(ref[k] for k in ("sheet", "part", "object_id"))
        if key not in known or key in seen:
            raise ReviewedExcelError("模型分组绘图证据不在已发送范围或重复")
        seen.add(key)
        if not any(
            part["sheet"] == ref["sheet"]
            and rectangles_cover(drawing_bounds(known[key]), [bounded_range(part["range"])])
            for part in group["fragments"]
        ):
            raise ReviewedExcelError("模型分组绘图证据不在片段内")
