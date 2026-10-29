"""Bounded and fair evidence selection from the complete inspected workbook."""

from collections import deque


def bounded_range(value: str) -> tuple[int, int, int, int]:
    from openpyxl.utils.cell import range_boundaries

    c1, r1, c2, r2 = range_boundaries(value)
    if (
        not isinstance(c1, int)
        or not isinstance(r1, int)
        or not isinstance(c2, int)
        or not isinstance(r2, int)
    ):
        raise ValueError("Evidence ranges must include both row and column coordinates")
    return c1, r1, c2, r2


def summarize_model_input(
    messages: list[dict],
    cells: list[dict],
    *,
    source_sha256: str,
    provider: str,
    model: str,
    max_output_tokens: int,
) -> dict:
    """Store verifiable references and a digest, never the full model input."""
    import hashlib
    import json

    encoded = json.dumps(
        messages, ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")
    return {
        "schema_version": 1,
        "source_sha256": source_sha256,
        "input_sha256": hashlib.sha256(encoded).hexdigest(),
        "input_bytes": len(encoded),
        "digest_encoding": "json-utf8-sortkeys-compact-v1",
        "provider": provider,
        "model": model,
        "max_output_tokens": max_output_tokens,
        "sent_cells": [
            {
                "sheet": cell["sheet"],
                "cell": cell["cell"],
                "kind": cell["kind"],
                "context_only": bool(cell.get("context_only")),
            }
            for cell in cells
        ],
    }


def candidate_sheets(review: dict, rules: list[dict]) -> list[dict]:
    """Find additional dense blocks, without treating blank rows alone as boundaries."""
    from openpyxl.utils.cell import (
        column_index_from_string,
        coordinate_from_string,
        get_column_letter,
    )

    overviews = {sheet["sheet"]: sheet for sheet in review["sheet_overview"]}
    output = []
    for sheet in rules:
        if len(sheet["regions"]) != 1:
            output.append(sheet)
            continue
        name = sheet["sheet"]
        cells = [
            (column_index_from_string(col), row)
            for cell in review["cells"]
            if cell["sheet"] == name
            for col, row in [coordinate_from_string(cell["cell"])]
        ]
        rows = sorted({row for _, row in cells})
        boundaries = []
        for previous, current in zip(rows, rows[1:]):
            if current - previous < 3:
                continue
            left = [(col, row) for col, row in cells if row < current]
            right = [(col, row) for col, row in cells if row >= current]
            if any(
                len(part) < 4
                or len({row for _, row in part}) < 2
                or len({col for col, _ in part}) < 2
                for part in (left, right)
            ):
                continue
            if any(
                mr1 < current <= mr2
                for merge in overviews[name].get("merged_regions", [])
                for _, mr1, _, mr2 in [bounded_range(merge)]
            ):
                continue
            boundaries.append(current)
        if not boundaries:
            output.append(sheet)
            continue
        max_col = max(1, overviews[name]["columns"])
        max_row = max(1, overviews[name]["rows"])
        starts = [1] + boundaries
        regions = []
        for index, start in enumerate(starts):
            end = starts[index + 1] - 1 if index + 1 < len(starts) else max_row
            region = dict(sheet["regions"][0])
            region["range"] = f"A{start}:{get_column_letter(max_col)}{end}"
            for field in ("evidence", "header_cells"):
                region[field] = [
                    ref
                    for ref in region[field]
                    if start <= coordinate_from_string(ref["cell"])[1] <= end
                ]
            region["needs_confirmation"] = True
            regions.append(region)
        output.append({**sheet, "regions": regions})
    return output


def build_structure_map(review: dict, rules: list[dict]) -> list[dict]:
    """Keep the local full index; later model payloads must still obey their budget."""
    by_sheet = {sheet["sheet"]: sheet for sheet in candidate_sheets(review, rules)}
    return [
        {
            **sheet,
            "regions": by_sheet[sheet["sheet"]]["regions"],
            "nonempty_cells": sum(cell["sheet"] == sheet["sheet"] for cell in review["cells"]),
        }
        for sheet in review["sheet_overview"]
    ]


def build_relationship_context(structure_map: list[dict], sent_cells: list[dict]) -> dict:
    """Return bounded coordinate observations; never assert a business relationship."""
    from openpyxl.utils.cell import (
        column_index_from_string,
        coordinate_from_string,
    )

    sent = {(cell["sheet"], cell["cell"]) for cell in sent_cells}
    sheets = []
    pair_budget, merge_budget, match_budget = 200, 100, 100
    for sheet in structure_map:
        name = sheet["sheet"]
        regions = sheet["regions"]
        pairs = []
        for left, right in zip(regions, regions[1:]):
            if pair_budget == 0:
                break
            lc1, lr1, lc2, lr2 = bounded_range(left["range"])
            rc1, rr1, rc2, rr2 = bounded_range(right["range"])
            pairs.append(
                {
                    "first": left["range"],
                    "second": right["range"],
                    "row_gap": max(0, rr1 - lr2 - 1),
                    "column_gap": max(0, rc1 - lc2 - 1),
                    "row_overlap": lr1 <= rr2 and rr1 <= lr2,
                    "column_overlap": lc1 <= rc2 and rc1 <= lc2,
                }
            )
            pair_budget -= 1
        merges = sheet.get("merged_regions", [])
        included_merges = merges[:merge_budget]
        merge_budget -= len(included_merges)
        refs: list[dict] = []
        total_refs = 0
        for region in regions:
            for field, kind in [
                ("evidence", "possible_title"),
                ("possible_shared_context", "possible_shared_note"),
            ]:
                for ref in region[field]:
                    entry = {
                        "region": region["range"],
                        "cell": ref["cell"],
                        "kind": kind,
                        "value_sent": (name, ref["cell"]) in sent,
                    }
                    total_refs += 1
                    if len(refs) < 100:
                        refs.append(entry)
        rows: dict[int, list[tuple[int, str, str]]] = {}
        for cell in sent_cells:
            if cell["sheet"] != name or cell["kind"] == "formula":
                continue
            col, row = coordinate_from_string(cell["cell"])
            value = " ".join(cell["value"][:160].casefold().split())
            if value and len(value) <= 80 and not value.replace(".", "", 1).isdigit():
                rows.setdefault(row, []).append(
                    (column_index_from_string(col), cell["cell"], value)
                )
        populated_rows = sorted(row for row, cells in rows.items() if len(cells) >= 2)
        row_matches = []
        total_matches = 0
        for index, left_row in enumerate(populated_rows):
            for right_row in populated_rows[index + 1 :]:
                same = [
                    (lcell, rcell, lcol == rcol)
                    for lcol, lcell, lv in rows[left_row]
                    for rcol, rcell, rv in rows[right_row]
                    if lv == rv
                ]
                if not same:
                    continue
                total_matches += 1
                if match_budget:
                    row_matches.append(
                        {
                            "first_row": left_row,
                            "second_row": right_row,
                            "matches": [
                                {"first_cell": a, "second_cell": b, "same_column": same_col}
                                for a, b, same_col in same[:20]
                            ],
                            "matches_omitted": max(0, len(same) - 20),
                        }
                    )
                    match_budget -= 1
        sheet_sent = sum(cell["sheet"] == name for cell in sent_cells)
        sheets.append(
            {
                "sheet": name,
                "candidate_adjacencies": pairs,
                "adjacencies_omitted": max(0, len(regions) - 1 - len(pairs)),
                "possible_context_refs": refs,
                "context_refs_omitted": max(0, total_refs - len(refs)),
                "merged_ranges": included_merges,
                "merged_omitted": len(merges) - len(included_merges),
                "sent_row_matches": row_matches,
                "row_matches_omitted": total_matches - len(row_matches),
                "sent_nonempty": sheet_sent,
                "unread_nonempty": max(0, sheet["nonempty_cells"] - sheet_sent),
            }
        )
    return {
        "kind": "structural_observations_only",
        "sheets": sheets,
        "limitation": "Coordinates and repeated values do not prove continuation, shared dependency, or a category; unseen values remain unknown.",
    }


def unresolved_fragments(
    candidate: dict, model_regions: list[dict], merges: list[str]
) -> list[dict]:
    """Retain disjoint unmodelled rectangles; never manufacture a split through a merge."""
    from openpyxl.utils.cell import get_column_letter

    fragments = [bounded_range(candidate["range"])]
    for region in model_regions:
        x1, y1, x2, y2 = bounded_range(region["range"])
        next_fragments = []
        for c1, r1, c2, r2 in fragments:
            ic1, ir1, ic2, ir2 = max(c1, x1), max(r1, y1), min(c2, x2), min(r2, y2)
            if ic1 > ic2 or ir1 > ir2:
                next_fragments.append((c1, r1, c2, r2))
                continue
            if r1 < ir1:
                next_fragments.append((c1, r1, c2, ir1 - 1))
            if ir2 < r2:
                next_fragments.append((c1, ir2 + 1, c2, r2))
            if c1 < ic1:
                next_fragments.append((c1, ir1, ic1 - 1, ir2))
            if ic2 < c2:
                next_fragments.append((ic2 + 1, ir1, c2, ir2))
        if len(next_fragments) > 500:
            return []
        fragments = next_fragments
    retained = []
    for c1, r1, c2, r2 in fragments:
        cut_merge = False
        for merged in merges:
            mc1, mr1, mc2, mr2 = bounded_range(merged)
            if (
                c1 <= mc2
                and mc1 <= c2
                and r1 <= mr2
                and mr1 <= r2
                and not (c1 <= mc1 <= mc2 <= c2 and r1 <= mr1 <= mr2 <= r2)
            ):
                cut_merge = True
                break
        if cut_merge:
            continue
        region = dict(candidate)
        region.update(
            {
                "range": f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}",
                "labels": ["unknown"],
                "label_source": "unassigned",
                "evidence": [],
                "header_cells": [],
                "merged_ranges": [],
                "needs_confirmation": True,
                "parent_heading": "",
                "dependency_status": "模型未覆盖，边界和共享说明待确认",
            }
        )
        retained.append(region)
    return retained


def flag_repeated_header_conflicts(regions: list[dict], sent_cells: list[dict]) -> list[dict]:
    """Flag adjacent incompatible model suggestions sharing a multi-cell row signature."""
    from openpyxl.utils.cell import (
        column_index_from_string,
        coordinate_from_string,
    )

    def row_signatures(region):
        c1, r1, c2, r2 = bounded_range(region["range"])
        rows = {}
        for cell in sent_cells:
            if cell["sheet"] != region["sheet"] or cell["kind"] == "formula":
                continue
            column, row = coordinate_from_string(cell["cell"])
            col = column_index_from_string(column)
            value = " ".join(cell["value"][:160].casefold().split())
            if (
                c1 <= col <= c2
                and r1 <= row <= r2
                and value
                and len(value) <= 80
                and not value.replace(".", "", 1).isdigit()
            ):
                rows.setdefault(row, set()).add((col, value))
        return {row: tokens for row, tokens in rows.items() if len(tokens) >= 2}

    import re

    values = {(cell["sheet"], cell["cell"]): cell["value"][:160] for cell in sent_cells}
    output = [dict(region) for region in regions]
    for region in output:
        evidence = " ".join(
            values.get((ref["sheet"], ref["cell"]), "") for ref in region["evidence"]
        )
        if re.search(
            r"不是|并非|並非|不属于|不屬於|不作为|不作為|\bnot\s+(?:a|an|the)\b|\bis\s+not\b",
            evidence,
            re.I,
        ):
            region["context_warning"] = "依据含否定描述，需结合上下文核对模型建议。"
    for current in output:
        if current["labels"] == ["unknown"]:
            continue
        c1, r1, c2, _ = bounded_range(current["range"])
        prior = []
        for other in regions:
            if (
                other is current
                or other["sheet"] != current["sheet"]
                or other["labels"] == ["unknown"]
            ):
                continue
            oc1, _or1, oc2, or2 = bounded_range(other["range"])
            if (
                0 <= r1 - or2 - 1 <= 3
                and min(c2, oc2) - max(c1, oc1) + 1 >= min(c2 - c1 + 1, oc2 - oc1 + 1) * 0.8
            ):
                prior.append((or2, other))
        if not prior:
            continue
        parent = max(prior, key=lambda entry: entry[0])[1]
        parent_rows, current_rows = row_signatures(parent), row_signatures(current)
        header_matches = [
            (pr, cr, sorted(old & new))
            for pr, old in parent_rows.items()
            for cr, new in current_rows.items()
            if len(old & new) >= 2
        ]
        if header_matches:
            parent_row, current_row, matched_columns = header_matches[0]
            if not (set(parent["labels"]) & set(current["labels"])):
                current["suggested_labels"] = current["labels"]
                current["labels"] = ["unknown"]
                current["context_warning"] = (
                    current.get("context_warning", "")
                    + "相邻区域出现重复表头且模型分类不同，可能是续表，需人工确认。"
                )
            current["continuation_context"] = {
                "parent_range": parent["range"],
                "parent_header_row": parent_row,
                "header_row": current_row,
                "matched_columns": [col for col, _ in matched_columns],
            }
    return output


def value_extent(region: dict, cells: list[dict]) -> str | None:
    """Bound occupied cells for display only; not a safe crop or layout boundary."""
    from openpyxl.utils.cell import (
        column_index_from_string,
        coordinate_from_string,
        get_column_letter,
    )

    c1, r1, c2, r2 = bounded_range(region["range"])
    points = []
    for cell in cells:
        if cell["sheet"] == region["sheet"]:
            column, row = coordinate_from_string(cell["cell"])
            col = column_index_from_string(column)
            if c1 <= col <= c2 and r1 <= row <= r2:
                points.append((col, row))
    if not points:
        return None
    return f"{get_column_letter(min(c for c, _ in points))}{min(r for _, r in points)}:{get_column_letter(max(c for c, _ in points))}{max(r for _, r in points)}"


def candidate_fragment_links(regions: list[dict], sent_cells: list[dict] | None = None) -> dict:
    """Suggest reviewable relations from shared dependencies or observed headers, never labels alone."""
    from openpyxl.utils.cell import (
        column_index_from_string,
        coordinate_from_string,
        get_column_letter,
    )

    cells = {(c["sheet"], c["cell"]): c for c in (sent_cells or [])}
    identity_fields = {
        "订单号": "number",
        "单号": "number",
        "报价单号": "number",
        "客户": "customer",
        "客户名称": "customer",
        "日期": "date",
        "order number": "number",
        "customer": "customer",
        "date": "date",
    }

    def identities(region):
        c1, r1, c2, r2 = bounded_range(region["range"])
        found = {}
        for (sheet, coordinate), cell in cells.items():
            if sheet != region["sheet"] or cell.get("kind") == "formula":
                continue
            column, row = coordinate_from_string(coordinate)
            col = column_index_from_string(column)
            if not (c1 <= col < c2 and r1 <= row <= r2):
                continue
            field = identity_fields.get(cell["value"].strip().casefold().rstrip(":："))
            value_cell = cells.get((sheet, f"{get_column_letter(col + 1)}{row}"))
            if field and value_cell and value_cell.get("kind") != "formula":
                value = value_cell["value"].strip().casefold()
                if value:
                    found.setdefault(field, []).append((value, value_cell["cell"]))
        return found

    identity_cache = [identities(region) for region in regions]
    links: list[dict] = []
    total = 0
    for index, first in enumerate(regions):
        for second_index in range(index + 1, len(regions)):
            second = regions[second_index]
            if first["sheet"] != second["sheet"]:
                continue
            left_refs = {(ref["sheet"], ref["cell"]) for ref in first.get("dependencies", [])}
            right_refs = {(ref["sheet"], ref["cell"]) for ref in second.get("dependencies", [])}
            shared = sorted(left_refs & right_refs)
            first_context = first.get("continuation_context") or {}
            second_context = second.get("continuation_context") or {}
            repeated = (
                first_context.get("parent_range") == second["range"]
                or second_context.get("parent_range") == first["range"]
            )
            if not shared and not repeated:
                continue
            left_identity, right_identity = identity_cache[index], identity_cache[second_index]
            matches: list[dict] = []
            conflicts: list[dict] = []
            for field in sorted(left_identity.keys() & right_identity.keys()):
                left, right = left_identity[field], right_identity[field]
                if len(left) != 1 or len(right) != 1:
                    continue
                item = {"field": field, "first_cell": left[0][1], "second_cell": right[0][1]}
                (matches if left[0][0] == right[0][0] else conflicts).append(item)
            total += 1
            if len(links) < 200:
                links.append(
                    {
                        "sheet": first["sheet"],
                        "first": first["range"],
                        "second": second["range"],
                        "shared_dependencies": [
                            {"sheet": sheet, "cell": cell} for sheet, cell in shared[:20]
                        ],
                        "dependencies_omitted": max(0, len(shared) - 20),
                        "repeated_header_observed": repeated,
                        "identity_matches": matches,
                        "identity_conflicts": conflicts,
                        "status": "identity_conflict" if conflicts else "candidate",
                    }
                )
    return {"items": links, "omitted": total - len(links)}


MAX_MODEL_CELLS = 500
MAX_MODEL_TEXT_BYTES = 96_000
MAX_VALUE_CHARS = 4000


def _model_cell_cost(cell: dict) -> int:
    return len(cell["value"].encode("utf-8")) + 64


def plan_complete_batches(review: dict, rules: list[dict]) -> list[list[dict]]:
    """Cover complete small workbooks or contiguous rows with bounded repeated context."""
    from openpyxl.utils.cell import (
        column_index_from_string,
        coordinate_from_string,
    )

    cells = review["cells"]
    if (
        len(cells) <= MAX_MODEL_CELLS
        and all(len(cell["value"]) <= MAX_VALUE_CHARS for cell in cells)
        and sum(map(_model_cell_cost, cells)) <= MAX_MODEL_TEXT_BYTES
    ):
        return [list(cells)]
    lookup = {(cell["sheet"], cell["cell"]): cell for cell in cells}
    batches: list[list[dict]] = []
    for sheet in candidate_sheets(review, rules):
        name = sheet["sheet"]
        for region in sheet["regions"]:
            start_batch = len(batches)
            c1, r1, c2, r2 = bounded_range(region["range"])
            grouped: dict[int, list[dict]] = {}
            for cell in cells:
                if cell["sheet"] != name:
                    continue
                col, row = coordinate_from_string(cell["cell"])
                if c1 <= column_index_from_string(col) <= c2 and r1 <= row <= r2:
                    grouped.setdefault(row, []).append(cell)
            contexts = []
            seen = set()
            for ref in region["evidence"] + region["possible_shared_context"]:
                key = (name, ref["cell"])
                if (
                    key in lookup
                    and key not in seen
                    and len(lookup[key]["value"]) <= MAX_VALUE_CHARS
                ):
                    contexts.append(lookup[key])
                    seen.add(key)
                    if len(contexts) == 10:
                        break
            current: list[dict] = []
            for row in sorted(grouped):
                row_cells = grouped[row]
                # An indivisible row beyond the hard budget remains explicitly unread.
                if (
                    len(row_cells) > MAX_MODEL_CELLS
                    or any(len(cell["value"]) > MAX_VALUE_CHARS for cell in row_cells)
                    or sum(map(_model_cell_cost, row_cells)) > MAX_MODEL_TEXT_BYTES
                ):
                    if current:
                        batches.append(current)
                        current = []
                    continue
                combined = current + row_cells
                context_keys = {(cell["sheet"], cell["cell"]) for cell in combined}
                missing_context = [
                    cell for cell in contexts if (cell["sheet"], cell["cell"]) not in context_keys
                ]
                if (
                    len(combined) + len(missing_context) > MAX_MODEL_CELLS
                    or sum(map(_model_cell_cost, combined + missing_context)) > MAX_MODEL_TEXT_BYTES
                ):
                    if current:
                        batches.append(current)
                    current = row_cells[:]
                else:
                    current = combined
            if current:
                batches.append(current)
            # Add context only if it fits; all primary row cells keep their order and are never displaced.
            for batch in batches[start_batch:]:
                if not batch:
                    continue
                batch_keys = {(cell["sheet"], cell["cell"]) for cell in batch}
                for context in contexts:
                    key = (context["sheet"], context["cell"])
                    if (
                        key not in batch_keys
                        and len(batch) < MAX_MODEL_CELLS
                        and sum(map(_model_cell_cost, batch)) + _model_cell_cost(context)
                        <= MAX_MODEL_TEXT_BYTES
                    ):
                        batch.append({**context, "context_only": True})
                        batch_keys.add(key)
    if not batches:
        raise ValueError("单行超出模型安全预算，未调用模型；请人工检查并分开文件")
    # Fair across sheets without disturbing contiguous order within any sheet.
    queues: dict[str, deque[list[dict]]] = {}
    for batch in batches:
        queues.setdefault(batch[0]["sheet"], deque()).append(batch)
    fair = []
    while any(queues.values()):
        for queue in queues.values():
            if queue:
                fair.append(queue.popleft())
    return fair
