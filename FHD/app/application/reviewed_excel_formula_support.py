"""Plan a bounded, source-coordinate formula dependency closure without changing Excel."""

import hashlib
import re
import warnings
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer, TokenizerError
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, get_column_letter

from app.application.elt.errors import ReviewedExcelError
from app.application.reviewed_excel_evidence import bounded_range

_SAFE_FUNCTIONS = {"SUM", "AVERAGE", "MIN", "MAX", "ROUND", "IF", "IFERROR", "COUNT", "COUNTA"}
_SUBTOTAL = re.compile(
    r"=SUBTOTAL\(\s*9\s*,\s*(\$?[A-Za-z]{1,3}\$?\d+:\$?[A-Za-z]{1,3}\$?\d+)\s*\)\Z",
    re.IGNORECASE,
)
_CELL_RANGE = re.compile(r"\$?[A-Za-z]{1,3}\$?\d+(?::\$?[A-Za-z]{1,3}\$?\d+)?\Z")


def _beyond_sheet_dimensions(col: int, row: int, sheet) -> bool:
    """True when a referenced cell lies outside the worksheet's used dimensions."""
    max_col = max(1, int(sheet.max_column or 1))
    max_row = max(1, int(sheet.max_row or 1))
    return col < 1 or row < 1 or col > max_col or row > max_row


def plan_formula_support(body: bytes, fragments: list[dict], *, max_cells: int = 20_000) -> dict:
    """Find required same-sheet cells and dependencies; do not rewrite or export a workbook.

    References to other sheets, external files, names, dynamic or unsupported
    functions are explicitly rejected until they can be preserved and verified.
    """
    if not fragments or len(fragments) > 20 or not 1 <= max_cells <= 20_000:
        raise ReviewedExcelError("公式依赖规划范围无效")
    if any(not isinstance(part, dict) or set(part) != {"sheet", "range"} for part in fragments):
        raise ReviewedExcelError("公式依赖片段无效")
    sheets = {part["sheet"] for part in fragments}
    if len(sheets) != 1:
        raise ReviewedExcelError("跨工作表文档暂不能验证")
    from app.application.reviewed_excel_template import inspect_workbook

    info = inspect_workbook(body)
    from app.application.reviewed_excel_template import _read_only_print_area_warnings

    with warnings.catch_warnings(record=True) as found:
        warnings.simplefilter("always")
        book = load_workbook(BytesIO(body), data_only=False, keep_links=True, rich_text=True)
    if found and _read_only_print_area_warnings(body, found) != info["read_only_print_areas"]:
        book.close()
        raise ReviewedExcelError("打印区域来源校验失败")
    try:
        sheet_name = next(iter(sheets))
        if sheet_name not in book.sheetnames:
            raise ReviewedExcelError("公式依赖来源工作表不存在")
        sheet = book[sheet_name]
        rectangles: list[tuple[int, int, int, int]] = []
        for part in fragments:
            try:
                bounds = bounded_range(part["range"])
                c1, r1, c2, r2 = bounds
                if not (1 <= c1 <= c2 <= sheet.max_column and 1 <= r1 <= r2 <= sheet.max_row):
                    raise ValueError("bounds")
            except (TypeError, ValueError):
                raise ReviewedExcelError("公式依赖片段坐标无效") from None
            rectangles.append(bounds)

        def selected(col: int, row: int) -> bool:
            return any(c1 <= col <= c2 and r1 <= row <= r2 for c1, r1, c2, r2 in rectangles)

        pending = [
            cell.coordinate
            for row in sheet.iter_rows()
            for cell in row
            if cell.data_type == "f"
            and isinstance(cell.column, int)
            and isinstance(cell.row, int)
            and selected(cell.column, cell.row)
        ]
        visited: set[str] = set()
        references: dict[str, list[str]] = {}
        outside: set[str] = set()
        inside: set[str] = set()
        proxies: set[str] = set()
        edges: dict[str, set[str]] = {}
        subtotals: set[str] = set()
        error_operand_warnings: list[dict] = []
        while pending:
            address = pending.pop()
            if address in visited:
                continue
            visited.add(address)
            value = sheet[address].value
            try:
                tokens = Tokenizer(value).items
            except (TypeError, TokenizerError):
                raise ReviewedExcelError("公式无法静态解析") from None
            has_subtotal = any(
                token.type == "FUNC"
                and token.subtype == "OPEN"
                and token.value[:-1].upper() == "SUBTOTAL"
                for token in tokens
            )
            if has_subtotal:
                match = _SUBTOTAL.fullmatch(value)
                if not match:
                    raise ReviewedExcelError("SUBTOTAL仅支持9和单个同表静态竖向范围")
                c1, r1, c2, r2 = bounded_range(match[1].replace("$", ""))
                if c1 != c2 or r1 > r2:
                    raise ReviewedExcelError("SUBTOTAL仅支持静态竖向范围")
                subtotals.add(address)
            direct = []
            edges[address] = set()
            for token in tokens:
                if token.type == "OPERAND" and token.subtype == "ERROR":
                    # Pre-existing dead refs (#REF!, etc.) are source-quality
                    # issues, not slice-boundary breaks. Allow with a receipt
                    # warning so derivation can proceed; do not chase them.
                    error_operand_warnings.append(
                        {
                            "sheet": sheet_name,
                            "cell": address,
                            "formula": value,
                            "token": token.value,
                            "message": "公式含错误引用（原件预存，警告放行）",
                        }
                    )
                    continue
                if (
                    token.type == "FUNC"
                    and token.subtype == "OPEN"
                    and token.value[:-1].upper() not in _SAFE_FUNCTIONS
                    and not (has_subtotal and token.value[:-1].upper() == "SUBTOTAL")
                ):
                    raise ReviewedExcelError("公式函数依赖尚无法保真")
                if token.type != "OPERAND" or token.subtype != "RANGE":
                    continue
                target = token.value
                if "!" in target or "[" in target:
                    raise ReviewedExcelError("跨工作表或外部公式依赖尚无法保真")
                if not _CELL_RANGE.fullmatch(target):
                    raise ReviewedExcelError("公式含名称或动态范围，暂不能保真")
                c1, r1, c2, r2 = bounded_range(target.replace("$", ""))
                if not (c1 <= c2 and r1 <= r2) or (c2 - c1 + 1) * (r2 - r1 + 1) > max_cells:
                    raise ReviewedExcelError("单个公式引用范围超出安全预算")
                direct.append(target)
                selected_refs: set[str] = set()
                includes_support = False
                beyond = False
                for row in range(r1, r2 + 1):
                    for col in range(c1, c2 + 1):
                        if _beyond_sheet_dimensions(col, row, sheet):
                            beyond = True
                            continue
                        ref = f"{get_column_letter(col)}{row}"
                        if selected(col, row):
                            inside.add(ref)
                            selected_refs.add(ref)
                        else:
                            outside.add(ref)
                            includes_support = True
                        if len(inside) + len(outside) > max_cells:
                            raise ReviewedExcelError("公式依赖总量超出安全预算")
                        if sheet[ref].data_type == "f":
                            if ref != address:
                                edges[address].add(ref)
                            if ref not in visited and ref != address:
                                pending.append(ref)
                if beyond:
                    error_operand_warnings.append(
                        {
                            "sheet": sheet_name,
                            "cell": address,
                            "formula": value,
                            "token": target,
                            "message": "公式引用超出工作表已用行列，越界部分已忽略（警告放行）",
                        }
                    )
                source_col, source_row = coordinate_from_string(address)
                if includes_support or not selected(
                    column_index_from_string(source_col), source_row
                ):
                    proxies.update(selected_refs)
            references[address] = direct
        # Circular / self-inclusive formulas: quarantine cells, do NOT abandon the
        # whole selection. Prefer rewriting self-inclusive aggregates (SUM/… that
        # include the formula cell); otherwise isolate the circular cell as text.
        circular_quarantine: list[dict] = []
        rewrites: dict[str, str] = {}

        def _self_inclusive_rewrite(address: str, formula: str) -> str | None:
            match = re.match(
                r"^=\s*(SUM|AVERAGE|MIN|MAX|COUNT|COUNTA)\s*\(\s*"
                r"(\$?[A-Za-z]{1,3}\$?\d+:\$?[A-Za-z]{1,3}\$?\d+)\s*\)\s*$",
                formula,
                re.IGNORECASE,
            )
            if not match:
                return None
            try:
                c1, r1, c2, r2 = bounded_range(match.group(2).replace("$", ""))
                acol, arow = coordinate_from_string(address)
                acol_i = column_index_from_string(acol)
            except (TypeError, ValueError):
                return None
            if not (c1 <= acol_i <= c2 and r1 <= arow <= r2):
                return None
            # Shrink the end that contains the formula cell (common self-SUM mistake).
            if c1 == c2 == acol_i:
                if arow == r2 and r2 > r1:
                    new_range = f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2 - 1}"
                elif arow == r1 and r2 > r1:
                    new_range = f"{get_column_letter(c1)}{r1 + 1}:{get_column_letter(c2)}{r2}"
                else:
                    return None
            elif r1 == r2 == arow:
                if acol_i == c2 and c2 > c1:
                    new_range = f"{get_column_letter(c1)}{r1}:{get_column_letter(c2 - 1)}{r2}"
                elif acol_i == c1 and c2 > c1:
                    new_range = f"{get_column_letter(c1 + 1)}{r1}:{get_column_letter(c2)}{r2}"
                else:
                    return None
            else:
                return None
            fn = match.group(1).upper()
            return f"={fn}({new_range})"

        # Pass 1: rewrite obvious self-inclusive aggregates and drop self-edges.
        for address in list(edges):
            value = sheet[address].value
            if not isinstance(value, str):
                continue
            rewritten = _self_inclusive_rewrite(address, value)
            if rewritten is None:
                # Also drop trivial self-loops even without rewrite.
                edges[address].discard(address)
                continue
            rewrites[address] = rewritten
            circular_quarantine.append(
                {
                    "sheet": sheet_name,
                    "cell": address,
                    "original_formula": value[:200],
                    "action": "rewrite_exclude_self",
                    "new_formula": rewritten,
                    "reason": "合计区间含自身单元格，已缩窄区间以隔离循环",
                }
            )
            edges[address].discard(address)
            # After rewrite, formula no longer refs itself as a formula edge target
            # within the old range end; clear edges to rebuild is unnecessary for DAG
            # check if self-loop was the only cycle contributor.

        def _cycle_nodes(graph: dict[str, set[str]]) -> set[str]:
            found: set[str] = set()
            done: set[str] = set()
            for start in graph:
                if start in done:
                    continue
                visiting: set[str] = set()
                stack = [(start, False)]
                path: list[str] = []
                while stack:
                    node, leaving = stack.pop()
                    if leaving:
                        visiting.remove(node)
                        if path and path[-1] == node:
                            path.pop()
                        done.add(node)
                        continue
                    if node in visiting:
                        # mark the cycle slice
                        if node in path:
                            found.update(path[path.index(node) :])
                        found.add(node)
                        continue
                    if node in done:
                        continue
                    visiting.add(node)
                    path.append(node)
                    stack.append((node, True))
                    stack.extend((child, False) for child in graph.get(node, set()))
            return found

        remaining_cycles = _cycle_nodes(edges) - set(rewrites)
        for address in sorted(remaining_cycles):
            value = sheet[address].value
            original = value if isinstance(value, str) else str(value)
            circular_quarantine.append(
                {
                    "sheet": sheet_name,
                    "cell": address,
                    "original_formula": original[:200],
                    "action": "isolate_as_text",
                    "new_formula": None,
                    "reason": "互相引用循环公式已隔离为文本，其余单元格照常导出",
                }
            )
            edges.pop(address, None)
            for targets in edges.values():
                targets.discard(address)
            # Drop from formula reference chase so install skips rewriting it as live formula
            references.pop(address, None)

        # Final DAG check — should be clean after quarantine.
        done: set[str] = set()
        for start in edges:
            if start in done:
                continue
            visiting: set[str] = set()
            stack = [(start, False)]
            while stack:
                node, leaving = stack.pop()
                if leaving:
                    visiting.remove(node)
                    done.add(node)
                    continue
                if node in visiting:
                    # Absolute last resort: isolate and continue
                    value = sheet[node].value
                    circular_quarantine.append(
                        {
                            "sheet": sheet_name,
                            "cell": node,
                            "original_formula": (value if isinstance(value, str) else str(value))[
                                :200
                            ],
                            "action": "isolate_as_text",
                            "new_formula": None,
                            "reason": "残留循环公式已隔离为文本，其余单元格照常导出",
                        }
                    )
                    edges.pop(node, None)
                    for targets in edges.values():
                        targets.discard(node)
                    references.pop(node, None)
                    visiting.clear()
                    stack.clear()
                    break
                if node in done:
                    continue
                visiting.add(node)
                stack.append((node, True))
                stack.extend((child, False) for child in edges.get(node, set()))
        # SUBTOTAL skips nested subtotals and filtered-out rows. Moving even a
        # transitive dependency to the support sheet can change those semantics.
        preserved: set[str] = set()
        pending = list(subtotals)
        while pending:
            address = pending.pop()
            if address in preserved:
                continue
            preserved.add(address)
            cell = sheet[address]
            if not selected(cell.column, cell.row):
                raise ReviewedExcelError("SUBTOTAL依赖必须完整保留于同一文档原坐标")
            for target in references.get(address, []):
                c1, r1, c2, r2 = bounded_range(target.replace("$", ""))
                pending.extend(
                    f"{get_column_letter(col)}{row}"
                    for row in range(r1, r2 + 1)
                    for col in range(c1, c2 + 1)
                    if not _beyond_sheet_dimensions(col, row, sheet)
                )
        return {
            "preserved_subtotal_cells": {
                address: sheet[address].value for address in sorted(preserved)
            },
            "subtotal_hidden_rows": {
                sheet[address].row: bool(sheet.row_dimensions[sheet[address].row].hidden)
                for address in preserved
            },
            "subtotal_filter": _filter_signature(sheet) if preserved else None,
            "source_sha256": hashlib.sha256(body).hexdigest(),
            "sheet": sheet_name,
            "fragments": fragments,
            "formula_references": references,
            "support_cells": sorted(outside),
            "selected_reference_cells": sorted(inside),
            "proxy_cells": sorted(proxies),
            "formula_error_operand_warnings": error_operand_warnings,
            "circular_quarantine": circular_quarantine,
            "circular_rewrites": rewrites,
        }
    finally:
        book.close()


def install_formula_support(book, plan: dict) -> str | None:
    """Install only planned support cells in memory; caller validates and saves the workbook.

    Selected references in a support range become formulas pointing at the
    editable main sheet, rather than frozen copies of the selected values.
    """
    from copy import copy

    from openpyxl.utils.cell import quote_sheetname

    verify_preserved_subtotals(book, plan)
    sheet = book[plan["sheet"]]
    # Apply circular quarantine before support-sheet wiring.
    for item in plan.get("circular_quarantine") or []:
        address = item.get("cell")
        if not isinstance(address, str) or not address.strip():
            continue
        action = item.get("action")
        if action == "rewrite_exclude_self" and item.get("new_formula"):
            sheet[address].value = item["new_formula"]
        elif action == "isolate_as_text":
            original = item.get("original_formula") or sheet[address].value
            sheet[address].value = f"[循环已隔离] {original}"
    for address, formula in (plan.get("circular_rewrites") or {}).items():
        if isinstance(address, str) and formula:
            sheet[address].value = formula
    support_cells = plan["support_cells"]
    if not support_cells:
        return None
    title = "依赖支撑"
    number = 1
    while title in book.sheetnames:
        number += 1
        title = f"依赖支撑_{number}"
    support = book.create_sheet(title)
    for address in support_cells:
        source = sheet[address]
        if source.hyperlink or source.comment:
            raise ReviewedExcelError("公式支撑单元格含未验证的链接或批注")
        target = support[address]
        target.value = source.value
        target._style = copy(source._style)
    for address in plan["proxy_cells"]:
        source = sheet[address]
        target = support[address]
        target.value = f"={quote_sheetname(sheet.title)}!{address}"
        target._style = copy(source._style)
    rectangles = [bounded_range(item["range"]) for item in plan["fragments"]]

    def selected(col, row):
        return any(c1 <= col <= c2 and r1 <= row <= r2 for c1, r1, c2, r2 in rectangles)

    for address in plan["formula_references"]:
        source = sheet[address]
        if not selected(source.column, source.row):
            continue
        tokens = Tokenizer(source.value).items
        changed = False
        for token in tokens:
            if token.type != "OPERAND" or token.subtype != "RANGE":
                continue
            c1, r1, c2, r2 = bounded_range(token.value.replace("$", ""))
            if any(
                not selected(col, row) for row in range(r1, r2 + 1) for col in range(c1, c2 + 1)
            ):
                token.value = f"{quote_sheetname(title)}!{token.value}"
                changed = True
        if changed:
            source.value = "=" + "".join(token.value for token in tokens)
    if book.calculation is None:
        from openpyxl.workbook.properties import CalcProperties

        book.calculation = CalcProperties()
    book.calculation.fullCalcOnLoad = True
    book.calculation.forceFullCalc = True
    verify_preserved_subtotals(book, plan)
    return title


def _filter_signature(sheet):
    from openpyxl.xml.functions import tostring

    return tostring(sheet.auto_filter.to_tree())


def verify_preserved_subtotals(book, plan: dict) -> None:
    """Fail closed if export changes any SUBTOTAL dependency or filter state."""
    cells = plan.get("preserved_subtotal_cells", {})
    if not cells:
        return
    if plan["sheet"] not in book.sheetnames:
        raise ReviewedExcelError("SUBTOTAL来源工作表未保留")
    sheet = book[plan["sheet"]]
    if (
        any(sheet[address].value != value for address, value in cells.items())
        or any(
            bool(sheet.row_dimensions[row].hidden) != hidden
            for row, hidden in plan["subtotal_hidden_rows"].items()
        )
        or _filter_signature(sheet) != plan["subtotal_filter"]
    ):
        raise ReviewedExcelError("SUBTOTAL原坐标依赖或隐藏/筛选状态未保真")
