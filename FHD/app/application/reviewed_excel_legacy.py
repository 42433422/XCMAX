"""Bounded pure-Python XLS derivatives, explicitly limited to data/template tests.

No office executable, macro execution, URL reader, corruption repair, or silent
formula/value conversion. Every derivative remains ineligible for publication.
"""

from __future__ import annotations

import hashlib
import struct
from copy import copy
from io import BytesIO
from xml.etree import ElementTree
from zipfile import ZipFile

import xlrd
from openpyxl import Workbook, load_workbook
from openpyxl.packaging.custom import StringProperty
from openpyxl.styles.numbers import BUILTIN_FORMATS
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from xlrd.compdoc import CompDoc, CompDocError

from app.application.etl.errors import ReviewedExcelError
from app.application.reviewed_excel_legacy_styles import assign_style, legacy_style
from app.application.reviewed_excel_template import MAX_BYTES, MAX_CELLS, inspect_workbook

OLE_MAGIC = bytes.fromhex("d0cf11e0a1b11ae1")
TEST_RESTRICTION = "仅供数据/模板测试；打印版式未验证，不得发布到企业模板库"
_PASSIVE = {
    0x809,
    0x0A,
    0x85,
    0x42,
    0x22,
    0xE0,
    0x31,
    0x41E,
    0x92,
    0x293,
    0xFC,
    0xFF,
    0x3C,
    0xFD,
    0x203,
    0x27E,
    0xBD,
    0x201,
    0xBE,
    0x205,
    0x200,
    0x208,
    0xE5,
    0x7D,
    0x225,
    0x55,
    0x23E,
    0x1D,
    0x99,
    0x3D,
    0x40,
    0x8D,
    0x5C,
    0xE1,
    0xC1,
    0xE2,
    0x161,
    0x13D,
    0x19,
    0x12,
    0x13,
    0x1AF,
    0x1BC,
    0xE,
    0x1B7,
    0xDA,
    0x160,
    0x8C,
    0xD,
    0xC,
    0xF,
    0x11,
    0x10,
    0x5F,
    0x2A,
    0x2B,
    0x82,
    0x80,
    0x81,
    0x14,
    0x15,
    0x83,
    0x84,
    0x26,
    0x27,
    0x28,
    0x29,
    0xA1,
    0x20B,
    0xD7,
}
# Known compatibility/style/view extensions are never treated as verified print
# content. Used XF extensions are rejected; unused ones remain only in original.
_EXTENSIONS = {
    0x1C0,
    0x1C1,
    0x87C,
    0x87D,
    0x892,
    0x863,
    0x896,
    0x88C,
    0x88E,
    0x89A,
    0x89C,
    0x88B,
    0x867,
}
_NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def _error(message):
    raise ReviewedExcelError("旧 XLS 转换未通过：" + message)


def inspect_legacy(body: bytes):
    """Reject active/unknown content and oversized grids before xlrd allocation."""
    if not body.startswith(OLE_MAGIC) or not 0 < len(body) <= MAX_BYTES:
        _error("文件不是受支持的 OLE/BIFF8 XLS 或超过 20 MB")
    book = None
    try:
        compound = CompDoc(body)
        allowed = {
            "Root Entry",
            "Workbook",
            "\x05SummaryInformation",
            "\x05DocumentSummaryInformation",
            "ETExtData",
        }
        entries = [entry for entry in compound.dirlist if entry.etype]
        if any(entry.name not in allowed for entry in entries):
            _error("存在宏、嵌入对象、加密或尚不支持的数据流")
        warnings = [
            TEST_RESTRICTION,
            "尚未验证 Windows/Excel/WPS 字体渲染及实际打印；不是完全保真转换",
            "窗口/冻结窗格/选区/缩放等视图元数据不迁移；原件完整保留",
        ]
        if any(entry.name == "ETExtData" for entry in entries):
            warnings.append("WPS ETExtData 扩展未迁入副本，仅完整保留于原始 XLS；其呈现影响未验证")
        stream = compound.get_named_stream("Workbook")
        if stream is None:
            _error("缺少 BIFF8 Workbook 数据流")
        position, section, declared, dimensions = 0, -1, 0, 0
        records: list[tuple[int, int, bytes]] = []
        extension_ids: set[int] = set()
        grid_bounds: dict[int, tuple[int, int]] = {}

        def admit_grid(row, col):
            old_row, old_col = grid_bounds.get(section, (0, 0))
            grid_bounds[section] = (max(row, old_row), max(col, old_col))
            if sum(r * c for r, c in grid_bounds.values()) > MAX_CELLS:
                _error("聚合工作表范围超过安全限制")

        while position + 4 <= len(stream):
            if len(records) >= 100_000:
                _error("BIFF 记录数量超过安全限制")
            opcode, length = struct.unpack_from("<HH", stream, position)
            if opcode == 0 and not any(stream[position:]):
                break
            end = position + 4 + length
            if end > len(stream) or opcode not in _PASSIVE | _EXTENSIONS:
                _error("含公式、外链、绘图、批注、验证规则或尚不支持的 BIFF 记录")
            data = stream[position + 4 : end]
            if opcode == 0x809:
                if len(data) < 4 or struct.unpack_from("<HH", data) not in {
                    (0x600, 5),
                    (0x600, 16),
                }:
                    _error("仅支持无宏的 BIFF8 普通工作表")
                if struct.unpack_from("<H", data, 2)[0] == 16:
                    section += 1
            if opcode == 0x85 and (len(data) < 6 or data[5] != 0):
                _error("宏表或图表页不支持")
            if opcode == 0x200:
                if len(data) < 14:
                    _error("工作表范围记录损坏")
                _, last_row, _, last_col = struct.unpack_from("<IIHH", data)
                if last_row > 65536 or last_col > 256:
                    _error("BIFF8 工作表范围无效")
                admit_grid(last_row, last_col)
                declared += last_row * last_col
                if declared > MAX_CELLS:
                    _error("工作表范围超过安全限制")
            if opcode in {0xFD, 0x203, 0x27E, 0xBD, 0x201, 0xBE, 0x205}:
                if len(data) < 6:
                    _error("单元格记录损坏")
                row, first_col = struct.unpack_from("<HH", data)
                col = first_col
                if opcode in {0xBD, 0xBE}:
                    col = struct.unpack_from("<H", data, len(data) - 2)[0]
                    stride = 6 if opcode == 0xBD else 2
                    if first_col > col or len(data) != 6 + stride * (col - first_col + 1):
                        _error("MULRK/MULBLANK 单元格记录结构损坏")
                elif len(data) != {0xFD: 10, 0x203: 14, 0x27E: 10, 0x201: 6, 0x205: 8}[opcode]:
                    _error("单元格记录长度无效")
                if first_col > 255 or col > 255:
                    _error("BIFF8 单元格列坐标无效")
                admit_grid(row + 1, col + 1)
                if (row + 1) * (col + 1) > MAX_CELLS:
                    _error("单元格坐标超过安全限制")
            if opcode == 0xE5:
                if len(data) < 2:
                    _error("合并区域记录损坏")
                count = struct.unpack_from("<H", data)[0]
                if len(data) != 2 + count * 8:
                    _error("合并区域记录长度无效")
                for offset in range(2, len(data), 8):
                    first_row, last_row, first_col, last_col = struct.unpack_from(
                        "<HHHH", data, offset
                    )
                    if first_row > last_row or first_col > last_col or last_col > 255:
                        _error("合并区域范围无效")
                    admit_grid(last_row + 1, last_col + 1)
            if opcode == 0x7D and (len(data) < 4 or struct.unpack_from("<H", data, 2)[0] > 255):
                _error("列范围记录无效")
            if opcode == 0x208:
                if len(data) < 16 or struct.unpack_from("<H", data)[0] >= MAX_CELLS:
                    _error("行格式范围超过安全限制")
                dimensions += 1
            elif opcode == 0x7D:
                first, last = struct.unpack_from("<HH", data)
                if first > last:
                    _error("列格式范围无效")
                dimensions += last - first + 1
            if dimensions > MAX_CELLS:
                _error("显式行列格式数量超过安全限制")
            if opcode in {0x12, 0x13} and any(data):
                _error("受保护的旧工作表暂不支持转换")
            if opcode in {0x14, 0x15} and data not in {b"", b"\0\0", b"\0\0\0"}:
                _error("非空页眉页脚尚不支持精确迁移")
            if opcode == 0x87D:
                if len(data) < 20:
                    _error("扩展格式记录损坏")
                extension_ids.add(struct.unpack_from("<H", data, 14)[0])
            records.append((section, opcode, data))
            position = end
        book = xlrd.open_workbook(file_contents=body, formatting_info=True, on_demand=True)
        if not 0 < book.nsheets <= 32 or sum(s.nrows * s.ncols for s in book.sheets()) > MAX_CELLS:
            _error("工作表数量或范围超过安全限制")
        used, text_bytes = {0}, 0
        for sheet in book.sheets():
            if sheet.default_additional_space_above or sheet.default_additional_space_below:
                _error("默认行额外留白不能精确映射")
            if sheet.horizontal_page_breaks or sheet.vertical_page_breaks:
                _error("手动分页尚不支持")
            for (row, col), runs in sheet.rich_text_runlist_map.items():
                font_id = book.xf_list[sheet.cell_xf_index(row, col)].font_index
                if any(font != font_id for _, font in runs):
                    _error("单元格内混合字体富文本尚不能完整验证")
                warnings.append(
                    f"{sheet.name}!{get_column_letter(col + 1)}{row + 1}：同字体的冗余富文本标记规范化，文字及字体保持不变"
                )
            if any(
                info.additional_space_above or info.additional_space_below
                for info in sheet.rowinfo_map.values()
            ):
                _error("显式行额外留白不能精确映射")
            used.update(i.xf_index for i in sheet.colinfo_map.values())
            used.update(i.xf_index for i in sheet.rowinfo_map.values() if i.xf_index >= 0)
            for row in sheet:
                for cell in row:
                    if cell.ctype not in {0, 1, 2, 4, 6}:
                        _error("日期、错误值或未知单元格类型暂不支持；不会转成字符串")
                    if cell.ctype:
                        used.add(cell.xf_index)
                    if cell.ctype == 1:
                        text_bytes += len(cell.value.encode("utf-8"))
                        if text_bytes > 8 * 1024 * 1024:
                            _error("展开后的单元格文本超过安全限制")
        if used & extension_ids:
            _error("使用中的单元格或行列格式包含尚不能验证的 XF 扩展")
        if any(opcode in _EXTENSIONS for _, opcode, _ in records):
            warnings.append(
                "未使用命名样式、兼容性/视图/主题/打印扩展未迁入副本；完整内容保留于原件"
            )
        return book, records, used, warnings
    except ReviewedExcelError:
        if book is not None:
            book.release_resources()
        raise
    except (
        xlrd.XLRDError,
        CompDocError,
        AssertionError,  # xlrd/CompDoc use assertions for malformed binary structures.
        ValueError,
        IndexError,
        KeyError,
        struct.error,
        TypeError,
    ) as exc:
        if book is not None:
            book.release_resources()
        raise ReviewedExcelError("旧 XLS 结构损坏或无法安全读取") from exc


def _print_settings(sheet, records, warnings):
    """Map understood scalar settings; custom paper dimensions remain unknown."""
    data = {opcode: payload for _, opcode, payload in records}
    if 0xA1 not in data:
        warnings.append(f"{sheet.title}：缺少完整打印设置，打印版式未验证")
        return
    try:
        paper, scale, start, fitw, fith, flags, dpi, vdpi, header, footer, copies = struct.unpack(
            "<HHhHHHHHddH", data[0xA1]
        )
        if flags & 4:
            _error("原件打印设置被标记为未定义，暂不支持")
        setup = sheet.page_setup
        setup.paperSize, setup.scale, setup.fitToWidth, setup.fitToHeight = paper, scale, fitw, fith
        setup.horizontalDpi, setup.verticalDpi, setup.copies = dpi, vdpi, copies
        setup.orientation = "portrait" if flags & 2 or flags & 64 else "landscape"
        setup.pageOrder = "overThenDown" if flags & 1 else "downThenOver"
        setup.blackAndWhite, setup.draft = bool(flags & 8), bool(flags & 16)
        setup.useFirstPageNumber = bool(flags & 128)
        if flags & 128:
            setup.firstPageNumber = start
        setup.cellComments = ("atEnd" if flags & 512 else "asDisplayed") if flags & 32 else "none"
        setup.errors = ["displayed", "blank", "dash", "NA"][(flags >> 10) & 3]
        for opcode, name in [(0x26, "left"), (0x27, "right"), (0x28, "top"), (0x29, "bottom")]:
            if opcode in data:
                setattr(sheet.page_margins, name, struct.unpack("<d", data[opcode])[0])
        sheet.page_margins.header, sheet.page_margins.footer = header, footer
        fit = bool(struct.unpack("<H", data.get(0x81, b"\0\0"))[0] & 256)
        sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=fit)
        for opcode, name in [
            (0x2A, "headings"),
            (0x2B, "gridLines"),
            (0x83, "horizontalCentered"),
            (0x84, "verticalCentered"),
        ]:
            if opcode in data:
                setattr(sheet.print_options, name, bool(struct.unpack("<H", data[opcode])[0]))
        if paper == 0 or paper >= 118:
            warnings.append(
                f"{sheet.title}：纸张代码 {paper} 的实际尺寸/驱动定义未验证，禁止作为已验收打印模板"
            )
    except (ValueError, struct.error) as exc:
        raise ReviewedExcelError("旧 XLS 打印设置无法明确读取") from exc


def _derive(book, records, used, warnings, source_sha):
    styles = {index: legacy_style(book, index) for index in used}
    workbook = Workbook()
    workbook.remove(workbook.active)
    normal = workbook._named_styles[0]
    assign_style(normal, styles[0])
    workbook._cell_styles[0] = copy(normal._style)
    workbook.properties.description = TEST_RESTRICTION
    workbook.custom_doc_props.append(
        StringProperty(name="ELTUsageRestriction", value=TEST_RESTRICTION)
    )
    workbook.custom_doc_props.append(StringProperty(name="ELTOriginalSHA256", value=source_sha))
    for index, source in enumerate(book.sheets()):
        sheet = workbook.create_sheet(source.name)
        sheet.sheet_format.zeroHeight = bool(source.default_row_hidden)
        sheet.sheet_state = {0: "visible", 1: "hidden", 2: "veryHidden"}[source.visibility]
        if source.default_row_height:
            sheet.sheet_format.defaultRowHeight = source.default_row_height / 20
        if source.standardwidth is not None:
            sheet.sheet_format.defaultColWidth = source.standardwidth / 256
        elif source.defcolwidth is not None:
            sheet.sheet_format.defaultColWidth = source.defcolwidth
        if source.defcolwidth is not None:
            sheet.sheet_format.baseColWidth = source.defcolwidth
        for row, info in source.rowinfo_map.items():
            dimension = sheet.row_dimensions[row + 1]
            dimension.height, dimension.hidden = info.height / 20, bool(info.hidden)
            dimension.outlineLevel, dimension.collapsed = (
                info.outline_level,
                bool(info.outline_group_starts_ends),
            )
            if info.xf_index >= 0:
                assign_style(dimension, styles[info.xf_index])
        for col, info in source.colinfo_map.items():
            dimension = sheet.column_dimensions[get_column_letter(col + 1)]
            dimension.width, dimension.hidden = info.width / 256, bool(info.hidden)
            dimension.outlineLevel, dimension.collapsed = info.outline_level, bool(info.collapsed)
            assign_style(dimension, styles[info.xf_index])
        for row in range(source.nrows):
            for col in range(source.ncols):
                cell = source.cell(row, col)
                if not cell.ctype:
                    continue
                target = sheet.cell(row + 1, col + 1)
                if cell.ctype == 1:
                    target.value, target.data_type = cell.value, "s"
                elif cell.ctype == 2:
                    target.value = cell.value
                elif cell.ctype == 4:
                    target.value = bool(cell.value)
                assign_style(target, styles[cell.xf_index])
                target.quotePrefix = bool(book.xf_list[cell.xf_index].lotus_123_prefix)
        # Preserve blank-cell style records inside merges instead of letting
        # openpyxl normalize/delete them while constructing the workbook.
        for first, last, left, right in source.merged_cells:
            sheet.merged_cells.add(
                f"{get_column_letter(left + 1)}{first + 1}:{get_column_letter(right)}{last}"
            )
        _print_settings(sheet, [r for r in records if r[0] == index], warnings)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue(), styles


def _validate_layout(source, sheet, styles, records):
    """Read back serialized dimensions/styles and independently decoded BIFF scalars."""
    if sheet.sheet_state != {0: "visible", 1: "hidden", 2: "veryHidden"}[source.visibility]:
        _error("工作表可见性改变")
    if sheet.sheet_format.zeroHeight != bool(source.default_row_hidden):
        _error("默认行隐藏状态改变")
    expected_width = (
        source.standardwidth / 256 if source.standardwidth is not None else source.defcolwidth
    )
    if (
        source.default_row_height
        and sheet.sheet_format.defaultRowHeight != source.default_row_height / 20
    ):
        _error("默认行高改变")
    if expected_width is not None and sheet.sheet_format.defaultColWidth != expected_width:
        _error("默认列宽改变")
    if source.defcolwidth is not None and sheet.sheet_format.baseColWidth != source.defcolwidth:
        _error("基础列宽改变")
    for is_row, entries in ((True, source.rowinfo_map), (False, source.colinfo_map)):
        for index, info in entries.items():
            dim = (
                sheet.row_dimensions[index + 1]
                if is_row
                else sheet.column_dimensions[get_column_letter(index + 1)]
            )
            size = dim.height if is_row else dim.width
            expected_size = info.height / 20 if is_row else info.width / 256
            collapsed = info.outline_group_starts_ends if is_row else info.collapsed
            if (size, dim.hidden, dim.outlineLevel, dim.collapsed) != (
                expected_size,
                bool(info.hidden),
                info.outline_level,
                bool(collapsed),
            ):
                _error("行列尺寸、隐藏或分组状态改变")
            if info.xf_index >= 0:
                actual = (
                    copy(dim.font),
                    copy(dim.fill),
                    copy(dim.border),
                    copy(dim.alignment),
                    dim.number_format,
                    copy(dim.protection),
                )
                if actual != styles[info.xf_index]:
                    _error("行列默认样式改变")
    data = {opcode: payload for _, opcode, payload in records}
    if 0xA1 not in data:
        return
    values = struct.unpack("<HHhHHHHHddH", data[0xA1])
    paper, scale, start, fitw, fith, flags, dpi, vdpi, header, footer, copies = values
    setup = sheet.page_setup
    expected = {
        "paperSize": paper,
        "scale": scale,
        "fitToWidth": fitw,
        "fitToHeight": fith,
        "horizontalDpi": dpi,
        "verticalDpi": vdpi,
        "copies": copies,
        "orientation": "portrait" if flags & 66 else "landscape",
        "pageOrder": "overThenDown" if flags & 1 else "downThenOver",
        "blackAndWhite": bool(flags & 8),
        "draft": bool(flags & 16),
        "useFirstPageNumber": bool(flags & 128),
        "cellComments": ("atEnd" if flags & 512 else "asDisplayed") if flags & 32 else None,
        "errors": ["displayed", "blank", "dash", "NA"][(flags >> 10) & 3],
    }
    if flags & 128:
        expected["firstPageNumber"] = start
    if any(getattr(setup, name) != value for name, value in expected.items()):
        _error("打印标量设置改变")
    margins = {"header": header, "footer": footer}
    for opcode, name in ((0x26, "left"), (0x27, "right"), (0x28, "top"), (0x29, "bottom")):
        if opcode in data:
            margins[name] = struct.unpack("<d", data[opcode])[0]
    if any(
        abs(getattr(sheet.page_margins, name) - value) > 1e-12 for name, value in margins.items()
    ):
        _error("打印边距改变")
    fit = bool(struct.unpack("<H", data.get(0x81, b"\0\0"))[0] & 256)
    if sheet.sheet_properties.pageSetUpPr.fitToPage != fit:
        _error("打印适配模式改变")
    for opcode, name in (
        (0x2A, "headings"),
        (0x2B, "gridLines"),
        (0x83, "horizontalCentered"),
        (0x84, "verticalCentered"),
    ):
        if opcode in data and getattr(sheet.print_options, name) != bool(
            struct.unpack("<H", data[opcode])[0]
        ):
            _error("打印选项改变")


def validate_conversion(book, converted, styles, records=()):
    review = inspect_workbook(converted)
    target = load_workbook(BytesIO(converted), data_only=False)
    cells, styled, normalized = 0, 0, 0
    try:
        if book.sheet_names() != target.sheetnames:
            _error("工作表名称或顺序改变")
        with ZipFile(BytesIO(converted)) as package:
            for index, (source, sheet) in enumerate(zip(book.sheets(), target)):
                xml = ElementTree.fromstring(package.read(f"xl/worksheets/sheet{index + 1}.xml"))
                raw = {c.attrib["r"]: c for c in xml.findall(".//m:sheetData/m:row/m:c", _NS)}
                expected = {
                    (r + 1, c + 1): source.cell(r, c)
                    for r in range(source.nrows)
                    for c in range(source.ncols)
                    if source.cell_type(r, c) not in {0, 6}
                }
                actual = {
                    (c.row, c.column): c.value for row in sheet for c in row if c.value is not None
                }
                if set(expected) != set(actual):
                    _error("单元格位置或数量改变")
                for address, cell in expected.items():
                    value = actual[address]
                    valid_type = (
                        type(value) is str
                        if cell.ctype == 1
                        else type(value) in {int, float}
                        if cell.ctype == 2
                        else type(value) is bool
                    )
                    if not valid_type or value != cell.value:
                        _error("单元格值或类型改变")
                cells += len(expected)
                merges = {
                    f"{get_column_letter(c + 1)}{r + 1}:{get_column_letter(d)}{b}"
                    for r, b, c, d in source.merged_cells
                }
                if merges != {str(r) for r in sheet.merged_cells.ranges}:
                    _error("合并区域改变")
                for row in range(source.nrows):
                    for col in range(source.ncols):
                        cell = source.cell(row, col)
                        if not cell.ctype:
                            continue
                        key = f"{get_column_letter(col + 1)}{row + 1}"
                        fmt = target._cell_styles[int(raw[key].attrib.get("s", "0"))]
                        number = (
                            target._number_formats[fmt.numFmtId - 164]
                            if fmt.numFmtId >= 164
                            else BUILTIN_FORMATS[fmt.numFmtId]
                        )
                        mapped = (
                            target._fonts[fmt.fontId],
                            target._fills[fmt.fillId],
                            target._borders[fmt.borderId],
                            target._alignments[fmt.alignmentId],
                            number,
                            target._protections[fmt.protectionId],
                        )
                        if mapped != styles[cell.xf_index] or bool(fmt.quotePrefix) != bool(
                            book.xf_list[cell.xf_index].lotus_123_prefix
                        ):
                            _error("单元格样式记录改变")
                        normalized += sheet.cell(row + 1, col + 1)._style != fmt
                        styled += 1
                _validate_layout(source, sheet, styles, [r for r in records if r[0] == index])
        return review, {
            "sheets": book.nsheets,
            "nonempty_cells": cells,
            "xml_cell_styles": styled,
            "merged_cells_normalized_on_readback": normalized,
            "checks": [
                "sheet_order",
                "values_types_positions",
                "merge_ranges",
                "raw_cell_style_records",
                "default_and_explicit_dimensions_styles_visibility",
                "supported_print_scalars_not_print_fidelity",
            ],
            "print_layout_verified": False,
        }
    finally:
        target.close()


def _convert_legacy_body(body: bytes, *, purpose: str):
    """Shared fail-closed XLS→XLSX convert used by archive consent and slice entry."""
    digest = hashlib.sha256(body).hexdigest()
    book, records, used, warnings = inspect_legacy(body)
    try:
        converted, styles = _derive(book, records, used, warnings, digest)
        if len(converted) > MAX_BYTES:
            _error("转换副本超过 20 MB")
        review, validation = validate_conversion(book, converted, styles, records)
        if validation["merged_cells_normalized_on_readback"]:
            warnings.append(
                f"{validation['merged_cells_normalized_on_readback']} 个合并内部样式在读回时会规范化；XML记录已保留，渲染仍须核对"
            )
        review["usage_restrictions"] = [TEST_RESTRICTION]
        review["conversion_receipt"] = {
            "source_sha256": digest,
            "derived_sha256": hashlib.sha256(converted).hexdigest(),
            "converter": "bounded xlrd/openpyxl data-test converter v1",
            "purpose": purpose,
            "print_layout_verified": False,
            "warnings": warnings,
            "validation": validation,
        }
        return converted, review
    except ReviewedExcelError:
        raise
    except (ValueError, KeyError, IndexError, TypeError, ElementTree.ParseError) as exc:
        raise ReviewedExcelError("旧 XLS 转换校验失败；不会生成正式模板") from exc
    finally:
        book.release_resources()


def convert_xls_to_xlsx_for_slice(body: bytes):
    """Convert OLE/BIFF8 XLS for classify→group→slice without archive HTTP consent.

    Reuses the same xlrd/openpyxl converter and fail-closed admission checks as
    prepare_archive_workbook. Returns (xlsx_bytes, conversion_receipt). Archive
    HTTP still requires legacy_test_only + matching SHA via prepare_archive_workbook.
    """
    if not body.startswith(OLE_MAGIC):
        _error("文件不是受支持的 OLE/BIFF8 XLS 或超过 20 MB")
    converted, review = _convert_legacy_body(body, purpose="slice_pipeline_auto_convert")
    return converted, review["conversion_receipt"]


def ensure_slice_workbook(body: bytes):
    """Return XLSX bytes for the slicing pipeline; auto-convert OLE XLS first.

    Non-OLE bodies pass through unchanged with receipt=None. OLE bodies are
    converted under the slice purpose and return the conversion receipt.
    """
    if not body.startswith(OLE_MAGIC):
        return body, None
    return convert_xls_to_xlsx_for_slice(body)


def prepare_archive_workbook(body: bytes, *, legacy_test_only=False, legacy_source_sha256=""):
    if not body.startswith(OLE_MAGIC):
        return body, inspect_workbook(body)
    digest = hashlib.sha256(body).hexdigest()
    if legacy_test_only is not True or legacy_source_sha256 != digest:
        _error("请明确确认仅数据/模板测试并绑定原件 SHA-256；打印未验证，不可正式发布")
    return _convert_legacy_body(body, purpose="data_template_test_only")
