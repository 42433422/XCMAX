"""Source workbook partition geometry and reviewed classification corrections."""

from __future__ import annotations

import hashlib
import json
from copy import deepcopy

from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

from app.application.elt.errors import ReviewedExcelError
from app.application.reviewed_excel_evidence import bounded_range

DOCUMENT_LABELS = frozenset(
    {
        "quotation",
        "delivery",
        "purchase",
        "attendance",
        "production",
        "process",
        "material",
        "inspection",
        "template",
        "unknown",
    }
)


def apply_classification_corrections(review: dict, latest: dict, labels: list) -> dict:
    """Validate user corrections against source geometry, preserving shared context conservatively."""
    from copy import deepcopy

    expected = {
        (sheet["sheet"], region["range"]): region
        for sheet in latest["sheets"]
        for region in sheet["regions"]
    }
    supplied: dict[tuple[str, str], list[dict]] = {}
    from openpyxl.utils.cell import (
        column_index_from_string,
        coordinate_from_string,
    )

    dimensions = {sheet["sheet"]: sheet for sheet in review["sheet_overview"]}
    rectangles: dict[str, list[tuple[int, int, int, int]]] = {}
    limit = max(200, len(expected))
    if not isinstance(labels, list) or len(labels) > limit:
        raise ReviewedExcelError(f"分类区域数量无效，最多{limit}个片段")
    for item in labels:
        if (
            not isinstance(item, dict)
            or set(item)
            not in ({"sheet", "range", "labels"}, {"sheet", "original_range", "range", "labels"})
            or not isinstance(item["sheet"], str)
            or not isinstance(item["range"], str)
            or ("original_range" in item and not isinstance(item["original_range"], str))
            or not isinstance(item["labels"], list)
            or not item["labels"]
            or any(
                not isinstance(label, str) or label not in DOCUMENT_LABELS
                for label in item["labels"]
            )
            or len(set(item["labels"])) != len(item["labels"])
            or ("unknown" in item["labels"] and len(item["labels"]) > 1)
        ):
            raise ReviewedExcelError("分类修正无效")
        key = (item["sheet"], item.get("original_range", item["range"]))
        if key not in expected:
            raise ReviewedExcelError("分类区域不属于该版本")
        try:
            c1, r1, c2, r2 = bounded_range(item["range"])
            size = dimensions[item["sheet"]]
            if not all(isinstance(x, int) for x in (c1, r1, c2, r2)) or not (
                1 <= c1 <= c2 <= size["columns"] and 1 <= r1 <= r2 <= size["rows"]
            ):
                raise ReviewedExcelError("bounds")
        except (ValueError, TypeError, KeyError):
            raise ReviewedExcelError("分类区域越界") from None
        for oc1, or1, oc2, or2 in rectangles.setdefault(item["sheet"], []):
            if c1 <= oc2 and oc1 <= c2 and r1 <= or2 and or1 <= r2:
                raise ReviewedExcelError("分类区域重叠")
        rectangles[item["sheet"]].append((c1, r1, c2, r2))
        for merged in size.get("merged_regions", []):
            mc1, mr1, mc2, mr2 = bounded_range(merged)
            if (
                c1 <= mc2
                and mc1 <= c2
                and r1 <= mr2
                and mr1 <= r2
                and not (c1 <= mc1 <= mc2 <= c2 and r1 <= mr1 <= mr2 <= r2)
            ):
                raise ReviewedExcelError("分类区域切断合并单元格")
        supplied.setdefault(key, []).append(item)
    if set(supplied) != set(expected):
        raise ReviewedExcelError("分类区域不完整")
    partition_changed = False
    corrected: dict[str, list[dict]] = {sheet["sheet"]: [] for sheet in latest["sheets"]}
    for key, original in expected.items():
        children = supplied[key]
        if len(children) > 1:
            partition_changed = True
            c1, r1, c2, r2 = bounded_range(original["range"])
            area = 0
            for child in children:
                sc1, sr1, sc2, sr2 = bounded_range(child["range"])
                if not (c1 <= sc1 <= sc2 <= c2 and r1 <= sr1 <= sr2 <= r2):
                    raise ReviewedExcelError("拆分片段必须位于原区域内")
                area += (sc2 - sc1 + 1) * (sr2 - sr1 + 1)
            # All final rectangles were checked for overlap above; equal area proves no holes.
            if area != (c2 - c1 + 1) * (r2 - r1 + 1):
                raise ReviewedExcelError("拆分必须完整覆盖原区域，不能遗漏行列")
        for item in children:
            region = deepcopy(original)
            if region["labels"] != item["labels"]:
                region["label_source"] = "user"
            region["labels"] = item["labels"]
            c1, r1, c2, r2 = bounded_range(item["range"])
            region["range"] = f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}"
            for field in ("original_range", "modified"):
                region.pop(field, None)
            if original["range"] != region["range"]:
                region["range_source"] = "user"

                def in_region(ref):
                    col, row = coordinate_from_string(ref["cell"])
                    return (
                        ref.get("sheet", key[0]) == key[0]
                        and c1 <= column_index_from_string(col) <= c2
                        and r1 <= row <= r2
                    )

                region["evidence"] = [ref for ref in original["evidence"] if in_region(ref)]
                region["header_cells"] = [ref for ref in original["header_cells"] if in_region(ref)]
                region["checked_cells"] = sum(
                    in_region(cell) for cell in review["cells"] if cell["sheet"] == key[0]
                )
                region["merged_ranges"] = [
                    merged
                    for merged in dimensions[key[0]].get("merged_regions", [])
                    for mc1, mr1, mc2, mr2 in [bounded_range(merged)]
                    if c1 <= mc1 <= mc2 <= c2 and r1 <= mr1 <= mr2 <= r2
                ]
                region.pop("continuation_context", None)
                region.pop("value_extent", None)
                region["parent_heading"] = ""
                region["dependency_status"] = "用户修改边界；共享说明和依赖需重新确认"
                region["needs_confirmation"] = True
            corrected[key[0]].append(region)
    for sheet in latest["sheets"]:
        sheet["regions"] = corrected[sheet["sheet"]]
        sheet["labels"] = list(
            dict.fromkeys(label for region in sheet["regions"] for label in region["labels"])
        )
    latest["partition_pending"] = partition_changed
    for coverage in latest.get("coverage", []):
        name = coverage["sheet"]
        coverage.setdefault("model_unassigned", list(coverage["unassigned"]))
        coverage["unassigned"] = []
        for cell in (cell for cell in review["cells"] if cell["sheet"] == name):
            col, row = coordinate_from_string(cell["cell"])
            number = column_index_from_string(col)
            if not any(
                c1 <= number <= c2 and r1 <= row <= r2
                for c1, r1, c2, r2 in rectangles.get(name, [])
            ):
                coverage["unassigned"].append(cell["cell"])
    return latest


def _normalize_atom_ids(data: dict, references: dict, review: dict) -> dict:
    from openpyxl.utils.cell import get_column_letter

    expected = {key for key, entry in references.items() if entry["primary"]}
    source = review.get("full_cells", review["cells"])
    used: set[str] = set()
    decoded: list[dict] = []
    by_primary: dict[str, dict] = {}
    for region in data["physical_regions"]:
        identifiers = region["atom_ids"]
        if any(key not in references for key in identifiers):
            raise ReviewedExcelError("源原子身份不在已发送事实中")
        if any(key not in expected and not references[key].get("blank") for key in identifiers):
            raise ReviewedExcelError("布局原子不是空白源合并区域，不能替代主体")
        if len(identifiers) != len(set(identifiers)) or used.intersection(identifiers):
            raise ReviewedExcelError("源原子重复分配，未修改模型答案")
        primary = region["primary_id"]
        if primary not in identifiers or primary not in expected:
            raise ReviewedExcelError("原子矩形必须选择自身主体原子")
        sheets = {references[key]["ref"]["sheet"] for key in identifiers}
        if len(sheets) != 1:
            raise ReviewedExcelError("源原子矩形不能跨工作表")
        sheet = next(iter(sheets))
        boxes = [bounded_range(references[key]["atom_range"]) for key in identifiers]
        c1, r1 = min(b[0] for b in boxes), min(b[1] for b in boxes)
        c2, r2 = max(b[2] for b in boxes), max(b[3] for b in boxes)
        selected = {
            (references[key]["ref"]["sheet"], references[key]["ref"]["cell"])
            for key in identifiers
            if key in expected
        }
        for cell in source:
            col, row = coordinate_from_string(cell["cell"])
            if (
                cell["sheet"] == sheet
                and c1 <= column_index_from_string(col) <= c2
                and r1 <= row <= r2
            ):
                if (sheet, cell["cell"]) not in selected:
                    raise ReviewedExcelError("原子矩形含未选主体、上下文或未发送内容，不能自动补齐")
        for overview in review["sheet_overview"]:
            if overview["sheet"] != sheet:
                continue
            for merge in overview.get("merged_regions", []):
                mc1, mr1, mc2, mr2 = bounded_range(merge)
                if c1 <= mc2 and mc1 <= c2 and r1 <= mr2 and mr1 <= r2:
                    if not (c1 <= mc1 and mc2 <= c2 and r1 <= mr1 and mr2 <= r2):
                        raise ReviewedExcelError("原子矩形切断合并单元格，不能自动扩框")
        dependencies = region["dependency_ids"]
        if len(dependencies) != len(set(dependencies)):
            raise ReviewedExcelError("模型依赖原子重复")
        area = f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}"
        decoded.append(
            {
                "sheet": sheet,
                "range": area,
                "labels": ["unknown"],
                "evidence": [deepcopy(references[primary]["ref"])],
                "dependencies": [deepcopy(references[key]["ref"]) for key in dependencies],
                "reason": region["reason"],
            }
        )
        by_primary[primary] = {"sheet": sheet, "range": area}
        used.update(identifiers)
    if used.intersection(expected) != expected:
        raise ReviewedExcelError("源主体原子未完整分配，不能自动补齐")
    documents = []
    for document in data["logical_documents"]:
        if any(key not in by_primary for key in document["fragments"]):
            raise ReviewedExcelError("文档引用未知主体片段编号")
        documents.append(
            {
                "document_id": document["document_id"],
                "fragments": [deepcopy(by_primary[key]) for key in document["fragments"]],
            }
        )
    return {
        "physical_regions": decoded,
        "logical_documents": documents,
        "layout_references": [
            deepcopy(references[key]["ref"]) for key in data["layout_references"]
        ],
    }


def source_ranges_schema(sheets: list[dict]) -> dict:
    properties: dict[str, dict] = {
        "sheet_name": {"type": "string", "enum": [sheet["sheet"] for sheet in sheets]}
    }
    for axis, dimension in (("row", "rows"), ("column", "columns")):
        for edge in ("start", "end"):
            properties[f"{axis}_{edge}"] = {
                "type": "integer",
                "minimum": 1,
                "maximum": max(sheet[dimension] for sheet in sheets),
            }
    area = {
        "type": "object",
        "required": list(properties),
        "additionalProperties": False,
        "properties": properties,
    }
    document = {
        "type": "object",
        "required": ["ranges", "reason"],
        "additionalProperties": False,
        "properties": {
            "ranges": {"type": "array", "minItems": 1, "maxItems": 20, "items": area},
            "reason": {"type": "string", "minLength": 1},
        },
    }
    return {
        "type": "object",
        "required": ["documents"],
        "additionalProperties": False,
        "properties": {
            "documents": {"type": "array", "minItems": 1, "maxItems": 100, "items": document}
        },
    }


def decode_source_range(region: dict, dimensions: dict) -> dict:
    from openpyxl.utils.cell import get_column_letter

    sheet = region["sheet_name"]
    r1, r2, c1, c2 = (region[key] for key in ("row_start", "row_end", "column_start", "column_end"))
    if (
        sheet not in dimensions
        or any(type(value) is not int for value in (r1, r2, c1, c2))
        or not (
            1 <= r1 <= r2 <= dimensions[sheet]["rows"]
            and 1 <= c1 <= c2 <= dimensions[sheet]["columns"]
        )
    ):
        raise ReviewedExcelError("模型边界非整数、倒序或超出原件范围，未修改答案")
    return {"sheet": sheet, "range": f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}"}


def internal_document_id(document: dict) -> str:
    """Storage identity only; never requested from the model or written to workbook data."""
    return (
        "source-"
        + hashlib.sha256(
            json.dumps(document["ranges"], sort_keys=True, ensure_ascii=False).encode()
        ).hexdigest()[:24]
    )


def _normalize_boundaries(data: dict, references: dict, review: dict) -> dict:
    if sum(len(document["ranges"]) for document in data["documents"]) > 100:
        raise ReviewedExcelError("模型范围总数超过100，未处理或修改答案")
    dimensions = {sheet["sheet"]: sheet for sheet in review["sheet_overview"]}
    primary = {
        (entry["ref"]["sheet"], entry["ref"]["cell"]): key
        for key, entry in references.items()
        if entry["primary"]
    }
    source = review.get("full_cells", review["cells"])
    seen: set[tuple[str, str]] = set()
    decoded, documents, proof = [], [], []
    for document in data["documents"]:
        parts = []
        for region in document["ranges"]:
            part = decode_source_range(region, dimensions)
            c1, r1, c2, r2 = bounded_range(part["range"])
            covered = []
            for cell in source:
                col, row = coordinate_from_string(cell["cell"])
                if (
                    cell["sheet"] == part["sheet"]
                    and c1 <= column_index_from_string(col) <= c2
                    and r1 <= row <= r2
                ):
                    key = (cell["sheet"], cell["cell"])
                    if key not in primary:
                        raise ReviewedExcelError("模型边界含上下文或未发送主体，不能裁剪或补齐")
                    covered.append(key)
            if not covered or seen.intersection(covered):
                raise ReviewedExcelError("模型边界主体为空或重复覆盖")
            decoded.append(
                {
                    **part,
                    "labels": ["unknown"],
                    "evidence": [{"sheet": covered[0][0], "cell": covered[0][1]}],
                    "dependencies": [],
                    "reason": document["reason"],
                }
            )
            proof.append(
                {**part, "source_cells": [cell for _, cell in covered], "count": len(covered)}
            )
            seen.update(covered)
            parts.append(part)
        documents.append({"document_id": internal_document_id(document), "fragments": parts})
    if seen != set(primary):
        raise ReviewedExcelError("模型边界未完整覆盖本批主体，不能自动补齐")
    references["_source_coverage"] = {
        "basis": "verified_original_source_rectangle",
        "document_id_origin": "internal_source_range_hash",
        "regions": proof,
    }
    return {"physical_regions": decoded, "logical_documents": documents, "layout_references": []}
