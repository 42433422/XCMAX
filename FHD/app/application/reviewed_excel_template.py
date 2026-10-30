"""Conservative, user-reviewed cleaning of a copy of an Excel workbook."""

from __future__ import annotations

import hashlib
import warnings
from io import BytesIO
from typing import Any
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.print_settings import PrintArea
from openpyxl.worksheet.worksheet import Worksheet

from app.application.etl.errors import ReviewedExcelError
from app.application.reviewed_excel_evidence import bounded_range
from app.utils.operational_errors import BOUNDARY_ERRORS

MAX_BYTES = 20 * 1024 * 1024
MAX_UNPACKED = 100 * 1024 * 1024
MAX_CELLS = 20_000


from app.application.reviewed_excel_classification import (
    classify_review_workbook as classify_review_workbook,
)
from app.application.reviewed_excel_classification import (
    classify_semantic_workbook as classify_semantic_workbook,
)


def _validate_template_schema(schema: dict, payload: object) -> bool:
    """Validate the complete schema subset used by template requests; fail closed otherwise."""
    permitted = {
        "type",
        "required",
        "properties",
        "additionalProperties",
        "items",
        "minItems",
        "maxItems",
        "minLength",
        "maxLength",
        "enum",
    }

    def supported(rule, depth=0):
        if not isinstance(rule, dict) or depth > 16 or set(rule) - permitted:
            return False
        if rule.get("type") not in {"object", "array", "string"}:
            return False
        props = rule.get("properties", {})
        if not isinstance(props, dict) or any(
            not isinstance(key, str) or not supported(value, depth + 1)
            for key, value in props.items()
        ):
            return False
        if "items" in rule and not supported(rule["items"], depth + 1):
            return False
        if "required" in rule and (
            not isinstance(rule["required"], list)
            or any(not isinstance(item, str) for item in rule["required"])
        ):
            return False
        if "additionalProperties" in rule and not isinstance(rule["additionalProperties"], bool):
            return False
        if "enum" in rule and not isinstance(rule["enum"], list):
            return False
        return all(
            key not in rule or (type(rule[key]) is int and rule[key] >= 0)
            for key in ("minItems", "maxItems", "minLength", "maxLength")
        )

    def matches(rule, value):
        kind = rule["type"]
        if kind == "object":
            if not isinstance(value, dict):
                return False
            properties = rule.get("properties", {})
            if any(key not in value for key in rule.get("required", [])):
                return False
            if rule.get("additionalProperties") is False and any(
                key not in properties for key in value
            ):
                return False
            if any(
                key in properties and not matches(properties[key], item)
                for key, item in value.items()
            ):
                return False
        elif kind == "array":
            if (
                not isinstance(value, list)
                or len(value) < rule.get("minItems", 0)
                or len(value) > rule.get("maxItems", len(value))
            ):
                return False
            if "items" in rule and any(not matches(rule["items"], item) for item in value):
                return False
        elif (
            not isinstance(value, str)
            or len(value) < rule.get("minLength", 0)
            or len(value) > rule.get("maxLength", len(value))
        ):
            return False
        return "enum" not in rule or value in rule["enum"]

    return bool(supported(schema) and matches(schema, payload))


def _template_mimo_complete(messages: list[dict], schema: dict, max_tokens: int):
    """Explicit template-only MiMo route; one bounded request, without telemetry payloads."""
    import asyncio
    import logging
    import os

    from app.application.etl import llm_assist
    from app.application.reviewed_excel_audit import ModelRequestNotSent
    from app.domain.shared_services.conversation.llm_adapter import OpenAICompatibleAdapter
    from app.infrastructure.llm.providers.openai_compatible_provider import XiaomiMimoProvider
    from app.infrastructure.llm.structured_output import extract_json
    from app.utils.operational_errors import BOUNDARY_ERRORS

    base = os.environ.get("FHD_TEMPLATE_MIMO_BASE_URL", "").strip().rstrip("/")
    model = os.environ.get("FHD_TEMPLATE_MIMO_MODEL", "").strip()
    if (
        llm_assist.etl_llm_mode() == "off"
        or base != "https://token-plan-cn.xiaomimimo.com/v1"
        or model not in {"mimo-v2.6-flash", "mimo-v2.6-pro"}
    ):
        raise ModelRequestNotSent("模板 MiMo 配置未完成或不受支持；未调用模型")
    adapter = OpenAICompatibleAdapter(provider="xiaomi", model=model, base_url=base)
    if not adapter.is_configured:
        raise ModelRequestNotSent("模板 MiMo 凭据未配置；未调用模型")
    provider = XiaomiMimoProvider(adapter)
    from app.application.reviewed_excel_audit import response_capture

    audit_callback = response_capture()

    async def call():
        try:
            return await asyncio.wait_for(
                provider.chat_completion(
                    messages,
                    temperature=0.0,
                    max_tokens=max_tokens,
                    thinking={"type": "disabled"},
                    **(
                        {"_reviewed_excel_audit_response": audit_callback} if audit_callback else {}
                    ),
                ),
                timeout=llm_assist.etl_llm_timeout_seconds(),
            )
        finally:
            await adapter.close()

    try:
        response = asyncio.run(call())
        choice = (response or {}).get("choices", [])[0]
        if choice.get("finish_reason") != "stop":
            raise ReviewedExcelError("incomplete")
        data = extract_json(str((choice.get("message") or {}).get("content") or ""))
        if data is None or not _validate_template_schema(schema, data):
            raise ReviewedExcelError("invalid")
    except BOUNDARY_ERRORS as exc:
        # Never log the provider exception, request, response or credential.
        logging.getLogger(__name__).info("template MiMo call failed: %s", type(exc).__name__)
        raise ReviewedExcelError("模型调用失败或结果不完整，可能已产生费用；未自动重试") from None
    return llm_assist.LlmAssistResult(used_llm=True, model=model, data=data, provider="mimo")


def _read_only_print_area_warnings(body: bytes, found: list) -> list[dict]:
    """Allow only verified static print-area warnings for read-only inspection."""
    import re

    from defusedxml import ElementTree

    prefix = "Print area cannot be set to Defined name: "
    messages = [str(item.message) for item in found]
    if any(
        item.category is not UserWarning
        or not message.startswith(prefix)
        or not message.endswith(".")
        for item, message in zip(found, messages, strict=True)
    ):
        raise ReviewedExcelError("工作簿含无法保证保真的扩展，暂不支持清理")
    ns = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with ZipFile(BytesIO(body)) as archive:
        root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
    sheets = [sheet.get("name") for sheet in root.findall("x:sheets/x:sheet", ns)]
    names = root.findall("x:definedNames/x:definedName", ns)
    entries = []
    pattern = re.compile(
        r"(?:'((?:[^']|'')+)'|([^'!,]+))!(\$?[A-Za-z]{1,3}\$?\d+:\$?[A-Za-z]{1,3}\$?\d+|\$?\d+:\$?\d+|\$?[A-Za-z]{1,3}:\$?[A-Za-z]{1,3})"
    )
    for message in messages:
        value = message[len(prefix) : -1]
        matching = [
            node for node in names if node.get("name") == "_xlnm.Print_Area" and node.text == value
        ]
        if len(matching) != 1:
            raise ReviewedExcelError("打印区域定义无法核验，暂不支持")
        raw_index = matching[0].get("localSheetId", "")
        if not raw_index.isdigit() or int(raw_index) >= len(sheets):
            raise ReviewedExcelError("打印区域来源工作表无效")
        sheet = sheets[int(raw_index)]
        offset = 0
        for match in pattern.finditer(value):
            if (
                match.start() != offset
                or (match.group(1) or match.group(2)).replace("''", "'") != sheet
            ):
                raise ReviewedExcelError("打印区域定义无法核验，暂不支持")
            parts = match.group(3).replace("$", "").split(":")
            bounds = []
            for part in parts:
                cell = re.fullmatch(r"([A-Za-z]+)?(\d+)?", part)
                if (
                    not cell
                    or (cell.group(1) and column_index_from_string(cell.group(1).upper()) > 16384)
                    or (cell.group(2) and not 1 <= int(cell.group(2)) <= 1048576)
                ):
                    raise ReviewedExcelError("打印区域坐标无效")
                bounds.append(
                    (
                        column_index_from_string(cell.group(1).upper()) if cell.group(1) else None,
                        int(cell.group(2)) if cell.group(2) else None,
                    )
                )
            if any(
                a is not None and b is not None and a > b
                for a, b in zip(bounds[0], bounds[1], strict=True)
            ):
                raise ReviewedExcelError("打印区域坐标无效")
            offset = match.end()
            if offset < len(value) and value[offset] == ",":
                offset += 1
            elif offset < len(value):
                raise ReviewedExcelError("打印区域定义无法核验，暂不支持")
        if not value or offset != len(value) or value.endswith(","):
            raise ReviewedExcelError("打印区域定义无法核验，暂不支持")
        entries.append({"sheet": sheet, "value": value})
    return entries


def inspect_workbook(body: bytes) -> dict:
    if not body or len(body) > MAX_BYTES:
        raise ReviewedExcelError("文件为空或超过 20 MB")
    try:
        with ZipFile(BytesIO(body)) as archive:
            entries = archive.infolist()
            if len(entries) > 5000 or sum(item.file_size for item in entries) > MAX_UNPACKED:
                raise ReviewedExcelError("工作簿解压后过大")
            if any(
                item.filename.startswith(("xl/embeddings/", "xl/vbaProject")) for item in entries
            ):
                raise ReviewedExcelError("含嵌入附件或宏，未读取或执行")
            unsupported = (
                "xl/drawings/",
                "xl/charts/",
                "xl/media/",
                "xl/embeddings/",
                "xl/activeX/",
                "xl/vbaProject",
                "xl/pivot",
                "xl/externalLinks/",
                "xl/slicers/",
                "xl/threadedComments/",
                "customXml/",
            )
            if any(item.filename.startswith(unsupported) for item in entries):
                raise ReviewedExcelError("工作簿含暂不支持保真的对象或扩展")
        with warnings.catch_warnings(record=True) as found:
            warnings.simplefilter("always")
            wb = load_workbook(BytesIO(body), data_only=False, keep_links=True, rich_text=True)
        read_only_print_areas = []
        if found:
            try:
                read_only_print_areas = _read_only_print_area_warnings(body, found)
            except BOUNDARY_ERRORS:
                wb.close()
                raise
    except (BadZipFile, KeyError, OSError, ParseError, InvalidFileException) as exc:
        raise ReviewedExcelError("无法读取 XLSX 工作簿") from exc
    cells: list[dict[str, Any]] = []
    layout: list[dict[str, Any]] = []
    try:
        if len(wb.worksheets) > 100:
            raise ReviewedExcelError("工作表超过 100 张，暂不能安全预览")
        grid_budget = 0
        for sheet in wb.worksheets:
            if len(sheet.merged_cells.ranges) > 5000:
                raise ReviewedExcelError("合并区域超过 5000 个，暂不能安全预览")
            if sheet.max_row * sheet.max_column > 100_000:
                raise ReviewedExcelError("工作表范围过大，暂不能安全预览")
            grid_budget += sheet.max_row * sheet.max_column
            if grid_budget > 250_000:
                raise ReviewedExcelError("工作簿总单元格范围过大，暂不能安全预览")
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell, MergedCell) or cell.value is None:
                        continue
                    if len(cells) >= MAX_CELLS:
                        raise ReviewedExcelError("非空单元格超过 20000 个，暂不能完整预览")
                    formula = cell.data_type == "f"
                    cells.append(
                        {
                            "sheet": sheet.title,
                            "cell": cell.coordinate,
                            "value": str(cell.value),
                            "kind": "formula" if formula else "review",
                            "reason": "公式保留" if formula else "未确定，默认保留；请逐格确认",
                        }
                    )
            if len(layout) < 20:
                max_row, max_col = min(sheet.max_row, 40), min(sheet.max_column, 20)
                merged = {}
                covered: set[tuple[int, int]] = set()
                for region in sheet.merged_cells.ranges:
                    if region.min_row > max_row or region.min_col > max_col:
                        continue
                    end_row, end_col = min(region.max_row, max_row), min(region.max_col, max_col)
                    merged[(region.min_row, region.min_col)] = (
                        end_row - region.min_row + 1,
                        end_col - region.min_col + 1,
                    )
                    covered.update(
                        (r, c)
                        for r in range(region.min_row, end_row + 1)
                        for c in range(region.min_col, end_col + 1)
                        if (r, c) != (region.min_row, region.min_col)
                    )
                rows = []
                for r in range(1, max_row + 1):
                    row_cells = []
                    for c in range(1, max_col + 1):
                        if (r, c) in covered:
                            continue
                        cell = sheet.cell(r, c)
                        span = merged.get((r, c), (1, 1))
                        rgb = (
                            cell.fill.fgColor.rgb
                            if cell.fill.patternType == "solid" and cell.fill.fgColor.type == "rgb"
                            else None
                        )
                        row_cells.append(
                            {
                                "cell": cell.coordinate,
                                "value": "" if cell.value is None else str(cell.value),
                                "formula": cell.data_type == "f",
                                "rowspan": span[0],
                                "colspan": span[1],
                                "bold": bool(cell.font.bold),
                                "italic": bool(cell.font.italic),
                                "font_size": cell.font.sz or 11,
                                "font_color": str(cell.font.color.rgb)[-6:]
                                if cell.font.color
                                and cell.font.color.type == "rgb"
                                and isinstance(cell.font.color.rgb, str)
                                else "",
                                "fill": str(rgb)[-6:] if isinstance(rgb, str) else "",
                                "align": cell.alignment.horizontal or "",
                            }
                        )
                    rows.append(
                        {
                            "height": sheet.row_dimensions[r].height
                            or sheet.sheet_format.defaultRowHeight
                            or 15,
                            "cells": row_cells,
                        }
                    )
                layout.append(
                    {
                        "sheet": sheet.title,
                        "rows": rows,
                        "widths": [
                            sheet.column_dimensions[get_column_letter(c)].width
                            or sheet.sheet_format.defaultColWidth
                            or 8.43
                            for c in range(1, max_col + 1)
                        ],
                        "total_rows": sheet.max_row,
                        "total_columns": sheet.max_column,
                        "truncated": sheet.max_row > max_row or sheet.max_column > max_col,
                    }
                )
        return {
            "sha256": hashlib.sha256(body).hexdigest(),
            "read_only_print_areas": read_only_print_areas,
            "cells": cells,
            "sheets": wb.sheetnames,
            "sheet_overview": [
                {
                    "sheet": s.title,
                    "rows": s.max_row,
                    "columns": s.max_column,
                    "merged_ranges": len(s.merged_cells.ranges),
                    "merged_regions": [str(item) for item in s.merged_cells.ranges],
                    "print_area": s.print_area,
                    "sheet_state": s.sheet_state,
                    "hidden_rows": [
                        int(key) for key, dim in s.row_dimensions.items() if dim.hidden
                    ],
                    "hidden_columns": [
                        str(key) for key, dim in s.column_dimensions.items() if dim.hidden
                    ],
                    "hidden_column_ranges": [
                        {
                            "start": dim.min or column_index_from_string(key),
                            "end": dim.max or dim.min or column_index_from_string(key),
                        }
                        for key, dim in s.column_dimensions.items()
                        if dim.hidden
                    ],
                }
                for s in wb.worksheets
            ],
            "layout": layout,
            "layout_truncated": len(wb.sheetnames) > len(layout),
            "model": "not_called",
        }
    finally:
        wb.close()


def _restore_print_areas(body: bytes, areas: list[dict]) -> bytes:
    """Restore validated OOXML print definitions dropped by openpyxl and verify them."""
    from xml.etree import ElementTree as ET
    from zipfile import ZIP_DEFLATED

    from defusedxml import ElementTree as SafeET

    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

    def tag(name):
        return f"{{{namespace}}}{name}"

    output = BytesIO()
    with ZipFile(BytesIO(body)) as old, ZipFile(output, "w", ZIP_DEFLATED) as new:
        for entry in old.infolist():
            data = old.read(entry.filename)
            if entry.filename == "xl/workbook.xml":
                root = SafeET.fromstring(data)
                sheets = [
                    sheet.get("name") for sheet in root.findall(f"{tag('sheets')}/{tag('sheet')}")
                ]
                names = root.find(tag("definedNames"))
                if names is None:
                    names = ET.Element(tag("definedNames"))
                    calc = root.find(tag("calcPr"))
                    root.insert(list(root).index(calc) if calc is not None else len(root), names)
                for area in areas:
                    if area["sheet"] not in sheets:
                        raise ReviewedExcelError("打印区域的工作表不存在")
                    index = str(sheets.index(area["sheet"]))
                    for node in list(names):
                        if (
                            node.get("name") == "_xlnm.Print_Area"
                            and node.get("localSheetId") == index
                        ):
                            names.remove(node)
                    node = ET.SubElement(
                        names,
                        tag("definedName"),
                        {"name": "_xlnm.Print_Area", "localSheetId": index},
                    )
                    node.text = area["value"]
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            new.writestr(entry, data)
    result = output.getvalue()
    restored = inspect_workbook(result)["read_only_print_areas"]
    if sorted((x["sheet"], x["value"]) for x in restored) != sorted(
        (x["sheet"], x["value"]) for x in areas
    ):
        raise ReviewedExcelError("打印区域写回后核验失败，已停止")
    return result


def derive_document_workbook(
    body: bytes, fragments: list[dict], *, print_area_mode: str = "source"
) -> bytes:
    """Create one isolated sheet at source coordinates; reject unprovable dependencies."""
    import re

    from openpyxl.formula.tokenizer import Tokenizer, TokenizerError

    info = inspect_workbook(body)
    if print_area_mode not in {"source", "selected"}:
        raise ReviewedExcelError("打印区域选择无效")
    if (
        not fragments
        or len(fragments) > 20
        or any(not isinstance(item, dict) or set(item) != {"sheet", "range"} for item in fragments)
    ):
        raise ReviewedExcelError("文档片段范围无效")
    sheets = {item["sheet"] for item in fragments}
    if len(sheets) != 1 or next(iter(sheets)) not in info["sheets"]:
        raise ReviewedExcelError("跨工作表文档暂不能保证保真")
    sheet_name = next(iter(sheets))
    special_areas = [area for area in info["read_only_print_areas"] if area["sheet"] == sheet_name]
    if special_areas and print_area_mode != "selected":
        raise ReviewedExcelError("原打印区域跨越文档边界，请明确选择所选片段作为派生打印区域")
    dimensions = next(sheet for sheet in info["sheet_overview"] if sheet["sheet"] == sheet_name)
    rectangles = []
    for item in fragments:
        try:
            bounds = bounded_range(item["range"])
            c1, r1, c2, r2 = bounds
            if not all(isinstance(x, int) for x in bounds) or not (
                1 <= c1 <= c2 <= dimensions["columns"] and 1 <= r1 <= r2 <= dimensions["rows"]
            ):
                raise ReviewedExcelError("bounds")
        except (TypeError, ValueError):
            raise ReviewedExcelError("文档片段坐标无效") from None
        rectangles.append(bounds)

    def contained(bounds):
        c1, r1, c2, r2 = bounds
        return any(x1 <= c1 <= c2 <= x2 and y1 <= r1 <= r2 <= y2 for x1, y1, x2, y2 in rectangles)

    def intersects(bounds):
        c1, r1, c2, r2 = bounds
        return any(c1 <= x2 and x1 <= c2 and r1 <= y2 and y1 <= r2 for x1, y1, x2, y2 in rectangles)

    from app.application.reviewed_excel_formula_support import (
        install_formula_support,
        plan_formula_support,
    )

    with warnings.catch_warnings(record=True) as found:
        warnings.simplefilter("always")
        wb = load_workbook(BytesIO(body), data_only=False, keep_links=True, rich_text=True)
    if found and _read_only_print_area_warnings(body, found) != info["read_only_print_areas"]:
        wb.close()
        raise ReviewedExcelError("打印区域来源校验失败")
    try:
        original_modified = wb.properties.modified
        ws = wb[sheet_name]
        if not isinstance(ws, Worksheet):
            raise ReviewedExcelError("所选对象不是工作表")
        if ws.sheet_state != "visible":
            raise ReviewedExcelError("隐藏工作表暂不能作为独立模板提取")
        if (
            wb.defined_names
            or ws.defined_names
            or ws.print_title_rows
            or ws.print_title_cols
            or ws.data_validations.count
            or len(ws.conditional_formatting)
            or ws.tables
            or ws.auto_filter.ref
        ):
            raise ReviewedExcelError("工作簿存在未验证的名称或区域依赖，暂不能提取")
        header_names = (
            "oddHeader",
            "oddFooter",
            "evenHeader",
            "evenFooter",
            "firstHeader",
            "firstFooter",
        )
        headers = {name: str(getattr(ws, name)) for name in header_names}
        print_settings = (str(ws.page_setup), str(ws.page_margins), str(ws.print_options))
        source_print = (
            [bounded_range(str(item)) for item in PrintArea.from_string(ws.print_area).ranges]
            if ws.print_area
            else []
        )
        if (
            source_print
            and any(not contained(bounds) for bounds in source_print)
            and print_area_mode != "selected"
        ):
            raise ReviewedExcelError("打印区域跨越所选文档，请明确选择所选片段作为派生打印区域")
        mapped_print = print_area_mode == "selected"
        if mapped_print:
            ws.print_area = [item["range"] for item in fragments]
        for merge in list(ws.merged_cells.ranges):
            bounds = bounded_range(str(merge))
            if intersects(bounds) and not contained(bounds):
                raise ReviewedExcelError("所选范围切断合并单元格")
            if not intersects(bounds):
                ws.unmerge_cells(str(merge))
        formula_plan = plan_formula_support(body, fragments)
        clear_cells = []
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue
                if not contained((cell.column, cell.row, cell.column, cell.row)):
                    clear_cells.append(cell)
                    continue
                if cell.hyperlink:
                    raise ReviewedExcelError("文档包含未核验的链接，暂不能提取")
                if cell.data_type == "f":
                    try:
                        tokens = Tokenizer(cell.value).items
                    except (TokenizerError, TypeError):
                        raise ReviewedExcelError("公式引用无法核验") from None
                    for token in tokens:
                        if token.type == "OPERAND" and token.subtype == "ERROR":
                            raise ReviewedExcelError("公式含错误引用，暂不能提取")
                        if (
                            token.type == "FUNC"
                            and token.subtype == "OPEN"
                            and token.value[:-1].upper()
                            not in {
                                "SUM",
                                "AVERAGE",
                                "MIN",
                                "MAX",
                                "ROUND",
                                "IF",
                                "IFERROR",
                                "COUNT",
                                "COUNTA",
                            }
                        ):
                            raise ReviewedExcelError("公式函数依赖无法核验")
                        if token.type == "OPERAND" and token.subtype == "RANGE":
                            if not re.fullmatch(
                                r"\$?[A-Za-z]{1,3}\$?\d+(?::\$?[A-Za-z]{1,3}\$?\d+)?", token.value
                            ):
                                raise ReviewedExcelError("公式含外部或名称引用，暂不能提取")
        support_sheet = install_formula_support(wb, formula_plan)
        support_values = (
            {
                address: wb[support_sheet][address].value
                for address in formula_plan["support_cells"] + formula_plan["proxy_cells"]
            }
            if support_sheet
            else {}
        )
        formula_values = {
            address: ws[address].value
            for address in formula_plan["formula_references"]
            if contained((ws[address].column, ws[address].row, ws[address].column, ws[address].row))
        }
        for cell in clear_cells:
            cell.value = None
            cell.comment = None
            cell.hyperlink = None
        for name in list(wb.sheetnames):
            if name not in {sheet_name, support_sheet}:
                del wb[name]
        wb.active = 0
        output = BytesIO()
        with warnings.catch_warnings(record=True) as found:
            warnings.simplefilter("always")
            wb.save(output)
        if found:
            raise ReviewedExcelError("文档保存存在未验证的内容，已停止")
        result = output.getvalue()
        checked = load_workbook(BytesIO(result), data_only=False, keep_links=True, rich_text=True)
        try:
            saved_sheet = checked[sheet_name]
            if not isinstance(saved_sheet, Worksheet):
                raise ReviewedExcelError("派生对象不是工作表")
            if any(
                saved_sheet[address].value != value for address, value in formula_values.items()
            ):
                raise ReviewedExcelError("派生公式写出后核验失败")
            if support_sheet and any(
                checked[support_sheet][address].value != value
                for address, value in support_values.items()
            ):
                raise ReviewedExcelError("支撑数据写出后核验失败")
            if headers != {
                name: str(getattr(saved_sheet, name)) for name in header_names
            } or print_settings != (
                str(saved_sheet.page_setup),
                str(saved_sheet.page_margins),
                str(saved_sheet.print_options),
            ):
                raise ReviewedExcelError("派生页眉或打印设置写出后核验失败")
            if mapped_print:
                actual = sorted(
                    bounded_range(str(item))
                    for item in PrintArea.from_string(saved_sheet.print_area).ranges
                )
                if actual != sorted(rectangles):
                    raise ReviewedExcelError("派生打印区域写出后核验失败")
        finally:
            checked.close()
        # ZIP member timestamps otherwise vary between saves, preventing a
        # verified re-derivation of exactly the same temporary document later.
        normalized = BytesIO()
        with ZipFile(BytesIO(result)) as source, ZipFile(normalized, "w") as target:
            for entry in source.infolist():
                entry.date_time = (1980, 1, 1, 0, 0, 0)
                contents = source.read(entry.filename)
                if entry.filename == "docProps/core.xml":
                    modified = (
                        original_modified.strftime("%Y-%m-%dT%H:%M:%SZ")
                        if original_modified
                        else "1980-01-01T00:00:00Z"
                    ).encode("ascii")
                    contents, count = re.subn(
                        rb"(<dcterms:modified\b[^>]*>)[^<]*(</dcterms:modified>)",
                        lambda match: match.group(1) + modified + match.group(2),
                        contents,
                    )
                    if count != 1:
                        raise ReviewedExcelError("派生文件元数据无法稳定重建")
                target.writestr(entry, contents)
        return normalized.getvalue()
    finally:
        wb.close()


def clean_workbook(body: bytes, expected_sha256: str, selections: list[dict]) -> bytes:
    inspection = inspect_workbook(body)
    if inspection["sha256"] != expected_sha256:
        raise ReviewedExcelError("文件与预览不一致，请重新预览")
    allowed = {
        (item["sheet"], item["cell"]) for item in inspection["cells"] if item["kind"] == "review"
    }
    if len(selections) > MAX_CELLS:
        raise ReviewedExcelError("选择的单元格过多")
    selected = set()
    for item in selections:
        if not isinstance(item, dict) or set(item) != {"sheet", "cell"}:
            raise ReviewedExcelError("清理范围无效")
        key = (item["sheet"], item["cell"])
        if key not in allowed or key in selected:
            raise ReviewedExcelError("清理范围包含无效、重复或受保护的单元格")
        selected.add(key)
    with warnings.catch_warnings(record=True) as found:
        warnings.simplefilter("always")
        wb = load_workbook(BytesIO(body), data_only=False, keep_links=True, rich_text=True)
    if found and _read_only_print_area_warnings(body, found) != inspection["read_only_print_areas"]:
        wb.close()
        raise ReviewedExcelError("打印区域来源校验失败")
    try:
        for sheet, address in selected:
            wb[sheet][address].value = None
        if any(item["kind"] == "formula" for item in inspection["cells"]):
            # openpyxl does not preserve formula caches. Ask spreadsheet software
            # to recalculate when the cleaned copy is opened.
            if wb.calculation is None:
                from openpyxl.workbook.properties import CalcProperties

                wb.calculation = CalcProperties()
            wb.calculation.fullCalcOnLoad = True
            wb.calculation.forceFullCalc = True
        output = BytesIO()
        with warnings.catch_warnings(record=True) as found:
            warnings.simplefilter("always")
            wb.save(output)
        if found:
            raise ReviewedExcelError("保存时发现无法保证保真的内容，已停止")
        result = output.getvalue()
        if inspection["read_only_print_areas"]:
            result = _restore_print_areas(result, inspection["read_only_print_areas"])
        return result
    finally:
        wb.close()


def make_proposals(review: dict) -> dict:
    """Recommend bounded, explainable candidates; uncertain cells stay unchanged."""
    import json
    from collections import defaultdict

    from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

    labels = {"客户", "客户名称", "客户姓名", "日期", "报价日期", "单号", "报价单号"}
    headings = {
        "产品名称",
        "商品名称",
        "项目名称",
        "項目名稱",
        "品名",
        "数量",
        "數量",
        "单价",
        "單價",
        "型号",
        "规格",
        "规格型号",
        "小计",
        "小計",
    }
    stops = (
        "合计",
        "合計",
        "总计",
        "總計",
        "条款",
        "條款",
        "说明",
        "說明",
        "备注",
        "備註",
        "税",
        "稅",
    )
    by_row: defaultdict[tuple[str, int], dict[int, dict[str, Any]]] = defaultdict(dict)
    for cell in review["cells"]:
        column_letter, row = coordinate_from_string(cell["cell"])
        by_row[(cell["sheet"], row)][column_index_from_string(column_letter)] = cell
    fixed = set()
    identity = []
    table = []
    evidence = []
    for (sheet, row), cells in by_row.items():
        for col, cell in cells.items():
            value = cell["value"].strip()
            if (
                cell["kind"] == "formula"
                or value in labels
                or value in headings
                or any(x in value for x in stops)
            ):
                fixed.add((sheet, cell["cell"]))
            if any(x in value for x in stops) and col + 1 in cells:
                fixed.add((sheet, cells[col + 1]["cell"]))
            if value in labels and col + 1 in cells:
                target = cells[col + 1]
                if target["kind"] != "formula":
                    identity.append({"sheet": sheet, "cell": target["cell"]})
        header_cols = [col for col, cell in cells.items() if cell["value"].strip() in headings]
        if len(header_cols) < 2:
            continue
        evidence.append({"sheet": sheet, "header_row": row, "columns": len(header_cols)})
        next_row = row + 1
        while (sheet, next_row) in by_row:
            values = by_row[(sheet, next_row)]
            if any(any(term in cell["value"] for term in stops) for cell in values.values()):
                break
            candidates = [values[col] for col in header_cols if col in values]
            if len(candidates) < 2:
                break
            for item in candidates:
                if item["kind"] != "formula":
                    table.append({"sheet": sheet, "cell": item["cell"]})
            next_row += 1

    def safe(items):
        seen = set()
        out = []
        for item in items:
            pair = (item["sheet"], item["cell"])
            if pair not in fixed and pair not in seen:
                seen.add(pair)
                out.append(item)
        return out

    table = safe(table)
    identity = safe(identity)
    candidates = [
        {
            "id": "preserve",
            "label": "保留原件内容",
            "clear": [],
            "explanation": "全部保留，适合先核对",
        },
        {
            "id": "table",
            "label": "清理识别到的明细",
            "clear": table,
            "explanation": "仅建议清理已识别表头下的连续明细；公式保留",
        },
        {
            "id": "table_identity",
            "label": "清理明细及单据信息",
            "clear": safe(table + identity),
            "explanation": "在明细基础上建议清理客户、日期或单号的相邻值",
        },
    ]
    version = hashlib.sha256(
        json.dumps(
            {"source": review["sha256"], "candidates": candidates},
            ensure_ascii=False,
            sort_keys=True,
        ).encode()
    ).hexdigest()
    return {
        "version": version,
        "parent_version": None,
        "candidates": candidates,
        "evidence": evidence,
        "protected": [{"sheet": s, "cell": c} for s, c in sorted(fixed)],
        "model": {"status": "not_called"},
        "uncertain": "未命中的单元格保留，建议请人工确认",
    }


def refine_proposal(
    review: dict,
    parent_version: str,
    feedback: str,
    *,
    owner_id: int,
    current: list | None = None,
    initial: bool = False,
) -> dict:
    """Use the existing ETL provider for a bounded, explicitly requested revision."""
    import json

    from app.application.etl import llm_assist
    from app.application.etl.llm_session_provider import bind_etl_llm_owner, reset_etl_llm_owner

    base = make_proposals(review)
    if len(parent_version) != 64 or not feedback.strip() or len(feedback) > 1000:
        raise ReviewedExcelError("版本或反馈无效")
    protected_keys = {(c["sheet"], c["cell"]) for c in base["protected"]}
    by_sheet = {
        name: [c for c in review["cells"] if c["sheet"] == name] for name in review["sheets"]
    }
    prioritized = [c for c in review["cells"] if (c["sheet"], c["cell"]) in protected_keys]
    sample = [cell for items in by_sheet.values() if items for cell in (items[0], items[-1])]
    for index in range(100):
        for items in by_sheet.values():
            if items:
                sample.append(items[round(index * (len(items) - 1) / 99)])
    unique: dict[tuple[str, str], dict[str, Any]] = {}
    for cell in prioritized + sample:
        unique.setdefault((cell["sheet"], cell["cell"]), cell)
    chosen = list(unique.values())[:100]
    bounded = [
        {"sheet": c["sheet"], "cell": c["cell"], "value": c["value"][:160], "kind": c["kind"]}
        for c in chosen
    ]
    schema = {
        "type": "object",
        "required": ["clear", "reason"],
        "properties": {
            "clear": {
                "type": "array",
                "items": {
                    "type": "object",
                    "required": ["sheet", "cell"],
                    "properties": {"sheet": {"type": "string"}, "cell": {"type": "string"}},
                    "additionalProperties": False,
                },
            },
            "reason": {"type": "string"},
        },
        "additionalProperties": False,
    }
    if initial:
        schema = {
            "type": "object",
            "required": ["candidates"],
            "properties": {
                "candidates": {"type": "array", "minItems": 3, "maxItems": 3, "items": schema}
            },
            "additionalProperties": False,
        }
    current = current or []
    if len(current) > 100 or any(
        not isinstance(item, dict) or set(item) != {"sheet", "cell"} for item in current
    ):
        raise ReviewedExcelError("当前候选超出模型输入范围，原候选未改变")
    bounded_coordinates = {(item["sheet"], item["cell"]) for item in bounded}
    if any((item["sheet"], item["cell"]) not in bounded_coordinates for item in current):
        raise ReviewedExcelError("当前候选包含样本范围外的单元格，不能交给模型修改")
    sent_protected = [
        item for item in base["protected"] if (item["sheet"], item["cell"]) in bounded_coordinates
    ]
    value_truncated = any(len(item["value"]) > 160 for item in chosen)
    dimensions = {item["sheet"]: item for item in review.get("sheet_overview", [])}
    overview = [
        {
            "sheet": name,
            **{
                key: value
                for key, value in dimensions.get(name, {}).items()
                if key != "merged_regions"
            },
            "nonempty_cells": len(items),
            "formulas": sum(c["kind"] == "formula" for c in items),
            "provided": sum(c["sheet"] == name for c in bounded),
            "uncovered": len(items) - sum(c["sheet"] == name for c in bounded),
        }
        for name, items in by_sheet.items()
    ]
    messages = [
        {
            "role": "system",
            "content": "Propose cells to clear in a copy of a spreadsheet. Preserve formulas, headers and fixed instructions. Use only exact sheet and cell coordinates supplied in cells; do not use ranges or operation objects. Return one JSON object matching output_schema exactly, with no extra properties. Each clear array contains individual {sheet, cell} objects; an empty array means retain all."
            + (
                " candidates must contain exactly three complete independent plans, each with its own clear array and reason. They are alternatives, not three operations in one plan; explain equivalence if evidence is insufficient."
                if initial
                else ""
            ),
        },
        {
            "role": "user",
            "content": json.dumps(
                {
                    "feedback": feedback,
                    "cells": bounded,
                    "current": current,
                    "protected": sent_protected,
                    "sheet_overview": overview,
                    "detected_headers": base["evidence"],
                    "truncated": len(review["cells"]) > 100,
                    "output_schema": schema,
                },
                ensure_ascii=False,
            ),
        },
    ]
    token = bind_etl_llm_owner(owner_id)
    try:
        import os

        configured_provider = os.environ.get("FHD_TEMPLATE_LLM_PROVIDER", "").strip().lower()
        if configured_provider == "mimo":
            result = _template_mimo_complete(messages, schema, 1400)
        elif configured_provider:
            raise ReviewedExcelError("模板模型供应商配置不受支持；未调用模型")
        else:
            result = llm_assist._complete(messages, schema=schema, max_tokens=1400)
    finally:
        reset_etl_llm_owner(token)
    if not result.used_llm or result.degraded:
        state = (
            "模型调用失败" if result.used_llm or result.degraded else "模型未调用（未配置或已关闭）"
        )
        raise ReviewedExcelError(f"{state}，自然语言修改不可用；原候选未改变")
    allowed = {(c["sheet"], c["cell"]) for c in bounded if c["kind"] != "formula"}
    protected = {(c["sheet"], c["cell"]) for c in base["protected"]}
    if initial and (not isinstance(result.data, dict) or set(result.data) != {"candidates"}):
        raise ReviewedExcelError("模型返回格式无效，原候选未改变")
    answers = result.data.get("candidates") if initial else [result.data]
    if not isinstance(answers, list) or len(answers) != (3 if initial else 1):
        raise ReviewedExcelError("模型未返回所需候选，原候选未改变")
    validated: list[dict[str, Any]] = []
    for answer in answers:
        if (
            not isinstance(answer, dict)
            or set(answer) != {"clear", "reason"}
            or not isinstance(answer.get("clear"), list)
            or not isinstance(answer.get("reason"), str)
        ):
            raise ReviewedExcelError("模型返回格式无效，原候选未改变")
        pairs = []
        for item in answer["clear"]:
            if (
                not isinstance(item, dict)
                or set(item) != {"sheet", "cell"}
                or not isinstance(item["sheet"], str)
                or not isinstance(item["cell"], str)
            ):
                raise ReviewedExcelError("模型返回坐标无效，原候选未改变")
            pair = (item["sheet"], item["cell"])
            if pair not in allowed or pair in protected or pair in pairs:
                raise ReviewedExcelError("模型建议包含未提供或受保护单元格，原候选未改变")
            pairs.append(pair)
        index = len(validated) + 1
        validated.append(
            {
                "id": f"model_initial_{index}" if initial else "model_revision",
                "label": f"模型候选 {index}" if initial else "模型反馈候选",
                "clear": answer["clear"],
                "explanation": answer["reason"][:300],
            }
        )
    if initial:
        base["candidates"] = validated
    else:
        base["candidates"][2] = validated[0]
    base["parent_version"] = parent_version
    base["version"] = hashlib.sha256(
        json.dumps(
            {
                "source": review["sha256"],
                "parent": parent_version,
                "feedback": feedback,
                "candidates": validated,
            },
            ensure_ascii=False,
            sort_keys=True,
        ).encode()
    ).hexdigest()
    from app.application.reviewed_excel_evidence import summarize_model_input

    input_summary = summarize_model_input(
        messages,
        bounded,
        source_sha256=review["sha256"],
        provider=configured_provider or "etl_default",
        model=result.model,
        max_output_tokens=1400,
    )
    base["model"] = {
        "status": "returned",
        "adopted": False,
        "cells_sent": len(bounded),
        "cells_total": len(review["cells"]),
        "truncated": len(review["cells"]) > 100,
        "values_truncated": value_truncated,
        "max_value_chars": 160,
        "input_summary": input_summary,
        "sheet_overview": overview,
    }
    return base
