"""Line-item normalization, stable note identity and product projection."""

from __future__ import annotations

import hashlib
import json
import re
from datetime import datetime
from typing import Any, cast

from app.application.shipment_etl_profile import (
    ShipmentEtlProfile,
)
from app.utils.operational_errors import RECOVERABLE_ERRORS

from .layout import _joined_row, _to_float, _to_int

_NON_PRODUCT_TOKENS = frozenset(
    {
        "january",
        "february",
        "march",
        "april",
        "may",
        "june",
        "july",
        "august",
        "september",
        "october",
        "november",
        "december",
        "title",
        "identifier",
        "subject",
        "description",
        "notes",
        "creator",
        "accession",
        "my title",
        "another title",
        "the best image ever",
    }
)


def _looks_like_non_product_token(value: str) -> bool:
    text = str(value or "").strip().lower()
    return bool(text) and text in _NON_PRODUCT_TOKENS


def _looks_like_titleish(value: str) -> bool:
    text = str(value or "").strip().lower()
    if not text:
        return False
    if text in _NON_PRODUCT_TOKENS:
        return True
    return bool(re.fullmatch("[a-z][a-z\\s\\-']{2,40}", text)) and (not re.search("\\d", text))


def _build_item_from_row(ws, row: int, mapping: dict[str, int]) -> dict[str, Any] | None:
    model = ""
    name = ""
    if "model_number" in mapping:
        model = str(ws.cell(row, mapping["model_number"]).value or "").strip()
    if "product_name" in mapping:
        name = str(ws.cell(row, mapping["product_name"]).value or "").strip()
    if not name and model and (not re.search("[A-Za-z0-9]", model)):
        name, model = (model, "")
    if not name and (not model):
        return None
    if _looks_like_non_product_token(name) or _looks_like_non_product_token(model):
        return None
    tins = (
        _to_int(ws.cell(row, mapping["quantity_tins"]).value) if "quantity_tins" in mapping else 0
    )
    tin_spec = _to_float(ws.cell(row, mapping["tin_spec"]).value) if "tin_spec" in mapping else 0.0
    qty_kg = (
        _to_float(ws.cell(row, mapping["quantity_kg"]).value) if "quantity_kg" in mapping else 0.0
    )
    unit_price = (
        _to_float(ws.cell(row, mapping["unit_price"]).value) if "unit_price" in mapping else 0.0
    )
    amount = _to_float(ws.cell(row, mapping["amount"]).value) if "amount" in mapping else 0.0
    if tins <= 0 and qty_kg <= 0 and (unit_price <= 0) and (amount <= 0):
        return None
    if 1990 <= tins <= 2035 and (not model) and _looks_like_titleish(name):
        return None
    if tin_spec <= 0 and tins > 0 and (qty_kg > 0):
        tin_spec = qty_kg / tins
    if qty_kg <= 0 and tins > 0 and (tin_spec > 0):
        qty_kg = tins * tin_spec
    if amount <= 0 and unit_price > 0 and (qty_kg > 0):
        amount = unit_price * qty_kg
    if tins <= 0 and qty_kg > 0:
        tins = 1
        if tin_spec <= 0:
            tin_spec = qty_kg
    return {
        "product_name": name or model,
        "model_number": model,
        "quantity_tins": max(0, tins),
        "tin_spec": tin_spec or 0.0,
        "spec_per_tin": tin_spec or 0.0,
        "quantity_kg": qty_kg,
        "unit_price": unit_price,
        "amount": amount,
        "quantity": max(1, tins) if tins else 1,
    }


def _parse_items(
    ws, header_row: int, mapping: dict[str, int], profile: ShipmentEtlProfile
) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    max_row = int(ws.max_row or 0)
    for row in range(header_row + 1, max_row + 1):
        joined = _joined_row(ws, row)
        if not joined:
            continue
        if profile.meta_patterns.stop_row.search(joined):
            break
        item = _build_item_from_row(ws, row, mapping)
        if item:
            items.append(item)
    return items


def note_fingerprint(note: dict[str, Any]) -> str:
    """内容指纹：同客户+单号+明细再导入可幂等跳过。"""
    payload = {
        "unit": str(note.get("unit_name") or "").strip(),
        "order": str(note.get("order_number") or "").strip(),
        "date": str(note.get("order_date") or "").strip(),
        "items": sorted(
            [
                {
                    "m": str(i.get("model_number") or "").strip().upper(),
                    "n": str(i.get("product_name") or "").strip(),
                    "q": float(i.get("quantity_tins") or i.get("quantity") or 0),
                    "k": float(i.get("quantity_kg") or 0),
                    "p": float(i.get("unit_price") or 0),
                }
                for i in note.get("items") or []
            ],
            key=lambda x: (x["m"], x["n"], x["q"], x["k"], x["p"]),
        ),
    }
    raw = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()[:28]


def _enrich_note(note: dict[str, Any]) -> dict[str, Any]:
    out = dict(note)
    out["sheet_name"] = str(out.get("sheet_name") or out.get("sheet") or "")
    out["sheet"] = out["sheet_name"] or str(out.get("sheet") or "")
    out["fingerprint"] = note_fingerprint(out)
    out["item_count"] = len(out.get("items") or [])
    out["total_amount"] = round(sum(float(i.get("amount") or 0) for i in out.get("items") or []), 2)
    return out


def _excel_date_to_str(value: Any, profile: ShipmentEtlProfile) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (int, float)):
        try:
            from openpyxl.utils.datetime import from_excel

            return cast("str", from_excel(value).strftime("%Y-%m-%d"))
        except RECOVERABLE_ERRORS:
            return str(value)
    text = str(value).strip()
    date_m = profile.meta_patterns.date.search(text)
    return date_m.group(1).replace(" ", "") if date_m else text


def _notes_to_product_records(
    notes: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str]] = set()
    for note in notes:
        unit = str(note.get("unit_name") or "").strip()
        for item in note.get("items") or []:
            model = str(item.get("model_number") or "").strip().upper()
            name = str(item.get("product_name") or "").strip()
            key = (unit, model, name)
            if key in seen:
                continue
            seen.add(key)
            records.append(
                {
                    "unit_name": unit,
                    "product_name": name,
                    "model_number": model,
                    "unit_price": float(item.get("unit_price") or 0),
                }
            )
    return records
