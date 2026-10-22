"""Workbook output and parse/write round trips through explicit profile adapters."""

from __future__ import annotations

import re
from collections.abc import Callable
from datetime import datetime
from pathlib import Path
from typing import TYPE_CHECKING, Any, Protocol, cast

from app.application.shipment_etl_profile import ShipmentEtlProfile

from .workbook import ParseNotes

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet

ResolveProfile = Callable[[ShipmentEtlProfile | None, str | None], ShipmentEtlProfile]


class WriteDelivery(Protocol):
    def __call__(
        self,
        notes: list[dict[str, Any]],
        output_path: str | Path,
        *,
        seller_title: str | None = None,
        profile_id: str | None = None,
        profile: ShipmentEtlProfile | None = None,
    ) -> dict[str, Any]: ...


def _write_delivery_item(
    sheet: Worksheet,
    row: int,
    item: dict[str, Any],
    columns: dict[str, Any],
    *,
    aliases: bool,
) -> None:
    values = {
        "model_number": item.get("model_number") or "",
        "product_name": item.get("product_name") or "",
        "quantity_tins": item.get("quantity_tins") or (item.get("quantity") if aliases else 0) or 0,
        "tin_spec": item.get("tin_spec") or (item.get("spec_per_tin") if aliases else 0) or 0,
        "quantity_kg": item.get("quantity_kg") or 0,
        "unit_price": item.get("unit_price") or 0,
        "amount": item.get("amount") or 0,
    }
    for field, value in values.items():
        if field in columns:
            sheet.cell(row, int(columns[field]), value)


def write_delivery(
    notes: list[dict[str, Any]],
    output_path: str | Path,
    *,
    seller_title: str | None = None,
    profile_id: str | None = None,
    profile: ShipmentEtlProfile | None = None,
    resolve_profile: ResolveProfile,
) -> dict[str, Any]:
    try:
        from openpyxl import Workbook
    except ImportError:
        return {"success": False, "message": "缺少 openpyxl，无法解析 Excel"}
    prof = resolve_profile(profile, profile_id)
    config = prof.write or {}
    title = str(seller_title if seller_title is not None else config.get("seller_title") or "")
    headers = list(config.get("header_row") or [])
    columns = dict(config.get("item_columns") or {})
    date_format = str(config.get("date_format") or "%Y-%m-%d")
    meta_template = str(
        config.get("meta_line_template") or "{unit} {contact} {order_date} {order_no}"
    )
    footer = str(config.get("footer_label") or "")
    default_sheet = str(config.get("default_sheet_name") or "Sheet1")
    sheet_prefix = str(config.get("sheet_name_prefix") or "S")
    path = Path(output_path).expanduser().resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    default = cast("Worksheet", workbook.active or workbook.create_sheet(default_sheet[:31]))
    created = 0
    try:
        for index, note in enumerate(notes or [], start=1):
            unit = str(note.get("unit_name") or f"客户{index}").strip()
            name = (
                str(note.get("sheet_name") or note.get("sheet") or unit)[:28]
                or f"{sheet_prefix}{index}"
            )
            name = re.sub(r"[\\/*?:\[\]]", "_", name)[:31]
            if index == 1:
                sheet = default
                sheet.title = name
            else:
                base, suffix = name, 1
                while name in workbook.sheetnames:
                    name = f"{base[:28]}_{suffix}"
                    suffix += 1
                sheet = workbook.create_sheet(name)
            sheet["A1"] = title
            sheet["A2"] = meta_template.format(
                unit=unit,
                contact=str(note.get("contact_person") or "").strip(),
                order_date=str(
                    note.get("order_date") or datetime.now().strftime(date_format)
                ).strip(),
                order_no=str(note.get("order_number") or f"TEST-{index:04d}").strip(),
            )
            for column, header in enumerate(headers, start=1):
                sheet.cell(3, column, header)
            last_row = 3
            for row, item in enumerate(note.get("items") or [], start=4):
                _write_delivery_item(sheet, row, item, columns, aliases=True)
                last_row = row
            if footer:
                sheet.cell(last_row + 2, 1, footer)
            created += 1
        if not created:
            default.title = default_sheet[:31]
            default["A1"] = title
            default["A2"] = str(config.get("demo_meta_line") or meta_template)
            for column, header in enumerate(headers, start=1):
                default.cell(3, column, header)
            _write_delivery_item(
                default, 4, dict(config.get("demo_item") or {}), columns, aliases=False
            )
            created = 1
        workbook.save(path)
    finally:
        workbook.close()
    return {
        "success": True,
        "file_path": str(path),
        "sheet_count": created,
        "profile_id": prof.id,
        "message": f"已生成送货单模板 {path.name}（{created} 张表）",
    }


def write_ledger(
    rows: list[dict[str, Any]],
    output_path: str | Path,
    *,
    sheet_name: str | None = None,
    unit_name: str | None = None,
    profile_id: str | None = None,
    profile: ShipmentEtlProfile | None = None,
    resolve_profile: ResolveProfile,
) -> dict[str, Any]:
    try:
        from openpyxl import Workbook
    except ImportError:
        return {"success": False, "message": "缺少 openpyxl，无法解析 Excel"}
    prof = resolve_profile(profile, profile_id)
    config = prof.write or {}
    default_sheet = str(sheet_name or config.get("ledger_sheet_name") or "ledger")
    resolved_unit = str(unit_name or config.get("ledger_default_unit") or "unit")
    headers = list(config.get("ledger_header_row") or [])
    columns = dict(config.get("ledger_item_columns") or {})
    samples = rows or list(config.get("ledger_sample_rows") or [])
    extra_sheet = str(config.get("ledger_extra_sheet") or "").strip()
    path = Path(output_path).expanduser().resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = cast("Worksheet", workbook.active)
    assert sheet is not None
    numeric_fields = {"quantity_tins", "tin_spec", "quantity_kg", "unit_price", "amount"}
    try:
        sheet.title = default_sheet[:31] or "ledger"
        for column, header in enumerate(headers, start=1):
            sheet.cell(1, column, header)
        for row_index, row in enumerate(samples, start=2):
            for field, column in columns.items():
                value = row.get(field)
                if value is None or value == "":
                    value = 0 if field in numeric_fields else ""
                sheet.cell(row_index, int(column), value)
        if extra_sheet and extra_sheet not in workbook.sheetnames:
            workbook.create_sheet(extra_sheet[:31])
        workbook.save(path)
    finally:
        workbook.close()
    return {
        "success": True,
        "file_path": str(path),
        "unit_name": resolved_unit,
        "row_count": len(samples),
        "profile_id": prof.id,
        "message": f"已生成出货流水模板 {path.name}",
    }


def regenerate(
    file_path: str | Path,
    output_path: str | Path,
    *,
    include_ledger: bool | str = "auto",
    profile_id: str | None = None,
    profile: ShipmentEtlProfile | None = None,
    resolve_profile: ResolveProfile,
    parse_notes: ParseNotes,
    write_notes: WriteDelivery,
) -> dict[str, Any]:
    prof = resolve_profile(profile, profile_id)
    source = parse_notes(file_path, include_ledger=include_ledger, profile=prof)
    if not source.get("success"):
        return source
    notes = source.get("notes") or []
    if not notes:
        return {"success": False, "message": "无可反推的单据", "error_code": "no_delivery_notes"}
    generated = write_notes(notes, output_path, profile=prof)
    if not generated.get("success"):
        return generated
    reparsed = parse_notes(output_path, include_ledger=False, profile=prof)
    return {
        "success": True,
        "source": source,
        "generated": generated,
        "reparsed": reparsed,
        "fingerprint_match": {note.get("fingerprint") for note in notes}
        == {note.get("fingerprint") for note in reparsed.get("notes") or []}
        if reparsed.get("success")
        else False,
        "profile_id": prof.id,
        "message": "模板反推完成",
    }
