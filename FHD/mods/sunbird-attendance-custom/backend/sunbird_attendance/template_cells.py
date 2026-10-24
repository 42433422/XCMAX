"""Cell text, number display and style preservation for attendance templates."""

from __future__ import annotations

import re
import unicodedata
from collections.abc import Iterable
from copy import copy

from openpyxl.utils import get_column_letter

from .template_schema import (
    _ATT_MARK,
    _CN_BLOCK_DIGITS,
    DETAIL_ARABIC_NUMBER_DISPLAY,
    DETAIL_MAX_COL,
    DETAIL_PERSON_BLOCK_ROWS,
    DETAIL_SUM_COL_START,
    DETAIL_TEMPLATE_SUMMARY_BEGIN_COL,
    BlockCellStyle,
)


def _chinese_block_label_to_int(text: str) -> int | None:
    """块内左侧标签常见「一…三十一」纯中文数字串 → 整数；含其它汉字则不改。"""
    if not text:
        return None
    text = unicodedata.normalize("NFKC", text).strip()
    if not text:
        return None
    allowed = set(_CN_BLOCK_DIGITS) | {"十"}
    if any(c not in allowed for c in text):
        return None
    d = _CN_BLOCK_DIGITS
    if len(text) == 1:
        if text in d:
            return d[text]
        if text == "十":
            return 10
        return None
    if len(text) == 2:
        if text[0] == "十" and text[1] in d:
            return 10 + d[text[1]]
        if text[1] == "十" and text[0] in d:
            return d[text[0]] * 10
    if len(text) == 3 and text[1] == "十" and (text[0] in d) and (text[2] in d):
        return d[text[0]] * 10 + d[text[2]]
    return None


def _plain_cell_text(value: object) -> str:
    """与 ``str()`` 相比，显式支持 openpyxl 的富文本单元格（拼接各段，便于匹配 ``18:30记加班``）。"""
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    try:
        from openpyxl.cell.rich_text import CellRichText

        if isinstance(value, CellRichText):
            return "".join(str(part) for part in value)
    except ImportError:
        pass
    return str(value)


def _replace_line_all_cn_numerals(ls: str) -> tuple[str, bool]:
    """把一行里每一段「仅由数字用字组成」的串按最长可解析规则换成阿拉伯数字（含无换行的「二一」）。"""
    if not ls:
        return (ls, False)
    if ls.strip() == "〇":
        return (ls, False)
    if ls[0] == "〇" and len(ls.strip()) > 1:
        return (ls, False)
    allowed = set(_CN_BLOCK_DIGITS) | {"十"}
    out: list[str] = []
    i = 0
    changed = False
    while i < len(ls):
        if ls[i] not in allowed:
            out.append(ls[i])
            i += 1
            continue
        j = i
        while j < len(ls) and ls[j] in allowed:
            j += 1
        run = ls[i:j]
        tup: tuple[int, int] | None = None
        for take in range(len(run), 0, -1):
            prefix = run[:take]
            n = _chinese_block_label_to_int(prefix)
            if n is None:
                continue
            if n == 0 and prefix == "〇":
                break
            tup = (n, take)
            break
        if tup is None:
            out.append(ls[i])
            i += 1
            continue
        n, take = tup
        out.append(str(n))
        i += take
        changed = True
    return ("".join(out), changed)


def _cell_text_replace_chinese_numerals(text: str) -> str | None:
    """多行、富文本拼成的单行、或「中文数+符号」混排：替换行内所有可解析的中文数字段。"""
    text = unicodedata.normalize("NFKC", text)
    lines = text.splitlines()
    out: list[str] = []
    changed = False
    for line in lines:
        ls = line.lstrip(" \t")
        indent_len = len(line) - len(ls)
        indent = line[:indent_len]
        if not ls:
            out.append(line)
            continue
        nl, _ = _replace_line_all_cn_numerals(ls)
        new_line = indent + nl
        if new_line != line:
            changed = True
        out.append(new_line)
    if not changed:
        return None
    return "\n".join(out)


def _merge_anchor_for_cell(
    row: int, col: int, overlaps: list[tuple[int, int, int, int]]
) -> tuple[int, int]:
    """在预筛选的合并矩形列表中解析 (row,col) 所属合并的左上角。"""
    for mr1, mc1, mr2, mc2 in overlaps:
        if mr1 <= row <= mr2 and mc1 <= col <= mc2:
            return (mr1, mc1)
    return (row, col)


def _normalize_chinese_numerals_in_rect(ws, r_lo: int, r_hi: int, col_lo: int, col_hi: int) -> None:
    """矩形 [r_lo..r_hi]×[col_lo..col_hi] 内：中文数字改为阿拉伯数字（合并格写左上角）。"""
    overlaps: list[tuple[int, int, int, int]] = []
    for rng in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = rng.bounds
        if max_row < r_lo or min_row > r_hi or max_col < col_lo or (min_col > col_hi):
            continue
        overlaps.append((min_row, min_col, max_row, max_col))
    processed_merge_tops: set[tuple[int, int]] = set()
    for row in range(r_lo, r_hi + 1):
        for col in range(col_lo, col_hi + 1):
            wr, wc = _merge_anchor_for_cell(row, col, overlaps)
            key = (wr, wc)
            if key in processed_merge_tops:
                continue
            processed_merge_tops.add(key)
            v = ws.cell(wr, wc).value
            if v is None:
                continue
            if isinstance(v, (int, float)) and type(v) is not bool:
                continue
            text = _plain_cell_text(v).strip().replace("\u3000", "").strip()
            if not text:
                continue
            text = unicodedata.normalize("NFKC", text).strip()
            if not text:
                continue
            if re.fullmatch("-?\\d+\\.", text):
                try:
                    coerced = float(text)
                except ValueError:
                    coerced = None
                if coerced is not None:
                    c = ws.cell(wr, wc)
                    c.value = coerced
                    if DETAIL_SUM_COL_START <= wc < DETAIL_TEMPLATE_SUMMARY_BEGIN_COL:
                        _force_arabic_number_format(c)
                    continue
            if text.startswith("="):
                continue
            if text in _ATT_MARK:
                continue
            n = _chinese_block_label_to_int(text)
            if n is not None:
                ws.cell(wr, wc).value = n
                continue
            mixed = _cell_text_replace_chinese_numerals(text)
            if mixed is not None:
                ws.cell(wr, wc).value = mixed


def _normalize_block_chinese_numerals(
    ws, block_top: int, col_lo: int = 1, col_hi: int = DETAIL_MAX_COL
) -> None:
    """每人 6 行块内 [col_lo..col_hi]：纯中文数字改为阿拉伯数字（合并格写左上角）。

    整格为考勤符号（如请假「〇」）的不改。合并区只预筛与本块矩形相交的范围，避免全表 merged 扫描卡死。
    """
    r_hi = block_top + DETAIL_PERSON_BLOCK_ROWS - 1
    _normalize_chinese_numerals_in_rect(ws, block_top, r_hi, col_lo, col_hi)


def _round_display(value: float) -> float:
    """一律保留一位小数的 float（如 2.0），避免整数在列格式下被显示成 ``2.``。"""
    return round(float(value), 1)


def _snapshot_block_cell_styles(ws, block_top: int) -> dict[tuple[int, int], BlockCellStyle]:
    """快照首块 6×DETAIL_MAX_COL 的字体、边框、对齐（供粘贴、清空与写入后恢复版式）。"""
    styles: dict[tuple[int, int], BlockCellStyle] = {}
    for dr in range(DETAIL_PERSON_BLOCK_ROWS):
        for c in range(1, DETAIL_MAX_COL + 1):
            cell = ws.cell(block_top + dr, c)
            try:
                font = copy(cell.font)
                border = copy(cell.border)
                alignment = copy(cell.alignment)
            except (TypeError, ValueError):
                font = border = alignment = None
            styles[dr + 1, c] = (font, border, alignment)
    return styles


def _apply_style_bundle(cell, bundle: BlockCellStyle | None) -> None:
    if not bundle:
        return
    for attribute, value in zip(("font", "border", "alignment"), bundle):
        if value is not None:
            setattr(cell, attribute, value)


def _force_arabic_number_format(cell, value: object | None = None) -> None:
    """单元格数字显示为常规阿拉伯数字（须配合列级 ``_strip_dbnum_column_styles``）。"""
    cell.number_format = DETAIL_ARABIC_NUMBER_DISPLAY


def _strip_dbnum_column_styles(ws, col_lo: int, col_hi: int) -> None:
    """去掉 ``column_dimensions`` 上的 ``[DBNum1]``，否则保存后仍按列样式显示中文数或 ``2.`` 等。"""
    for c in range(col_lo, col_hi + 1):
        letter = get_column_letter(c)
        dim = ws.column_dimensions.get(letter)
        if dim is None:
            continue
        fmt = str(dim.number_format or "")
        if "DBNum" in fmt or "dbnum" in fmt.lower():
            dim.number_format = DETAIL_ARABIC_NUMBER_DISPLAY


def clear_template_blocks(
    ws,
    base_rows: Iterable[int],
    proto_styles: dict[tuple[int, int], BlockCellStyle] | None = None,
) -> None:
    for base_row in base_rows:
        for row_idx in range(base_row, base_row + 6):
            for col_idx in range(5, DETAIL_TEMPLATE_SUMMARY_BEGIN_COL):
                cell = ws.cell(row_idx, col_idx)
                cell.value = None
                if proto_styles:
                    rel_r = row_idx - base_row + 1
                    st = proto_styles.get((rel_r, col_idx))
                    if st:
                        _apply_style_bundle(cell, st)
