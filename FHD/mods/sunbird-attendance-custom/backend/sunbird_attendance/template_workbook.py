"""Load template workbooks and rebuild their employee roster blocks."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from .template_cells import (
    _apply_style_bundle,
    _normalize_block_chinese_numerals,
    _snapshot_block_cell_styles,
)
from .template_schema import (
    DETAIL_HEADER_ROWS,
    DETAIL_PERSON_BLOCK_ROWS,
    DETAIL_SIDE_SUMMARY_NIGHT_COLS,
    DETAIL_SUM_COL_START,
    DETAIL_TEMPLATE_SUMMARY_BEGIN_COL,
    BlockCellStyle,
)


def _ensure_template_workbook(output_path: Path, template_path: Path | None = None) -> Workbook:
    if template_path is not None:
        return load_workbook(template_path)
    if output_path.exists():
        return load_workbook(output_path)
    wb = Workbook()
    ws = wb.worksheets[0]
    ws.title = "明细"
    ws["A1"] = "未提供模板，已生成基础明细页"
    return wb


def open_output_workbook(output_path: Path, template_path: Path | None = None) -> Workbook:
    return _ensure_template_workbook(output_path, template_path)


def _snapshot_first_person_block(
    ws, block_top: int
) -> tuple[dict[tuple[int, int], object], list[tuple[int, int, int, int]]]:
    """复制首个人员块（6 行）的单元格值与合并区域（相对块内行 1-6）。

    不含 ``DETAIL_TEMPLATE_SUMMARY_BEGIN_COL`` 及以右列，避免每人块粘贴时覆盖侧栏 SUMIF/夜班公式。
    """
    cap = DETAIL_TEMPLATE_SUMMARY_BEGIN_COL - 1
    vals: dict[tuple[int, int], object] = {}
    for dr in range(DETAIL_PERSON_BLOCK_ROWS):
        for c in range(1, cap + 1):
            vals[dr + 1, c] = ws.cell(block_top + dr, c).value
    rel_merges: list[tuple[int, int, int, int]] = []
    for rng in list(ws.merged_cells.ranges):
        min_col, min_row, max_col, max_row = rng.bounds
        if (
            min_row >= block_top
            and max_row <= block_top + DETAIL_PERSON_BLOCK_ROWS - 1
            and (min_col >= 1)
            and (max_col <= cap)
        ):
            rel_merges.append((min_col, min_row - block_top + 1, max_col, max_row - block_top + 1))
    return (vals, rel_merges)


def _paste_one_person_block(
    ws,
    block_top: int,
    proto_vals: dict[tuple[int, int], object],
    rel_merges: list[tuple[int, int, int, int]],
    department: str,
    nature: str,
    name: str,
    proto_styles: dict[tuple[int, int], BlockCellStyle] | None = None,
) -> None:
    for (dr, c), v in proto_vals.items():
        if c >= DETAIL_TEMPLATE_SUMMARY_BEGIN_COL:
            continue
        tgt = ws.cell(block_top + dr - 1, c)
        tgt.value = None if c >= 5 else v
        if proto_styles:
            _apply_style_bundle(tgt, proto_styles.get((dr, c)))
    for col, value in enumerate((department, nature, name), 1):
        cell = ws.cell(block_top, col, value)
        cell.data_type = "s"
    for min_c, r1, max_c, r2 in rel_merges:
        ref = f"{get_column_letter(min_c)}{block_top + r1 - 1}:{get_column_letter(max_c)}{block_top + r2 - 1}"
        try:
            ws.merge_cells(ref)
        except ValueError:
            pass
    _normalize_block_chinese_numerals(
        ws,
        block_top,
        DETAIL_SUM_COL_START,
        DETAIL_TEMPLATE_SUMMARY_BEGIN_COL - 1,
    )


def rebuild_detail_sheet_person_blocks(
    ws,
    people: list[tuple[str, str, str]],
    *,
    header_rows: int = DETAIL_HEADER_ROWS,
    prototype_block_top: int = 4,
) -> None:
    """Rebuild whole six-row blocks; preserve prototype merges and night formulas."""
    if not people:
        return
    if ws.max_row < prototype_block_top + DETAIL_PERSON_BLOCK_ROWS - 1:
        raise ValueError("明细表过短：缺少模板首个人员 6 行块")
    night_tpl = tuple(
        ws.cell(prototype_block_top, col).value for col in DETAIL_SIDE_SUMMARY_NIGHT_COLS
    )
    proto_vals, rel_merges = _snapshot_first_person_block(ws, prototype_block_top)
    proto_styles = _snapshot_block_cell_styles(ws, prototype_block_top)
    last = ws.max_row
    for merged in list(ws.merged_cells.ranges):
        if merged.max_row > header_rows:
            ws.unmerge_cells(str(merged))
    if last > header_rows:
        ws.delete_rows(header_rows + 1, last - header_rows)
    for i, (dept, nature, name) in enumerate(people):
        top = header_rows + 1 + i * DETAIL_PERSON_BLOCK_ROWS
        _paste_one_person_block(ws, top, proto_vals, rel_merges, dept, nature, name, proto_styles)
    first_body = header_rows + 1
    for col, val in zip(DETAIL_SIDE_SUMMARY_NIGHT_COLS, night_tpl):
        if val not in (None, ""):
            ws.cell(first_body, col).value = val
