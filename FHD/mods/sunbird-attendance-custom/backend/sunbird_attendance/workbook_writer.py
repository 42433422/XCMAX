"""Own the template workbook from validation through atomic output publication."""

from __future__ import annotations

from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from .attendance_month import (
    _aggregate_employee_records,
    _build_monthly_rows_for_template,
    _filter_records_to_template_roster,
)
from .detail_sheet import write_detail_sheet
from .monthly_sheet import write_monthly_sheet
from .parser import AttendanceDayRecord
from .roster_identity import PersonnelRoster, identify_roster_records, restore_display_names
from .template_profiles import build_template_profiles
from .template_workbook import open_output_workbook, rebuild_detail_sheet_person_blocks

WorkbookOpener = Callable[[Path, Path | None], Workbook]


def _retain_detail_and_monthly_sheets(workbook: Workbook) -> None:
    """Keep the detail and linked monthly totals in the delivered workbook."""
    if "明细" not in workbook.sheetnames:
        active = workbook.active
        if not isinstance(active, Worksheet):
            raise ValueError("明细页不是有效的工作表")
        active.title = "明细"
    for name in list(workbook.sheetnames):
        if name not in {"明细", "月度统计"}:
            del workbook[name]


def _publish_workbook(workbook: Workbook, output: Path) -> None:
    """Replace a previous delivery only after a complete workbook has been saved."""
    output.parent.mkdir(parents=True, exist_ok=True)
    with NamedTemporaryFile(
        dir=output.parent, prefix=f".{output.name}-", suffix=".tmp", delete=False
    ) as stream:
        temporary = Path(stream.name)
    try:
        workbook.save(temporary)
        temporary.replace(output)
    finally:
        temporary.unlink(missing_ok=True)


@dataclass(frozen=True)
class AttendanceWorkbookWriter:
    open_workbook: WorkbookOpener = open_output_workbook

    def write(
        self,
        records: list[AttendanceDayRecord],
        out: Path,
        template: Path | None,
        month_label: str,
        personnel_roster: PersonnelRoster | None,
    ) -> dict[str, Any]:
        display: dict[str, str] = {}
        if personnel_roster:
            records, personnel_roster, display = identify_roster_records(records, personnel_roster)
        workbook = self.open_workbook(out, template)
        try:
            detail_ws = workbook["明细"] if "明细" in workbook.sheetnames else workbook.active
            if not isinstance(detail_ws, Worksheet):
                raise ValueError("明细页不是有效的工作表")
            if personnel_roster:
                rebuild_detail_sheet_person_blocks(detail_ws, personnel_roster)
            template_profiles = build_template_profiles(detail_ws)
            if not template_profiles:
                return {
                    "success": False,
                    "error": "明细页未解析到任何员工块，请检查固定模板或人员管理名单",
                }
            filtered = _filter_records_to_template_roster(records, template_profiles)
            if not filtered and not personnel_roster:
                return {
                    "success": False,
                    "error": "钉钉数据与模板明细中的姓名无交集：请核对模板人员名单与「每日统计」姓名列是否一致（含空格/全半角）。",
                }
            employees, analysis_rows = _aggregate_employee_records(
                filtered, template_profiles=template_profiles
            )
            monthly_rows = _build_monthly_rows_for_template(employees, personnel_roster, detail_ws)
            template_result = write_detail_sheet(workbook, employees, month_label=month_label)
            write_monthly_sheet(workbook, monthly_rows, link_detail_side_totals=True)
            _retain_detail_and_monthly_sheets(workbook)
            restore_display_names(workbook, display)
            _publish_workbook(workbook, out)
            return {
                "success": True,
                "output": str(out),
                "month": month_label,
                "rows_used_for_template": len(filtered),
                "rows_stats": len(analysis_rows),
                "employees_total": len(employees),
                "employees_matched": template_result.matched_employee_count,
                "unmatched_names": [
                    display.get(name, name) for name in template_result.unmatched_employee_names
                ],
                "personnel_roster_count": len(personnel_roster) if personnel_roster else 0,
                "output_sheet_names": list(workbook.sheetnames),
            }
        finally:
            workbook.close()


def write_records(
    records: list[AttendanceDayRecord],
    out: Path,
    template: Path | None,
    month_label: str,
    personnel_roster: PersonnelRoster | None,
) -> dict[str, Any]:
    return AttendanceWorkbookWriter().write(records, out, template, month_label, personnel_roster)
