"""Business outcome assertions independent of planner action labels."""

from __future__ import annotations

from typing import Any


def check_spreadsheets(execution: dict[str, Any], assertions: list[dict]) -> tuple[bool, str]:
    from io import BytesIO

    from openpyxl import load_workbook

    from app.application.agent_orchestrator.artifact_files import read_verified_spreadsheet

    for assertion in assertions:
        accepted = False
        for artifact in execution.get("artifacts", []):
            if artifact.get("artifact_type") != "file":
                continue
            content = read_verified_spreadsheet(execution["run_id"], artifact)
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            try:
                if assertion["sheet"] not in workbook.sheetnames:
                    continue
                if "cells" in assertion:
                    cells = assertion["cells"]
                    if not cells:
                        raise ValueError("spreadsheet cell assertions must not be empty")
                    sheet = workbook[assertion["sheet"]]
                    accepted = all(
                        sheet[address].value == expected for address, expected in cells.items()
                    )
                    if accepted:
                        break
                    continue
                rows = list(workbook[assertion["sheet"]].values)
                if not rows or len(rows) - 1 != assertion["count"]:
                    continue
                records = [dict(zip(rows[0], row)) for row in rows[1:]]
                expected = assertion["includes"]
                if not expected:
                    raise ValueError("spreadsheet assertions require expected cell values")
                accepted = all(
                    any(
                        all(key in row and row[key] == value for key, value in item.items())
                        for row in records
                    )
                    for item in expected
                )
                if accepted:
                    break
            finally:
                workbook.close()
        if not accepted:
            return False, f"spreadsheet content differs: {assertion['sheet']}"
    return True, ""


def check_returned_records(execution: dict[str, Any], assertions: list[dict]) -> tuple[bool, str]:
    for assertion in assertions:
        matching = [
            step
            for step in execution.get("steps", [])
            if step.get("status") == "completed"
            and step.get("tool_id") == assertion["tool_id"]
            and step.get("action") == assertion["action"]
        ]
        accepted = False
        for step in matching:
            value = step.get("output")
            for key in assertion.get("path", ["data"]):
                value = value.get(key) if isinstance(value, dict) else None
            if "equals" in assertion:
                expected_value = assertion["equals"]
                if value == expected_value and isinstance(value, bool) == isinstance(
                    expected_value, bool
                ):
                    accepted = True
                    break
                continue
            if not isinstance(value, list):
                continue
            if "count" in assertion and len(value) != assertion["count"]:
                continue
            expected = assertion.get("includes") or []
            if not expected:
                raise ValueError("returned-record assertions need expected record contents")
            if all(
                any(
                    isinstance(row, dict)
                    and all(row.get(key) == field for key, field in item.items())
                    for row in value
                )
                for item in expected
            ):
                accepted = True
                break
        if not accepted:
            return False, f"returned records differ: {assertion['tool_id']}.{assertion['action']}"
    return True, ""


def seed_records(fixtures: list[dict]) -> None:
    from app.application.business_db_write_verification import _model_config
    from app.db import SessionLocal
    from app.infrastructure.tenant_scope import current_tenant_id

    with SessionLocal() as db:
        for fixture in fixtures:
            model, fields, _ = _model_config(fixture["entity"])
            values = {fields.get(key, key): value for key, value in fixture["values"].items()}
            if "tenant_id" in values:
                raise ValueError("fixture may not override the isolated trial tenant")
            db.add(model(tenant_id=current_tenant_id(), **values))
        db.commit()


def seed_sales_period_orders(rows: list[dict]) -> None:
    """Seed explicit month-boundary amounts in the trial's already-isolated DB."""
    from calendar import monthrange
    from datetime import datetime, timedelta
    from decimal import Decimal

    from app.db import SessionLocal
    from app.db.models.sales import SalesOrder, SalesOrderItem
    from app.infrastructure.tenant_scope import current_tenant_id

    now = datetime.now()
    first = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    last = first.replace(
        day=monthrange(first.year, first.month)[1],
        hour=23,
        minute=59,
        second=59,
        microsecond=999999,
    )
    instants = {
        "month_start": first,
        "month_end": last,
        "previous_month_end": first - timedelta(microseconds=1),
    }
    with SessionLocal() as db:
        for index, row in enumerate(rows):
            instant = instants[row["at"]]
            amount = Decimal(str(row["amount"]))
            order = SalesOrder(
                tenant_id=current_tenant_id(),
                order_no=f"BENCH-PERIOD-{index}",
                customer_name="月报测试客户",
                created_at=instant,
                total_amount=amount,
            )
            db.add(order)
            db.flush()
            db.add(
                SalesOrderItem(
                    tenant_id=current_tenant_id(),
                    order_id=order.id,
                    product_name="月报测试产品",
                    quantity=1,
                    unit_price=amount,
                    amount=amount,
                )
            )
        db.commit()


def seed_ledger_period_entries(rows: list[dict]) -> None:
    from calendar import monthrange
    from datetime import date, timedelta
    from decimal import Decimal

    from app.db import SessionLocal
    from app.db.models.accounting import JournalEntry, JournalEntryLine
    from app.infrastructure.tenant_scope import current_tenant_id

    first = date.today().replace(day=1)
    dates = {
        "month_start": first,
        "month_end": first.replace(day=monthrange(first.year, first.month)[1]),
        "previous_month_end": first - timedelta(days=1),
    }
    with SessionLocal() as db:
        for index, row in enumerate(rows):
            amount = Decimal(str(row["amount"]))
            entry = JournalEntry(
                tenant_id=current_tenant_id(),
                entry_no=f"BENCH-LEDGER-{index}",
                journal_date=dates[row["at"]],
                status="posted",
                debit_total=amount,
                credit_total=amount,
            )
            entry.lines = [
                JournalEntryLine(
                    tenant_id=current_tenant_id(),
                    account_code="1001",
                    account_name="库存现金",
                    debit=amount,
                    credit=0,
                ),
                JournalEntryLine(
                    tenant_id=current_tenant_id(),
                    account_code="6001",
                    account_name="主营业务收入",
                    debit=0,
                    credit=amount,
                ),
            ]
            db.add(entry)
        db.commit()


def seed_inventory_quantities(rows: list[dict]) -> None:
    from app.db import SessionLocal
    from app.db.models.inventory import InventoryLedger, Warehouse
    from app.db.models.product import Product
    from app.infrastructure.tenant_scope import current_tenant_id

    if not rows:
        return
    with SessionLocal() as db:
        warehouse = Warehouse(tenant_id=current_tenant_id(), code="BENCH-STOCK", name="测试仓库")
        db.add(warehouse)
        db.flush()
        for row in rows:
            product = Product(
                tenant_id=current_tenant_id(),
                name=row["model_number"],
                model_number=row["model_number"],
            )
            db.add(product)
            db.flush()
            db.add(
                InventoryLedger(
                    tenant_id=current_tenant_id(),
                    product_id=product.id,
                    warehouse_id=warehouse.id,
                    quantity=row["quantity"],
                    available_quantity=row["quantity"],
                    reserved_quantity=0,
                )
            )
        db.commit()


def assertion_model_config(entity: str):
    if entity in {"inventory_ledgers", "inventory_transactions"}:
        from app.db.models import InventoryLedger, InventoryTransaction

        return (InventoryLedger if entity == "inventory_ledgers" else InventoryTransaction), {}, ()
    if entity in {"sales_orders", "sales_order_items"}:
        from app.db.models.sales import SalesOrder, SalesOrderItem

        return (SalesOrder if entity == "sales_orders" else SalesOrderItem), {}, ()
    if entity == "financial_transactions":
        from app.db.models.finance import FinancialTransaction

        return FinancialTransaction, {}, ()
    from app.application.business_db_write_verification import _model_config

    return _model_config(entity)
