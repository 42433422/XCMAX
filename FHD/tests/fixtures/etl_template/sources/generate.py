"""Generate project-owned synthetic Excel fixtures; never use customer data."""

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font


def sample():
    wb = Workbook()
    ws = wb.active
    ws.title = "报价"
    ws.merge_cells("A1:C1")
    ws["A1"] = "固定说明：付款后交货"
    ws["A2"] = "客户"
    ws["B2"] = "示例客户"
    ws["A3"] = "数量"
    ws["B3"] = 42
    ws["C3"] = "=B3*2"
    ws["A4"] = "模糊说明"
    ws.column_dimensions["B"].width = 27
    ws.row_dimensions[3].height = 31
    ws.print_options.horizontalCentered = True
    ws.print_area = "A1:C4"
    ws["B3"].number_format = "0.00"
    ws["A1"].font = Font(bold=True)
    out = BytesIO()
    wb.save(out)
    wb.close()
    return out.getvalue()


def large_quotation(rows=601):
    wb = Workbook()
    ws = wb.active
    ws.title = "报价"
    for row in range(1, rows + 1):
        ws.cell(row, 1, f"记录{row}")
    ws["A1"] = "报价单"
    out = BytesIO()
    wb.save(out)
    wb.close()
    return out.getvalue()


if __name__ == "__main__":
    Path(__file__).with_name("synthetic_quotation.xlsx").write_bytes(sample())
