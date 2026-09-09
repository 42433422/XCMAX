"""Export bytes survive the API bridge and stay bound to their account/Mod."""

import hashlib
import importlib
import json
import logging
from io import BytesIO

import pytest
from fastapi.responses import Response
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.application.aiopen import api_artifacts
from app.application.aiopen.service import _tool_api_call, generate_api_key
from app.db.models.user import Session
from tests.test_application.test_aiopen_api_execution import application as application
from tests.test_application.test_aiopen_api_execution import caller
from tests.test_application.test_aiopen_api_execution import host as host
from tests.test_application.test_aiopen_screen_identity import request_for


@pytest.fixture
def export_app(application, monkeypatch, tmp_path):
    app, _ = application
    monkeypatch.setattr(
        "app.utils.path_io.path_utils.get_data_dir", lambda: str(tmp_path / "private-data")
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["客户", "数量"])
    for index in range(200):
        sheet.append([f"客户-{index}", index])
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    content = buffer.getvalue()

    @app.get("/api/test/export")
    def export():
        return Response(
            content,
            media_type="application/octet-stream",
            headers={"Content-Disposition": "attachment; filename*=UTF-8''%E5%AE%A2%E6%88%B7.xlsx"},
        )

    return app, content


def test_exact_binary_download_and_cross_account_denial(export_app):
    app, content = export_app
    key = generate_api_key(request=request_for(**{"X-Session-Id": "login-3"}))["key"]
    with caller({"X-AIOPEN-Key": key}):
        result = _tool_api_call(app, {"path": "/api/test/export"})
    assert result["success"], result
    receipt = result["artifacts"][0]
    assert receipt["name"] == "客户.xlsx"
    assert (
        receipt["size"] == len(content) and receipt["sha256"] == hashlib.sha256(content).hexdigest()
    )
    assert "private-data" not in json.dumps(result)
    importlib.reload(api_artifacts)  # Receipt must not depend on process-local maps.
    with TestClient(app) as client:
        for headers in ({}, {"X-Session-Id": "login-4"}, {"X-Session-Id": "login-5"}):
            assert client.get(receipt["download_url"], headers=headers).status_code == 404
        response = client.get(receipt["download_url"], headers={"X-AIOPEN-Key": key})
        assert response.status_code == 200 and response.content == content
        downloaded = load_workbook(BytesIO(response.content), read_only=True)
        assert downloaded.active["A201"].value == "客户-199"
        assert downloaded.active["B201"].value == 199
        downloaded.close()
        assert response.headers["x-content-sha256"] == receipt["sha256"]
        assert response.headers["cache-control"] == "private, no-store"
        assert response.headers["content-disposition"].startswith("attachment;")


def test_download_rechecks_mod_entitlement_and_expiry(export_app, host):
    app, _ = export_app
    with caller({"X-Session-Id": "login-3"}):
        result = _tool_api_call(app, {"path": "/api/test/export", "mod_id": "mod-a"})
    receipt = result["artifacts"][0]
    with host.begin() as db:
        db.get(Session, 3).entitled_mod_ids_json = "[]"
    with TestClient(app) as client:
        assert client.get(receipt["uri"], headers={"X-Session-Id": "login-3"}).status_code == 404
    with host.begin() as db:
        db.get(Session, 3).entitled_mod_ids_json = '["mod-a"]'
    manifest = api_artifacts._root() / f"{receipt['artifact_id']}.json"
    metadata = json.loads(manifest.read_text())
    metadata["expires_at"] = 0
    manifest.write_text(json.dumps(metadata))
    with TestClient(app) as client:
        assert client.get(receipt["uri"], headers={"X-Session-Id": "login-3"}).status_code == 404
    with caller({"X-Session-Id": "login-3"}):
        assert _tool_api_call(app, {"path": "/api/test/export"})["success"]
    assert not manifest.exists()
    assert not manifest.with_suffix(".bin").exists()


def test_corrupted_or_oversized_export_never_reports_complete(export_app, monkeypatch):
    app, _ = export_app
    with caller({"X-Session-Id": "login-3"}):
        result = _tool_api_call(app, {"path": "/api/test/export"})
    receipt = result["artifacts"][0]
    (api_artifacts._root() / f"{receipt['artifact_id']}.bin").write_bytes(b"damaged")
    with TestClient(app) as client:
        assert client.get(receipt["uri"], headers={"X-Session-Id": "login-3"}).status_code == 404
    monkeypatch.setattr(api_artifacts, "MAX_BYTES", 1)
    with caller({"X-Session-Id": "login-3"}):
        result = _tool_api_call(app, {"path": "/api/test/export"})
    assert not result["success"] and result["code"] == "API_EXPORT_FAILED"


def test_export_receipt_attaches_to_actual_owner_task(export_app, monkeypatch, host):
    from app.application.agent_orchestrator.chat_trace import create_chat_trace_run
    from app.application.agent_orchestrator.run_sql_repository import SQLAlchemyAgentRunRepository
    from app.fastapi_routes.aiopen_route_support import trace_tool_call

    app, _ = export_app
    with caller({"X-Session-Id": "login-3"}):
        result = _tool_api_call(app, {"path": "/api/test/export"})
    repo = SQLAlchemyAgentRunRepository(session_factory=host)
    monkeypatch.setattr(
        "app.application.agent_orchestrator.chat_trace.get_agent_run_repository", lambda: repo
    )
    run_id = trace_tool_call(
        route="/api/aiopen/invoke",
        channel="aiopen_invoke",
        tool="api_call",
        args={"path": "/api/test/export"},
        result=result,
        user_id="forged-account",
        create_trace_run=create_chat_trace_run,
        logger=logging.getLogger(__name__),
    )
    run = SQLAlchemyAgentRunRepository(session_factory=host).get(run_id)
    assert run is not None and run.user_id == "3"
    assert run.metadata["runtime_context"]["tenant_id"] == "7"
    assert len(run.artifacts) == 1
    assert run.artifacts[0].uri == result["artifacts"][0]["uri"]
    assert run.artifacts[0].metadata["sha256"] == result["artifacts"][0]["sha256"]


@pytest.mark.parametrize(
    "content,media",
    [
        (b"\xff\x00\x91", None),
        ("完整文本".encode() * 3000, "text/plain"),
        (b"a,b\n1,2", "text/csv"),
    ],
)
def test_unlabelled_binary_and_long_text_are_not_truncated(export_app, content, media):
    app, _ = export_app

    @app.get("/api/test/unlabelled")
    def unlabelled():
        return Response(content, media_type=media)

    with caller({"X-Session-Id": "login-3"}):
        result = _tool_api_call(app, {"path": "/api/test/unlabelled"})
    assert result["success"] and result["artifacts"][0]["size"] == len(content)
    with TestClient(app) as client:
        response = client.get(result["artifacts"][0]["uri"], headers={"X-Session-Id": "login-3"})
        assert response.content == content
