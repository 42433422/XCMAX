"""Customer exchange entrypoints use real account, workbook and ETL services."""

import importlib.util
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.application.aiopen.api_artifacts import read_api_export
from app.application.aiopen.service import AIOPEN_STATE, _tool_api_call
from app.db.models.purchase_unit import PurchaseUnit
from app.infrastructure.tenant_scope import tenant_scope
from tests.test_application.test_aiopen_api_artifacts import export_app as export_app
from tests.test_application.test_aiopen_api_execution import application as application
from tests.test_application.test_aiopen_api_execution import caller
from tests.test_application.test_aiopen_api_execution import host as host
from tests.test_application.test_aiopen_etl_roundtrip import await_run, call
from tests.test_application.test_aiopen_etl_roundtrip import etl_app as etl_app


@pytest.mark.parametrize("via_mod", [False, True])
def test_customer_import_preview_and_execution_keep_target_mod(
    etl_app, application, monkeypatch, via_mod
):
    app, _, _ = etl_app
    _, factories = application
    prefix = "/api/customers"
    mod_id = "mod-a" if via_mod else ""
    if via_mod:
        path = (
            Path(__file__).resolve().parents[2]
            / "mods/xcagi-erp-domain-bridge/backend/blueprints.py"
        )
        spec = importlib.util.spec_from_file_location("customer_exchange_bridge_fixture", path)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        module.register_fastapi_routes(app, "mod-a")
        prefix = "/api/mod/mod-a/customers"
        monkeypatch.setitem(AIOPEN_STATE["whitelist"], prefix, True)
    with caller({"X-Session-ID": "login-3"}):
        artifact = _tool_api_call(app, {"path": "/api/test/export"})["artifacts"][0]
        result = _tool_api_call(
            app,
            {
                "path": prefix + "/import",
                "method": "POST",
                "mod_id": mod_id,
                "files": [{"field": "file", "artifact_id": artifact["artifact_id"]}],
            },
        )
        assert result["success"] and result["status_code"] == 202, result
        assert result["data"]["data"]["requires_confirmation"] is True
        run_id = result["data"]["data"]["run_id"]
        from app.request_active_mod_ctx import (
            reset_request_active_mod_id,
            set_request_active_mod_id,
        )

        token = set_request_active_mod_id(mod_id)
        try:
            run = await_run(app, run_id, "preview_ready")
            assert run["summary"]["new"] == 200, run
            for factory in factories.values():
                with tenant_scope(7), factory() as db:
                    assert db.query(PurchaseUnit).count() == 1
            call(app, f"/api/etl/runs/{run_id}/execute", method="POST", body={"confirmed": True})
            await_run(app, run_id, "completed")
        finally:
            reset_request_active_mod_id(token)
    for name, factory in factories.items():
        with tenant_scope(7), factory() as db:
            expected = 201 if name == (mod_id or "host-business") else 1
            assert db.query(PurchaseUnit).count() == expected


def test_actual_customer_export_is_owned_unique_and_not_shadowed(etl_app, application):
    app, _, factory = etl_app
    with tenant_scope(7), factory.begin() as db:
        row = db.get(PurchaseUnit, 7)
        row.unit_name = '=HYPERLINK("https://invalid.test")'
        row.contact_phone = "0013800000000"
    with caller({"X-Session-ID": "login-3"}):
        first = _tool_api_call(app, {"path": "/api/customers/export"})
        second = _tool_api_call(app, {"path": "/api/customers/export"})
        assert first["success"] and second["success"], (first, second)
        one, two = first["artifacts"][0], second["artifacts"][0]
        assert one["name"] != two["name"]
        content, _ = read_api_export(one["artifact_id"])
        workbook = load_workbook(BytesIO(content))
        rows = list(workbook.active.values)
        assert len(rows) == 2 and rows[1][0] == 7
        assert rows[1][1].startswith("'=") and rows[1][3] == "0013800000000"
        assert workbook.active["B2"].data_type == "s"
        workbook.close()
        missing = _tool_api_call(app, {"path": "/api/customers/export?template_id=missing"})
        assert not missing["success"] and missing["status_code"] == 400, missing
    with caller({"X-Session-ID": "login-5"}):
        other = _tool_api_call(app, {"path": "/api/customers/export"})
        assert other["success"], other
        raw, _ = read_api_export(other["artifacts"][0]["artifact_id"])
        workbook = load_workbook(BytesIO(raw))
        assert list(workbook.active.values)[1][1] == "host-business-8"
        workbook.close()
