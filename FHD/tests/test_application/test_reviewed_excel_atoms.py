"""Source-atom rectangles: ordinary internal blanks are allowed, hidden content is not."""

import json
from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from app.application.etl.errors import ReviewedExcelError
from app.application.reviewed_excel_classification import _classify_semantic_batch
from app.application.reviewed_excel_template import inspect_workbook


def source(cells, merges=(), *, second=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "O'Brien ! 表"
    for address in cells:
        sheet[address] = "source"
    for area in merges:
        sheet.merge_cells(area)
    if second:
        book.create_sheet("第二")["A1"] = "source"
    stream = BytesIO()
    book.save(stream)
    return inspect_workbook(stream.getvalue())


def region(*atoms, primary=None, dependencies=()):
    return {
        "atom_ids": list(atoms),
        "primary_id": primary or atoms[0],
        "dependency_ids": list(dependencies),
        "reason": "synthetic source membership",
    }


def answer(*regions, documents=(), layout=()):
    return {
        "physical_regions": list(regions),
        "logical_documents": list(documents),
        "layout_references": list(layout),
    }


def execute(monkeypatch, review, raw):
    original = deepcopy(raw)
    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    monkeypatch.setattr(
        "app.application.reviewed_excel_template._template_mimo_complete",
        lambda *args: SimpleNamespace(data=raw, model="synthetic-no-network"),
    )
    try:
        return _classify_semantic_batch(review, reference_ids=True)
    finally:
        assert raw == original


def test_atoms_allow_existing_internal_blank_but_never_add_outer_padding(monkeypatch):
    review = source(["A1", "C1", "E1"])
    raw = answer(
        region("S1:A1", "S1:C1"),
        region("S1:E1"),
        documents=[{"document_id": "d", "fragments": ["S1:A1", "S1:E1"]}],
    )
    result = execute(monkeypatch, review, raw)
    assert [r["range"] for r in result["regions"]] == ["A1:C1", "E1:E1"]
    assert result["logical_documents"][0]["fragments"][0]["range"] == "A1:C1"
    assert result["input_summary"]["reference_response"] == raw
    assert result["input_summary"]["reference_index"]["S1:C1"]["atom_range"] == "C1"


@pytest.mark.parametrize("mode", ["unselected", "context", "unread"])
def test_atom_rectangle_rejects_nonempty_holes_without_filling_them(monkeypatch, mode):
    review = source(["A1", "B1", "C1"])
    if mode != "unselected":
        review["full_cells"] = deepcopy(review["cells"])
        if mode == "context":
            review["cells"][1]["context_only"] = True
        else:
            review["cells"].pop(1)
    raw = answer(region("S1:A1", "S1:C1"))
    with pytest.raises(ReviewedExcelError, match="未选主体、上下文或未发送"):
        execute(monkeypatch, review, raw)


@pytest.mark.parametrize(
    "raw,error",
    [
        (answer(region("S1:A1", "S1:Z99")), "白名单"),
        (answer(region("S1:A1", "S1:A1")), "重复分配"),
        (answer(region("S1:A1"), region("S1:A1", "S1:B1")), "重复分配"),
        (answer(region("S1:A1")), "未完整分配"),
        (answer(region("S1:A1", primary="S1:B1"), region("S1:B1")), "自身主体"),
        (answer(region("S1:A1", "S1:B1", dependencies=["S1:A1", "S1:A1"])), "依赖原子重复"),
    ],
)
def test_atom_references_fail_closed(monkeypatch, raw, error):
    with pytest.raises(ReviewedExcelError, match=error):
        execute(monkeypatch, source(["A1", "B1"]), raw)


def test_atoms_reject_cross_sheet_and_context_as_primary(monkeypatch):
    review = source(["A1"], second=True)
    with pytest.raises(ReviewedExcelError, match="跨工作表"):
        execute(monkeypatch, review, answer(region("S1:A1", "S2:A1")))
    review = source(["A1", "C1"])
    review["full_cells"] = deepcopy(review["cells"])
    review["cells"][1]["context_only"] = True
    with pytest.raises(ReviewedExcelError, match="白名单"):
        execute(monkeypatch, review, answer(region("S1:A1", "S1:C1")))
    result = execute(monkeypatch, review, answer(region("S1:A1", dependencies=["S1:C1"])))
    assert result["regions"][0]["dependencies"][0]["cell"] == "C1"


def test_merge_anchor_atom_decodes_full_extent_and_preserves_audit(monkeypatch):
    review = source(["B2", "D3"], ["D3:E3"])
    raw = answer(region("S1:B2", "S1:D3"))
    result = execute(monkeypatch, review, raw)
    assert result["regions"][0]["range"] == "B2:E3"
    assert result["input_summary"]["reference_index"]["S1:D3"]["atom_range"] == "D3:E3"
    forged = deepcopy(raw)
    forged["physical_regions"][0]["range"] = "B2:D3"
    with pytest.raises(ReviewedExcelError, match="白名单"):
        execute(monkeypatch, review, forged)


def test_blank_merge_must_be_explicit_to_extend_a_candidate(monkeypatch):
    review = source(["A1", "D3"], ["C2:E2"])
    with pytest.raises(ReviewedExcelError, match="切断合并.*不能自动扩框"):
        execute(monkeypatch, review, answer(region("S1:A1", "S1:D3")))
    raw = answer(region("S1:A1", "S1:D3", "M1:C2:E2"), layout=["M1:C2:E2"])
    result = execute(monkeypatch, review, raw)
    assert result["regions"][0]["range"] == "A1:E3"
    assert result["layout_references"] == [{"sheet": "O'Brien ! 表", "range": "C2:E2"}]
    assert result["input_summary"]["reference_index"]["M1:C2:E2"]["layout"] is True


def test_blank_layout_cannot_smuggle_unselected_content_or_form_a_primary(monkeypatch):
    review = source(["A1", "C1"], ["A2:C2"])
    with pytest.raises(ReviewedExcelError, match="未选主体"):
        execute(monkeypatch, review, answer(region("S1:A1", "M1:A2:C2"), region("S1:C1")))
    with pytest.raises(ReviewedExcelError, match="自身主体"):
        execute(monkeypatch, review, answer(region("M1:A2:C2", primary="S1:A1")))


def test_nonblank_merge_layout_id_is_not_a_substitute_for_primary(monkeypatch):
    review = source(["A1"], ["A1:C1"])
    with pytest.raises(ReviewedExcelError, match="白名单"):
        execute(monkeypatch, review, answer(region("S1:A1", "M1:A1:C1")))
    result = execute(monkeypatch, review, answer(region("S1:A1"), layout=["M1:A1:C1"]))
    assert result["regions"][0]["range"] == "A1:C1"


@pytest.mark.parametrize(
    "fragments,error", [(["S1:Z99"], "未知主体"), (["S1:A1", "S1:A1"], "重复")]
)
def test_documents_reference_only_decoded_region_primary_ids(monkeypatch, fragments, error):
    with pytest.raises(ReviewedExcelError, match=error):
        execute(
            monkeypatch,
            source(["A1"]),
            answer(region("S1:A1"), documents=[{"document_id": "d", "fragments": fragments}]),
        )


def test_atom_contract_has_no_free_geometry_and_keeps_existing_region_limits():
    prepared = _classify_semantic_batch(
        source(["A1"], ["B2:C2"]), reference_ids=True, prepare_only=True
    )
    payload = json.loads(prepared["messages"][1]["content"])
    schema = prepared["schema"]["properties"]
    props = schema["physical_regions"]["items"]["properties"]
    assert not {"sheet", "range", "evidence_ids"} & props.keys()
    assert props["atom_ids"]["items"]["enum"] == ["S1:A1", "M1:B2:C2"]
    assert props["atom_ids"]["maxItems"] == len(props["atom_ids"]["items"]["enum"])
    assert schema["physical_regions"]["maxItems"] == 100
    assert schema["logical_documents"]["items"]["properties"]["fragments"]["maxItems"] == 100
    assert payload["output_schema"] == prepared["schema"]


def test_unsent_nonempty_merge_anchor_is_never_offered_as_blank_atom():
    review = source(["A1", "D2"], ["D2:E2"])
    review["full_cells"] = deepcopy(review["cells"])
    review["cells"] = review["cells"][:1]
    request = _classify_semantic_batch(review, reference_ids=True, prepare_only=True)
    props = request["schema"]["properties"]
    atoms = props["physical_regions"]["items"]["properties"]["atom_ids"]["items"]["enum"]
    assert atoms == ["S1:A1"]
    assert props["layout_references"]["items"]["enum"] == ["M1:D2:E2"]


def test_atom_rectangles_still_reject_overlap_only_in_source_blank_cells(monkeypatch):
    review = source(["A2", "C2", "B1", "B3"])
    with pytest.raises(ReviewedExcelError, match="重叠"):
        execute(monkeypatch, review, answer(region("S1:A2", "S1:C2"), region("S1:B1", "S1:B3")))


def test_layout_atoms_do_not_reduce_existing_500_primary_capacity(monkeypatch):
    from openpyxl.utils.cell import get_column_letter

    cells = [f"{get_column_letter(col)}{row}" for row in range(2, 22) for col in range(1, 26)]
    book = Workbook()
    book.active.title = "S"
    for cell in cells:
        book.active[cell] = "x"
    book.active.merge_cells("Z1:AA1")
    stream = BytesIO()
    book.save(stream)
    review = inspect_workbook(stream.getvalue())
    atoms = [f"S1:{cell}" for cell in cells] + ["M1:Z1:AA1"]
    assert len(review["cells"]) == 500
    result = execute(monkeypatch, review, answer(region(*atoms)))
    assert result["regions"][0]["range"] == "A1:AA21"
