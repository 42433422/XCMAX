from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.application.etl.errors import ReviewedExcelError
from app.application.reviewed_excel_auto_slice import derive_classified_documents
from app.application.reviewed_excel_template import inspect_workbook


def workbook():
    book = Workbook()
    sheet = book.active
    sheet.title = "One"
    sheet["A1"] = "第一份"
    sheet.merge_cells("A1:B1")
    sheet["A2"] = 7
    sheet["B2"] = "=A2*2"
    sheet["A5"] = "第二份"
    other = book.create_sheet("Two")
    other["A1"] = "第三份"
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def classification(body):
    return {
        "source_sha256": inspect_workbook(body)["sha256"],
        "model_status": "returned",
        "regions": [
            {"sheet": "One", "range": "A1:B2", "dependencies": []},
            {"sheet": "One", "range": "A5:B5", "dependencies": []},
            {"sheet": "Two", "range": "A1:A1", "dependencies": []},
        ],
        "suggested_documents": [
            {"fragments": [{"sheet": "One", "range": "A1:B2"}]},
            {"fragments": [{"sheet": "One", "range": "A5:B5"}]},
            {"fragments": [{"sheet": "Two", "range": "A1:A1"}]},
        ],
    }


def test_ungrouped_regions_block_and_model_groups_preserve_formula_and_merge():
    body = workbook()
    uncertain = classification(body)
    uncertain["suggested_documents"] = []
    with pytest.raises(ReviewedExcelError, match="独立单据关系"):
        derive_classified_documents(body, uncertain)
    result = derive_classified_documents(body, classification(body))
    assert len(result) == 3
    assert all(item["basis"] == "model_group" and not item["confirmed"] for item in result)
    first = load_workbook(BytesIO(result[0]["workbook"]))
    assert first.sheetnames == ["One"]
    assert first["One"]["B2"].value == "=A2*2"
    assert "A1:B1" in {str(r) for r in first["One"].merged_cells.ranges}
    assert first["One"]["A5"].value is None
    first["One"]["A2"] = 9
    stream = BytesIO()
    first.save(stream)
    assert load_workbook(BytesIO(stream.getvalue()))["One"]["A2"].value == 9


def test_model_group_is_not_claimed_user_confirmed():
    body = workbook()
    result = classification(body)
    result["suggested_documents"] = [
        {"fragments": [{"sheet": "One", "range": "A1:B2"}, {"sheet": "One", "range": "A5:B5"}]},
        {"fragments": [{"sheet": "Two", "range": "A1:A1"}]},
    ]
    slices = derive_classified_documents(body, result)
    assert len(slices) == 2 and slices[0]["basis"] == "model_group" and not slices[0]["confirmed"]
    assert load_workbook(BytesIO(slices[0]["workbook"]))["One"]["A5"].value == "第二份"


def test_rejects_external_dependency_duplicate_group_and_source_mismatch():
    body = workbook()
    result = classification(body)
    result["regions"][0]["dependencies"] = [{"sheet": "One", "cell": "A5"}]
    with pytest.raises(ReviewedExcelError, match="外部"):
        derive_classified_documents(body, result)
    result = classification(body)
    result["suggested_documents"] = [{"fragments": [{"sheet": "One", "range": "A1:B2"}]}] * 2
    with pytest.raises(ReviewedExcelError, match="重复"):
        derive_classified_documents(body, result)
    result = classification(body)
    result["source_sha256"] = "wrong"
    with pytest.raises(ReviewedExcelError, match="原件"):
        derive_classified_documents(body, result)


def test_rejects_cross_sheet_model_group_and_cut_merged_cell():
    body = workbook()
    result = classification(body)
    result["suggested_documents"] = [
        {"fragments": [{"sheet": "One", "range": "A1:B2"}, {"sheet": "Two", "range": "A1:A1"}]},
        {"fragments": [{"sheet": "One", "range": "A5:B5"}]},
    ]
    with pytest.raises(ReviewedExcelError, match="跨工作表"):
        derive_classified_documents(body, result)
    result = classification(body)
    result["regions"][0]["range"] = "A1:A2"
    result["suggested_documents"][0]["fragments"][0]["range"] = "A1:A2"
    with pytest.raises(ReviewedExcelError, match="合并单元格"):
        derive_classified_documents(body, result)


def test_chinese_names_are_bounded_unique_and_neutral_if_no_title():
    from app.application.reviewed_excel_auto_slice import _slice_filename

    used = set()
    part = [{"sheet": "One", "range": "A1:B2"}]
    cells = [{"sheet": "One", "cell": "A1", "value": "生产单"}]
    assert _slice_filename(part, cells, used) == "生产单.xlsx"
    assert _slice_filename(part, cells, used) == "生产单_2.xlsx"
    assert _slice_filename(part, [], used) == "切片_1.xlsx"
    assert _slice_filename(part, [], used) == "切片_2.xlsx"
    assert (
        _slice_filename(part, [{"sheet": "One", "cell": "A1", "value": "../客户/生产单"}], used)
        == "切片_3.xlsx"
    )


@pytest.mark.parametrize("count,accepted", [(19999, True), (20000, True), (20001, False)])
def test_workbook_nonempty_cell_limit_and_full_batch_coverage(count, accepted):
    from app.application.reviewed_excel_classification import plan_semantic_batches

    book = Workbook()
    sheet = book.active
    for row in range(1, count + 1):
        sheet.cell(row, 1, row)
    stream = BytesIO()
    book.save(stream)
    if not accepted:
        with pytest.raises(ReviewedExcelError, match="超过 20000"):
            inspect_workbook(stream.getvalue())
        return
    review = inspect_workbook(stream.getvalue())
    assert len(review["cells"]) == count
    batches, budget = plan_semantic_batches(review, 4)
    primary = [
        (c["sheet"], c["cell"]) for batch in batches for c in batch if not c.get("context_only")
    ]
    assert len(primary) == len(set(primary)) == count
    from app.application.reviewed_excel_evidence import MAX_MODEL_CELLS
    assert all(len(batch) <= MAX_MODEL_CELLS for batch in batches)
    assert budget == 4


def test_second_stage_requires_complete_coverage_and_evidence_for_each_fragment():
    from app.application.reviewed_excel_grouping import prepare_grouping, validate_grouping

    body = workbook()
    review = inspect_workbook(body)
    classification_result = classification(body)
    classification_result["coverage"] = [
        {"sheet": "One", "unread": [], "unassigned": [], "merged_omitted": 0},
        {"sheet": "Two", "unread": [], "unassigned": [], "merged_omitted": 0},
    ]
    for region in classification_result["regions"]:
        region.update(labels=["unknown"], evidence=[], reason="需要核对")
    proposal = prepare_grouping(review, classification_result)
    parts = [{"sheet": "One", "range": "A1:B2"}, {"sheet": "One", "range": "A5:B5"}]
    group = {
        "document_id": "doc1",
        "fragments": parts,
        "reason": "模型建议续表",
        "evidence": [{"sheet": "One", "cell": "A1"}, {"sheet": "One", "cell": "A5"}],
    }
    valid = validate_grouping(proposal, {"groups": [group]})
    assert valid[0]["status"] == "model_suggested" and valid[0]["fragments"] == parts
    with pytest.raises(ReviewedExcelError, match="片段内证据"):
        validate_grouping(proposal, {"groups": [{**group, "evidence": group["evidence"][:1]}]})
    with pytest.raises(ReviewedExcelError, match="来源无效或重复"):
        validate_grouping(proposal, {"groups": [group, {**group, "document_id": "doc2"}]})
    classification_result["coverage"][0]["unread"] = ["A2"]
    with pytest.raises(ReviewedExcelError, match="覆盖不完整"):
        prepare_grouping(review, classification_result)


def test_classified_document_export_preserves_formula_support():
    book = Workbook()
    sheet = book.active
    sheet.title = "生产"
    sheet["A1"] = "生产单"
    sheet["A2"] = "=B2*C2"
    sheet["B2"] = 3
    sheet["C2"] = 5
    sheet["D2"] = "无关信息"
    stream = BytesIO()
    book.save(stream)
    body = stream.getvalue()
    result = derive_classified_documents(
        body,
        {
            "source_sha256": inspect_workbook(body)["sha256"],
            "model_status": "returned",
            "regions": [{"sheet": "生产", "range": "A1:B2", "dependencies": []}],
            "suggested_documents": [{"fragments": [{"sheet": "生产", "range": "A1:B2"}]}],
        },
    )
    saved = load_workbook(BytesIO(result[0]["workbook"]), data_only=False)
    assert result[0]["filename"] == "生产单.xlsx"
    assert saved.sheetnames == ["生产", "依赖支撑"]
    assert saved["生产"]["A2"].value == "=B2*'依赖支撑'!C2"
    assert saved["依赖支撑"]["C2"].value == 5
    assert saved["依赖支撑"]["D2"].value is None


def test_unresolved_open_fragment_cannot_be_exported_as_a_finished_document():
    body = workbook()
    result = classification(body)
    result["regions"][0]["open_fragment"] = True
    result["regions"][0]["model_original_range"] = "A1:B5"
    with pytest.raises(ReviewedExcelError, match="未读边界"):
        derive_classified_documents(body, result)


def test_prepare_grouping_includes_middle_structure_and_cross_fragment_formulas():
    """Final grouping input must expose middle rows and cross-fragment formulas, not only ends."""
    from app.application.reviewed_excel_grouping import prepare_grouping

    book = Workbook()
    sheet = book.active
    sheet.title = "综合"
    sheet["A103"] = "表头甲"
    sheet["B103"] = "表头乙"
    for row in range(104, 125):
        sheet[f"A{row}"] = f"上段-{row}"
        sheet[f"B{row}"] = row
    sheet["A125"] = "续表标题"
    sheet["B125"] = "数量"
    for row in range(126, 140):
        sheet[f"A{row}"] = f"下段-{row}"
        sheet[f"B{row}"] = row
    sheet["C139"] = "=SUM(B104:B124)"
    stream = BytesIO()
    book.save(stream)
    body = stream.getvalue()
    review = inspect_workbook(body)
    classification = {
        "source_sha256": review["sha256"],
        "coverage": [
            {
                "sheet": "综合",
                "unread": [],
                "unassigned": [],
                "merged_omitted": 0,
            }
        ],
        "regions": [
            {
                "sheet": "综合",
                "range": "A103:K124",
                "evidence": [{"sheet": "综合", "cell": "A103"}],
                "reason": "上段",
            },
            {
                "sheet": "综合",
                "range": "A125:G203",
                "evidence": [{"sheet": "综合", "cell": "A125"}],
                "reason": "下段",
            },
        ],
    }
    proposal = prepare_grouping(review, classification)
    cells = {(c["sheet"], c["cell"]) for c in proposal["cells"]}
    # Middle rows from the upper fragment must be present, not only A103 and the last row.
    assert ("综合", "A103") in cells
    assert any(cell.startswith("A11") or cell.startswith("A12") for _, cell in cells)
    assert ("综合", "A125") in cells
    assert ("综合", "C139") in cells
    structure = {(item["sheet"], item["range"]): item for item in proposal["fragment_structure"]}
    assert structure[("综合", "A103:K124")]["nonempty_cells"] >= 2
    assert structure[("综合", "A125:G203")]["has_formulas"] is True
    links = proposal["cross_fragment_formulas"]
    assert any(
        link["from_fragment"] == "A125:G203" and link["to_fragment"] == "A103:K124"
        for link in links
    )
    # Evidence is descriptive only: prepare_grouping must not invent merged groups.
    assert len(proposal["fragments"]) == 2
