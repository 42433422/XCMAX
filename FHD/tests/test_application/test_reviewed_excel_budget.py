"""Synthetic-only transport budgeting regressions; no provider credentials/calls."""

import json
from copy import deepcopy
from io import BytesIO

import pytest
from openpyxl import Workbook

from app.application.elt.errors import ReviewedExcelError
from app.application.reviewed_excel_audit import request_digest
from app.application.reviewed_excel_budget import (
    BUDGET_VERSION,
    MAX_OUTPUT_TOKENS,
    MAX_REQUEST_BYTES,
    RequestBudgetExceeded,
    fit_transport_batches,
    output_budget,
    request_bytes,
    validate_checkpoint_budget,
    validate_remaining_audit_budget,
)
from app.application.reviewed_excel_classification import _classify_semantic_batch
from app.application.reviewed_excel_evidence import (
    MAX_MODEL_CELLS,
    MAX_MODEL_TEXT_BYTES,
    _model_cell_cost,
    plan_semantic_batches,
)
from app.application.reviewed_excel_template import inspect_workbook


def workbook_review(rows=120, *, formulas=True):
    book = Workbook()
    sheet = book.active
    sheet.title = "Synthetic"
    for row in range(1, rows + 1):
        sheet.cell(
            row, 1, "=SUM(" + ",".join(f"B{i}" for i in range(1, 26)) + ")" if formulas else "x"
        )
    stream = BytesIO()
    book.save(stream)
    return inspect_workbook(stream.getvalue())


def test_complete_request_includes_dense_formulas_schema_and_preserves_all_cells(monkeypatch):
    monkeypatch.delenv("FHD_TEMPLATE_LLM_PROVIDER", raising=False)
    # 300 dense formula rows still exceed the raised ~384KB request envelope as one call,
    # while staying under MAX_MODEL_TEXT_BYTES so transport packing (not cell-count) splits.
    review = workbook_review(300)
    assert sum(map(_model_cell_cost, review["cells"])) < MAX_MODEL_TEXT_BYTES
    with pytest.raises(RequestBudgetExceeded):
        _classify_semantic_batch(review, prepare_only=True, reference_ids=True)
    original = deepcopy(review)
    batches, rounds = plan_semantic_batches(review, 4, neutral=True)
    assert 1 < len(batches) <= 99 and rounds == 4
    found = []
    for batch in batches:
        prepared = _classify_semantic_batch(
            {**review, "cells": batch, "full_cells": review["cells"]},
            prepare_only=True,
            reference_ids=True,
        )
        assert request_bytes(prepared) <= MAX_REQUEST_BYTES
        payload = json.loads(prepared["messages"][1]["content"])
        assert payload["integrity_constraints"]["formula_references"]
        assert payload["output_schema"] == prepared["schema"]
        found.extend(c["cell"] for c in batch if not c.get("context_only"))
    assert sorted(found) == sorted(c["cell"] for c in review["cells"])
    assert len(found) == len(set(found))
    assert review == original


def test_single_row_overhead_budget_rejects_before_provider(monkeypatch):
    review = workbook_review(1)
    monkeypatch.setattr("app.application.reviewed_excel_budget.MAX_REQUEST_BYTES", 100)
    monkeypatch.setattr(
        "app.application.reviewed_excel_template._template_mimo_complete",
        lambda *a: pytest.fail("must not call provider"),
    )
    with pytest.raises(RequestBudgetExceeded, match="单行"):
        plan_semantic_batches(review, neutral=True)


def test_oversized_primary_row_is_never_silently_dropped():
    review = workbook_review(2, formulas=False)
    review["cells"][0]["value"] = "x" * 100_000
    with pytest.raises(ReviewedExcelError, match="单行"):
        plan_semantic_batches(review, neutral=True)


def test_prepare_and_execution_use_identical_bounded_output_budget(monkeypatch):
    review = workbook_review(180, formulas=False)
    prepared = _classify_semantic_batch(review, prepare_only=True, reference_ids=True)
    assert 4000 < prepared["max_tokens"] <= MAX_OUTPUT_TOKENS
    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")

    def capture(messages, schema, tokens):
        assert {"messages": messages, "schema": schema, "max_tokens": tokens} == prepared
        raise RuntimeError("synthetic provider boundary")

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", capture)
    with pytest.raises(RuntimeError, match="synthetic provider boundary"):
        _classify_semantic_batch(
            review, reference_ids=True, expected_request_sha256=request_digest(prepared)
        )


def test_output_budget_is_bounded_and_ignores_repeated_context():
    assert output_budget([{}]) == 4000
    assert output_budget([{}] * 1000) == MAX_OUTPUT_TOKENS
    assert output_budget([{}] * 90 + [{"context_only": True}] * 400) == output_budget([{}] * 90)


@pytest.mark.parametrize("old", [{"results": []}, {"budget_version": "previous", "results": []}])
def test_old_or_different_budget_checkpoint_rejected_unchanged(old):
    before = deepcopy(old)
    with pytest.raises(ReviewedExcelError, match="原结果与审计保留"):
        validate_checkpoint_budget(old)
    assert old == before
    validate_checkpoint_budget({**old, "budget_version": BUDGET_VERSION})


def test_audit_reserves_group_and_counts_failed_and_not_sent():
    state = {"model_audit": [{"state": "failed"}, {"state": "not_sent"}]}
    checkpoint = {"results": [], "auto_group": True, "phase": "ready"}
    validate_remaining_audit_budget(state, checkpoint, 97)
    before = deepcopy(state)
    with pytest.raises(ReviewedExcelError, match="100条"):
        validate_remaining_audit_budget(state, checkpoint, 98)
    assert state == before


def test_plan_has_hard_audit_limit(monkeypatch):
    review = workbook_review(100, formulas=False)
    for i, cell in enumerate(review["cells"]):
        cell["sheet"] = f"S{i}"
    review["sheets"] = [c["sheet"] for c in review["cells"]]
    monkeypatch.setattr(
        "app.application.reviewed_excel_classification._classify_semantic_batch", lambda *a, **k: {}
    )
    with pytest.raises(RequestBudgetExceeded, match="99批"):
        fit_transport_batches(review, [[c] for c in review["cells"]])


@pytest.mark.parametrize("round_budget", [0, 5])
def test_per_round_budget_still_one_to_four(round_budget):
    with pytest.raises(ReviewedExcelError, match="1 至 4"):
        plan_semantic_batches(workbook_review(1), round_budget, neutral=True)


@pytest.mark.parametrize("value", ["0", "181", "nan", "inf", "not-a-number"])
def test_template_timeout_rejects_invalid_configuration(value, monkeypatch):
    from app.application.reviewed_excel_audit import ModelRequestNotSent
    from app.application.reviewed_excel_budget import template_timeout_seconds

    monkeypatch.setenv("FHD_TEMPLATE_LLM_TIMEOUT_SECONDS", value)
    with pytest.raises(ModelRequestNotSent):
        template_timeout_seconds()


def test_template_timeout_is_separate_from_etl_row_limit(monkeypatch):
    from app.application.reviewed_excel_budget import template_timeout_seconds

    monkeypatch.delenv("FHD_TEMPLATE_LLM_TIMEOUT_SECONDS", raising=False)
    monkeypatch.setenv("FHD_ETL_LLM_TIMEOUT_SECONDS", "2")
    assert template_timeout_seconds() == 120
    monkeypatch.setenv("FHD_TEMPLATE_LLM_TIMEOUT_SECONDS", "180")
    assert template_timeout_seconds() == 180


def test_actual_prefix_request_sizes_are_monotone_with_fixed_edge_context():
    review = workbook_review(24)
    cells = review["cells"]
    edge_context = cells[:5] + cells[-5:]
    sizes = []
    for count in range(1, len(cells) + 1):
        selected = cells[:count]
        keys = {c["cell"] for c in selected}
        batch = selected + [
            {**c, "context_only": True} for c in edge_context if c["cell"] not in keys
        ]
        request = _classify_semantic_batch(
            {**review, "cells": batch, "full_cells": cells}, prepare_only=True, reference_ids=True
        )
        sizes.append(request_bytes(request))
    assert sizes == sorted(sizes)
    assert len(set(sizes)) == len(sizes)


def test_old_cell_chunk_boundaries_do_not_change_neutral_packing():
    review = workbook_review(300, formulas=False)
    cells = review["cells"]
    old_chunks = [cells[i : i + 75] for i in range(0, len(cells), 75)]
    combined = fit_transport_batches(review, [cells])
    repacked = fit_transport_batches(review, old_chunks)
    assert combined == repacked
    assert len(repacked) < len(old_chunks)
    assert sum(not c.get("context_only") for b in repacked for c in b) == 300


def test_full_width_primary_row_is_not_rejected_for_optional_context(monkeypatch):
    from openpyxl.utils.cell import get_column_letter

    width = MAX_MODEL_CELLS
    review = {
        "sheets": ["S"],
        "cells": [
            {"sheet": "S", "cell": f"{get_column_letter(c)}2", "value": "x", "kind": "review"}
            for c in range(1, width + 1)
        ]
        + [{"sheet": "S", "cell": "A3", "value": "tail", "kind": "review"}],
    }
    monkeypatch.setattr(
        "app.application.reviewed_excel_classification._classify_semantic_batch", lambda *a, **k: {}
    )
    fitted = fit_transport_batches(review, [review["cells"][:width], review["cells"][width:]])
    assert len(fitted) == 2
    assert sum(not c.get("context_only") for c in fitted[0]) == width
    assert sum(not c.get("context_only") for b in fitted for c in b) == width + 1
    assert all(len(b) <= MAX_MODEL_CELLS for b in fitted)


def saved_plan(review, batches):
    from app.application.reviewed_excel_plan import store_coordinate_plan

    checkpoint = {
        "budget_version": BUDGET_VERSION,
        "source_sha256": review["sha256"],
        "version": "synthetic-version",
        "max_calls": 2,
        "total": len(batches),
    }
    store_coordinate_plan(checkpoint, batches)
    return checkpoint


def test_coordinate_checkpoint_roundtrip_is_value_free_without_replanning(monkeypatch):
    from app.application.reviewed_excel_plan import restore_coordinate_plan

    review = workbook_review(120)
    batches, _ = plan_semantic_batches(review, 2, neutral=True)
    checkpoint = saved_plan(review, batches)
    assert "=SUM" not in json.dumps(checkpoint)
    monkeypatch.setattr(
        "app.application.reviewed_excel_classification._classify_semantic_batch",
        lambda *a, **k: pytest.fail("restore must not prepare the entire plan"),
    )
    rebuilt = restore_coordinate_plan(review, checkpoint)
    assert [[(c["cell"], bool(c.get("context_only"))) for c in b] for b in rebuilt] == [
        [(c["cell"], bool(c.get("context_only"))) for c in b] for b in batches
    ]
    assert [c["value"] for b in rebuilt for c in b] == [c["value"] for b in batches for c in b]


@pytest.mark.parametrize("tamper", ["missing", "digest", "source", "unknown", "duplicate", "omit"])
def test_coordinate_checkpoint_tampering_is_rejected(tamper):
    from app.application.reviewed_excel_plan import _digest, restore_coordinate_plan

    review = workbook_review(3, formulas=False)
    checkpoint = saved_plan(review, [review["cells"]])
    if tamper == "missing":
        checkpoint.pop("coordinate_plan")
    elif tamper == "digest":
        checkpoint["plan_digest"] = "0" * 64
    elif tamper == "source":
        checkpoint["source_sha256"] = "changed"
    else:
        if tamper == "unknown":
            checkpoint["coordinate_plan"][0][0][1] = "Z999"
        elif tamper == "duplicate":
            checkpoint["coordinate_plan"][0].append(checkpoint["coordinate_plan"][0][0])
        else:
            checkpoint["coordinate_plan"][0].pop()
        # Even a recomputed hash cannot authorize invalid coverage/coordinates.
        checkpoint["plan_digest"] = _digest(checkpoint)
    with pytest.raises(ReviewedExcelError, match="坐标计划"):
        restore_coordinate_plan(review, checkpoint)


def test_optional_large_context_never_excludes_valid_primary_rows():
    book = Workbook()
    sheet = book.active
    for row, count, length in [(1, 5, 3500), (2, 30, 2000), (3, 5, 3500)]:
        for col in range(1, count + 1):
            sheet.cell(row, col, "x" * length)
    stream = BytesIO()
    book.save(stream)
    review = inspect_workbook(stream.getvalue())
    for row in [1, 2, 3]:
        chosen = [c for c in review["cells"] if c["cell"].endswith(str(row))]
        prepared = _classify_semantic_batch(
            {**review, "cells": chosen, "full_cells": review["cells"]},
            prepare_only=True,
            reference_ids=True,
        )
        assert request_bytes(prepared) <= MAX_REQUEST_BYTES
    batches, _ = plan_semantic_batches(review, neutral=True)
    primary = [c["cell"] for b in batches for c in b if not c.get("context_only")]
    assert len(primary) == len(set(primary)) == 40
    for batch in batches:
        request = _classify_semantic_batch(
            {**review, "cells": batch, "full_cells": review["cells"]},
            prepare_only=True,
            reference_ids=True,
        )
        assert request_bytes(request) <= MAX_REQUEST_BYTES


def test_merge_closed_band_uses_safe_column_cuts_without_rewriting_source(monkeypatch):
    from openpyxl.utils.cell import (
        column_index_from_string,
        coordinate_from_string,
        range_boundaries,
    )

    from app.application import reviewed_excel_classification as classification

    book = Workbook()
    sheet = book.active
    for row in (1, 2):
        for col in range(1, 5):
            sheet.cell(row, col, "source")
    sheet.merge_cells("A1:A2")
    sheet.merge_cells("D1:D2")
    output = BytesIO()
    book.save(output)
    review = inspect_workbook(output.getvalue())
    original = deepcopy(review)
    real_prepare = classification._classify_semantic_batch

    def bounded_prepare(current, **kwargs):
        assert kwargs == {"prepare_only": True, "reference_ids": True}
        if len(current["cells"]) > 3:
            raise RequestBudgetExceeded("synthetic transport limit")
        return real_prepare(current, **kwargs)

    monkeypatch.setattr(classification, "_classify_semantic_batch", bounded_prepare)
    batches = fit_transport_batches(review, [review["cells"]])
    assert len(batches) == 2
    assert review == original
    found = [cell["cell"] for batch in batches for cell in batch]
    assert len(found) == len(set(found)) == len(review["cells"])
    for batch in batches:
        coords = [coordinate_from_string(cell["cell"]) for cell in batch]
        left, right = (
            min(column_index_from_string(c) for c, _ in coords),
            max(column_index_from_string(c) for c, _ in coords),
        )
        # Both rows are supplied in each horizontal subset; no vertical merge is split.
        assert {row for _, row in coords} == {1, 2}
        for area in ("A1:A2", "D1:D2"):
            c1, _, c2, _ = range_boundaries(area)
            assert not (left <= c2 and c1 <= right) or left <= c1 <= c2 <= right
        prepared = real_prepare(
            {**review, "cells": batch, "full_cells": review["cells"]},
            prepare_only=True,
            reference_ids=True,
        )
        payload = json.loads(prepared["messages"][1]["content"])
        for cell in payload["cells"]:
            if cell["cell"] in {"A1", "D1"}:
                assert cell["required_extent"] == f"{cell['cell']}:{cell['cell'][0]}2"
            else:
                assert "required_extent" not in cell


def test_indivisible_merge_closure_rejects_before_provider(monkeypatch):
    from app.application import reviewed_excel_classification as classification

    book = Workbook()
    sheet = book.active
    # Nonoverlapping source merges form a pinwheel with no safe row or column cut.
    for area in ("A1:C1", "D1:D3", "B4:D4", "A2:A4", "B2:C3"):
        sheet[area.split(":")[0]] = "source"
        sheet.merge_cells(area)
    output = BytesIO()
    book.save(output)
    review = inspect_workbook(output.getvalue())

    def reject(current, **kwargs):
        assert kwargs == {"prepare_only": True, "reference_ids": True}
        if len(current["cells"]) > 1:
            raise RequestBudgetExceeded("synthetic transport limit")
        pytest.fail("an indivisible source block must not be split")

    monkeypatch.setattr(classification, "_classify_semantic_batch", reject)
    monkeypatch.setattr(
        "app.application.reviewed_excel_template._template_mimo_complete",
        lambda *a: pytest.fail("must not call provider"),
    )
    with pytest.raises(RequestBudgetExceeded, match="合并闭包无安全二维分割"):
        fit_transport_batches(review, [review["cells"]])


@pytest.mark.parametrize("reference_ids", [False, True])
def test_required_extent_is_source_derived_not_a_business_region(reference_ids, monkeypatch):
    book = Workbook()
    sheet = book.active
    sheet.title = "O'Brien ! 表"
    sheet["A1"] = "source"
    sheet.merge_cells("A1:B2")
    sheet["D1"] = "separate source"
    output = BytesIO()
    book.save(output)
    review = inspect_workbook(output.getvalue())
    review["cells"][0]["required_extent"] = "A1:Z99"
    review["cells"][1]["required_extent"] = "D1:Z99"
    original = deepcopy(review)
    request = _classify_semantic_batch(review, prepare_only=True, reference_ids=reference_ids)
    cells = json.loads(request["messages"][1]["content"])["cells"]
    assert cells[0]["required_extent"] == "A1:B2"
    assert "required_extent" not in cells[1]
    assert [c["cell"] for c in cells] == ["A1", "D1"]  # Blank subordinate cells are not evidence.
    assert (
        "do not prove one business document"
        if reference_ids
        else "not a business boundary or a separate document"
    ) in request["messages"][0]["content"]
    assert review == original

    if reference_ids:
        from app.application.elt.llm_assist_runtime import LlmAssistResult

        raw = {
            "physical_regions": [
                {
                    "sheet": sheet.title,
                    "range": "A1:B1",
                    "primary_id": "S1:A1",
                    "evidence_ids": ["S1:A1"],
                    "dependency_ids": [],
                    "reason": "source",
                }
            ],
            "logical_documents": [],
            "layout_references": [],
        }
        untouched = deepcopy(raw)
        monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
        monkeypatch.setattr(
            "app.application.reviewed_excel_template._template_mimo_complete",
            lambda *args: LlmAssistResult(used_llm=True, data=raw),
        )
        with pytest.raises(ReviewedExcelError, match="白名单"):
            _classify_semantic_batch(review, reference_ids=True)
        assert raw == untouched  # The source extent never repairs the model's returned range.
