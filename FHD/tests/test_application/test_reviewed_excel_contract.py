import json
from io import BytesIO

import pytest
from openpyxl import Workbook

from app.application.etl.errors import ReviewedExcelError
from app.application.reviewed_excel_classification import _classify_semantic_batch
from app.application.reviewed_excel_grouping import group_semantic_fragments
from app.application.reviewed_excel_reference_ids import (
    ATOM_IDENTIFY_PROMPT as IDENTIFY_PROMPT,
)
from app.application.reviewed_excel_reference_ids import (
    contract_hashes,
    verify_contract,
)
from app.application.reviewed_excel_template import inspect_workbook


@pytest.fixture(autouse=True)
def archived_v7_atom_contract(monkeypatch):
    from app.application import reviewed_excel_reference_ids as contracts

    monkeypatch.setattr(contracts, "_boundary_contract", contracts._atom_contract)


def test_independent_contract_is_identical_for_unseen_structures(monkeypatch):
    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    hashes = []
    for sheet_name, address, merged in [
        ("甲", "B3", False),
        ("测试", "C8", True),
        ("Other", "F5", False),
    ]:
        book = Workbook()
        sheet = book.active
        sheet.title = sheet_name
        sheet[address] = "synthetic title"
        sheet["A12"] = "synthetic data"
        if merged:
            sheet.merge_cells("C8:D8")
        stream = BytesIO()
        book.save(stream)
        review = inspect_workbook(stream.getvalue())
        prepared = _classify_semantic_batch(review, prepare_only=True, reference_ids=True)
        payload = json.loads(prepared["messages"][1]["content"])
        assert prepared["messages"][0]["content"] == IDENTIFY_PROMPT
        assert (
            not {"rule_hints", "structure_map", "relationship_context", "labels"} & payload.keys()
        )
        assert (
            "labels"
            not in prepared["schema"]["properties"]["physical_regions"]["items"]["properties"]
        )
        hashes.append(
            verify_contract(prepared["messages"][0]["content"], prepared["schema"], "identify")
        )
    assert hashes[0] == hashes[1] == hashes[2]
    assert (
        hashes[0]["schema_template_sha256"]
        == "ffab0241efd5e7fba5087b4ca32774b3ba3b2ccb95124709f624d095ddd0adc0"
    )
    with pytest.raises(ReviewedExcelError, match="契约"):
        verify_contract(IDENTIFY_PROMPT + " changed", prepared["schema"], "identify")


def test_grouping_contract_has_stable_template_across_model_fragments():
    fingerprints = []
    for sheet, cell in [("Alpha", "B5"), ("乙", "C9")]:
        review = {
            "sha256": "synthetic",
            "sheets": [sheet],
            "cells": [{"sheet": sheet, "cell": cell, "value": "synthetic", "kind": "text"}],
        }
        classification = {
            "source_sha256": "synthetic",
            "coverage": [{"sheet": sheet, "unread": [], "unassigned": [], "merged_omitted": 0}],
            "regions": [
                {
                    "sheet": sheet,
                    "range": f"{cell}:{cell}",
                    "evidence": [{"sheet": sheet, "cell": cell}],
                }
            ],
        }
        request = group_semantic_fragments(review, classification, prepare_only=True)
        fingerprints.append(contract_hashes(request["schema"], "group"))
    assert fingerprints[0] == fingerprints[1]
    assert (
        fingerprints[0]["schema_template_sha256"]
        == "06548a8c4a41b87333a2e19dba43320f23ead21993ff7884ff884c54c60699e2"
    )


def test_identification_and_grouping_keep_workbook_prefix_when_batch_skips_first_sheet():
    book = Workbook()
    book.active.title = "First"
    second = book.create_sheet("O'Brien ! 表")
    second["C17"] = "synthetic"
    stream = BytesIO()
    book.save(stream)
    review = inspect_workbook(stream.getvalue())
    batch_review = {**review, "full_cells": review["cells"], "cells": review["cells"]}
    identify = _classify_semantic_batch(batch_review, prepare_only=True, reference_ids=True)
    identify_payload = json.loads(identify["messages"][1]["content"])
    prefixes = [
        {"sheet": "First", "prefix": "S1"},
        {"sheet": "O'Brien ! 表", "prefix": "S2"},
    ]
    assert identify_payload["source_sheet_prefixes"] == prefixes
    assert [cell["evidence_id"] for cell in identify_payload["cells"]] == ["S2:C17"]
    classification = {
        "source_sha256": review["sha256"],
        "coverage": [
            {"sheet": name, "unread": [], "unassigned": [], "merged_omitted": 0}
            for name in ["First", "O'Brien ! 表"]
        ],
        "regions": [
            {
                "sheet": "O'Brien ! 表",
                "range": "C17",
                "evidence": [{"sheet": "O'Brien ! 表", "cell": "C17"}],
            }
        ],
    }
    grouping = group_semantic_fragments(review, classification, prepare_only=True)
    grouping_payload = json.loads(grouping["messages"][1]["content"])
    assert "source_sheet_prefixes" not in grouping_payload
    assert [cell["sheet"] for cell in grouping_payload["cells"]] == ["O'Brien ! 表"]


def source_protocol_review():
    book = Workbook()
    first = book.active
    first.title = "O'Brien ! 表"
    first["A1"] = "=Z99"
    first.merge_cells("A1:A2")
    second = book.create_sheet("其他")
    second["C1"] = "source"
    second.merge_cells("C1:C2")
    stream = BytesIO()
    book.save(stream)
    return inspect_workbook(stream.getvalue())


@pytest.mark.parametrize("reference_ids", [False, True])
def test_final_source_protocol_defines_layout_documents_and_context(reference_ids):
    request = _classify_semantic_batch(
        source_protocol_review(), prepare_only=True, reference_ids=reference_ids
    )
    prompt = request["messages"][0]["content"]
    if reference_ids:
        for clause in (
            "smallest bounding rectangle",
            "Ordinary source blank cells inside",
            "no padding outside",
            "Every nonempty source cell inside",
            "Invalid selections are rejected, never expanded",
            "logical_documents.fragments contains the primary_id",
            "not a coverage ledger",
        ):
            assert clause in prompt
    else:
        for clause in (
            "logical_documents.fragments must copy exact sheet/range pairs",
            "Assign each fragment to at most one document",
            "layout_references contains only exact sheet/range pairs from sheets[].merged",
            "not a merge-coverage ledger",
            "context_only cells may be dependencies, never primary or classification evidence",
            "Any region intersecting a merged range must contain that entire merged range",
            "never cut a merge or borrow unread nonempty cells",
            "A formula target is available evidence only if it is explicitly present in cells",
        ):
            assert clause in prompt
    payload = json.loads(request["messages"][1]["content"])
    assert payload["output_schema"] == request["schema"]
    layout = request["schema"]["properties"]["layout_references"]
    fragments = request["schema"]["properties"]["logical_documents"]["items"]["properties"][
        "fragments"
    ]["items"]
    assert layout["items"] is not fragments
    if reference_ids:
        assert fragments == {"type": "string"}
        assert layout["items"]["enum"] == ["M1:A1:A2", "M2:C1:C2"]
    else:
        assert "enum" not in fragments["properties"]["range"]
        assert layout["items"]["properties"]["range"]["enum"] == ["A1:A2", "C1:C2"]
        assert layout["items"]["properties"]["sheet"]["enum"] == ["O'Brien ! 表", "其他"]
    if reference_ids:
        assert (
            "S1:Z99"
            not in request["schema"]["properties"]["physical_regions"]["items"]["properties"][
                "dependency_ids"
            ]["items"]["enum"]
        )
        assert payload["integrity_constraints"]["formula_references"][0]["references"] == [
            {"sheet": "O'Brien ! 表", "range": "Z99"}
        ]


def test_layout_source_enums_do_not_replace_sheet_range_pair_guard(monkeypatch):
    from app.application.etl.llm_assist_runtime import LlmAssistResult
    from app.application.reviewed_excel_template import _validate_template_schema

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    review = source_protocol_review()
    request = _classify_semantic_batch(review, prepare_only=True, reference_ids=False)
    data = {
        "physical_regions": [
            {
                "sheet": sheet,
                "range": area,
                "labels": ["unknown"],
                "evidence": [{"sheet": sheet, "cell": identifier.split(":", 1)[1]}],
                "dependencies": [],
                "reason": "source",
            }
            for sheet, area, identifier in [
                ("O'Brien ! 表", "A1:A2", "S1:A1"),
                ("其他", "C1:C2", "S2:C1"),
            ]
        ],
        "logical_documents": [],
        "layout_references": [{"sheet": "O'Brien ! 表", "range": "C1:C2"}],
    }
    assert _validate_template_schema(request["schema"], data)
    monkeypatch.setattr(
        "app.application.reviewed_excel_template._template_mimo_complete",
        lambda *args: LlmAssistResult(used_llm=True, data=data),
    )
    with pytest.raises(ReviewedExcelError, match="布局引用"):
        _classify_semantic_batch(review, reference_ids=False)
    data["layout_references"][0]["range"] = "A1:A2"
    assert len(_classify_semantic_batch(review, reference_ids=False)["regions"]) == 2
    data["layout_references"] = []
    assert _classify_semantic_batch(review, reference_ids=False)["layout_references"] == []
    data["logical_documents"] = [
        {"document_id": "d", "fragments": [{"sheet": "O'Brien ! 表", "range": "A1:A2"}]},
        {"document_id": "e", "fragments": [{"sheet": "O'Brien ! 表", "range": "A1:A2"}]},
    ]
    with pytest.raises(ReviewedExcelError, match="重复"):
        _classify_semantic_batch(review, reference_ids=False)
    data["logical_documents"].pop()
    data["logical_documents"][0]["fragments"][0]["range"] = "A1:C2"
    with pytest.raises(ReviewedExcelError, match="片段来源"):
        _classify_semantic_batch(review, reference_ids=False)


def test_empty_merge_source_allows_only_empty_layout_and_preserves_contract():
    from app.application.reviewed_excel_template import _validate_template_schema

    book = Workbook()
    book.active["A1"] = "source"
    stream = BytesIO()
    book.save(stream)
    request = _classify_semantic_batch(
        inspect_workbook(stream.getvalue()), prepare_only=True, reference_ids=True
    )
    layout = request["schema"]["properties"]["layout_references"]
    assert layout["maxItems"] == 0
    assert _validate_template_schema(layout, [])
    assert not _validate_template_schema(layout, [{"sheet": "Sheet", "range": "A1"}])
    assert verify_contract(IDENTIFY_PROMPT, request["schema"], "identify") == verify_contract(
        IDENTIFY_PROMPT,
        _classify_semantic_batch(source_protocol_review(), prepare_only=True, reference_ids=True)[
            "schema"
        ],
        "identify",
    )


@pytest.mark.parametrize(
    "mutation",
    ["bound", "bound_boolean", "required", "document_length", "evidence_role", "layout_role"],
)
def test_source_normalization_does_not_hide_protocol_mutations(mutation):
    schema = _classify_semantic_batch(
        source_protocol_review(), prepare_only=True, reference_ids=True
    )["schema"]
    props = schema["properties"]
    if mutation == "bound":
        props["layout_references"]["maxItems"] = 99
    elif mutation == "bound_boolean":
        props["layout_references"]["maxItems"] = False
    elif mutation == "required":
        props["layout_references"]["items"]["required"] = ["sheet"]
    elif mutation == "document_length":
        props["logical_documents"]["items"]["properties"]["document_id"]["maxLength"] = 81
    else:
        props["physical_regions"]["items"]["properties"]["atom_ids"]["items"]["enum"].append(
            "M1:Z99:AA99" if mutation == "layout_role" else "S1:Z99"
        )
    with pytest.raises(ReviewedExcelError, match="契约"):
        verify_contract(IDENTIFY_PROMPT, schema, "identify")
