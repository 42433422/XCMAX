from io import BytesIO
from zipfile import ZipFile

import pytest

from app.application.etl.errors import ReviewedExcelError
from app.application.reviewed_excel_drawing import extract_drawing_evidence


def package(*, objects=None, drawing_rels=None):
    marker = "<x:from><x:col>0</x:col><x:colOff>1</x:colOff><x:row>0</x:row><x:rowOff>2</x:rowOff></x:from>"
    shape = '<x:sp><x:nvSpPr><x:cNvPr id="1" name="shape"/><x:cNvSpPr/></x:nvSpPr><x:spPr/><x:txBody><a:bodyPr vertOverflow="clip" horzOverflow="clip"/><a:p><a:r><a:t>标题</a:t></a:r></a:p></x:txBody></x:sp>'
    if objects is None:
        objects = f'<x:oneCellAnchor>{marker}<x:ext cx="100" cy="200"/>{shape}<x:clientData/></x:oneCellAnchor>'
    files = {
        "[Content_Types].xml": '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Override PartName="/xl/drawings/drawing1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>',
        "xl/workbook.xml": '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="表" r:id="r1"/></sheets></workbook>',
        "xl/_rels/workbook.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="r1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/></Relationships>',
        "xl/worksheets/sheet1.xml": '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><drawing r:id="r2"/></worksheet>',
        "xl/worksheets/_rels/sheet1.xml.rels": '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="r2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" Target="../drawings/drawing1.xml"/></Relationships>',
        "xl/drawings/drawing1.xml": f'<x:wsDr xmlns:x="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">{objects}</x:wsDr>',
    }
    if drawing_rels is not None:
        files["xl/drawings/_rels/drawing1.xml.rels"] = drawing_rels
    output = BytesIO()
    with ZipFile(output, "w") as z:
        for path, content in files.items():
            z.writestr(path, content)
    output.seek(0)
    return output


def test_text_shape_identity_and_anchor():
    with ZipFile(package()) as z:
        items = extract_drawing_evidence(z)
    assert items == [
        {
            "sheet": "表",
            "part": "xl/drawings/drawing1.xml",
            "object_id": "1",
            "kind": "sp",
            "anchor_type": "oneCellAnchor",
            "from": {"col": 0, "colOff": 1, "row": 0, "rowOff": 2},
            "to": None,
            "extent": {"cx": 100, "cy": 200},
            "text": "标题",
        }
    ]


def test_unknown_drawing_and_external_relationship_rejected():
    with ZipFile(package(objects="<x:absoluteAnchor/>")) as z:
        with pytest.raises(ReviewedExcelError, match="锚点类型|未知对象结构"):
            extract_drawing_evidence(z)
    rel = '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="r1" Type="x" Target="https://example.invalid/" TargetMode="External"/></Relationships>'
    with ZipFile(package(drawing_rels=rel)) as z:
        with pytest.raises(ReviewedExcelError, match="外部目标"):
            extract_drawing_evidence(z)


def test_duplicate_object_identity_rejected():
    with ZipFile(package()) as z:
        from xml.etree import ElementTree as ET

        root = ET.fromstring(z.read("xl/drawings/drawing1.xml"))
        element = ET.tostring(root[0], encoding="unicode")
    with ZipFile(package(objects=element + element)) as z:
        with pytest.raises(ReviewedExcelError, match="编号无效"):
            extract_drawing_evidence(z)


def test_planner_accounts_for_multiple_drawings_at_same_cell_and_fails_on_uncovered_row():
    from app.application.reviewed_excel_drawing import plan_drawing_batches
    from app.application.reviewed_excel_evidence import MAX_MODEL_CELLS

    # Two drawings on A1 add +2 to the count; width must exceed MAX so packing needs 2 batches.
    n_cells = MAX_MODEL_CELLS
    cells = [
        {"sheet": "表", "cell": f"A{i}", "value": "v", "kind": "review"}
        for i in range(1, n_cells + 1)
    ]
    drawings = [
        {
            "sheet": "表",
            "part": "xl/drawings/drawing1.xml",
            "object_id": str(i),
            "from": {"row": 0},
            "text": "标题",
        }
        for i in range(2)
    ]
    review = {
        "sheets": ["表"],
        "cells": cells,
        "drawing_evidence": drawings,
        "sheet_overview": [
            {"sheet": "表", "drawing_row_boundaries": list(range(1, n_cells + 2))}
        ],
    }
    batches = plan_drawing_batches(review)
    assert len(batches) == 2
    assert sum(len(c.get("drawing_evidence", [])) for b in batches for c in b) == 2
    assert all(
        sum(1 + len(c.get("drawing_evidence", [])) for c in b) <= MAX_MODEL_CELLS for b in batches
    )
    assert cells[0].get("drawing_evidence") is None
    review["drawing_evidence"] = [{**drawings[0], "from": {"row": n_cells + 100}}]
    with pytest.raises(ReviewedExcelError, match="没有单元格证据"):
        plan_drawing_batches(review)


def test_classifier_sends_distinct_drawing_evidence_and_rejects_forged_reference(monkeypatch):
    import json

    from app.application.etl.llm_assist_runtime import LlmAssistResult
    from app.application.reviewed_excel_classification import (
        _classify_semantic_batch,
        plan_semantic_batches,
    )
    from app.application.reviewed_excel_template import inspect_workbook
    from tests.fixtures.etl_template.sources.generate import sample

    review = inspect_workbook(sample())
    drawing = {
        "sheet": "报价",
        "part": "xl/drawings/drawing1.xml",
        "object_id": "17",
        "kind": "sp",
        "anchor_type": "oneCellAnchor",
        "from": {"col": 0, "colOff": 0, "row": 0, "rowOff": 0},
        "to": None,
        "extent": {"cx": 100, "cy": 100},
        "text": "报价单",
        "last_col": 1,
        "last_row": 1,
    }
    review["drawing_evidence"] = [drawing]
    review["sheet_overview"][0]["drawing_row_boundaries"] = list(range(1, 6))
    batches, _ = plan_semantic_batches(review, 4)
    current_id = ["17"]
    region_range = ["A1:C4"]

    def fake(messages, schema, max_tokens):
        request = json.loads(messages[1]["content"])
        assert request["drawing_evidence"] == [drawing]
        assert "drawing_evidence" in schema["properties"]["physical_regions"]["items"]["required"]
        region = {
            "sheet": "报价",
            "range": region_range[0],
            "labels": ["quotation"],
            "evidence": [{"sheet": "报价", "cell": "A2"}],
            "dependencies": [],
            "reason": "绘图标题",
            "drawing_evidence": [
                {"sheet": "报价", "part": drawing["part"], "object_id": current_id[0]}
            ],
        }
        return LlmAssistResult(
            used_llm=True,
            model="test",
            data={"physical_regions": [region], "logical_documents": [], "layout_references": []},
        )

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    batch_review = {**review, "cells": batches[0], "full_cells": review["cells"]}
    result = _classify_semantic_batch(batch_review)
    assert result["input_summary"]["sent_drawing_objects"][0]["object_id"] == "17"
    assert result["regions"][0]["drawing_evidence"][0]["object_id"] == "17"
    current_id[0] = "forged"
    with pytest.raises(ReviewedExcelError, match="不在.*已发送范围"):
        _classify_semantic_batch(batch_review)
    current_id[0] = "17"
    region_range[0] = "A2:C4"
    with pytest.raises(ReviewedExcelError, match="绘图证据不在区域内"):
        _classify_semantic_batch(batch_review)


def synthetic_workbook_with_shape(*, extra_rows=False):
    from xml.etree import ElementTree as ET

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "表"
    ws["A1"] = "考勤表"
    ws["A2"] = "姓名"
    ws["B2"] = "日期"
    if extra_rows:
        ws["A4"] = "末尾"
        ws.print_area = "A1:B2"
    original = BytesIO()
    wb.save(original)
    with ZipFile(package()) as fixture:
        drawing = fixture.read("xl/drawings/drawing1.xml")
        drawing = (
            drawing.replace(b"oneCellAnchor", b"twoCellAnchor")
            .replace(b"<x:colOff>1</x:colOff>", b"<x:colOff>0</x:colOff>")
            .replace(b"<x:rowOff>2</x:rowOff>", b"<x:rowOff>0</x:rowOff>")
            .replace(
                b'<x:ext cx="100" cy="200"/>',
                b"<x:to><x:col>0</x:col><x:colOff>0</x:colOff><x:row>1</x:row><x:rowOff>0</x:rowOff></x:to>",
            )
        )
        rels = fixture.read("xl/worksheets/_rels/sheet1.xml.rels")
    output = BytesIO()
    spreadsheet = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    relationships = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    contenttypes = "http://schemas.openxmlformats.org/package/2006/content-types"
    with ZipFile(BytesIO(original.getvalue())) as old, ZipFile(output, "w") as new:
        for item in old.infolist():
            content = old.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(content)
                ET.SubElement(root, f"{{{spreadsheet}}}drawing", {f"{{{relationships}}}id": "r2"})
                content = ET.tostring(root, encoding="utf-8")
            if item.filename == "[Content_Types].xml":
                root = ET.fromstring(content)
                ET.SubElement(
                    root,
                    f"{{{contenttypes}}}Override",
                    {
                        "PartName": "/xl/drawings/drawing1.xml",
                        "ContentType": "application/vnd.openxmlformats-officedocument.drawing+xml",
                    },
                )
                content = ET.tostring(root, encoding="utf-8")
            new.writestr(item, content)
        new.writestr("xl/drawings/drawing1.xml", drawing)
        new.writestr("xl/worksheets/_rels/sheet1.xml.rels", rels)
    return output.getvalue()


def test_synthetic_shape_real_api_upload_analyze_download(tmp_path, monkeypatch):
    from app.application.etl.llm_assist_runtime import LlmAssistResult
    from app.application.reviewed_excel_template import inspect_workbook
    from tests.test_application.test_reviewed_excel_template import review_store

    fixture = review_store.__wrapped__(tmp_path, monkeypatch)
    store = next(fixture)
    source = synthetic_workbook_with_shape()

    def model(messages, schema, max_tokens):
        import json

        data = json.loads(messages[1]["content"])
        d = data["drawing_evidence"][0]
        return LlmAssistResult(
            used_llm=True,
            model="test",
            data={
                "physical_regions": [
                    {
                        "sheet": "表",
                        "range": "A1:B2",
                        "labels": ["attendance"],
                        "evidence": [{"sheet": "表", "cell": "A1"}],
                        "dependencies": [],
                        "reason": "考勤表",
                        "drawing_evidence": [{k: d[k] for k in ("sheet", "part", "object_id")}],
                    }
                ],
                "logical_documents": [
                    {"document_id": "d1", "fragments": [{"sheet": "表", "range": "A1:B2"}]}
                ],
                "layout_references": [],
            },
        )

    try:
        monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
        monkeypatch.setattr(
            "app.application.reviewed_excel_template._template_mimo_complete", model
        )
        assert len(inspect_workbook(source, allow_drawings=True)["drawing_evidence"]) == 1
        upload = store.client.post(
            "/api/templates/reviewed-excel/classify",
            data={"expected_actor": "1:7"},
            files={"files": ("synthetic.xlsx", source)},
        )
        assert upload.status_code == 200, upload.text
        entry = upload.json()["files"][0]
        assert entry["status"] == "classified", upload.text
        path = "/api/templates/reviewed-excel/classifications/" + entry["classification_id"]
        result = store.client.post(
            path + "/analyze", data={"version": entry["version"], "sample_authorized": "true"}
        )
        assert result.status_code == 200, result.text
        output = result.json()
        assert output["coverage"][0]["drawing_unread"] == []
        assert output["coverage"][0]["drawing_unassigned"] == []
        listing = store.client.get(path + "/auto-slices/" + output["version"])
        assert listing.status_code == 200, listing.text
        assert len(listing.json()["slices"]) == 1
        download = store.client.get(path + "/auto-slices/" + output["version"] + "/0")
        assert download.status_code == 200
        assert (
            inspect_workbook(download.content)["drawing_evidence"]
            == inspect_workbook(source)["drawing_evidence"]
        )
    finally:
        try:
            next(fixture)
        except StopIteration:
            pass


def test_grouping_drawing_identity_and_coverage():
    from app.application.reviewed_excel_grouping import prepare_grouping, validate_grouping

    drawing = {
        "sheet": "表",
        "part": "xl/drawings/drawing1.xml",
        "object_id": "1",
        "text": "考勤表",
        "from": {"col": 0, "row": 0},
        "last_col": 1,
        "last_row": 1,
    }
    review = {
        "sha256": "source",
        "sheets": ["表"],
        "cells": [{"sheet": "表", "cell": "A1", "value": "考勤", "kind": "review"}],
        "drawing_evidence": [drawing],
    }
    result = {
        "source_sha256": "source",
        "coverage": [
            {
                "unread": [],
                "unassigned": [],
                "merged_omitted": 0,
                "drawing_unread": [],
                "drawing_unassigned": [],
            }
        ],
        "regions": [
            {
                "sheet": "表",
                "range": "A1:B2",
                "labels": ["attendance"],
                "reason": "标题",
                "evidence": [],
            }
        ],
    }
    group_input = prepare_grouping(review, result)
    ref = {k: drawing[k] for k in ("sheet", "part", "object_id")}
    response = {
        "groups": [
            {
                "document_id": "d1",
                "fragments": [{"sheet": "表", "range": "A1:B2"}],
                "evidence": [{"sheet": "表", "cell": "A1"}],
                "drawing_evidence": [ref],
                "reason": "同一标题",
            }
        ]
    }
    assert validate_grouping(group_input, response)[0]["drawing_evidence"] == [ref]
    response["groups"][0]["drawing_evidence"][0]["object_id"] = "wrong"
    with pytest.raises(ReviewedExcelError, match="不在.*已发送范围"):
        validate_grouping(group_input, response)
    result["coverage"][0]["drawing_unread"] = [ref]
    with pytest.raises(ReviewedExcelError, match="覆盖不完整"):
        prepare_grouping(review, result)


def test_partial_drawing_slice_preserves_objects_and_rejects_crossing():
    from xml.etree import ElementTree as ET

    from openpyxl import load_workbook

    from app.application.reviewed_excel_template import derive_document_workbook, inspect_workbook

    source = synthetic_workbook_with_shape()
    out = BytesIO()
    with ZipFile(BytesIO(source)) as old, ZipFile(out, "w") as new:
        for entry in old.infolist():
            content = old.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(content)
                ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
                data = root.find(f"{{{ns}}}sheetData")
                row = ET.SubElement(data, f"{{{ns}}}row", {"r": "4"})
                cell = ET.SubElement(row, f"{{{ns}}}c", {"r": "A4", "t": "inlineStr"})
                inline = ET.SubElement(cell, f"{{{ns}}}is")
                ET.SubElement(inline, f"{{{ns}}}t").text = "末尾"
                content = ET.tostring(root, encoding="utf-8")
            if entry.filename == "xl/drawings/drawing1.xml":
                root = ET.fromstring(content)
                clone = ET.fromstring(ET.tostring(root[0]))
                x = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
                clone.find(f".//{{{x}}}cNvPr").set("id", "2")
                clone.find(f"{{{x}}}from/{{{x}}}row").text = "2"
                clone.find(f"{{{x}}}to/{{{x}}}row").text = "3"
                root.append(clone)
                content = ET.tostring(root, encoding="utf-8")
            new.writestr(entry, content)
    source = out.getvalue()
    before = inspect_workbook(source, allow_drawings=True)
    assert len(before["drawing_evidence"]) == 2
    result = derive_document_workbook(source, [{"sheet": "表", "range": "A1:B2"}])
    after = inspect_workbook(result, allow_drawings=True)
    assert [x["object_id"] for x in after["drawing_evidence"]] == ["1"]
    assert load_workbook(BytesIO(result))["表"]["A4"].value is None
    assert all(
        after["drawing_evidence"][0][k] == before["drawing_evidence"][0][k]
        for k in ("text", "from", "extent", "kind")
    )
    crossing = {**before, "drawing_evidence": [{**before["drawing_evidence"][0], "last_row": 3}]}
    from app.application.reviewed_excel_drawing import select_drawing_objects

    with pytest.raises(ReviewedExcelError, match="切断绘图对象"):
        select_drawing_objects(crossing, "表", [(1, 1, 2, 2)])


def test_boundary_candidates_keep_merges_and_shapes_together():
    from openpyxl import Workbook

    from app.application.reviewed_excel_drawing import drawing_row_boundaries

    ws = Workbook().active
    ws["A6"] = "tail"
    ws.merge_cells("A2:B3")
    obj = {"from": {"row": 3, "rowOff": 0, "colOff": 0}, "to": {"row": 5, "rowOff": 0, "colOff": 0}}
    boundaries = drawing_row_boundaries(ws, [obj])
    assert 3 not in boundaries
    assert obj["last_row"] == 6
    assert 5 not in boundaries and 6 not in boundaries
    assert 1 in boundaries and 7 in boundaries


def test_auto_slice_preserves_source_print_area_or_rejects_without_implicit_consent(monkeypatch):
    from io import BytesIO

    from openpyxl import load_workbook

    from app.application.reviewed_excel_auto_slice import derive_classified_documents
    from app.application.reviewed_excel_template import inspect_workbook

    source = synthetic_workbook_with_shape(extra_rows=True)
    review = inspect_workbook(source, allow_drawings=True)
    regions = [
        {"sheet": "表", "range": "A1:B2", "dependencies": []},
        {"sheet": "表", "range": "A3:B4", "dependencies": []},
    ]
    classification = {
        "source_sha256": review["sha256"],
        "model_status": "returned",
        "regions": regions,
        "coverage": [
            {
                "unread": [],
                "unassigned": [],
                "merged_omitted": 0,
                "drawing_unread": [],
                "drawing_unassigned": [],
            }
        ],
        "suggested_documents": [
            {"fragments": [{"sheet": r["sheet"], "range": r["range"]}]} for r in regions
        ],
    }
    from app.application import reviewed_excel_auto_slice as auto_slice

    original_derive = auto_slice.derive_document_workbook
    modes = []

    def checked_derive(body, fragments, *, print_area_mode, object_receipt=None):
        modes.append(print_area_mode)
        return original_derive(
            body, fragments, print_area_mode=print_area_mode, object_receipt=object_receipt
        )

    monkeypatch.setattr(auto_slice, "derive_document_workbook", checked_derive)
    with pytest.raises(ReviewedExcelError, match="打印区域跨越所选文档"):
        derive_classified_documents(source, classification)
    assert modes == ["source", "source"]
    assert inspect_workbook(source)["sha256"] == review["sha256"]

    # An explicitly grouped document covering the entire source print area still
    # exports normally, without selecting a new print range or changing source.
    classification["suggested_documents"] = [
        {"fragments": [{"sheet": r["sheet"], "range": r["range"]} for r in regions]}
    ]
    slices = derive_classified_documents(source, classification)
    assert len(slices) == 1
    assert slices[0]["print_area_mode"] == "source"
    assert modes == ["source", "source", "source"]
    assert load_workbook(BytesIO(slices[0]["workbook"]))["表"].print_area == "'表'!$A$1:$B$2"
    assert (
        len(inspect_workbook(slices[0]["workbook"], allow_drawings=True)["drawing_evidence"]) == 1
    )


def test_nonzero_row_offset_is_not_a_proven_boundary():
    from openpyxl import Workbook

    from app.application.reviewed_excel_drawing import drawing_row_boundaries, set_drawing_columns

    ws = Workbook().active
    ws["A6"] = "tail"
    ws.sheet_format.zeroHeight = True
    shape = {
        "sheet": "Sheet",
        "from": {"row": 0, "rowOff": 0, "col": 0, "colOff": 0},
        "to": None,
        "extent": None,
    }
    set_drawing_columns(ws, [shape])
    assert shape["geometry_error"]
    with pytest.raises(ReviewedExcelError, match="尚未核验"):
        drawing_row_boundaries(ws, [shape])


def rewrite_drawing(body, change):
    out = BytesIO()
    with ZipFile(BytesIO(body)) as old, ZipFile(out, "w") as new:
        for item in old.infolist():
            content = old.read(item.filename)
            if item.filename == "xl/drawings/drawing1.xml":
                content = change(content)
            new.writestr(item, content)
    return out.getvalue()


def assert_archive_roundtrip(store, body):
    import hashlib

    from sqlalchemy import text

    with store.engine.begin() as conn:
        conn.execute(
            text(
                "INSERT INTO etl_archive_quotas (tenant_id, limit_bytes, reserved_bytes) VALUES (1, 1000000, 0)"
            )
        )
    receipt = store.client.post(
        "/api/templates/reviewed-excel/archive-classify",
        data={"expected_actor": "1:7", "request_key": "drawing-archive"},
        files={"file": ("synthetic.xlsx", body)},
    )
    assert receipt.status_code == 200, receipt.text
    item = receipt.json()
    assert item["model_status"] == "not_called" and not item["confirmed"]
    version = item["archive_version_id"]
    downloaded = store.client.get(
        f"/api/templates/reviewed-excel/archive-versions/{version}/download"
    )
    assert downloaded.status_code == 200
    assert hashlib.sha256(downloaded.content).digest() == hashlib.sha256(body).digest()
    return item


def test_archive_accepts_unsupported_geometry_without_allowing_model_or_derivation(
    tmp_path, monkeypatch
):
    from app.application.reviewed_excel_drawing import plan_drawing_batches
    from app.application.reviewed_excel_template import derive_document_workbook, inspect_workbook
    from tests.test_application.test_reviewed_excel_template import review_store

    fixture = review_store.__wrapped__(tmp_path, monkeypatch)
    store = next(fixture)
    monkeypatch.setattr(
        "app.application.reviewed_excel_template._template_mimo_complete",
        lambda *_: pytest.fail("unexpected model call"),
    )
    body = rewrite_drawing(
        synthetic_workbook_with_shape(),
        lambda b: b.replace(b"<x:colOff>0</x:colOff>", b"<x:colOff>1</x:colOff>"),
    )
    try:
        assert_archive_roundtrip(store, body)
        review = inspect_workbook(body)
        # twoCell marker bounds remain usable with sub-cell offsets.
        assert "geometry_error" not in review["drawing_evidence"][0]
        assert review["drawing_evidence"][0]["last_row"] >= 1
        plan_drawing_batches(review)
        # oneCell without extent still fail-closes
        bare = {
            "sheet": "表",
            "from": {"row": 0, "col": 0, "colOff": 0, "rowOff": 0},
            "to": None,
            "extent": None,
        }
        from openpyxl import Workbook

        from app.application.reviewed_excel_drawing import set_drawing_columns

        set_drawing_columns(Workbook().active, [bare])
        assert "geometry_error" in bare
    finally:
        try:
            next(fixture)
        except StopIteration:
            pass


def test_two_dimensional_selection_rejects_horizontal_cut_and_holes():
    from app.application.reviewed_excel_drawing import select_drawing_objects

    obj = {"sheet": "表", "from": {"row": 1, "col": 1}, "last_row": 4, "last_col": 4}
    info = {"drawing_evidence": [obj]}
    assert select_drawing_objects(info, "表", [(2, 2, 4, 4)]) == [obj]
    assert select_drawing_objects(info, "表", [(2, 2, 2, 4), (3, 2, 4, 4)]) == [obj]
    assert select_drawing_objects(info, "表", [(5, 2, 6, 4)]) == []
    for boxes in [[(2, 2, 3, 4)], [(2, 2, 4, 2), (2, 4, 4, 4)], [(2, 2, 2, 4), (4, 2, 4, 4)]]:
        with pytest.raises(ReviewedExcelError, match="切断"):
            select_drawing_objects(info, "表", boxes)


def test_nonzero_transform_is_readable_but_not_derivable():
    from app.application.reviewed_excel_template import derive_document_workbook, inspect_workbook

    # Ordinary xfrm is allowed; unmeasured visual overflow (effectLst) still fail-closes.
    body = rewrite_drawing(
        synthetic_workbook_with_shape(),
        lambda b: b.replace(
            b"<x:spPr/>",
            b"<x:spPr><a:effectLst/></x:spPr>",
        ),
    )
    assert inspect_workbook(body)["drawing_evidence"][0]["geometry_error"]
    with pytest.raises(ReviewedExcelError, match="独立变换|视觉外溢"):
        derive_document_workbook(body, [{"sheet": "表", "range": "A1:B2"}])


@pytest.mark.parametrize(
    "change",
    [
        lambda b: b.replace(b"<x:from>", b"<x:from><x:col>99</x:col>"),
        lambda b: b.replace(b"<x:spPr/>", b"<x:spPr><a:unrecognizedGeometry/></x:spPr>"),
        lambda b: b.replace(b"<a:t>", b"<x:t>").replace(b"</a:t>", b"</x:t>"),
        lambda b: b.replace(b"<a:r>", b'<a:fld id="x" type="datetime">').replace(
            b"</a:r>", b"</a:fld>"
        ),
    ],
)
def test_malformed_or_unknown_xml_is_not_supported(change):
    from app.application.reviewed_excel_template import inspect_workbook

    with pytest.raises(ReviewedExcelError):
        inspect_workbook(rewrite_drawing(synthetic_workbook_with_shape(), change))


def test_paragraphs_and_line_breaks_survive_evidence_extraction():
    from app.application.reviewed_excel_template import inspect_workbook

    body = rewrite_drawing(
        synthetic_workbook_with_shape(),
        lambda b: b.replace(
            b"</a:r></a:p>",
            b"</a:r><a:br/><a:r><a:t>B</a:t></a:r></a:p><a:p><a:r><a:t>C</a:t></a:r></a:p>",
        ),
    )
    assert inspect_workbook(body)["drawing_evidence"][0]["text"] == "标题\nB\nC"


def test_drawing_derivation_keeps_formula_support_sheet():
    from xml.etree import ElementTree as ET

    from openpyxl import load_workbook

    from app.application.reviewed_excel_template import derive_document_workbook, inspect_workbook

    source = synthetic_workbook_with_shape(extra_rows=True)
    out = BytesIO()
    with ZipFile(BytesIO(source)) as old, ZipFile(out, "w") as new:
        for item in old.infolist():
            data = old.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(data)
                ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
                cell = next(c for c in root.iter(f"{{{ns}}}c") if c.get("r") == "B2")
                cell.attrib.pop("t", None)
                for child in list(cell):
                    cell.remove(child)
                ET.SubElement(cell, f"{{{ns}}}f").text = "A4"
                data = ET.tostring(root)
            new.writestr(item, data)
    result = derive_document_workbook(out.getvalue(), [{"sheet": "表", "range": "A1:B2"}])
    workbook = load_workbook(BytesIO(result))
    assert workbook["表"]["B2"].value == "='依赖支撑'!A4"
    assert workbook["依赖支撑"]["A4"].value == "末尾"
    assert inspect_workbook(result)["drawing_evidence"][0]["text"] == "标题"


def test_independent_drawing_ids_and_contract_reject_forged_duplicate_or_outside(monkeypatch):
    import json

    from app.application.etl.llm_assist_runtime import LlmAssistResult
    from app.application.reviewed_excel_classification import _classify_semantic_batch
    from app.application.reviewed_excel_drawing import plan_drawing_batches
    from app.application.reviewed_excel_reference_ids import DRAWING_IDENTIFY_PROMPT
    from app.application.reviewed_excel_template import inspect_workbook

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    review = inspect_workbook(synthetic_workbook_with_shape())
    review = {**review, "full_cells": review["cells"], "cells": plan_drawing_batches(review)[0]}
    request = _classify_semantic_batch(review, reference_ids=True, prepare_only=True)
    payload = json.loads(request["messages"][1]["content"])
    assert request["messages"][0]["content"] == DRAWING_IDENTIFY_PROMPT
    assert not {"rule_hints", "structure_map", "relationship_context", "labels"} & payload.keys()
    assert payload["drawing_evidence"][0]["drawing_id"] == "D0001"
    properties = request["schema"]["properties"]["physical_regions"]["items"]["properties"]
    assert "labels" not in properties and properties["drawing_ids"]["items"]["enum"] == ["D0001"]
    assert all(key.startswith("S1:") for key in properties["primary_id"]["enum"])
    ids, area, primary = ["D0001"], "A1:B2", "S1:A1"

    def model(messages, schema, max_tokens):
        return LlmAssistResult(
            used_llm=True,
            model="test",
            data={
                "physical_regions": [
                    {
                        "sheet": "表",
                        "range": area,
                        "primary_id": primary,
                        "evidence_ids": [primary],
                        "dependency_ids": [],
                        "drawing_ids": ids,
                        "reason": "source",
                    }
                ],
                "logical_documents": [],
                "layout_references": [],
            },
        )

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", model)
    accepted = _classify_semantic_batch(review, reference_ids=True)
    assert accepted["regions"][0]["drawing_evidence"][0]["object_id"] == "1"
    ids = ["D9999"]
    with pytest.raises(ReviewedExcelError, match="白名单"):
        _classify_semantic_batch(review, reference_ids=True)
    ids = ["D0001", "D0001"]
    with pytest.raises(ReviewedExcelError, match="重复"):
        _classify_semantic_batch(review, reference_ids=True)
    ids, area, primary = ["D0001"], "B2:B2", "S1:B2"
    with pytest.raises(ReviewedExcelError, match="绘图证据不在区域内"):
        _classify_semantic_batch(review, reference_ids=True)


@pytest.mark.parametrize(
    "replacement",
    [
        b'<a:bodyPr vertOverflow="clip" horzOverflow="clip"><a:spAutoFit/></a:bodyPr>',
        b'<a:bodyPr vertOverflow="clip" horzOverflow="clip" rot="900000"/>',
        b'<a:bodyPr vertOverflow="clip" horzOverflow="clip"/><a:bodyPr vertOverflow="overflow"/>',
    ],
)
def test_text_autofit_rotation_or_duplicate_layout_cannot_be_derived(replacement):
    from app.application.reviewed_excel_template import derive_document_workbook

    body = rewrite_drawing(
        synthetic_workbook_with_shape(),
        lambda b: b.replace(b'<a:bodyPr vertOverflow="clip" horzOverflow="clip"/>', replacement),
    )
    with pytest.raises(ReviewedExcelError):
        derive_document_workbook(body, [{"sheet": "表", "range": "A1:B2"}])


def test_relationship_type_requires_exact_ooxml_namespace():
    from app.application.reviewed_excel_template import inspect_workbook

    source = synthetic_workbook_with_shape()
    out = BytesIO()
    with ZipFile(BytesIO(source)) as old, ZipFile(out, "w") as new:
        for item in old.infolist():
            data = old.read(item.filename)
            if item.filename == "xl/worksheets/_rels/sheet1.xml.rels":
                data = data.replace(
                    b"http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing",
                    b"https://example.invalid/not-ooxml/drawing",
                )
            new.writestr(item, data)
    with pytest.raises(ReviewedExcelError, match="关系不受支持"):
        inspect_workbook(out.getvalue())


def test_drawing_content_type_is_not_silently_repaired():
    from app.application.reviewed_excel_template import inspect_workbook

    source = synthetic_workbook_with_shape()
    out = BytesIO()
    with ZipFile(BytesIO(source)) as old, ZipFile(out, "w") as new:
        for item in old.infolist():
            data = old.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.drawing+xml",
                    b"application/unknown-unsupported",
                )
            new.writestr(item, data)
    with pytest.raises(ReviewedExcelError, match="内容类型声明"):
        inspect_workbook(out.getvalue())


def test_grouping_drawing_id_contract_and_duplicate_rejection(monkeypatch):
    from app.application.etl.llm_assist_runtime import LlmAssistResult
    from app.application.reviewed_excel_grouping import group_semantic_fragments
    from app.application.reviewed_excel_template import inspect_workbook

    review = inspect_workbook(synthetic_workbook_with_shape())
    classification = {
        "source_sha256": review["sha256"],
        "coverage": [
            {
                "sheet": "表",
                "unread": [],
                "unassigned": [],
                "merged_omitted": 0,
                "drawing_unread": [],
                "drawing_unassigned": [],
            }
        ],
        "regions": [{"sheet": "表", "range": "A1:B2", "evidence": [{"sheet": "表", "cell": "A1"}]}],
    }
    request = group_semantic_fragments(review, classification, prepare_only=True)
    assert request["schema"]["properties"]["groups"]["items"]["properties"]["drawing_ids"]["items"][
        "enum"
    ] == ["D0001"]
    ids = ["D0001"]

    def model(*args):
        return LlmAssistResult(
            used_llm=True,
            model="test",
            data={
                "groups": [
                    {
                        "document_id": "d1",
                        "fragment_ids": ["R0001"],
                        "evidence_ids": ["S1:A1"],
                        "drawing_ids": ids,
                        "reason": "source",
                    }
                ]
            },
        )

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", model)
    groups = group_semantic_fragments(review, classification)
    assert groups[0]["drawing_evidence"][0]["object_id"] == "1"
    ids = ["D0001", "D0001"]
    with pytest.raises(ReviewedExcelError, match="重复"):
        group_semantic_fragments(review, classification)
