from io import BytesIO

import pytest
from openpyxl import Workbook

from app.application.elt.errors import ReviewedExcelError
from app.application.reviewed_excel_formula_support import plan_formula_support


def workbook(values):
    book = Workbook()
    sheet = book.active
    sheet.title = "表"
    for address, value in values.items():
        sheet[address] = value
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def test_dependency_closure_preserves_formula_graph_and_limits_support_to_required_cells():
    body = workbook({"A1": "=B1+C1", "B1": "=D1*2", "C1": 3, "D1": 4, "E1": "无关"})
    result = plan_formula_support(body, [{"sheet": "表", "range": "A1:C1"}])
    assert result["support_cells"] == ["D1"]
    assert result["selected_reference_cells"] == ["B1", "C1"]
    assert result["formula_references"] == {"A1": ["B1", "C1"], "B1": ["D1"]}
    assert "E1" not in result["support_cells"]


def test_cycles_external_refs_unsupported_functions_and_large_ranges_fail_closed():
    part = [{"sheet": "表", "range": "A1:A1"}]
    # Mutual cycles are quarantined (isolate), not whole-segment fail-closed.
    mutual = plan_formula_support(
        workbook({"A1": "=B1", "B1": "=A1", "C1": 1}), [{"sheet": "表", "range": "A1:C1"}]
    )
    assert mutual["circular_quarantine"]
    assert {item["action"] for item in mutual["circular_quarantine"]} <= {
        "isolate_as_text",
        "rewrite_exclude_self",
    }
    with pytest.raises(ReviewedExcelError, match="外部"):
        plan_formula_support(workbook({"A1": "='[外部.xlsx]表'!A1"}), part)
    with pytest.raises(ReviewedExcelError, match="函数依赖"):
        plan_formula_support(workbook({"A1": '=INDIRECT("B1")'}), part)
    with pytest.raises(ReviewedExcelError, match="范围超出安全预算"):
        plan_formula_support(workbook({"A1": "=SUM(B1:D1)"}), part, max_cells=2)


def test_self_inclusive_sum_rewritten_and_segment_exports():
    from openpyxl import load_workbook

    from app.application.reviewed_excel_template import derive_document_workbook

    body = workbook(
        {
            "A1": "购货明细",
            "C2": 1,
            "C3": 2,
            "C4": 3,
            "C5": "=SUM(C2:C5)",
            "D2": 10,
        }
    )
    plan = plan_formula_support(body, [{"sheet": "表", "range": "A1:D5"}])
    assert plan["circular_quarantine"]
    assert any(item["action"] == "rewrite_exclude_self" for item in plan["circular_quarantine"])
    assert plan["circular_rewrites"]["C5"] == "=SUM(C2:C4)"
    result = derive_document_workbook(body, [{"sheet": "表", "range": "A1:D5"}])
    saved = load_workbook(__import__("io").BytesIO(result), data_only=False)
    assert saved["表"]["C5"].value == "=SUM(C2:C4)"


def test_support_sheet_keeps_only_needed_values_and_links_back_to_editable_main_cells():
    from openpyxl import load_workbook

    from app.application.reviewed_excel_formula_support import install_formula_support

    original = workbook({"A1": "=SUM(B1:D1)", "B1": 2, "C1": 3, "D1": 4, "E1": "不相关"})
    fragments = [{"sheet": "表", "range": "A1:B1"}]
    plan = plan_formula_support(original, fragments)
    assert plan["support_cells"] == ["C1", "D1"] and plan["proxy_cells"] == ["B1"]
    book = load_workbook(BytesIO(original))
    name = install_formula_support(book, plan)
    assert name == "依赖支撑"
    assert book["表"]["A1"].value == "=SUM('依赖支撑'!B1:D1)"
    assert book[name]["B1"].value == "='表'!B1"
    assert book[name]["C1"].value == 3 and book[name]["D1"].value == 4
    assert book[name]["E1"].value is None
    assert book.calculation.fullCalcOnLoad and book.calculation.forceFullCalc
    stream = BytesIO()
    book.save(stream)
    saved = load_workbook(BytesIO(stream.getvalue()))
    assert saved["表"]["A1"].value == "=SUM('依赖支撑'!B1:D1)"
    assert saved[name]["B1"].value == "='表'!B1"
    assert saved[name]["E1"].value is None


def test_actual_derivation_preserves_needed_formula_support_only():
    from openpyxl import load_workbook

    from app.application.reviewed_excel_template import derive_document_workbook

    original = workbook({"A1": "=SUM(B1:D1)", "B1": 2, "C1": 3, "D1": 4, "E1": "不相关"})
    result = derive_document_workbook(original, [{"sheet": "表", "range": "A1:B1"}])
    saved = load_workbook(BytesIO(result), data_only=False)
    assert saved.sheetnames == ["表", "依赖支撑"]
    assert saved["表"]["A1"].value == "=SUM('依赖支撑'!B1:D1)"
    assert saved["依赖支撑"]["B1"].value == "='表'!B1"
    assert saved["依赖支撑"]["C1"].value == 3
    assert saved["依赖支撑"]["D1"].value == 4
    assert saved["依赖支撑"]["E1"].value is None
    assert saved["表"]["C1"].value is None


def test_subtotal_nested_formulas_keep_original_coordinates_across_fragment_union():
    from openpyxl import load_workbook

    from app.application.reviewed_excel_formula_support import (
        install_formula_support,
        verify_preserved_subtotals,
    )

    original = workbook(
        {
            "A1": 2,
            "A2": 3,
            "A3": "=SUBTOTAL(9,A1:A2)",
            "A4": "=SUBTOTAL(9,$A$1:$A$3)",
            "B1": "=C1",
            "C1": 7,
        }
    )
    fragments = [{"sheet": "表", "range": "A1:B2"}, {"sheet": "表", "range": "A3:A4"}]
    plan = plan_formula_support(original, fragments)
    assert plan["support_cells"] == ["C1"]
    assert set(plan["preserved_subtotal_cells"]) == {"A1", "A2", "A3", "A4"}
    book = load_workbook(BytesIO(original))
    install_formula_support(book, plan)
    output = BytesIO()
    book.save(output)
    saved = load_workbook(BytesIO(output.getvalue()))
    verify_preserved_subtotals(saved, plan)
    assert saved["表"]["A4"].value == "=SUBTOTAL(9,$A$1:$A$3)"
    assert saved["表"]["A3"].value == "=SUBTOTAL(9,A1:A2)"
    assert saved["表"]["B1"].value == "='依赖支撑'!C1"


@pytest.mark.parametrize(
    "formula",
    [
        "=SUBTOTAL(109,A1:A2)",
        "=SUBTOTAL(9,A1:B1)",
        "=SUBTOTAL(9,A1,A2)",
        "=SUBTOTAL(9,A1)",
        "=SUBTOTAL(9,表!A1:A2)",
        '=SUBTOTAL(9,INDIRECT("A1:A2"))',
        "=SUBTOTAL(9,A1:A2)+1",
        "=SUM(SUBTOTAL(9,A1:A2))",
        "=SUBTOTAL(9,#REF!)",
    ],
)
def test_unsupported_subtotal_forms_fail_closed(formula):
    with pytest.raises(ReviewedExcelError):
        plan_formula_support(
            workbook({"A1": 2, "A2": 3, "A3": formula}), [{"sheet": "表", "range": "A1:A3"}]
        )


@pytest.mark.parametrize(
    "values, selected",
    [
        ({"A1": 1, "A2": 2, "A3": "=SUBTOTAL(9,A1:A2)"}, "A2:A3"),
        ({"A1": "=B1", "A2": 2, "A3": "=SUBTOTAL(9,A1:A2)", "B1": 1}, "A1:A3"),
        ({"A1": 1, "A2": 2, "A3": "=SUBTOTAL(9,A1:A2)", "B3": "=A3"}, "B3:B3"),
    ],
)
def test_subtotal_cannot_use_support_cells_even_transitively(values, selected):
    with pytest.raises(ReviewedExcelError, match="SUBTOTAL依赖必须完整保留"):
        plan_formula_support(workbook(values), [{"sheet": "表", "range": selected}])


@pytest.mark.parametrize("change", ["value", "formula", "hidden", "filter", "sheet"])
def test_subtotal_export_guard_rejects_changed_source_semantics(change):
    from openpyxl import load_workbook

    from app.application.reviewed_excel_formula_support import verify_preserved_subtotals

    book = load_workbook(
        BytesIO(workbook({"A1": 1, "A2": "=SUBTOTAL(9,A1:A1)", "A3": "=SUBTOTAL(9,A1:A2)"}))
    )
    sheet = book["表"]
    sheet.row_dimensions[1].hidden = True
    sheet.auto_filter.ref = "A1:A3"
    sheet.auto_filter.add_filter_column(0, ["1"])
    output = BytesIO()
    book.save(output)
    plan = plan_formula_support(output.getvalue(), [{"sheet": "表", "range": "A1:A3"}])
    verify_preserved_subtotals(book, plan)
    if change == "value":
        sheet["A1"] = 99
    elif change == "formula":
        sheet["A2"] = "=SUM(A1:A1)"
    elif change == "hidden":
        sheet.row_dimensions[1].hidden = False
    elif change == "filter":
        sheet.auto_filter.ref = "A1:A2"
    else:
        sheet.title = "不同表"
    with pytest.raises(ReviewedExcelError, match="SUBTOTAL"):
        verify_preserved_subtotals(book, plan)


def test_preexisting_error_operands_are_allowed_with_warning_receipt():
    """Dead #REF! in source SUM formulas must not hard-reject derivation."""
    body = workbook(
        {
            "A1": "=SUM(#REF!)",
            "A2": 1,
            "B1": "=SUM(A2,#REF!)",
            "C1": 9,
        }
    )
    result = plan_formula_support(body, [{"sheet": "表", "range": "A1:B1"}])
    warnings = result["formula_error_operand_warnings"]
    assert {item["cell"] for item in warnings} == {"A1", "B1"}
    assert all(item["message"].startswith("公式含错误引用") for item in warnings)
    assert result["formula_references"]["A1"] == []
    assert result["formula_references"]["B1"] == ["A2"]
    assert result["selected_reference_cells"] == []
    assert result["support_cells"] == ["A2"]


def test_derive_surfaces_error_operand_warnings_on_receipt():
    from app.application.reviewed_excel_template import derive_document_workbook

    original = workbook({"A1": "=SUM(#REF!)", "A2": 3, "B1": 1})
    receipt: dict = {}
    result = derive_document_workbook(
        original, [{"sheet": "表", "range": "A1:B1"}], object_receipt=receipt
    )
    assert result
    assert receipt["formula_error_operand_warnings"][0]["cell"] == "A1"
    assert receipt["formula_error_operand_warnings"][0]["formula"] == "=SUM(#REF!)"
