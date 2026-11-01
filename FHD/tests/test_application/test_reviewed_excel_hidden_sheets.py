"""Hidden source sheets remain archived, but are not standalone template choices."""

import hashlib
import json
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.application.elt.errors import ReviewedExcelError
from app.application.reviewed_excel_classification import classify_review_workbook
from app.application.reviewed_excel_template import derive_document_workbook, inspect_workbook
from tests.test_application.test_reviewed_excel_template import review_store as review_store


def hidden_first_workbook(state):
    workbook = Workbook()
    hidden = workbook.active
    hidden.title = "Internal"
    hidden.append(["Synthetic internal", 1])
    hidden.sheet_state = state
    visible = workbook.create_sheet("Visible")
    visible.append(["Synthetic visible", 2])
    workbook.active = 1
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


@pytest.mark.parametrize("state", ["hidden", "veryHidden"])
def test_classification_marks_hidden_unprocessable_without_omitting_or_unhiding(state):
    body = hidden_first_workbook(state)
    info = inspect_workbook(body)
    sheets = classify_review_workbook(info)
    assert [sheet["sheet"] for sheet in sheets] == ["Internal", "Visible"]
    assert sheets[0]["sheet_state"] == state and sheets[0]["processable"] is False
    assert "隐藏工作表" in sheets[0]["processing_message"]
    assert sheets[1]["sheet_state"] == "visible" and sheets[1]["processable"] is True
    with pytest.raises(ReviewedExcelError, match="隐藏工作表"):
        derive_document_workbook(body, [{"sheet": "Internal", "range": "A1:B1"}])
    derived = derive_document_workbook(body, [{"sheet": "Visible", "range": "A1:B1"}])
    assert load_workbook(BytesIO(derived)).sheetnames == ["Visible"]
    assert inspect_workbook(body)["sha256"] == hashlib.sha256(body).hexdigest()
    original = load_workbook(BytesIO(body))
    assert original["Internal"].sheet_state == state
    original.close()


def test_archive_classification_exposes_visibility_and_hidden_preview_stays_400(review_store):
    body = hidden_first_workbook("hidden")
    client = review_store.client
    response = client.post(
        "/api/templates/reviewed-excel/archive-classify",
        data={"expected_actor": "1:7", "request_key": "synthetic-hidden-first"},
        files={"file": ("synthetic.xlsx", body)},
    )
    assert response.status_code == 200, response.text
    classification = response.json()
    sheets = classification["sheets"]
    assert sheets[0]["processable"] is False and sheets[1]["processable"] is True
    base = f"/api/templates/reviewed-excel/classifications/{classification['classification_id']}"
    labels = [
        {"sheet": sheet["sheet"], "range": region["range"], "labels": region["labels"]}
        for sheet in sheets
        for region in sheet["regions"]
    ]
    confirmed = client.post(
        base + "/confirm",
        data={"version": classification["version"], "labels_json": json.dumps(labels)},
    )
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["sheets"][0]["processable"] is False
    params = {
        "expected_actor": "1:7",
        "classification_id": classification["classification_id"],
        "classification_version": confirmed.json()["version"],
        "region_sheet": "Internal",
        "region_range": "A1:B1",
    }
    blocked = client.post(
        "/api/templates/reviewed-excel/preview",
        data=params,
        files={"file": ("synthetic.xlsx", body)},
    )
    assert blocked.status_code == 400 and "隐藏工作表" in blocked.text
    assert inspect_workbook(body)["sha256"] == hashlib.sha256(body).hexdigest()
