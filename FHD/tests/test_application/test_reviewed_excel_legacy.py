"""Pure-Python legacy conversion: real fixtures, consent, provenance and limits."""

import hashlib
import json
import struct
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import sessionmaker

from app.application import reviewed_excel_legacy as legacy
from app.application.etl.errors import ReviewedExcelError
from app.application.etl.service_support import dump_json, load_json
from app.db.models.etl import EtlRun, EtlUpload
from app.db.models.etl_archive import EtlArchiveQuota, EtlArchiveVersion
from app.fastapi_routes.reviewed_excel_template import _source
from app.infrastructure.tenant_scope import tenant_scope
from tests.fixtures.etl_template.legacy_biff import fixture, record
from tests.test_application.test_reviewed_excel_template import review_store as _review_store

review_store = _review_store


def consent(body):
    return {"legacy_test_only": True, "legacy_source_sha256": hashlib.sha256(body).hexdigest()}


def convert(body):
    return legacy.prepare_archive_workbook(body, **consent(body))


def archive(store, body=None, limit=1000000):
    body = body or fixture()
    client = store.client
    assert (
        client.post(
            "/api/templates/reviewed-excel/archive-quota",
            data={"expected_actor": "1:7", "limit_bytes": limit},
        ).status_code
        == 200
    )
    response = client.post(
        "/api/templates/reviewed-excel/archive-classify",
        data={"expected_actor": "1:7", "request_key": "legacy-test", **consent(body)},
        files={"file": ("legacy.xls", body)},
    )
    assert response.status_code == 200, response.text
    item = response.json()
    derived = client.get(
        "/api/templates/reviewed-excel/classifications/" + item["classification_id"] + "/source"
    )
    assert derived.status_code == 200
    return item, derived.content


def test_real_python_conversion_and_original_unchanged():
    body = fixture()
    before = hashlib.sha256(body).hexdigest()
    output, review = convert(body)
    assert load_workbook(BytesIO(output)).active["A1"].value == 7
    receipt = review["conversion_receipt"]
    assert receipt["source_sha256"] == before == hashlib.sha256(body).hexdigest()
    assert receipt["derived_sha256"] == hashlib.sha256(output).hexdigest()
    assert receipt["validation"]["nonempty_cells"] == 1
    assert (
        receipt["purpose"] == "data_template_test_only"
        and receipt["print_layout_verified"] is False
    )
    assert review["usage_restrictions"] and receipt["warnings"]


@pytest.mark.parametrize(
    "options",
    [{}, {"legacy_test_only": True}, {"legacy_test_only": True, "legacy_source_sha256": "wrong"}],
)
def test_explicit_consent_is_bound_to_original(options):
    with pytest.raises(ReviewedExcelError, match="明确确认"):
        legacy.prepare_archive_workbook(fixture(), **options)


@pytest.mark.parametrize("opcode", [0x6, 0x2F, 0x1AE, 0x23, 0x5D, 0xEC, 0x1B8, 0x18, 0x1B0, 0xFFFF])
def test_unsupported_active_content_is_rejected_before_writing(opcode, monkeypatch):
    monkeypatch.setattr(legacy, "_derive", lambda *a: pytest.fail("no output before admission"))
    with pytest.raises(ReviewedExcelError):
        convert(fixture(record(opcode)))


@pytest.mark.parametrize("name", ["VBA", "ObjectPool", "EncryptedPackage"])
def test_unknown_stream_rejected(name):
    with pytest.raises(ReviewedExcelError):
        convert(fixture(stream_name=name))


def test_wrong_format_and_huge_grid_rejected(monkeypatch):
    with pytest.raises(ReviewedExcelError):
        legacy.inspect_legacy(b"not XLS")
    body = fixture(record(0x200, struct.pack("<IIHHH", 0, 65536, 0, 256, 0)))
    monkeypatch.setattr(legacy.xlrd, "open_workbook", lambda **kw: pytest.fail("before allocation"))
    with pytest.raises(ReviewedExcelError, match="范围超过"):
        legacy.inspect_legacy(body)


@pytest.mark.parametrize("weight,bold", [(400, False), (700, True)])
def test_font_weight_not_unused_bold_flag_is_authoritative(weight, bold):
    output, _ = convert(fixture(font_weight=weight))
    assert load_workbook(BytesIO(output)).active["A1"].font.b is bold


def test_unrepresentable_font_weight_rejected():
    with pytest.raises(ReviewedExcelError, match="字重"):
        convert(fixture(font_weight=600))


def test_formula_looking_text_remains_literal():
    output, _ = convert(fixture(text="=1+1"))
    cell = load_workbook(BytesIO(output), data_only=False).active["A1"]
    assert cell.value == "=1+1" and cell.data_type == "s"


def test_custom_paper_has_specific_warning_not_print_acceptance():
    setup = struct.pack("<HHhHHHHHddH", 256, 100, 1, 1, 1, 2, 203, 203, 0.17, 0.3, 1)
    output, review = convert(fixture(record(0xA1, setup)))
    assert load_workbook(BytesIO(output)).active.page_setup.paperSize == 256
    assert any("256" in warning for warning in review["conversion_receipt"]["warnings"])
    assert review["conversion_receipt"]["print_layout_verified"] is False


def test_loss_during_serialization_is_rejected(monkeypatch):
    original = legacy._derive

    def changed(*args):
        body, styles = original(*args)
        wb = load_workbook(BytesIO(body))
        wb.active["A1"] = 8
        out = BytesIO()
        wb.save(out)
        return out.getvalue(), styles

    monkeypatch.setattr(legacy, "_derive", changed)
    with pytest.raises(ReviewedExcelError, match="单元格值"):
        convert(fixture())


def test_archive_separate_hashes_quota_and_persisted_restrictions(review_store):
    item, derived = archive(review_store)
    assert item["original_sha256"] == hashlib.sha256(fixture()).hexdigest()
    assert item["sha256"] == hashlib.sha256(derived).hexdigest()
    assert item["usage_restrictions"]
    original = review_store.client.get(
        "/api/templates/reviewed-excel/archive-versions/" + item["archive_version_id"] + "/download"
    )
    assert (
        original.content == fixture() and "original.xls" in original.headers["content-disposition"]
    )
    with tenant_scope(1), sessionmaker(bind=review_store.engine)() as db:
        run = db.get(EtlRun, item["classification_id"])
        assert _source(db, run, 1, 7) == derived
        assert db.get(EtlArchiveQuota, 1).reserved_bytes == len(fixture()) + len(derived)
        assert db.get(EtlUpload, run.upload_id).expires_at is None
        assert load_json(run.summary_json, {})["usage_restrictions"]
        from fastapi import HTTPException

        with pytest.raises(HTTPException):
            _source(db, run, 1, 99)
        Path(db.get(EtlUpload, run.upload_id).storage_path).write_bytes(b"tampered")
        with pytest.raises(HTTPException):
            _source(db, run, 1, 7)


def test_existing_classification_recovery_needs_no_new_conversion_or_consent(
    review_store, monkeypatch
):
    item, _ = archive(review_store)
    monkeypatch.setattr(
        legacy, "inspect_legacy", lambda *a: pytest.fail("must reuse persisted derivative")
    )
    response = review_store.client.post(
        "/api/templates/reviewed-excel/archive-versions/"
        + item["archive_version_id"]
        + "/classify",
        data={"expected_actor": "1:7"},
    )
    assert (
        response.status_code == 200
        and response.json()["conversion_receipt"] == item["conversion_receipt"]
    )
    replay = review_store.client.post(
        "/api/templates/reviewed-excel/archive-classify",
        data={"expected_actor": "1:7", "request_key": "legacy-test"},
        files={"file": ("legacy.xls", fixture())},
    )
    assert replay.status_code == 200 and replay.json()["sha256"] == item["sha256"]


def test_no_consent_never_creates_archive(review_store):
    response = review_store.client.post(
        "/api/templates/reviewed-excel/archive-classify",
        data={"expected_actor": "1:7", "request_key": "no-consent"},
        files={"file": ("legacy.xls", fixture())},
    )
    assert response.status_code == 400
    with tenant_scope(1), sessionmaker(bind=review_store.engine)() as db:
        assert db.query(EtlArchiveVersion).count() == db.query(EtlRun).count() == 0


def test_http_preview_download_allowed_but_formal_save_blocked_by_server_lineage(review_store):
    item, derived = archive(review_store)
    client = review_store.client
    url = "/api/templates/reviewed-excel/classifications/" + item["classification_id"]
    labels = [
        {"sheet": s["sheet"], "range": r["range"], "labels": ["quotation"]}
        for s in item["sheets"]
        for r in s["regions"]
    ]
    confirmed = client.post(
        url + "/confirm",
        data={
            "expected_actor": "1:7",
            "version": item["version"],
            "labels_json": json.dumps(labels),
        },
    )
    assert confirmed.status_code == 200, confirmed.text
    data = {
        "expected_actor": "1:7",
        "classification_id": item["classification_id"],
        "classification_version": confirmed.json()["version"],
        "region_sheet": labels[0]["sheet"],
        "region_range": labels[0]["range"],
    }
    response = client.post(
        "/api/templates/reviewed-excel/preview",
        data=data,
        files={"file": ("derived.xlsx", derived)},
    )
    assert response.status_code == 200, response.text
    preview = response.json()
    assert preview["usage_restrictions"]
    download = client.get(
        f"/api/templates/reviewed-excel/reviews/{preview['review_id']}/versions/{preview['proposal']['version']}/candidates/preserve/download"
    )
    assert download.status_code == 200
    # Remove mutable latest state to prove the immutable server summary is enough.
    with tenant_scope(1), sessionmaker(bind=review_store.engine)() as db:
        run = db.get(EtlRun, item["classification_id"])
        state = load_json(run.draft_json, {})
        state["versions"][-1].pop("usage_restrictions", None)
        state["versions"][-1].pop("conversion_receipt", None)
        run.draft_json = dump_json(state)
        db.commit()
    saved = client.post(
        "/api/templates/reviewed-excel/save",
        data={
            "review_id": preview["review_id"],
            "proposal_version": preview["proposal"]["version"],
            "candidate_id": "preserve",
            "selections": "[]",
            "name": "must not publish",
            "publish_to_tenant": "true",
        },
    )
    assert saved.status_code == 409 and "打印版式未验证" in saved.text
    recent = client.get("/api/templates/reviewed-excel/classifications").json()["files"][0]
    assert (
        recent["usage_restrictions"]
        and recent["conversion_receipt"]["purpose"] == "data_template_test_only"
    )


def test_normal_xlsx_path_unchanged():
    from tests.fixtures.etl_template.sources.generate import sample

    body = sample()
    out, review = legacy.prepare_archive_workbook(body)
    assert out == body and not review["usage_restrictions"] and "conversion_receipt" not in review


@pytest.mark.parametrize(
    "field", ["defaultRowHeight", "baseColWidth", "scale", "left", "fitToPage"]
)
def test_layout_scalar_loss_rejected(field, monkeypatch):
    original = legacy._derive

    def changed(*args):
        body, styles = original(*args)
        wb = load_workbook(BytesIO(body))
        sheet = wb.active
        if field in {"defaultRowHeight", "baseColWidth"}:
            setattr(sheet.sheet_format, field, 19)
        elif field == "scale":
            sheet.page_setup.scale = 87
        elif field == "left":
            sheet.page_margins.left = 1.9
        else:
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
        out = BytesIO()
        wb.save(out)
        return out.getvalue(), styles

    monkeypatch.setattr(legacy, "_derive", changed)
    setup = struct.pack("<HHhHHHHHddH", 9, 100, 1, 1, 1, 2, 203, 203, 0.17, 0.3, 1)
    records = (
        record(0xA1, setup)
        + record(0x26, struct.pack("<d", 0.61))
        + record(0x225, struct.pack("<HH", 0, 300))
        + record(0x55, struct.pack("<H", 8))
    )
    with pytest.raises(ReviewedExcelError, match="改变"):
        convert(fixture(records))


def test_boolean_type_is_preserved_not_stringified():
    body = fixture(record(0x205, struct.pack("<HHHBB", 0, 0, 0, 1, 0)))
    output, _ = convert(body)
    value = load_workbook(BytesIO(output)).active["A1"].value
    assert value is True


def test_error_cell_rejected_before_derive(monkeypatch):
    body = fixture(record(0x205, struct.pack("<HHHBB", 0, 0, 0, 7, 1)))
    monkeypatch.setattr(legacy, "_derive", lambda *a: pytest.fail("reject before output"))
    with pytest.raises(ReviewedExcelError, match="错误值"):
        convert(body)


def test_default_hidden_rows_preserved_and_extra_spacing_rejected():
    body = fixture(record(0x225, struct.pack("<HH", 2, 300)))
    output, _ = convert(body)
    assert load_workbook(BytesIO(output)).active.sheet_format.zeroHeight is True
    with pytest.raises(ReviewedExcelError, match="额外留白"):
        convert(fixture(record(0x225, struct.pack("<HH", 4, 300))))


@pytest.mark.parametrize(
    "extra",
    [
        record(0xE5, struct.pack("<HHHHH", 1, 0, 65535, 0, 255)),
        record(0x203, struct.pack("<HHHd", 19999, 0, 0, 1))
        + record(0x203, struct.pack("<HHHd", 0, 99, 0, 2)),
        record(0x208, struct.pack("<HHHHHHHH", 65535, 0, 1, 300, 0, 0, 0, 0)),
    ],
)
def test_merge_combined_axes_and_row_dimensions_rejected_before_allocation(extra, monkeypatch):
    monkeypatch.setattr(legacy.xlrd, "open_workbook", lambda **kw: pytest.fail("before allocation"))
    with pytest.raises(ReviewedExcelError, match="范围"):
        convert(fixture(extra))


@pytest.mark.parametrize("flag", [1 << 28, 1 << 29])
def test_explicit_row_extra_spacing_rejected(flag):
    body = fixture(record(0x208, struct.pack("<HHHHHHI", 0, 0, 1, 300, 0, 0, flag)))
    with pytest.raises(ReviewedExcelError, match="显式行额外留白"):
        convert(body)


def test_python_conversion_dependencies_declared_and_frozen_modules_collected():
    import ast
    import tomllib

    root = Path(__file__).resolve().parents[2]
    project = tomllib.loads((root / "pyproject.toml").read_text())
    for name in ("xlrd", "openpyxl"):
        assert any(
            dep.startswith(name + "=") or dep.startswith(name + ">")
            for dep in project["project"]["optional-dependencies"]["server-api"]
        )
        for requirements in ("requirements.txt", "requirements-server-api.txt"):
            assert any(
                line.startswith(name + "=") or line.startswith(name + ">")
                for line in (root / requirements).read_text().splitlines()
            )
    spec = ast.parse((root / "scripts/package/xcagi_backend.spec").read_text())
    collected = set()
    for node in ast.walk(spec):
        if (
            isinstance(node, ast.For)
            and isinstance(node.target, ast.Name)
            and node.target.id == "module"
            and isinstance(node.iter, ast.List)
        ):
            collected.update(
                item.value for item in node.iter.elts if isinstance(item, ast.Constant)
            )
    assert {"xlrd", "openpyxl"} <= collected


@pytest.mark.parametrize(
    "extra",
    [
        record(0x203, struct.pack("<HHHd", 0, 256, 0, 1)),
        record(0xBE, struct.pack("<HHHH", 0, 0, 0, 4)),
        record(0xBE, struct.pack("<HHHH", 0, 4, 0, 0)),
        record(0xBD, struct.pack("<HHHIH", 0, 0, 0, 0, 4)),
        record(0xBD, struct.pack("<HHHIH", 0, 256, 0, 0, 256)),
        record(0x203, struct.pack("<HHH", 0, 0, 0)),
        record(0x200, struct.pack("<IIHHH", 0, 1, 0, 257, 0)),
    ],
)
def test_malformed_cell_records_rejected_before_parser(extra, monkeypatch):
    monkeypatch.setattr(legacy.xlrd, "open_workbook", lambda **kw: pytest.fail("before parser"))
    with pytest.raises(ReviewedExcelError):
        convert(fixture(extra))


def test_xlrd_structural_assertion_is_business_error(monkeypatch):
    def damaged(**kwargs):
        raise AssertionError("malformed BIFF structure")

    monkeypatch.setattr(legacy.xlrd, "open_workbook", damaged)
    with pytest.raises(ReviewedExcelError, match="结构损坏") as caught:
        convert(fixture())
    assert isinstance(caught.value.__cause__, AssertionError)


def test_valid_mulblank_and_mulrk_records_remain_supported():
    blanks = record(0xBE, struct.pack("<HHHHH", 1, 0, 0, 0, 1))
    numbers = record(0xBD, struct.pack("<HHHIHIH", 0, 0, 0, (7 << 2) | 2, 0, (8 << 2) | 2, 1))
    output, _ = convert(fixture(blanks + numbers))
    sheet = load_workbook(BytesIO(output)).active
    assert sheet["A1"].value == 7 and sheet["B1"].value == 8
    assert sheet["A2"].value is None and sheet["B2"].value is None


def test_convert_xls_to_xlsx_for_slice_enters_inspect_path():
    """Slice helper converts without archive consent and feeds inspect/fit_transport."""
    from app.application.reviewed_excel_budget import fit_transport_batches
    from app.application.reviewed_excel_classification import plan_semantic_batches
    from app.application.reviewed_excel_template import inspect_workbook

    body = fixture()
    converted, receipt = legacy.convert_xls_to_xlsx_for_slice(body)
    assert receipt["purpose"] == "slice_pipeline_auto_convert"
    assert receipt["source_sha256"] == hashlib.sha256(body).hexdigest()
    assert receipt["derived_sha256"] == hashlib.sha256(converted).hexdigest()
    assert receipt["print_layout_verified"] is False
    assert receipt["warnings"]
    review = inspect_workbook(converted)
    assert review["sheets"]
    assert len(review["cells"]) >= 1
    batches, _ = plan_semantic_batches(review, max_calls=4, neutral=True)
    fitted = fit_transport_batches(review, batches)
    assert fitted and sum(len(b) for b in fitted) == len(review["cells"])


def test_ensure_slice_workbook_passthrough_and_convert():
    from tests.fixtures.etl_template.sources.generate import sample

    xlsx = sample()
    out, receipt = legacy.ensure_slice_workbook(xlsx)
    assert out == xlsx and receipt is None
    body = fixture()
    converted, receipt = legacy.ensure_slice_workbook(body)
    assert receipt["purpose"] == "slice_pipeline_auto_convert"
    assert converted != body and converted[:2] == b"PK"


def test_slice_convert_still_fail_closed_on_unsupported(monkeypatch):
    monkeypatch.setattr(legacy, "_derive", lambda *a: pytest.fail("no output before admission"))
    with pytest.raises(ReviewedExcelError):
        legacy.convert_xls_to_xlsx_for_slice(fixture(record(0x6)))
