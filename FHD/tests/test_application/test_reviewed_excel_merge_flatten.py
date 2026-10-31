"""Export flattens merges cut by sticky row-band selections (国圣-style)."""

from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.application.reviewed_excel_template import derive_document_workbook


def _body_with_vertical_merges():
    book = Workbook()
    ws = book.active
    ws.title = "借原材料明细"
    ws["A1"] = "标题"
    ws.merge_cells("A1:H1")
    for row in range(3, 8):
        ws.cell(row, 1, value=f"row{row}")
        ws.cell(row, 2, value=row)
    ws["G3"] = "合并备注G"
    ws["H3"] = "合并备注H"
    ws.merge_cells("G3:G7")
    ws.merge_cells("H3:H7")
    buf = BytesIO()
    book.save(buf)
    return buf.getvalue()


def test_parent_range_containing_merges_still_ok():
    body = _body_with_vertical_merges()
    out = derive_document_workbook(
        body, [{"sheet": "借原材料明细", "range": "A1:H7"}], print_area_mode="selected"
    )
    wb = load_workbook(BytesIO(out))
    assert "借原材料明细" in wb.sheetnames


def test_row_band_cutting_vertical_merges_flattens_instead_of_raising():
    body = _body_with_vertical_merges()
    # Sticky ledger style children that previously raised 所选范围切断合并单元格.
    for area in ("A1:H3", "A4:H4", "A5:H5", "A6:H6", "A7:H7"):
        out = derive_document_workbook(
            body, [{"sheet": "借原材料明细", "range": area}], print_area_mode="selected"
        )
        wb = load_workbook(BytesIO(out))
        ws = wb["借原材料明细"]
        # Flattened cells inside the selection keep the former merge corner value.
        if area == "A4:H4":
            assert ws["G4"].value == "合并备注G"
            assert ws["H4"].value == "合并备注H"
        wb.close()
