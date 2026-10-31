"""Synthetic merge-metadata coverage: identities, never a count-only claim."""

import json
from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from app.application.etl.errors import ReviewedExcelError
from app.application.reviewed_excel_classification import (
    _classify_semantic_batch,
    combine_semantic_batches,
)
from app.application.reviewed_excel_grouping import prepare_grouping
from app.application.reviewed_excel_template import inspect_workbook


def merged_source(count=4):
    book = Workbook()
    sheet = book.active
    sheet.title = "Synthetic"
    for row in range(1, count + 1):
        sheet.cell(row, 1, "synthetic")
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    stream = BytesIO()
    book.save(stream)
    body = stream.getvalue()
    return body, inspect_workbook(body)


def classify_part(monkeypatch, review, cells):
    from app.application import reviewed_excel_template as template

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    rows = [int(cell["cell"][1:]) for cell in cells]
    raw = {
        "documents": [
            {
                "ranges": [
                    {
                        "sheet_name": "Synthetic",
                        "row_start": min(rows),
                        "row_end": max(rows),
                        "column_start": 1,
                        "column_end": 2,
                    }
                ],
                "reason": "synthetic evidence",
            }
        ]
    }
    monkeypatch.setattr(
        template,
        "_template_mimo_complete",
        lambda *args: SimpleNamespace(data=deepcopy(raw), model="synthetic"),
    )
    return _classify_semantic_batch(
        {**review, "cells": cells, "full_cells": review["cells"]}, reference_ids=True
    )


def test_101_merges_are_sent_and_exactly_covered(monkeypatch):
    _, review = merged_source(101)
    prepared = _classify_semantic_batch(review, prepare_only=True, reference_ids=True)
    payload = json.loads(prepared["messages"][1]["content"])
    source_merges = review["sheet_overview"][0]["merged_regions"]
    assert len(payload["sheets"][0]["merged"]) == 101
    assert set(payload["sheets"][0]["merged"]) == set(source_merges)
    assert payload["sheets"][0]["merged_omitted"] == 0
    part = classify_part(monkeypatch, review, review["cells"])
    assert set(part["coverage"][0]["merged_ranges"]) == set(source_merges)
    combined = combine_semantic_batches(review, [review["cells"]], [part], 1)
    coverage = combined["coverage"][0]
    assert coverage["merged_sent"] == 101
    assert coverage["merged_omitted"] == 0
    assert coverage["unread"] == coverage["unassigned"] == []


@pytest.mark.parametrize("duplicate", [False, True])
def test_merge_coverage_uses_identity_union_across_batches(monkeypatch, duplicate):
    _, review = merged_source()
    batches = [review["cells"][:2], review["cells"][2:]]
    parts = [classify_part(monkeypatch, review, batch) for batch in batches]
    merges = review["sheet_overview"][0]["merged_regions"]
    if not duplicate:
        parts[0]["coverage"][0]["merged_ranges"] = merges[:2]
        parts[1]["coverage"][0]["merged_ranges"] = merges[2:]
    combined = combine_semantic_batches(review, batches, parts, 1)
    assert combined["coverage"][0]["merged_sent"] == 4
    assert combined["coverage"][0]["merged_omitted"] == 0


def test_missing_merge_identity_remains_omitted_and_blocks_grouping(monkeypatch):
    _, review = merged_source()
    part = classify_part(monkeypatch, review, review["cells"])
    part["coverage"][0]["merged_ranges"] = part["coverage"][0]["merged_ranges"][:-1]
    # The stale count must not hide the missing identity.
    part["coverage"][0]["merged_sent"] = 4
    combined = combine_semantic_batches(review, [review["cells"]], [part], 1)
    assert combined["coverage"][0]["merged_sent"] == 3
    assert combined["coverage"][0]["merged_omitted"] == 1
    with pytest.raises(ReviewedExcelError, match="覆盖不完整"):
        prepare_grouping(review, combined)


def test_unknown_merge_identity_is_rejected(monkeypatch):
    _, review = merged_source()
    part = classify_part(monkeypatch, review, review["cells"])
    part["coverage"][0]["merged_ranges"].append("C1:D1")
    with pytest.raises(ReviewedExcelError):
        combine_semantic_batches(review, [review["cells"]], [part], 1)


def test_legacy_count_without_identities_cannot_claim_merge_coverage(monkeypatch):
    _, review = merged_source()
    part = classify_part(monkeypatch, review, review["cells"])
    part["coverage"][0].pop("merged_ranges")
    part["coverage"][0]["merged_sent"] = 4
    combined = combine_semantic_batches(review, [review["cells"]], [part], 1)
    assert combined["coverage"][0]["merged_sent"] == 0
    assert combined["coverage"][0]["merged_omitted"] == 4


def test_complete_merge_metadata_over_96k_is_rejected_before_provider(monkeypatch):
    from app.application import reviewed_excel_template as template
    from app.application.reviewed_excel_budget import RequestBudgetExceeded

    _, review = merged_source(1)
    # Legal, disjoint source geometry; only one nonempty cell. The rejection must
    # be the complete request envelope, not the 500-cell or per-cell guards.
    review["sheet_overview"][0]["rows"] = 5000
    review["sheet_overview"][0]["merged_regions"] = [f"A{row}:B{row}" for row in range(1, 5001)]
    assert len(review["cells"]) == 1
    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    calls = []

    def forbidden_provider(*args):
        calls.append(args)
        pytest.fail("oversized metadata must never reach the provider")

    monkeypatch.setattr(template, "_template_mimo_complete", forbidden_provider)
    for prepare_only in (True, False):
        with pytest.raises(RequestBudgetExceeded, match="完整模型请求超过安全预算"):
            _classify_semantic_batch(review, prepare_only=prepare_only, reference_ids=True)
    assert calls == []
