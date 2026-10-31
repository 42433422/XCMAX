"""Synthetic range-rule boundaries; no customer workbook or business values."""

from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import Rule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.styles.differential import DifferentialStyle

from app.application.etl.errors import ReviewedExcelError
from app.application.reviewed_excel_range_objects import (
    _WPS_FOLLOW,
    S,
    plan_range_objects,
    verify_range_objects,
)
from app.application.reviewed_excel_template import derive_document_workbook
from tests.test_application.test_reviewed_excel_template import review_store as _review_store

review_store = _review_store


def rewrite(body, part, change):
    out = BytesIO()
    with ZipFile(BytesIO(body)) as source, ZipFile(out, "w") as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == part:
                root = ET.fromstring(data)
                change(root)
                data = ET.tostring(root)
            target.writestr(entry, data)
    return out.getvalue()


def sample(*, area="A1:A4 B1:B4", filter_area="A1:C4", duplicate=True, outside=True, formula=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "Synthetic"
    for row in sheet.iter_rows(min_row=1, max_row=8, min_col=1, max_col=4):
        for cell in row:
            cell.value = cell.row
    style = DifferentialStyle(
        font=Font(color="FFFF0000"), fill=PatternFill("solid", fgColor="FFFFFF00")
    )
    rule = Rule(
        type="expression" if formula else "duplicateValues",
        dxf=style,
        priority=3,
        formula=["A1>0"] if formula else [],
    )
    sheet.conditional_formatting.add(area, rule)
    if duplicate:
        sheet.conditional_formatting.add(area, Rule(type="duplicateValues", dxf=style, priority=7))
    if outside:
        sheet.conditional_formatting.add(
            "D6:D8", Rule(type="duplicateValues", dxf=style, priority=9)
        )
    if filter_area:
        sheet.auto_filter.ref = filter_area
    out = BytesIO()
    book.save(out)
    if not filter_area:
        return out.getvalue()

    def wps(root):
        af = root.find(f"{{{S}}}autoFilter")
        af.set(_WPS_FOLLOW, "0")
        ET.SubElement(af, f"{{{S}}}extLst")

    return rewrite(out.getvalue(), "xl/worksheets/sheet1.xml", wps)


def fragments(*ranges):
    return [{"sheet": "Synthetic", "range": r} for r in ranges]


def test_complete_union_preserves_duplicate_rules_priorities_dxf_and_wps_filter():
    body = sample()
    selection = fragments("A1:C2", "A3:C4")
    receipt = {}
    output = derive_document_workbook(body, selection, object_receipt=receipt)
    verify_range_objects(output, plan_range_objects(body, "Synthetic", selection))
    sheet = load_workbook(BytesIO(output)).active
    rules = [r for cf in sheet.conditional_formatting for r in sheet.conditional_formatting[cf]]
    assert [r.priority for r in rules] == [3, 7]
    assert all(r.type == "duplicateValues" and r.dxf.font.color.rgb == "FFFF0000" for r in rules)
    assert receipt["excluded"] == [{"type": "duplicateValues", "ranges": ["D6:D8"], "priority": 9}]
    assert [r["type"] for r in receipt["retained"]] == [
        "duplicateValues",
        "duplicateValues",
        "autoFilter",
    ]
    with ZipFile(BytesIO(output)) as archive:
        af = ET.fromstring(archive.read("xl/worksheets/sheet1.xml")).find(f"{{{S}}}autoFilter")
        assert af.attrib == {"ref": "A1:C4", _WPS_FOLLOW: "0"}
        assert len(af) == 1 and af[0].tag == f"{{{S}}}extLst"
    assert sheet["D6"].value is None


@pytest.mark.parametrize(
    "selection", [fragments("A1:C3"), fragments("A1:A4"), fragments("A1:C1", "A3:C4")]
)
def test_partial_multirange_comparison_or_hole_is_not_clipped(selection):
    with pytest.raises(ReviewedExcelError, match="跨越|真实数据|空白上下文"):
        derive_document_workbook(sample(), selection)


def test_filter_crosses_boundary_even_when_conditional_rule_is_inside():
    with pytest.raises(ReviewedExcelError, match="未选尾列|空白上下文|真实数据"):
        derive_document_workbook(sample(area="A1:A4", filter_area="A1:D4"), fragments("A1:C4"))


def test_fully_outside_filter_and_rules_are_explicitly_excluded():
    receipt = {}
    body = sample(area="D6:D8", filter_area="A6:D8", outside=False)
    output = derive_document_workbook(body, fragments("A1:C4"), object_receipt=receipt)
    sheet = load_workbook(BytesIO(output)).active
    assert not sheet.auto_filter.ref and not len(sheet.conditional_formatting)
    assert receipt["retained"] == []
    assert len(receipt["excluded"]) == 3
    assert receipt["excluded"][-1] == {"type": "autoFilter", "ranges": ["A6:D8"]}


def test_formula_rule_stays_rejected():
    with pytest.raises(ReviewedExcelError, match="无公式"):
        derive_document_workbook(sample(formula=True), fragments("A1:C4"))


@pytest.mark.parametrize(
    "kind", ["rule_extension", "filter_extension", "filter_attribute", "filter_condition"]
)
def test_unknown_extensions_and_unverified_filter_conditions_rejected(kind):
    def change(root):
        af = root.find(f"{{{S}}}autoFilter")
        if kind == "rule_extension":
            ET.SubElement(
                root.find(f"{{{S}}}conditionalFormatting/{{{S}}}cfRule"), f"{{{S}}}extLst"
            )
        elif kind == "filter_extension":
            ET.SubElement(af[0], f"{{{S}}}ext", uri="unknown")
        elif kind == "filter_attribute":
            af.set("{urn:unknown}option", "1")
        else:
            ET.SubElement(af, f"{{{S}}}filterColumn", colId="0")

    body = rewrite(sample(), "xl/worksheets/sheet1.xml", change)
    with pytest.raises(ReviewedExcelError):
        plan_range_objects(body, "Synthetic", fragments("A1:C4"))


@pytest.mark.parametrize("kind", ["rule_priority", "filter_property", "style", "missing_rule"])
def test_final_readback_detects_xml_and_dxf_tampering(kind):
    body = sample()
    selection = fragments("A1:C4")
    plan = plan_range_objects(body, "Synthetic", selection)
    output = derive_document_workbook(body, selection)

    def change(root):
        if kind == "style":
            root.find(f"{{{S}}}dxfs/{{{S}}}dxf/{{{S}}}font/{{{S}}}color").set("rgb", "FF00FF00")
        elif kind == "filter_property":
            root.find(f"{{{S}}}autoFilter").attrib.pop(_WPS_FOLLOW)
        elif kind == "missing_rule":
            cf = root.find(f"{{{S}}}conditionalFormatting")
            cf.remove(cf[0])
        else:
            root.find(f"{{{S}}}conditionalFormatting/{{{S}}}cfRule").set("priority", "1")

    mutated = rewrite(
        output, "xl/styles.xml" if kind == "style" else "xl/worksheets/sheet1.xml", change
    )
    with pytest.raises(ReviewedExcelError, match="读回"):
        verify_range_objects(mutated, plan)


def test_dtd_and_duplicate_zip_members_rejected_before_rule_read():
    body = sample()
    with ZipFile(BytesIO(body)) as source:
        out = BytesIO()
        with ZipFile(out, "w") as target:
            for entry in source.infolist():
                target.writestr(entry, source.read(entry.filename))
            with pytest.warns(UserWarning, match="Duplicate"):
                target.writestr("xl/styles.xml", source.read("xl/styles.xml"))
    with pytest.raises(ReviewedExcelError, match="ZIP"):
        plan_range_objects(out.getvalue(), "Synthetic", fragments("A1:C4"))


def empty_tail_sample(mutation=None):
    book = load_workbook(BytesIO(sample(filter_area="A1:D4")))
    sheet = book.active
    for row in range(1, 5):
        sheet.cell(row, 4).value = None
    if mutation == "value":
        sheet["D2"] = "synthetic"
    elif mutation == "formula":
        sheet["D2"] = "=1+1"
    elif mutation == "style":
        sheet["D2"].font = Font(bold=True)
    elif mutation == "merge":
        sheet.merge_cells("D1:D2")
    elif mutation == "row_style":
        sheet.row_dimensions[2].font = Font(bold=True)
    elif mutation == "row_border":
        sheet.row_dimensions[2].border = Border(bottom=Side(style="thin"))
    elif mutation == "row_fill":
        sheet.row_dimensions[2].fill = PatternFill("solid", fgColor="FFFFFF00")
    elif mutation == "row_alignment":
        sheet.row_dimensions[2].alignment = Alignment(horizontal="center")
    elif mutation == "row_number_format":
        sheet.row_dimensions[2].number_format = "0.00"
    elif mutation == "row_protection":
        sheet.row_dimensions[2].protection = Protection(locked=False)
    elif mutation == "column_style":
        sheet.column_dimensions["D"].font = Font(bold=True)
    elif mutation == "comment":
        from openpyxl.comments import Comment

        sheet["D2"].comment = Comment("synthetic", "test")
    elif mutation == "conditional_rule":
        sheet.conditional_formatting.add("D1:D4", Rule(type="duplicateValues", priority=11))
    out = BytesIO()
    book.save(out)

    def wps(root):
        af = root.find(f"{{{S}}}autoFilter")
        af.set(_WPS_FOLLOW, "0")
        ET.SubElement(af, f"{{{S}}}extLst")

    return rewrite(out.getvalue(), "xl/worksheets/sheet1.xml", wps)


def test_filter_only_verified_empty_tail_preserves_node_without_expanding_business_range():
    source = empty_tail_sample()
    chosen = fragments("A1:C4")
    receipt = {}
    output = derive_document_workbook(source, chosen, object_receipt=receipt)
    verify_range_objects(output, plan_range_objects(source, "Synthetic", chosen))
    item = next(x for x in receipt["retained"] if x["type"] == "autoFilter")
    assert item["ranges"] == ["A1:D4"] and item["verified_empty_cells"] == 4
    assert item["blank_tail_bounds"] == [4, 1, 4, 4]
    assert "without_expanding_business_fragments" in item["reason"]
    assert chosen == fragments("A1:C4")
    before, after = load_workbook(BytesIO(source)).active, load_workbook(BytesIO(output)).active
    assert dict(before.sheet_format) == dict(after.sheet_format)
    assert str(after.auto_filter.ref) == "A1:D4"
    assert all(after.cell(row, 4).value is None for row in range(1, 5))


@pytest.mark.parametrize(
    "mutation",
    [
        "value",
        "formula",
        "style",
        "merge",
        "row_border",
        "row_fill",
        "row_alignment",
        "row_number_format",
        "row_protection",
        "column_style",
        "comment",
        "conditional_rule",
    ],
)
def test_empty_filter_tail_exception_rejects_any_meaningful_dependency(mutation):
    with pytest.raises(ReviewedExcelError, match="尾列|空白上下文|绘图"):
        derive_document_workbook(empty_tail_sample(mutation), fragments("A1:C4"))


def test_filter_tail_proof_rejects_real_drawing_object():
    from tests.test_application.test_reviewed_excel_drawing import synthetic_workbook_with_shape

    body = synthetic_workbook_with_shape()

    def shift_drawing(root):
        for node in root.iter():
            if node.tag.endswith("}col"):
                node.text = str(int(node.text) + 3)

    body = rewrite(body, "xl/drawings/drawing1.xml", shift_drawing)

    def add_filter(root):
        ET.SubElement(root, f"{{{S}}}autoFilter", ref="A1:E2")

    body = rewrite(body, "xl/worksheets/sheet1.xml", add_filter)
    with pytest.raises(ReviewedExcelError, match="绘图"):
        plan_range_objects(body, "表", [{"sheet": "表", "range": "A1:C2"}])


def test_filter_tail_dimensions_tampering_is_rejected():
    source = empty_tail_sample()
    chosen = fragments("A1:C4")
    plan = plan_range_objects(source, "Synthetic", chosen)
    output = derive_document_workbook(source, chosen)

    def change(root):
        root.find(f"{{{S}}}sheetFormatPr").set("defaultColWidth", "20")

    with pytest.raises(ReviewedExcelError, match="尺寸读回"):
        verify_range_objects(rewrite(output, "xl/worksheets/sheet1.xml", change), plan)


def test_subtotal_export_keeps_recursive_dependencies_hidden_filter_and_wps_metadata():
    book = load_workbook(BytesIO(empty_tail_sample()))
    sheet = book.active
    sheet["C4"] = "=SUBTOTAL(9,A1:A3)"
    sheet["A3"] = "=A1+A2"
    sheet.row_dimensions[2].hidden = True
    out = BytesIO()
    book.save(out)

    def wps(root):
        af = root.find(f"{{{S}}}autoFilter")
        af.set(_WPS_FOLLOW, "0")
        ET.SubElement(af, f"{{{S}}}extLst")

    source = rewrite(out.getvalue(), "xl/worksheets/sheet1.xml", wps)
    output = derive_document_workbook(source, fragments("A1:C4"))
    saved = load_workbook(BytesIO(output)).active
    assert saved["C4"].value == "=SUBTOTAL(9,A1:A3)"
    assert saved["A3"].value == "=A1+A2"
    assert saved.row_dimensions[2].hidden is True
    assert saved.auto_filter.ref == "A1:D4"
    verify_range_objects(output, plan_range_objects(source, "Synthetic", fragments("A1:C4")))


def test_cleaning_is_explicitly_blocked_only_for_raw_filter_metadata():
    from app.application.reviewed_excel_template import clean_workbook, inspect_workbook

    source = sample()
    derived = derive_document_workbook(source, fragments("A1:C4"))
    with pytest.raises(ReviewedExcelError, match="切片可下载；清理/模板保存尚未支持"):
        clean_workbook(derived, inspect_workbook(derived)["sha256"], [])
    ordinary = load_workbook(BytesIO(source))
    out = BytesIO()
    ordinary.save(out)  # Synthetic control without WPS attributes/extension.
    normal = out.getvalue()
    assert clean_workbook(normal, inspect_workbook(normal)["sha256"], [])


def test_dtd_xml_is_rejected_as_business_error():
    from app.application.reviewed_excel_range_objects import _xml

    with pytest.raises(ReviewedExcelError, match="XML"):
        _xml(b'<!DOCTYPE test [<!ENTITY x "data">]><test>&x;</test>')


def test_archive_auto_slice_http_download_returns_verified_metadata_and_receipt(review_store):
    from sqlalchemy.orm import sessionmaker

    from app.application.etl.service_support import dump_json, load_json
    from app.db.models.etl import EtlRun
    from app.infrastructure.tenant_scope import tenant_scope
    from tests.test_application.test_reviewed_excel_legacy import archive

    body = sample()
    item, _ = archive(review_store, body)
    with tenant_scope(1), sessionmaker(bind=review_store.engine)() as db:
        run = db.get(EtlRun, item["classification_id"])
        state = load_json(run.draft_json, {})
        latest = state["versions"][-1]
        latest["model_status"] = "returned"
        latest["sheets"] = [
            {
                "sheet": "Synthetic",
                "regions": [{"range": "A1:D8", "range_source": "model", "dependencies": []}],
            }
        ]
        latest["coverage"] = [
            {
                "sheet": "Synthetic",
                "total": 32,
                "sent": 32,
                "unread": [],
                "unassigned": [],
                "merged_omitted": 0,
            }
        ]
        latest["suggested_documents"] = [{"fragments": fragments("A1:D8")}]
        run.draft_json = dump_json(state)
        db.commit()
    url = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}/auto-slices/{item['version']}"
    response = review_store.client.get(url)
    assert response.status_code == 200, response.text
    receipt = response.json()["slices"][0]["range_objects_receipt"]
    assert receipt["excluded"] == []
    assert len(receipt["retained"]) == 4
    downloaded = review_store.client.get(url + "/0")
    assert downloaded.status_code == 200, downloaded.text
    verify_range_objects(
        downloaded.content, plan_range_objects(body, "Synthetic", fragments("A1:D8"))
    )
    original = review_store.client.get(
        f"/api/templates/reviewed-excel/archive-versions/{item['archive_version_id']}/download"
    )
    assert original.content == body


def test_inherited_row_font_only_is_kept_with_exact_geometry_and_font_readback():
    body = empty_tail_sample("row_style")
    receipt = {}
    selection = fragments("A1:C4")
    output = derive_document_workbook(body, selection, object_receipt=receipt)
    filter_record = next(r for r in receipt["retained"] if r["type"] == "autoFilter")
    assert filter_record["preserved_row_font_only_rows"] == [2]
    original, saved = load_workbook(BytesIO(body)).active, load_workbook(BytesIO(output)).active
    assert original.row_dimensions[2]._style == saved.row_dimensions[2]._style
    assert ET.tostring(original.row_dimensions[2].font.to_tree()) == ET.tostring(
        saved.row_dimensions[2].font.to_tree()
    )
    plan = plan_range_objects(body, "Synthetic", selection)

    def alter_font(root):
        fonts = root.find(f"{{{S}}}fonts")
        ET.SubElement(fonts[1], f"{{{S}}}i", val="1")

    with pytest.raises(ReviewedExcelError, match="读回"):
        verify_range_objects(rewrite(output, "xl/styles.xml", alter_font), plan)

    def alter_height(root):
        row = next(r for r in root.findall(f"{{{S}}}sheetData/{{{S}}}row") if r.get("r") == "2")
        row.set("ht", "31")
        row.set("customHeight", "1")

    with pytest.raises(ReviewedExcelError, match="读回"):
        verify_range_objects(rewrite(output, "xl/worksheets/sheet1.xml", alter_height), plan)


def test_http_template_preview_refuses_raw_filter_metadata_without_blocking_plain_xlsx(
    review_store,
):
    source = sample()
    response = review_store.client.post(
        "/api/templates/reviewed-excel/preview",
        data={"expected_actor": "1:7"},
        files={"file": ("synthetic.xlsx", source)},
    )
    assert response.status_code == 400
    assert "清理/模板保存尚未支持" in response.text
    wb = load_workbook(BytesIO(source))
    out = BytesIO()
    wb.save(out)
    ordinary = review_store.client.post(
        "/api/templates/reviewed-excel/preview",
        data={"expected_actor": "1:7"},
        files={"file": ("ordinary.xlsx", out.getvalue())},
    )
    assert ordinary.status_code == 200, ordinary.text


def test_conditional_format_blank_context_is_preserved_without_expanding_selection():
    """CF may keep a verified blank right-hand context; business selection stays unchanged."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Synthetic"
    for row in range(1, 5):
        for col in range(1, 4):
            sheet.cell(row, col, row)
        sheet.cell(row, 4).value = None
    style = DifferentialStyle(font=Font(color="FFFF0000"))
    sheet.conditional_formatting.add("A1:D4", Rule(type="duplicateValues", dxf=style, priority=3))
    out = BytesIO()
    book.save(out)
    body = out.getvalue()
    chosen = fragments("A1:C4")
    receipt = {}
    output = derive_document_workbook(body, chosen, object_receipt=receipt)
    verify_range_objects(output, plan_range_objects(body, "Synthetic", chosen))
    retained = [r for r in receipt["retained"] if r["type"] == "duplicateValues"]
    assert retained and retained[0]["ranges"] == ["A1:D4"]
    assert "blank_conditional_context" in retained[0]["reason"]
    assert chosen == fragments("A1:C4")
    saved = load_workbook(BytesIO(output)).active
    assert all(saved.cell(row, 4).value is None for row in range(1, 5))


def test_conditional_format_with_unselected_real_data_is_rejected():
    body = sample(area="A1:D4", filter_area=None, outside=False, duplicate=False)
    with pytest.raises(ReviewedExcelError, match="真实数据|空白上下文|跨越"):
        derive_document_workbook(body, fragments("A1:C4"))
