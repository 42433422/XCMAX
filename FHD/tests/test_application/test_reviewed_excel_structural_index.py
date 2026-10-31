"""Tests for general workbook structural index (no customer hardcodes)."""

from __future__ import annotations

import json
from io import BytesIO

from openpyxl import Workbook

from app.application.reviewed_excel_reference_ids import neutralize_messages
from app.application.reviewed_excel_structural_index import (
    SCHEMA_VERSION,
    attach_structural_index_to_draft,
    build_workbook_structural_index,
    compress_structural_index_for_model,
    prefer_structural_classification_if_undersegmented,
)


def _sample_workbook() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "出货台账"
    ws["A1"] = "厂名"
    ws["B1"] = "日期"
    ws["C1"] = "单号"
    ws["D1"] = "型号"
    # excel-ish serial dates + order keys
    rows = [(45000, 10), (45000, 10), (45001, 11), (45002, 12), (45003, 13), (45003, 13), (45004, 14), (45005, 15), (45006, 16)]
    for i, (serial, order) in enumerate(rows, start=2):
        ws[f"A{i}"] = "甲厂"
        ws[f"B{i}"] = serial
        ws[f"C{i}"] = order
        ws[f"D{i}"] = f"M{i}"
    hidden = wb.create_sheet("报价隐表")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "2026-01-01报价"
    hidden["A5"] = "2026-02-01报价"
    hidden["B2"] = "100"
    period = wb.create_sheet("期间块")
    period["A1"] = "某厂2026年01月01日-01月31日 购货明细"
    period["A10"] = "某厂2026年02月01日-02月28日 购货明细"
    period["B2"] = 1
    xref = wb.create_sheet("汇总")
    xref["A1"] = "=出货台账!C2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_index_schema_is_generic_and_includes_hidden():
    index = build_workbook_structural_index(_sample_workbook())
    assert index["schema_version"] == SCHEMA_VERSION
    assert index["stats"]["sheet_count"] == 4
    assert index["stats"]["hidden_sheet_count"] == 1
    names = {s["sheet"]: s for s in index["sheets"]}
    assert names["报价隐表"]["sheet_state"] == "hidden"
    assert names["报价隐表"]["block_title_rows"]
    assert names["期间块"]["candidate_row_ranges"]
    assert any(r["basis"] == "block_title" for r in names["期间块"]["candidate_row_ranges"])
    assert names["出货台账"]["ledger_run_hint_count"] >= 3
    assert any(r["basis"] == "ledger_key_run" for r in names["出货台账"]["candidate_row_ranges"])
    assert sum(1 for r in names["出货台账"]["candidate_row_ranges"] if r["basis"] == "ledger_key_run") >= 3
    assert index["cross_sheet_edges"]
    blob = json.dumps(index, ensure_ascii=False)
    assert "澜宇" not in blob
    assert "165" not in blob


def test_compress_and_draft_attach():
    index = build_workbook_structural_index(_sample_workbook())
    compact = compress_structural_index_for_model(index, budget_bytes=8_000)
    assert compact["encoded_bytes"] <= 8_000
    draft = attach_structural_index_to_draft({}, index)
    assert draft["workbook_structural_index"]["schema_version"] == SCHEMA_VERSION
    assert draft["workbook_structural_index_full_sha256"] == index["source_sha256"]


def test_neutralize_keeps_structural_index_drops_structure_map():
    payload = {
        "sheets": [{"sheet": "表", "merged": []}],
        "cells": [{"sheet": "表", "cell": "A1", "value": "x", "kind": "review"}],
        "batch_coverage": [{"sheet": "表", "total_nonempty": 1}],
        "structure_map": [{"sheet": "表"}],
        "rule_hints": [],
        "relationship_context": {},
        "labels": ["a"],
        "workbook_structural_index": {"schema_version": SCHEMA_VERSION, "sheets": []},
    }
    messages = [
        {"role": "system", "content": ""},
        {"role": "user", "content": json.dumps(payload, ensure_ascii=False)},
    ]
    neutralize_messages(messages)
    kept = json.loads(messages[1]["content"])
    assert "workbook_structural_index" in kept
    assert kept["workbook_structural_index"]["schema_version"] == SCHEMA_VERSION
    assert "structure_map" not in kept
    assert "rule_hints" not in kept
    assert "labels" not in kept



def test_prefer_structural_when_strictly_more_atoms():
    """Generic rule: promote when struct_n > model_n; keep model when equal or model has more."""
    review = {
        "sheet_overview": [{"sheet": "期间块", "sheet_state": "visible", "merged_regions": []}],
        "workbook_structural_index": {
            "schema_version": SCHEMA_VERSION,
            "sheets": [
                {
                    "sheet": "期间块",
                    "sheet_state": "visible",
                    "columns": 4,
                    "candidate_row_ranges": [
                        {"start_row": 1, "end_row": 9, "basis": "block_title"},
                        {"start_row": 10, "end_row": 20, "basis": "block_title"},
                        {"start_row": 21, "end_row": 30, "basis": "block_title"},
                    ],
                }
            ],
        },
        "sha256": "x",
    }
    # model under-segments (1 region vs 3 structural) → promote
    under = {"regions": [{"sheet": "期间块", "range": "A1:D30"}]}
    promoted = prefer_structural_classification_if_undersegmented(review, under)
    assert promoted.get("workbook_structural_index_source") == "layout_dateblock_fallback"
    assert len(promoted["regions"]) == 3

    # equal counts → keep model
    equal = {
        "regions": [
            {"sheet": "期间块", "range": "A1:D9"},
            {"sheet": "期间块", "range": "A10:D20"},
            {"sheet": "期间块", "range": "A21:D30"},
        ]
    }
    kept = prefer_structural_classification_if_undersegmented(review, equal)
    assert kept is equal

    # model has more → keep model
    over = {
        "regions": [
            {"sheet": "期间块", "range": "A1:D5"},
            {"sheet": "期间块", "range": "A6:D9"},
            {"sheet": "期间块", "range": "A10:D20"},
            {"sheet": "期间块", "range": "A21:D30"},
        ]
    }
    kept2 = prefer_structural_classification_if_undersegmented(review, over)
    assert kept2 is over

    # missing classification → structural
    from_none = prefer_structural_classification_if_undersegmented(review, None)
    assert len(from_none["regions"]) == 3


def test_additional_return_note_absorbs_form_meta():
    """Mid-sheet 退货单 heading + following 购货单位 meta = one block_title atom.

    Keeps 退货单 as a real title word without inventing a phantom extra document
    relative to a sheet that already has a prior delivery-note form.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "客户表"
    ws["A1"] = "成都某公司送货单"
    ws["A2"] = "购货单位（乙方）：甲厂    联系人：张    日期：2026年01月01日   订单编号：26-1"
    ws["A10"] = "销售协议 | 1.甲方所供上列产品均属检验合格产品，送货单一式四联"
    ws["A12"] = "3.本送货单一式四联，白色联为存根联"
    ws["A20"] = "成都某公司退货单"
    ws["A21"] = "购货单位（乙方）：甲厂    联系人：张    日期：2025年08月01日   订单编号：25-1"
    ws["B22"] = "item"
    buf = BytesIO()
    wb.save(buf)
    index = build_workbook_structural_index(buf.getvalue())
    sheet = next(s for s in index["sheets"] if s["sheet"] == "客户表")
    title_cands = [c for c in sheet["candidate_row_ranges"] if c.get("basis") == "block_title"]
    starts = [c["start_row"] for c in title_cands]
    # First form keeps fine grain (送货单 + meta + agreement lines as separate cues).
    assert 1 in starts and 2 in starts
    # Second form: 退货单 present as heading, meta absorbed (no separate start at 21).
    assert 20 in starts
    assert 21 not in starts
    ret = next(c for c in title_cands if c["start_row"] == 20)
    assert ret["end_row"] >= 21
