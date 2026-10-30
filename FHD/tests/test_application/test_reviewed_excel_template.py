import hashlib
import json
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, text
from sqlalchemy.orm import sessionmaker

from app.application.reviewed_excel_template import clean_workbook, inspect_workbook
from app.fastapi_routes import reviewed_excel_template as routes
from app.infrastructure.auth.dependencies import get_logged_in_user
from tests.fixtures.etl_template.sources.generate import large_quotation, sample


def input_digest(messages):
    return hashlib.sha256(
        json.dumps(messages, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode(
            "utf-8"
        )
    ).hexdigest()


def test_clean_formula_workbook_without_calculation_properties():
    wb = Workbook()
    wb.active["A1"] = 12
    wb.active["B1"] = "=A1*2"
    wb.calculation = None  # calcPr is optional in a valid XLSX workbook.
    output = BytesIO()
    wb.save(output)
    wb.close()
    source = output.getvalue()
    info = inspect_workbook(source)
    cleaned = clean_workbook(source, info["sha256"], [{"sheet": "Sheet", "cell": "A1"}])
    result = load_workbook(BytesIO(cleaned))
    try:
        assert result.active["A1"].value is None
        assert result.active["B1"].value == "=A1*2"
        assert result.calculation.fullCalcOnLoad is True
        assert result.calculation.forceFullCalc is True
    finally:
        result.close()


def test_clean_copy_preserves_layout_and_ambiguous_cells():
    fixture = Path(__file__).parents[1] / "fixtures" / "etl_template"
    expected = json.loads((fixture / "expected" / "quotation.json").read_text())
    source = (fixture / "sources" / "synthetic_quotation.xlsx").read_bytes()
    info = inspect_workbook(source)
    selected = [{"sheet": expected["sheet"], "cell": cell} for cell in expected["instance_cells"]]
    cleaned = clean_workbook(source, info["sha256"], selected)
    a, b = load_workbook(BytesIO(source)), load_workbook(BytesIO(cleaned))
    layout = info["layout"][0]
    assert layout["widths"][1] == 27 and layout["rows"][2]["height"] == 31
    assert layout["rows"][0]["cells"][0]["colspan"] == 3
    assert layout["rows"][0]["cells"][0]["bold"] is True
    ws, original = b.active, a.active
    assert ws["B2"].value is None and ws["B3"].value is None
    assert ws["A1"].value == original["A1"].value and ws["A4"].value == "模糊说明"
    assert ws["C3"].value == "=B3*2" and str(ws.merged_cells) == str(original.merged_cells)
    assert ws.column_dimensions["B"].width == 27 and ws.row_dimensions[3].height == 31
    assert ws["B3"].number_format == "0.00" and ws.print_area == original.print_area
    assert ws["A1"].font.bold is True
    assert ws.print_options.horizontalCentered and original["B3"].value == 42
    with pytest.raises(ValueError):
        clean_workbook(source, info["sha256"], [{"sheet": "报价", "cell": "C3"}])
    with pytest.raises(ValueError):
        clean_workbook(source, "wrong", selected)


def test_archived_document_rebuild_is_byte_stable_across_save_times(monkeypatch):
    from datetime import datetime, timezone
    from zipfile import ZipFile

    from app.application.reviewed_excel_template import derive_document_workbook

    original = sample()
    fragments = [{"sheet": "报价", "range": "A1:C4"}]
    clock = [datetime(2026, 10, 29, 1, 2, 4)]
    observed = []

    class SaveClock:
        @staticmethod
        def now(tz=None):
            return clock[0].replace(tzinfo=tz)

    def zip_time(*args):
        observed.append(clock[0])
        return clock[0].timetuple()

    monkeypatch.setattr(
        "openpyxl.writer.excel.datetime", SimpleNamespace(datetime=SaveClock, timezone=timezone)
    )
    monkeypatch.setattr("zipfile.time", SimpleNamespace(localtime=zip_time, time=lambda: 0))
    first = derive_document_workbook(original, fragments)
    clock[0] = datetime(2030, 12, 31, 23, 58, 58)
    second = derive_document_workbook(original, fragments)
    assert len(set(observed)) == 2
    assert first == second
    with ZipFile(BytesIO(second)) as archived:
        assert all(item.date_time == (1980, 1, 1, 0, 0, 0) for item in archived.infolist())


def test_dates_hidden_sheets_print_and_formula_cache_are_explicit():
    from datetime import datetime
    from zipfile import ZIP_DEFLATED, ZipFile

    wb = Workbook()
    ws = wb.active
    ws.title = "报价"
    ws.merge_cells("A1:B1")
    ws["A1"] = "固定说明"
    ws["A2"] = datetime(2026, 10, 28)
    ws["A2"].number_format = "yyyy-mm-dd"
    ws["B2"] = 42
    ws["C2"] = "=B2*2"
    ws.column_dimensions["B"].width = 25
    ws.row_dimensions[2].height = 29
    ws.print_area = "A1:C2"
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = "landscape"
    hidden = wb.create_sheet("固定规则")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "不可删除的条款"
    wb.calculation.fullCalcOnLoad = False
    raw = BytesIO()
    wb.save(raw)
    wb.close()
    # openpyxl does not calculate formulas: create a realistic source cache.
    original = BytesIO()
    with (
        ZipFile(BytesIO(raw.getvalue())) as source_zip,
        ZipFile(original, "w", ZIP_DEFLATED) as target_zip,
    ):
        for item in source_zip.infolist():
            data = source_zip.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"<f>B2*2</f><v></v>", b"<f>B2*2</f><v>84</v>")
            target_zip.writestr(item, data)
    source = original.getvalue()
    assert load_workbook(BytesIO(source), data_only=True).active["C2"].value == 84
    info = inspect_workbook(source)
    result = clean_workbook(source, info["sha256"], [{"sheet": "报价", "cell": "B2"}])
    unchanged = load_workbook(BytesIO(source), data_only=False)
    cleaned = load_workbook(BytesIO(result), data_only=False)
    out = cleaned["报价"]
    assert unchanged["报价"]["B2"].value == 42 and out["B2"].value is None
    assert out["A2"].value == datetime(2026, 10, 28) and out["A2"].number_format == "yyyy-mm-dd"
    assert out["C2"].value == "=B2*2" and str(out.merged_cells) == "A1:B1"
    assert out.column_dimensions["B"].width == 25 and out.row_dimensions[2].height == 29
    assert out.print_area == unchanged["报价"].print_area and out.print_options.horizontalCentered
    assert out.page_setup.orientation == "landscape"
    assert (
        cleaned["固定规则"].sheet_state == "hidden"
        and cleaned["固定规则"]["A1"].value == "不可删除的条款"
    )
    assert load_workbook(BytesIO(source), data_only=True).active["C2"].value == 84
    assert load_workbook(BytesIO(result), data_only=True).active["C2"].value is None
    assert cleaned.calculation.fullCalcOnLoad is True


def test_template_mimo_is_explicit_scoped_and_single_call(monkeypatch, caplog):
    import json

    from app.application.reviewed_excel_template import make_proposals, refine_proposal
    from app.infrastructure.llm.providers.openai_compatible_provider import XiaomiMimoProvider

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    monkeypatch.setenv("FHD_TEMPLATE_MIMO_BASE_URL", "https://token-plan-cn.xiaomimimo.com/v1")
    monkeypatch.setenv("FHD_TEMPLATE_MIMO_MODEL", "mimo-v2.6-flash")
    monkeypatch.setenv("MIMO_API_KEY", "synthetic-placeholder-not-a-secret")
    monkeypatch.setenv("FHD_ETL_LLM", "on")
    calls = []

    async def fake(self, messages, **kwargs):
        calls.append((self._adapter.model_name, self._adapter._base_url, kwargs))
        candidates = [{"clear": [], "reason": "保留待确认"} for _ in range(3)]
        return {
            "choices": [
                {
                    "finish_reason": "stop",
                    "message": {"content": json.dumps({"candidates": candidates})},
                }
            ]
        }

    monkeypatch.setattr(XiaomiMimoProvider, "chat_completion", fake)
    review = inspect_workbook(sample())
    proposal = refine_proposal(
        review, make_proposals(review)["version"], "生成候选", owner_id=7, initial=True
    )
    assert proposal["model"]["status"] == "returned" and len(proposal["candidates"]) == 3
    assert len(calls) == 1 and calls[0][0] == "mimo-v2.6-flash"
    assert calls[0][1] == "https://token-plan-cn.xiaomimimo.com/v1"
    assert calls[0][2]["thinking"] == {"type": "disabled"}
    assert "synthetic-placeholder-not-a-secret" not in caplog.text

    async def fail(*args, **kwargs):
        calls.append("failed")
        raise RuntimeError("429 synthetic-placeholder-not-a-secret")

    monkeypatch.setattr(XiaomiMimoProvider, "chat_completion", fail)
    with pytest.raises(ValueError, match="未自动重试"):
        refine_proposal(
            review, make_proposals(review)["version"], "生成候选", owner_id=7, initial=True
        )
    assert calls.count("failed") == 1

    assert "synthetic-placeholder-not-a-secret" not in caplog.text

    async def slow(*args, **kwargs):
        import asyncio

        calls.append("timeout")
        await asyncio.sleep(1)

    monkeypatch.setattr(XiaomiMimoProvider, "chat_completion", slow)
    monkeypatch.setattr("app.application.etl.llm_assist.etl_llm_timeout_seconds", lambda: 0.01)
    with pytest.raises(ValueError, match="未自动重试"):
        refine_proposal(
            review, make_proposals(review)["version"], "生成候选", owner_id=7, initial=True
        )
    assert calls.count("timeout") == 1
    monkeypatch.setenv("FHD_TEMPLATE_MIMO_BASE_URL", "https://unapproved.invalid/v1")
    with pytest.raises(ValueError, match="未调用"):
        refine_proposal(
            review, make_proposals(review)["version"], "生成候选", owner_id=7, initial=True
        )
    assert calls.count("failed") == 1


def test_template_mimo_preview_requires_permission_and_consent(monkeypatch, review_store):
    import json

    from app.infrastructure.llm.providers.openai_compatible_provider import XiaomiMimoProvider

    for name, value in {
        "FHD_TEMPLATE_LLM_PROVIDER": "mimo",
        "FHD_TEMPLATE_MIMO_BASE_URL": "https://token-plan-cn.xiaomimimo.com/v1",
        "FHD_TEMPLATE_MIMO_MODEL": "mimo-v2.6-flash",
        "MIMO_API_KEY": "synthetic-placeholder",
        "FHD_ETL_LLM": "on",
    }.items():
        monkeypatch.setenv(name, value)
    calls = []

    async def fake(*args, **kwargs):
        calls.append(1)
        candidates = [{"clear": [], "reason": "保留待确认"} for _ in range(3)]
        return {
            "choices": [
                {
                    "finish_reason": "stop",
                    "message": {"content": json.dumps({"candidates": candidates})},
                }
            ]
        }

    monkeypatch.setattr(XiaomiMimoProvider, "chat_completion", fake)
    client = review_store.client
    url = "/api/templates/reviewed-excel/preview"
    files = {"file": ("synthetic.xlsx", sample())}
    data = {"expected_actor": "1:7", "model_opt_in": "true"}
    assert client.post(url, files=files, data=data).status_code == 400 and not calls
    review_store.actor["manage"] = False
    assert (
        client.post(url, files=files, data={**data, "sample_authorized": "true"}).status_code == 403
        and not calls
    )
    review_store.actor["manage"] = True
    response = client.post(url, files=files, data={**data, "sample_authorized": "true"})
    assert response.status_code == 200 and len(calls) == 1
    assert response.json()["proposal"]["model"]["status"] == "returned"


def test_mimo_fenced_three_plans_and_strict_nested_contract(monkeypatch):
    import json

    from app.application.reviewed_excel_template import make_proposals, refine_proposal
    from app.infrastructure.llm.providers.openai_compatible_provider import XiaomiMimoProvider

    for name, value in {
        "FHD_TEMPLATE_LLM_PROVIDER": "mimo",
        "FHD_TEMPLATE_MIMO_BASE_URL": "https://token-plan-cn.xiaomimimo.com/v1",
        "FHD_TEMPLATE_MIMO_MODEL": "mimo-v2.6-flash",
        "MIMO_API_KEY": "synthetic-placeholder",
        "FHD_ETL_LLM": "on",
    }.items():
        monkeypatch.setenv(name, value)
    responses = []
    seen = []

    async def fake(self, messages, **kwargs):
        seen.append(messages)
        return {"choices": [{"finish_reason": "stop", "message": {"content": responses.pop(0)}}]}

    monkeypatch.setattr(XiaomiMimoProvider, "chat_completion", fake)
    review = inspect_workbook(sample())
    version = make_proposals(review)["version"]
    plans = [
        {"clear": [], "reason": "保留"},
        {"clear": [{"sheet": "报价", "cell": "B2"}], "reason": "清除客户"},
        {
            "clear": [{"sheet": "报价", "cell": "B2"}, {"sheet": "报价", "cell": "B3"}],
            "reason": "清除实例",
        },
    ]
    responses.append("```json\n" + json.dumps({"candidates": plans}, ensure_ascii=False) + "\n```")
    output = refine_proposal(review, version, "生成候选", owner_id=7, initial=True)
    assert [p["clear"] for p in output["candidates"]] == [p["clear"] for p in plans]
    sent = json.loads(seen[0][1]["content"])
    assert sent["output_schema"]["properties"]["candidates"]["minItems"] == 3
    assert (
        sent["output_schema"]["properties"]["candidates"]["items"]["additionalProperties"] is False
    )
    assert "independent plans" in seen[0][0]["content"]

    wrong_operation = [
        {
            "sheet": "报价",
            "range": "B2:B3",
            "action": "clear",
            "reason": "示例",
            "equivalent": False,
        }
    ] * 3
    invalid = [
        {"candidates": wrong_operation},
        {"candidates": [{"reason": "缺清除列表"}, *plans[1:]]},
        {"candidates": [{**plans[0], "extra": True}, *plans[1:]]},
        {"candidates": plans, "extra": True},
        {
            "candidates": [
                {"clear": [{"sheet": "报价", "cell": "C3"}], "reason": "公式"},
                *plans[1:],
            ]
        },
        {
            "candidates": [
                {
                    "clear": [{"sheet": "报价", "cell": "B2", "action": "clear"}],
                    "reason": "多余属性",
                },
                *plans[1:],
            ]
        },
    ]
    for item in invalid:
        responses.append(json.dumps(item, ensure_ascii=False))
        with pytest.raises(ValueError):
            refine_proposal(review, version, "生成候选", owner_id=7, initial=True)


@pytest.fixture
def review_store(tmp_path, monkeypatch):
    from fastapi import HTTPException

    from app.application.etl import service_uploads
    from app.db import session as db_session
    from app.db.etl_bootstrap import ensure_sqlite_etl_bootstrap
    from app.domain.shared_services.document_templates import crud
    from app.fastapi_routes.reviewed_excel import archive as archive_routes
    from app.fastapi_routes.template_api import router as product_template_router
    from app.infrastructure.templates import tenant_scope as template_scope

    engine = create_engine(
        f"sqlite:///{tmp_path / 'templates.db'}", connect_args={"check_same_thread": False}
    )
    with engine.begin() as conn:
        conn.execute(
            text(
                "CREATE TABLE templates (id INTEGER PRIMARY KEY, template_key TEXT, template_name TEXT, template_type TEXT, original_file_path TEXT, analyzed_data TEXT, editable_config TEXT, zone_config TEXT, merged_cells_config TEXT, style_config TEXT, business_rules TEXT, is_active INTEGER, tenant_id INTEGER)"
            )
        )
        conn.execute(
            text("CREATE TABLE template_usage_log (template_id INTEGER, action TEXT, result TEXT)")
        )
    ensure_sqlite_etl_bootstrap(engine=engine, swallow_errors=False)
    monkeypatch.setattr(db_session, "SessionLocal", sessionmaker(bind=engine))
    monkeypatch.setattr(crud, "_ensure_template_tables_ready", lambda: None)
    monkeypatch.setattr(template_scope, "ensure_templates_tenant_column", lambda: None)
    monkeypatch.setattr(service_uploads, "get_app_data_dir", lambda: str(tmp_path))
    monkeypatch.setattr(routes, "get_app_data_dir", lambda: str(tmp_path))
    monkeypatch.setattr(archive_routes, "get_app_data_dir", lambda: str(tmp_path))
    actor = {"id": 7, "tenant": 1, "manage": True, "read": True, "admin": True}

    def authorize(permission):
        if not actor[permission]:
            raise HTTPException(status_code=403, detail="权限已撤销")
        return SimpleNamespace(id=actor["id"], tenant_id=actor["tenant"])

    def recheck(_request, expected):
        user = authorize("manage")
        if (user.tenant_id, user.id) != expected:
            raise HTTPException(status_code=403, detail="身份已改变")

    monkeypatch.setattr(routes, "_recheck", recheck)
    monkeypatch.setattr(archive_routes, "_recheck", recheck)

    app = FastAPI()
    app.include_router(product_template_router)
    app.dependency_overrides[routes._manage] = lambda: authorize("manage")
    app.dependency_overrides[routes._read] = lambda: authorize("read")
    app.dependency_overrides[routes._archive_admin] = lambda: authorize("admin")
    with TestClient(app) as client:
        yield SimpleNamespace(client=client, actor=actor, engine=engine, root=tmp_path)
    engine.dispose()


def test_classification_multiple_files_sheets_neutral_and_no_write(review_store):
    wb = Workbook()
    ws = wb.active
    ws.title = "综合"
    ws["A1"] = "报价单"
    ws["A2"] = "采购订单"
    unknown = wb.create_sheet("其他")
    unknown["A1"] = "没有识别标记"
    output = BytesIO()
    wb.save(output)
    body = output.getvalue()
    client = review_store.client
    response = client.post(
        "/api/templates/reviewed-excel/classify",
        files=[
            ("files", ("非报价名称.xlsx", body)),
            ("files", ("客户.xlsx", sample())),
            ("files", ("bad.csv", b"x")),
        ],
        data={"expected_actor": "1:7"},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["files"][0]["sheets"][0]["labels"] == ["quotation", "purchase"]
    assert len(data["files"][0]["sheets"][0]["regions"]) == 1
    assert data["files"][0]["sheets"][0]["regions"][0]["labels"] == ["quotation", "purchase"]
    assert data["files"][0]["sheets"][0]["evidence"][0]["cell"] == "A1"
    assert data["files"][0]["sheets"][1]["labels"] == ["unknown"]
    assert data["files"][1]["sheets"][0]["labels"] == ["unknown"]
    assert data["files"][2]["status"] == "unsupported"
    assert "不表示语义完整理解" in data["scope"]
    with review_store.engine.connect() as conn:
        assert (
            conn.execute(
                text("SELECT count(*) FROM etl_runs WHERE target_type='template_classification'")
            ).scalar_one()
            == 2
        )
        assert conn.execute(text("SELECT count(*) FROM templates")).scalar_one() == 0


def test_production_archive_slices_keep_thirteen_rows_without_business_import(
    review_store, monkeypatch
):
    wb = Workbook()
    ws = wb.active
    ws.title = "生产单"
    for row in [
        ["生产单"],
        ["工艺"],
        ["工序", "说明"],
        ["切割", "合成说明"],
        ["焊接", "合成说明"],
        ["原料"],
        ["材料", "数量"],
        ["合成钢材", 3],
        ["合成板材", 4],
        ["检验说明"],
        ["项目", "要求"],
        ["外观", "合成规范"],
        ["尺寸", "合成公差"],
    ]:
        ws.append(row)
    output = BytesIO()
    wb.save(output)
    body = output.getvalue()
    monkeypatch.setattr(
        "app.application.reviewed_excel_template._template_mimo_complete",
        lambda *_: pytest.fail("未授权模型被调用"),
    )
    with review_store.engine.begin() as conn:
        conn.execute(
            text(
                "INSERT INTO etl_archive_quotas (tenant_id, limit_bytes, reserved_bytes) VALUES (1, 1000000, 0)"
            )
        )
    client = review_store.client
    receipt = client.post(
        "/api/templates/reviewed-excel/archive-classify",
        data={"expected_actor": "1:7", "request_key": "production-thirteen-rows"},
        files={"file": ("生产单.xlsx", body)},
    )
    assert receipt.status_code == 200, receipt.text
    item = receipt.json()
    regions = item["sheets"][0]["regions"]
    assert [(r["range"], r["labels"]) for r in regions] == [
        ("A1:B5", ["production", "process"]),
        ("A6:B9", ["material"]),
        ("A10:B13", ["inspection"]),
    ]
    assert item["model_status"] == "not_called" and not item["confirmed"]
    labels = [{"sheet": "生产单", "range": r["range"], "labels": r["labels"]} for r in regions]
    confirmed = client.post(
        f"/api/templates/reviewed-excel/classifications/{item['classification_id']}/confirm",
        data={"version": item["version"], "labels_json": json.dumps(labels)},
    ).json()
    assert confirmed["confirmed"]
    for index, region in enumerate(regions):
        result = client.post(
            "/api/templates/reviewed-excel/preview",
            data={
                "expected_actor": "1:7",
                "classification_id": item["classification_id"],
                "classification_version": confirmed["version"],
                "region_sheet": "生产单",
                "region_range": region["range"],
            },
            files={"file": ("生产单.xlsx", body)},
        )
        assert result.status_code == 200, result.text
        review = result.json()
        assert len(review["proposal"]["candidates"]) == 3
        saved = client.post(
            "/api/templates/reviewed-excel/save",
            data={
                "review_id": review["review_id"],
                "proposal_version": review["proposal"]["version"],
                "candidate_id": "preserve",
                "selections": "[]",
                "name": f"生产片段{index}",
                "publish_to_tenant": "true",
            },
        )
        assert saved.status_code == 200, saved.text
    original = client.get(
        f"/api/templates/reviewed-excel/archive-versions/{item['archive_version_id']}/download"
    )
    assert original.status_code == 200 and original.content == body
    with review_store.engine.connect() as conn:
        assert conn.execute(text("SELECT count(*) FROM templates")).scalar_one() == 3
        assert (
            conn.execute(
                text(
                    "SELECT count(*) FROM etl_runs WHERE target_type NOT IN ('template_classification', 'template_review')"
                )
            ).scalar_one()
            == 0
        )


def test_production_words_require_title_structure_and_do_not_cut_merges():
    from app.application.reviewed_excel_classification import classify_review_workbook

    wb = Workbook()
    ws = wb.active
    for row in [
        ["字段", "正文"],
        ["生产单", "正文偶然提词"],
        ["工艺", "原料"],
        ["检验说明", "正文描述"],
    ]:
        ws.append(row)
    output = BytesIO()
    wb.save(output)
    result = classify_review_workbook(inspect_workbook(output.getvalue()))[0]
    assert result["labels"] == ["unknown"] and len(result["regions"]) == 1
    ws.delete_rows(1, ws.max_row)
    for row in [
        ["工艺"],
        ["工序", "要求"],
        ["焊接", "共享合并值"],
        ["合成工序", "规范"],
        ["原料"],
        ["材料", "数量"],
        ["合成材料", 1],
    ]:
        ws.append(row)
    ws.merge_cells("B3:B6")
    output = BytesIO()
    wb.save(output)
    result = classify_review_workbook(inspect_workbook(output.getvalue()))[0]
    assert len(result["regions"]) == 1
    assert result["regions"][0]["merged_ranges"] == ["B3:B6"]


@pytest.mark.parametrize("archived", [False, True])
@pytest.mark.parametrize("combine", [False, True])
def test_initial_rule_continuation_requires_explicit_selection_and_replays_state(
    review_store, monkeypatch, archived, combine
):
    wb = Workbook()
    ws = wb.active
    ws.title = "工艺"
    for row in [["工艺说明"], ["工序", "要求"], ["焊接", "合成规范"], ["复核", "合成记录"]] * 2:
        ws.append(row)
    output = BytesIO()
    wb.save(output)
    body = output.getvalue()
    monkeypatch.setattr(
        "app.application.reviewed_excel_template._template_mimo_complete",
        lambda *_: pytest.fail("未授权模型被调用"),
    )
    client = review_store.client
    upload = {"file": ("工艺.xlsx", body)}
    request = {"expected_actor": "1:7", "request_key": "initial-continuation"}
    if archived:
        with review_store.engine.begin() as conn:
            conn.execute(
                text(
                    "INSERT INTO etl_archive_quotas (tenant_id, limit_bytes, reserved_bytes) VALUES (1, 1000000, 0)"
                )
            )
        response = client.post(
            "/api/templates/reviewed-excel/archive-classify", data=request, files=upload
        )
        assert response.status_code == 200, response.text
        item = response.json()
        replay = client.post(
            "/api/templates/reviewed-excel/archive-classify", data=request, files=upload
        )
        assert replay.status_code == 200 and replay.json() == item
    else:
        response = client.post(
            "/api/templates/reviewed-excel/classify",
            data=request,
            files={"files": ("工艺.xlsx", body)},
        )
        assert response.status_code == 200, response.text
        item = response.json()["files"][0]
    assert not item["confirmed"] and item["model_status"] == "not_called"
    regions = item["sheets"][0]["regions"]
    assert [r["range"] for r in regions] == ["A1:B4", "A5:B8"]
    assert regions[1]["continuation_context"]["parent_range"] == "A1:B4"
    links = item["fragment_links"]["items"]
    assert len(links) == 1 and links[0]["repeated_header_observed"] and not links[0]["selected"]
    url = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"
    history = client.get(url).json()
    assert (
        len(history["versions"]) == 1
        and history["versions"][-1]["fragment_links"] == item["fragment_links"]
    )
    labels = [{"sheet": "工艺", "range": r["range"], "labels": r["labels"]} for r in regions]
    edge = [{key: links[0][key] for key in ["sheet", "first", "second"]}] if combine else []
    response = client.post(
        url + "/confirm",
        data={
            "version": item["version"],
            "labels_json": json.dumps(labels),
            "links_json": json.dumps(edge),
            "independent_confirmed": "true",
        },
    )
    assert response.status_code == 200, response.text
    confirmed = response.json()
    assert confirmed["confirmed"] and len(confirmed["logical_documents"]) == int(combine)
    assert confirmed["fragment_links"]["items"][0]["selected"] is combine
    if archived:
        replay = client.post(
            "/api/templates/reviewed-excel/archive-classify", data=request, files=upload
        )
        assert replay.status_code == 200, replay.text
        for key in [
            "version",
            "sheets",
            "confirmed",
            "fragment_links",
            "logical_documents",
            "partition_pending",
            "confirmed_document_count",
            "document_count_status",
        ]:
            assert replay.json()[key] == confirmed[key]
        assert replay.json()["independent_confirmed"] is True
        recovered = client.post(
            f"/api/templates/reviewed-excel/archive-versions/{item['archive_version_id']}/classify",
            data={"expected_actor": "1:7"},
        )
        assert recovered.status_code == 200 and recovered.json() == replay.json()
        original = client.get(
            f"/api/templates/reviewed-excel/archive-versions/{item['archive_version_id']}/download"
        )
        assert original.status_code == 200 and original.content == body
        with review_store.engine.connect() as conn:
            assert conn.execute(text("SELECT count(*) FROM etl_archive_assets")).scalar_one() == 1
            assert conn.execute(text("SELECT count(*) FROM etl_archive_versions")).scalar_one() == 1
            assert (
                conn.execute(
                    text(
                        "SELECT count(*) FROM etl_runs WHERE target_type='template_classification'"
                    )
                ).scalar_one()
                == 1
            )
    preview = client.post(
        "/api/templates/reviewed-excel/preview",
        data={
            "expected_actor": "1:7",
            "classification_id": item["classification_id"],
            "classification_version": confirmed["version"],
            "region_sheet": "工艺",
            "region_range": "A1:B4",
        },
        files=upload,
    )
    assert preview.status_code == 200, preview.text
    value = preview.json()
    assert len(value["selected_region"]["fragments"]) == (2 if combine else 1)
    candidate = client.get(
        f"/api/templates/reviewed-excel/reviews/{value['review_id']}/versions/{value['proposal']['version']}/candidates/preserve/download"
    )
    assert candidate.status_code == 200, candidate.text
    sheet = load_workbook(BytesIO(candidate.content))["工艺"]
    assert sheet["B4"].value == "合成记录"
    assert sheet["B8"].value == ("合成记录" if combine else None)


def test_initial_rule_continuation_identity_conflict_cannot_be_selected(review_store):
    wb = Workbook()
    ws = wb.active
    for number in ["PO-1", "PO-2"]:
        for row in [["工艺说明"], ["工序", "要求"], ["订单号", number], ["合成工序", "规范"]]:
            ws.append(row)
    output = BytesIO()
    wb.save(output)
    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("工艺.xlsx", output.getvalue())},
    ).json()["files"][0]
    link = item["fragment_links"]["items"][0]
    assert link["status"] == "identity_conflict" and not link["selected"]
    labels = [
        {"sheet": s["sheet"], "range": r["range"], "labels": r["labels"]}
        for s in item["sheets"]
        for r in s["regions"]
    ]
    url = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"
    response = client.post(
        url + "/confirm",
        data={
            "version": item["version"],
            "labels_json": json.dumps(labels),
            "links_json": json.dumps([{key: link[key] for key in ["sheet", "first", "second"]}]),
        },
    )
    assert response.status_code == 400 and "身份冲突" in response.text
    assert len(client.get(url).json()["versions"]) == 1


@pytest.mark.parametrize("initial_label", ["template", "unknown"])
@pytest.mark.parametrize("archived", [False, True])
def test_manual_split_requires_new_version_second_confirmation_and_scoped_access(
    review_store, initial_label, archived
):
    wb = Workbook()
    ws = wb.active
    ws.title = "原稿"
    for row in [["合成凭证"], ["品名", "数量"], ["甲", 1], ["乙", 2], ["丙", 3], ["丁", 4]] * 2:
        ws.append(row)
    output = BytesIO()
    wb.save(output)
    body = output.getvalue()
    client = review_store.client
    if archived:
        with review_store.engine.begin() as conn:
            conn.execute(
                text(
                    "INSERT INTO etl_archive_quotas (tenant_id, limit_bytes, reserved_bytes) VALUES (1, 1000000, 0)"
                )
            )
        item = client.post(
            "/api/templates/reviewed-excel/archive-classify",
            data={"expected_actor": "1:7", "request_key": "manual-split-recovery"},
            files={"file": ("原稿.xlsx", body)},
        ).json()
    else:
        item = client.post(
            "/api/templates/reviewed-excel/classify",
            data={"expected_actor": "1:7"},
            files={"files": ("原稿.xlsx", body)},
        ).json()["files"][0]
    url = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"
    labels = [
        {"sheet": "原稿", "original_range": "A1:B12", "range": area, "labels": [initial_label]}
        for area in ["A1:B6", "A7:B12"]
    ]
    data = {
        "version": item["version"],
        "labels_json": json.dumps(labels),
        "links_json": "[]",
        "independent_confirmed": "true",
    }
    review_store.actor["id"] = 8
    assert client.post(url + "/confirm", data=data).status_code == 404
    review_store.actor.update(id=7, tenant=2)
    assert client.post(url + "/confirm", data=data).status_code == 404
    review_store.actor["tenant"] = 1
    response = client.post(url + "/confirm", data=data)
    assert response.status_code == 200, response.text
    split = response.json()
    assert not split["confirmed"] and split["partition_pending"]
    assert "再次确认" in split["message"] and split["logical_documents"] == []
    assert split["confirmed_document_count"] is None
    if archived:
        replay = client.post(
            "/api/templates/reviewed-excel/archive-classify",
            data={"expected_actor": "1:7", "request_key": "manual-split-recovery"},
            files={"file": ("原稿.xlsx", body)},
        )
        assert replay.status_code == 200, replay.text
        for key in [
            "version",
            "sheets",
            "confirmed",
            "partition_pending",
            "fragment_links",
            "logical_documents",
        ]:
            assert replay.json()[key] == split[key]
        assert replay.json()["independent_confirmed"] is False
    assert all(
        "original_range" not in region and "modified" not in region
        for region in split["sheets"][0]["regions"]
    )
    assert client.post(url + "/confirm", data=data).status_code == 409
    preview_data = {
        "expected_actor": "1:7",
        "classification_id": item["classification_id"],
        "classification_version": split["version"],
        "region_sheet": "原稿",
        "region_range": "A1:B6",
    }
    assert (
        client.post(
            "/api/templates/reviewed-excel/preview",
            data=preview_data,
            files={"file": ("原稿.xlsx", body)},
        ).status_code
        == 409
    )
    if initial_label == "unknown":
        assert split["fragment_links"]["items"] == []
        corrections = [
            {"sheet": "原稿", "range": r["range"], "labels": ["template"]}
            for r in split["sheets"][0]["regions"]
        ]
        corrected = client.post(
            url + "/confirm",
            data={"version": split["version"], "labels_json": json.dumps(corrections)},
        )
        assert corrected.status_code == 200, corrected.text
        split = corrected.json()
        assert split["partition_pending"] and not split["confirmed"]
        preview_data["classification_version"] = split["version"]
    candidates = split["fragment_links"]["items"]
    assert (
        len(candidates) == 1
        and candidates[0]["repeated_header_observed"]
        and not candidates[0]["selected"]
    )
    selected = [{key: candidates[0][key] for key in ["sheet", "first", "second"]}]
    labels = [
        {"sheet": "原稿", "range": region["range"], "labels": region["labels"]}
        for region in split["sheets"][0]["regions"]
    ]
    confirmed = client.post(
        url + "/confirm",
        data={
            "version": split["version"],
            "labels_json": json.dumps(labels),
            "links_json": json.dumps(selected),
        },
    )
    assert confirmed.status_code == 200, confirmed.text
    value = confirmed.json()
    assert (
        value["confirmed"]
        and not value["partition_pending"]
        and len(value["logical_documents"]) == 1
    )
    result = client.post(
        "/api/templates/reviewed-excel/preview",
        data={**preview_data, "classification_version": value["version"]},
        files={"file": ("原稿.xlsx", body)},
    )
    assert result.status_code == 200, result.text
    assert len(result.json()["selected_region"]["fragments"]) == 2


@pytest.mark.parametrize(
    "areas,merge,reason",
    [
        (["A1:B5", "A7:B12"], False, "完整覆盖"),
        (["A1:B7", "A7:B12"], False, "重叠"),
        (["A1:C6", "A7:B12"], False, "越界"),
        (["A1:A12", "B1:B12"], True, "合并单元格"),
    ],
)
def test_manual_split_rejects_missing_cells_overlaps_bounds_and_merges(
    review_store, areas, merge, reason
):
    wb = Workbook()
    ws = wb.active
    for row in range(1, 13):
        ws.append([f"合成{row}", row])
    if merge:
        ws.merge_cells("A1:B1")
    output = BytesIO()
    wb.save(output)
    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("原稿.xlsx", output.getvalue())},
    ).json()["files"][0]
    url = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"
    labels = [
        {"sheet": "Sheet", "original_range": "A1:B12", "range": area, "labels": ["template"]}
        for area in areas
    ]
    result = client.post(
        url + "/confirm", data={"version": item["version"], "labels_json": json.dumps(labels)}
    )
    assert result.status_code == 400 and reason in result.text
    assert len(client.get(url).json()["versions"]) == 1


def test_manual_split_checks_overlap_with_other_original_regions():
    from copy import deepcopy

    from app.application.reviewed_excel_classification import (
        apply_classification_corrections,
        classify_review_workbook,
    )

    wb = Workbook()
    ws = wb.active
    for row in [
        ["报价单"],
        ["品名", "数量"],
        ["甲", 1],
        ["乙", 2],
        ["采购订单"],
        ["品名", "数量"],
        ["丙", 3],
        ["丁", 4],
    ]:
        ws.append(row)
    output = BytesIO()
    wb.save(output)
    review = inspect_workbook(output.getvalue())
    latest = {"sheets": classify_review_workbook(review)}
    labels = [
        {"sheet": "Sheet", "original_range": "A1:B4", "range": area, "labels": ["quotation"]}
        for area in ["A1:B2", "A3:B4"]
    ]
    with pytest.raises(ValueError, match="重叠"):
        apply_classification_corrections(
            review,
            deepcopy(latest),
            labels
            + [
                {
                    "sheet": "Sheet",
                    "range": "A4:B8",
                    "original_range": "A5:B8",
                    "labels": ["purchase"],
                }
            ],
        )
    labels[1]["range"] = "A3:B5"
    with pytest.raises(ValueError, match="原区域内"):
        apply_classification_corrections(
            review,
            deepcopy(latest),
            labels
            + [
                {
                    "sheet": "Sheet",
                    "range": "A6:B8",
                    "original_range": "A5:B8",
                    "labels": ["purchase"],
                }
            ],
        )


def test_existing_large_classification_can_confirm_without_expanding_regions(review_store):
    wb = Workbook()
    ws = wb.active
    for _ in range(201):
        for row in [["报价单"], ["品名", "数量"], ["合成商品", 1]]:
            ws.append(row)
    output = BytesIO()
    wb.save(output)
    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("既有多单.xlsx", output.getvalue())},
    ).json()["files"][0]
    regions = item["sheets"][0]["regions"]
    assert len(regions) == 201
    labels = [
        {"sheet": "Sheet", "range": region["range"], "labels": region["labels"]}
        for region in regions
    ]
    url = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}/confirm"
    response = client.post(
        url, data={"version": item["version"], "labels_json": json.dumps(labels)}
    )
    assert response.status_code == 200 and response.json()["confirmed"]
    expanded = [
        {"sheet": "Sheet", "original_range": "A1:B3", "range": area, "labels": ["quotation"]}
        for area in ["A1:B1", "A2:B3"]
    ] + labels[1:]
    rejected = client.post(
        url, data={"version": response.json()["version"], "labels_json": json.dumps(expanded)}
    )
    assert rejected.status_code == 400 and "最多201" in rejected.text


def test_archive_classification_receipt_survives_expired_temporary_upload(review_store):
    from datetime import UTC, datetime, timedelta

    from app.db.models.etl import EtlRun
    from app.infrastructure.tenant_scope import tenant_scope

    client = review_store.client
    body = sample()
    data = {"expected_actor": "1:7", "request_key": "synthetic-request"}
    url = "/api/templates/reviewed-excel/archive-classify"
    assert client.post(url, data=data, files={"file": ("sample.xlsx", body)}).status_code == 400
    with review_store.engine.begin() as conn:
        conn.execute(
            text(
                "INSERT INTO etl_archive_quotas (tenant_id, limit_bytes, reserved_bytes) VALUES (1, 1000000, 0)"
            )
        )
    response = client.post(url, data=data, files={"file": ("sample.xlsx", body)})
    assert response.status_code == 200, response.text
    receipt = response.json()
    blocked = client.post(
        f"/api/templates/reviewed-excel/archive-assets/{receipt['asset_id']}/state",
        data={"expected_actor": "1:7", "state": "trashed"},
    )
    assert blocked.status_code == 409 and "依赖" in blocked.json()["detail"]
    assert receipt["sha256"] == hashlib.sha256(body).hexdigest()
    again = client.post(url, data=data, files={"file": ("sample.xlsx", body)})
    assert (
        again.status_code == 200
        and again.json()["classification_id"] == receipt["classification_id"]
    )
    with review_store.engine.begin() as conn:
        conn.execute(
            text(
                "UPDATE etl_uploads SET expires_at=:expiry WHERE id=(SELECT upload_id FROM etl_runs WHERE id=:id)"
            ),
            {
                "expiry": (datetime.now(UTC) - timedelta(days=1)).isoformat(),
                "id": receipt["classification_id"],
            },
        )
    with tenant_scope(1), sessionmaker(bind=review_store.engine)() as db:
        assert routes._source(db, db.get(EtlRun, receipt["classification_id"]), 1, 7) == body
    source_url = (
        f"/api/templates/reviewed-excel/classifications/{receipt['classification_id']}/source"
    )
    restored = client.get(source_url)
    assert restored.status_code == 200 and restored.content == body
    assert restored.headers["cache-control"] == "no-store"
    review_store.actor["tenant"] = 2
    assert (
        client.get(
            f"/api/templates/reviewed-excel/classifications/{receipt['classification_id']}"
        ).status_code
        == 404
    )
    assert client.get(source_url).status_code == 404
    review_store.actor["tenant"] = 1
    labels = [
        {"sheet": sheet["sheet"], "range": region["range"], "labels": region["labels"]}
        for sheet in receipt["sheets"]
        for region in sheet["regions"]
    ]
    confirm = client.post(
        f"/api/templates/reviewed-excel/classifications/{receipt['classification_id']}/confirm",
        data={"version": receipt["version"], "labels_json": json.dumps(labels)},
    )
    assert confirm.status_code == 200, confirm.text
    region = confirm.json()["sheets"][0]["regions"][0]
    preview = client.post(
        "/api/templates/reviewed-excel/preview",
        data={
            "expected_actor": "1:7",
            "classification_id": receipt["classification_id"],
            "classification_version": confirm.json()["version"],
            "region_sheet": "报价",
            "region_range": region["range"],
        },
        files={"file": ("sample.xlsx", body)},
    )
    assert preview.status_code == 200, preview.text
    result = preview.json()
    with review_store.engine.begin() as conn:
        conn.execute(
            text(
                "UPDATE etl_uploads SET expires_at=:expiry WHERE id=(SELECT upload_id FROM etl_runs WHERE id=:id)"
            ),
            {
                "expiry": (datetime.now(UTC) - timedelta(days=1)).isoformat(),
                "id": result["review_id"],
            },
        )
    candidate = result["proposal"]["candidates"][0]
    recovered = client.get(
        f"/api/templates/reviewed-excel/reviews/{result['review_id']}/versions/{result['proposal']['version']}/candidates/{candidate['id']}/download"
    )
    assert recovered.status_code == 200, recovered.text
    assert load_workbook(BytesIO(recovered.content))["报价"]["C3"].value == "=B3*2"


def test_archive_quota_requires_admin_and_never_drops_below_reserved(review_store):
    client = review_store.client
    url = "/api/templates/reviewed-excel/archive-quota"
    review_store.actor["admin"] = False
    assert client.get(url).status_code == 403
    assert client.post(url, data={"expected_actor": "1:7", "limit_bytes": 1000}).status_code == 403
    review_store.actor["admin"] = True
    assert client.get(url).json()["configured"] is False
    assert (
        client.post(url, data={"expected_actor": "1:7", "limit_bytes": 1000000}).status_code == 200
    )
    body = sample()
    item = client.post(
        "/api/templates/reviewed-excel/archive-classify",
        data={"expected_actor": "1:7", "request_key": "quota-test"},
        files={"file": ("synthetic.xlsx", body)},
    )
    assert item.status_code == 200
    used = client.get(url).json()["reserved_bytes"]
    assert used == len(body)
    assert (
        client.post(url, data={"expected_actor": "1:7", "limit_bytes": used - 1}).status_code == 409
    )


def test_archive_recover_complete_pending_and_reversible_trash(review_store):
    from app.application.etl.archive_store import create_original
    from app.db.models.etl_archive import EtlArchiveQuota, EtlArchiveVersion
    from app.infrastructure.tenant_scope import tenant_scope

    with tenant_scope(1), sessionmaker(bind=review_store.engine)() as db:
        db.add(EtlArchiveQuota(tenant_id=1, limit_bytes=1000000, reserved_bytes=0))
        db.commit()
        item = create_original(
            db,
            owner=7,
            name="sample.xlsx",
            body=sample(),
            root=review_store.root / "etl" / "archive",
            request_key="recover-test",
        )
        asset_id, version_id = item.asset_id, item.id
        db.get(EtlArchiveVersion, version_id).status = "pending"
        db.commit()
    client = review_store.client
    state_url = f"/api/templates/reviewed-excel/archive-assets/{asset_id}/state"
    assert (
        client.post(state_url, data={"expected_actor": "1:7", "state": "trashed"}).status_code
        == 409
    )

    assert (
        client.post(
            f"/api/templates/reviewed-excel/archive-versions/{version_id}/recover",
            data={"expected_actor": "1:7"},
        ).status_code
        == 200
    )
    wb = Workbook()
    wb.active["A1"] = "第二个版本"
    stream = BytesIO()
    wb.save(stream)
    wb.close()
    second_bytes = stream.getvalue()
    version_url = f"/api/templates/reviewed-excel/archive-assets/{asset_id}/versions"
    second = client.post(
        version_url,
        data={"expected_actor": "1:7", "request_key": "second-version"},
        files={"file": ("new.xlsx", second_bytes)},
    )
    assert second.status_code == 200, second.text
    assert second.json()["version"] == 2
    repeated = client.post(
        version_url,
        data={"expected_actor": "1:7", "request_key": "second-version"},
        files={"file": ("new.xlsx", second_bytes)},
    )
    assert (
        repeated.status_code == 200 and repeated.json()["version_id"] == second.json()["version_id"]
    )
    assert (
        client.get(
            f"/api/templates/reviewed-excel/archive-versions/{second.json()['version_id']}/download"
        ).content
        == second_bytes
    )
    review_store.actor["tenant"] = 2
    assert (
        client.get(
            f"/api/templates/reviewed-excel/archive-versions/{second.json()['version_id']}/download"
        ).status_code
        == 404
    )
    review_store.actor["tenant"] = 1
    assert (
        client.post(state_url, data={"expected_actor": "1:7", "state": "trashed"}).json()["state"]
        == "trashed"
    )
    assert (
        client.get("/api/templates/reviewed-excel/archive-assets").json()["assets"][0]["state"]
        == "trashed"
    )
    assert (
        client.post(state_url, data={"expected_actor": "1:7", "state": "active"}).json()["state"]
        == "active"
    )
    classify_url = (
        f"/api/templates/reviewed-excel/archive-versions/{second.json()['version_id']}/classify"
    )
    classified = client.post(classify_url, data={"expected_actor": "1:7"})
    assert classified.status_code == 200, classified.text
    assert classified.json()["sha256"] == hashlib.sha256(second_bytes).hexdigest()
    assert (
        client.post(classify_url, data={"expected_actor": "1:7"}).json()["classification_id"]
        == classified.json()["classification_id"]
    )
    assert (
        client.post(state_url, data={"expected_actor": "1:7", "state": "trashed"}).status_code
        == 409
    )


def test_archive_classify_and_withdraw_do_not_both_succeed(review_store):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier

    from app.application.etl.archive_store import create_original
    from app.db.models.etl_archive import EtlArchiveQuota
    from app.infrastructure.tenant_scope import tenant_scope

    with tenant_scope(1), sessionmaker(bind=review_store.engine)() as db:
        db.add(EtlArchiveQuota(tenant_id=1, limit_bytes=1000000, reserved_bytes=0))
        db.commit()
        item = create_original(
            db,
            owner=7,
            name="sample.xlsx",
            body=sample(),
            root=review_store.root / "etl" / "archive",
            request_key="race-test",
        )
        asset_id, version_id = item.asset_id, item.id
    barrier = Barrier(2)

    def submit(path, data):
        barrier.wait()
        return review_store.client.post(path, data=data).status_code

    with ThreadPoolExecutor(max_workers=2) as pool:
        classify = pool.submit(
            submit,
            f"/api/templates/reviewed-excel/archive-versions/{version_id}/classify",
            {"expected_actor": "1:7"},
        )
        withdraw = pool.submit(
            submit,
            f"/api/templates/reviewed-excel/archive-assets/{asset_id}/state",
            {"expected_actor": "1:7", "state": "trashed"},
        )
        statuses = (classify.result(), withdraw.result())
    assert statuses in {(200, 409), (400, 200), (409, 200)}, statuses


def test_classification_mixed_regions_and_embedded_attachment_are_conservative(review_store):
    from zipfile import ZIP_DEFLATED, ZipFile

    wb = Workbook()
    ws = wb.active
    ws.title = "综合"
    ws["A1"] = "报价单"
    ws["A2"] = "日期"
    ws["B2"] = "报价内容"
    ws["A6"] = "采购订单"
    ws["B7"] = "采购内容"
    ws["A11"] = "考勤表"
    ws["B12"] = "考勤内容"
    ws["A14"] = "共享图例"
    wb.create_sheet("未知")["A1"] = "无法判断"
    out = BytesIO()
    wb.save(out)
    body = out.getvalue()
    merged = Workbook()
    ms = merged.active
    ms.title = "边界"
    ms["A1"] = "报价单"
    ms["A6"] = "采购订单"
    ms.merge_cells("B4:B7")
    ms["B4"] = "跨区身份"
    shared = BytesIO()
    merged.save(shared)
    embedded = BytesIO()
    with ZipFile(BytesIO(body)) as source, ZipFile(embedded, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            target.writestr(item, source.read(item.filename))
        target.writestr("xl/embeddings/object1.bin", b"synthetic never executed")
    response = review_store.client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files=[
            ("files", ("mixed.xlsx", body)),
            ("files", ("embedded.xlsx", embedded.getvalue())),
            ("files", ("merged.xlsx", shared.getvalue())),
        ],
    )
    assert response.status_code == 200
    records = response.json()["files"]
    regions = records[0]["sheets"][0]["regions"]
    assert [r["labels"] for r in regions] == [["quotation"], ["purchase"], ["attendance"]]
    assert [r["range"] for r in regions] == ["A1:B5", "A6:B10", "A11:B14"]
    assert all(r["needs_confirmation"] and "未分析" in r["dependency_status"] for r in regions)
    assert regions[0]["header_cells"] == [{"cell": "A2", "value": "日期"}]
    assert all(any(c["cell"] == "A14" for c in r["possible_shared_context"]) for r in regions)
    assert records[0]["sheets"][1]["regions"][0]["labels"] == ["unknown"]
    assert records[1]["status"] == "unsupported" and "未读取或执行" in records[1]["message"]
    blocked = records[2]["sheets"][0]["regions"]
    assert len(blocked) == 1 and blocked[0]["labels"] == ["quotation", "purchase"]
    assert blocked[0]["merged_ranges"] == ["B4:B7"] and blocked[0]["needs_confirmation"]


def test_semantic_classification_persists_versions_and_links_source(review_store, monkeypatch):
    import json
    from datetime import UTC, datetime, timedelta

    from app.application.etl.llm_assist_runtime import LlmAssistResult

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")

    wb = Workbook()
    ws = wb.active
    ws.title = "混合"
    ws["A1"] = "报价单"
    ws["A2"] = "产品名称"
    ws["B3"] = "样例"
    ws["E1"] = "采购订单"
    ws["E2"] = "供应商"
    ws["G4"] = "样例"
    ws["A5"] = "共享图例"
    wb.create_sheet("未知")["A1"] = "未分类"
    output = BytesIO()
    wb.save(output)
    body = output.getvalue()
    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("mixed.xlsx", body)},
    ).json()["files"][0]
    cid, version = item["classification_id"], item["version"]
    url = f"/api/templates/reviewed-excel/classifications/{cid}"
    sent = []
    sent_messages = []
    regions = [
        {
            "sheet": "混合",
            "range": "A1:C4",
            "labels": ["quotation"],
            "evidence": [{"sheet": "混合", "cell": "A1"}],
            "dependencies": [{"sheet": "混合", "cell": "A5"}],
            "reason": "报价",
        },
        {
            "sheet": "混合",
            "range": "D1:G4",
            "labels": ["purchase"],
            "evidence": [{"sheet": "混合", "cell": "E1"}],
            "dependencies": [{"sheet": "混合", "cell": "A5"}],
            "reason": "采购",
        },
    ]

    def fake(messages, schema, max_tokens):
        sent.append(json.loads(messages[1]["content"]))
        sent_messages.append(messages)
        return LlmAssistResult(used_llm=True, data={"regions": regions})

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    assert client.post(url + "/analyze", data={"version": version}).status_code == 400 and not sent
    response = client.post(url + "/analyze", data={"version": version, "sample_authorized": "true"})
    assert response.status_code == 200, response.text
    analyzed = response.json()
    assert analyzed["model_status"] == "returned" and len(analyzed["sheets"][0]["regions"]) == 3
    assert analyzed["sheets"][0]["regions"][2]["range"] == "A5:G5"
    assert analyzed["sheets"][0]["regions"][2]["labels"] == ["unknown"]
    assert analyzed["sheets"][0]["regions"][0]["dependencies"] == [{"sheet": "混合", "cell": "A5"}]
    assert analyzed["fragment_links"]["items"][0]["status"] == "candidate"
    assert analyzed["sheets"][1]["regions"][0]["labels"] == ["unknown"]
    assert analyzed["coverage"][0]["unassigned"] == ["A5"]
    assert analyzed["input_summaries"][0]["input_sha256"] == input_digest(sent_messages[0])
    assert analyzed["input_summaries"][0]["source_sha256"] == hashlib.sha256(body).hexdigest()
    assert "input_messages" not in json.dumps(client.get(url).json(), ensure_ascii=False)
    with review_store.engine.connect() as conn:
        persisted = conn.execute(
            text("SELECT draft_json FROM etl_runs WHERE id=:id"), {"id": cid}
        ).scalar_one()
    assert '"input_messages"' not in persisted
    assert all(len(c["value"]) <= 160 for c in sent[0]["cells"])
    corrected = [
        {"sheet": sheet["sheet"], "range": region["range"], "labels": region["labels"]}
        for sheet in analyzed["sheets"]
        for region in sheet["regions"]
    ]
    overlap = [dict(entry) for entry in corrected]
    overlap[0] = {**overlap[0], "original_range": "A1:C4", "range": "A1:E4"}
    assert (
        client.post(
            url + "/confirm",
            data={"version": analyzed["version"], "labels_json": json.dumps(overlap)},
        ).status_code
        == 400
    )
    corrected[0] = {**corrected[0], "original_range": "A1:C4", "range": "A1:B4"}
    corrected[0]["labels"] = ["quotation", "template"]
    confirmed = client.post(
        url + "/confirm",
        data={"version": analyzed["version"], "labels_json": json.dumps(corrected)},
    )
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["fragment_links"]["items"] == []
    assert (
        client.post(
            url + "/confirm",
            data={"version": analyzed["version"], "labels_json": json.dumps(corrected)},
        ).status_code
        == 409
    )
    history = client.get(url).json()
    assert len(history["versions"]) == 3 and history["versions"][-1]["sheets"][0]["regions"][0][
        "labels"
    ] == ["quotation", "template"]
    assert history["versions"][-1]["sheets"][0]["regions"][0]["range"] == "A1:B4"
    assert history["versions"][-1]["sheets"][0]["regions"][0]["label_source"] == "user"
    assert history["versions"][-1]["sheets"][0]["regions"][0]["range_source"] == "user"
    assert history["versions"][-1]["coverage"][1]["model_unassigned"] == ["A1"]
    assert (
        client.get("/api/templates/reviewed-excel/classifications").json()["files"][0]["latest"][
            "confirmed"
        ]
        is True
    )
    from app.application.etl.errors import EtlConflict
    from app.application.etl.service import EtlService
    from app.infrastructure.tenant_scope import tenant_scope

    with tenant_scope(1), sessionmaker(bind=review_store.engine)() as db:
        with pytest.raises(EtlConflict):
            EtlService().execute(
                db, run_id=cid, owner_user_id=7, confirmed=True, valid_rows_only=False
            )
    preview_data = {
        "expected_actor": "1:7",
        "classification_id": cid,
        "classification_version": confirmed.json()["version"],
        "region_sheet": "混合",
        "region_range": "A1:B4",
    }
    preview = client.post(
        "/api/templates/reviewed-excel/preview",
        data=preview_data,
        files={"file": ("mixed.xlsx", body)},
    )
    assert preview.status_code == 400 and "依赖外部" in preview.json()["detail"]
    with tenant_scope(1), sessionmaker(bind=review_store.engine)() as db:
        history_ids = [item["id"] for item in EtlService().list_runs(db, owner_user_id=7)]
        assert cid not in history_ids
    assert (
        client.post(
            "/api/templates/reviewed-excel/preview",
            data=preview_data,
            files={"file": ("other.xlsx", sample())},
        ).status_code
        == 409
    )
    review_store.actor["tenant"] = 2
    assert client.get(url).status_code == 404
    review_store.actor["tenant"] = 1
    with review_store.engine.begin() as conn:
        conn.execute(
            text(
                "UPDATE etl_uploads SET expires_at=:expiry WHERE id=(SELECT upload_id FROM etl_runs WHERE id=:id)"
            ),
            {"expiry": (datetime.now(UTC) - timedelta(days=1)).isoformat(), "id": cid},
        )
    assert (
        client.post(
            url + "/analyze",
            data={"version": confirmed.json()["version"], "sample_authorized": "true"},
        ).status_code
        == 410
    )
    assert len(client.get(url).json()["versions"]) == 3


def test_semantic_classification_rejects_unsafe_model_regions(monkeypatch):
    from app.application.etl.llm_assist_runtime import LlmAssistResult
    from app.application.reviewed_excel_template import classify_semantic_workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "报价"
    ws["A1"] = "报价单"
    ws.merge_cells("B2:C2")
    ws["B2"] = "跨列身份"
    ws["D4"] = "图例"
    out = BytesIO()
    wb.save(out)
    review = inspect_workbook(out.getvalue())
    monkeypatch.delenv("FHD_TEMPLATE_LLM_PROVIDER", raising=False)
    with pytest.raises(ValueError, match="未进行模型分类"):
        classify_semantic_workbook(review)
    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    valid = {
        "sheet": "报价",
        "range": "A1:D4",
        "labels": ["quotation"],
        "evidence": [{"sheet": "报价", "cell": "A1"}],
        "dependencies": [],
        "reason": "报价",
    }
    invalid = [
        [{**valid, "range": "A1:B4"}],
        [{**valid, "range": "A1:E4"}],
        [valid, valid],
        [{**valid, "evidence": [{"sheet": "报价", "cell": "Z99"}]}],
        [{**valid, "labels": ["customer_products"]}],
        [{**valid, "extra": True}],
    ]
    for regions in invalid:
        monkeypatch.setattr(
            "app.application.reviewed_excel_template._template_mimo_complete",
            lambda *a, **k: LlmAssistResult(used_llm=True, data={"regions": regions}),
        )
        with pytest.raises(ValueError):
            classify_semantic_workbook(review)


def test_semantic_classification_uses_bounded_fenced_model_output(monkeypatch):
    import json

    from app.application.reviewed_excel_template import classify_semantic_workbook
    from app.infrastructure.llm.providers.openai_compatible_provider import XiaomiMimoProvider

    for name, value in {
        "FHD_TEMPLATE_LLM_PROVIDER": "mimo",
        "FHD_TEMPLATE_MIMO_BASE_URL": "https://token-plan-cn.xiaomimimo.com/v1",
        "FHD_TEMPLATE_MIMO_MODEL": "mimo-v2.6-flash",
        "MIMO_API_KEY": "synthetic-placeholder",
        "FHD_ETL_LLM": "on",
    }.items():
        monkeypatch.setenv(name, value)
    wb = Workbook()
    ws = wb.active
    ws.title = "报价"
    ws["A1"] = "报价单"
    ws["A2"] = "合成长文本" * 40
    for row in range(3, 130):
        ws.cell(row, 1, f"合成数据{row}")
    out = BytesIO()
    wb.save(out)
    review = inspect_workbook(out.getvalue())
    seen = []

    async def fake(self, messages, **kwargs):
        seen.append(json.loads(messages[1]["content"]))
        data = {
            "physical_regions": [
                {
                    "sheet": "报价",
                    "range": "A1:A5",
                    "labels": ["quotation"],
                    "evidence": [{"sheet": "报价", "cell": "A1"}],
                    "dependencies": [],
                    "reason": "报价标题",
                }
            ],
            "logical_documents": [],
            "layout_references": [],
        }
        return {
            "choices": [
                {
                    "finish_reason": "stop",
                    "message": {
                        "content": "```json\n" + json.dumps(data, ensure_ascii=False) + "\n```"
                    },
                }
            ]
        }

    monkeypatch.setattr(XiaomiMimoProvider, "chat_completion", fake)
    result = classify_semantic_workbook(review)
    assert len(seen) == 1 and len(seen[0]["cells"]) == 129
    assert (
        next(cell["value"] for cell in seen[0]["cells"] if cell["cell"] == "A2")
        == "合成长文本" * 40
    )
    assert result["coverage"][0]["total"] == 129 and not result["coverage"][0]["unread"]
    assert not result["coverage"][0]["value_truncated"]
    assert "A129" in result["coverage"][0]["unassigned"]


def test_semantic_stale_response_cannot_overwrite_confirmed_correction(review_store, monkeypatch):
    import json

    from app.application.etl.llm_assist_runtime import LlmAssistResult

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")

    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("sample.xlsx", sample())},
    ).json()["files"][0]
    url = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"
    labels = [
        {"sheet": sheet["sheet"], "range": region["range"], "labels": ["template"]}
        for sheet in item["sheets"]
        for region in sheet["regions"]
    ]

    def fake(*args, **kwargs):
        confirmed = client.post(
            url + "/confirm", data={"version": item["version"], "labels_json": json.dumps(labels)}
        )
        assert confirmed.status_code == 200
        return LlmAssistResult(
            used_llm=True,
            data={
                "regions": [
                    {
                        "sheet": "报价",
                        "range": "A1:C4",
                        "labels": ["unknown"],
                        "evidence": [],
                        "dependencies": [],
                        "reason": "待确认",
                    }
                ]
            },
        )

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    response = client.post(
        url + "/analyze", data={"version": item["version"], "sample_authorized": "true"}
    )
    assert response.status_code == 409
    history = client.get(url).json()["versions"]
    assert len(history) == 2 and history[-1]["confirmed"] is True
    assert history[-1]["sheets"][0]["regions"][0]["labels"] == ["template"]


def test_manual_region_correction_cannot_split_merge(review_store):
    import json

    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("sample.xlsx", sample())},
    ).json()["files"][0]
    region = item["sheets"][0]["regions"][0]
    correction = [
        {
            "sheet": "报价",
            "original_range": region["range"],
            "range": "A1:B4",
            "labels": ["unknown"],
        }
    ]
    url = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"
    response = client.post(
        url + "/confirm", data={"version": item["version"], "labels_json": json.dumps(correction)}
    )
    assert response.status_code == 400 and "合并单元格" in response.text
    assert len(client.get(url).json()["versions"]) == 1


def test_save_replay_scope_expiry_and_history(review_store):
    from datetime import UTC, datetime, timedelta

    store = review_store
    client = store.client
    upload = {
        "file": (
            "sample.xlsx",
            sample(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    store.actor["id"] = 8
    assert (
        client.post(
            "/api/templates/reviewed-excel/preview", files=upload, data={"expected_actor": "1:7"}
        ).status_code
        == 409
    )
    store.actor["id"] = 7
    response = client.post(
        "/api/templates/reviewed-excel/preview", files=upload, data={"expected_actor": "1:7"}
    )
    assert response.status_code == 200, response.text
    preview = response.json()
    review_id = preview["review_id"]
    for candidate, rendered in zip(
        preview["proposal"]["candidates"], preview["proposal"]["result_previews"], strict=True
    ):
        url = f"/api/templates/reviewed-excel/reviews/{review_id}/versions/{preview['proposal']['version']}/candidates/{candidate['id']}/download"
        downloaded = client.get(url)
        assert downloaded.status_code == 200
        assert inspect_workbook(downloaded.content)["layout"] == rendered["layout"]
        workbook = load_workbook(BytesIO(downloaded.content), data_only=False)
        assert workbook.active["C3"].value == "=B3*2"
    store.actor["id"] = 8
    assert client.get(url).status_code == 404
    store.actor["id"] = 7
    data = {
        "review_id": review_id,
        "name": "合成模板",
        "proposal_version": preview["proposal"]["version"],
        "candidate_id": "preserve",
        "publish_to_tenant": "true",
        "selections": '[{"sheet":"报价","cell":"B3"}]',
        "proposal_history": "forged",
    }
    history_url = f"/api/templates/reviewed-excel/reviews/{review_id}"
    before = client.get(history_url).json()
    assert (
        client.post(
            "/api/templates/reviewed-excel/save", data={**data, "publish_to_tenant": "false"}
        ).status_code
        == 400
    )
    assert before["versions"][0]["version"] == data["proposal_version"]
    wrong = {**data, "proposal_version": "0" * 64}
    assert client.post("/api/templates/reviewed-excel/save", data=wrong).status_code == 409
    wrong = {**data, "candidate_id": "forged"}
    assert client.post("/api/templates/reviewed-excel/save", data=wrong).status_code == 409
    wrong = {**data, "selections": '[{"sheet":"报价","cell":"C3"}]'}
    assert client.post("/api/templates/reviewed-excel/save", data=wrong).status_code == 400
    store.actor["id"] = 8
    assert client.get(history_url).status_code == 404
    assert client.post("/api/templates/reviewed-excel/save", data=data).status_code == 404
    store.actor.update(id=7, tenant=2)
    assert client.get(history_url).status_code == 404
    store.actor.update(tenant=1, manage=False)
    assert client.post("/api/templates/reviewed-excel/save", data=data).status_code == 403
    store.actor.update(manage=True, read=False)
    assert client.get(history_url).status_code == 403
    store.actor["read"] = True
    with store.engine.begin() as conn:
        row = conn.execute(
            text(
                "SELECT u.storage_path, u.id FROM etl_uploads u JOIN etl_runs r ON r.upload_id=u.id WHERE r.id=:id"
            ),
            {"id": review_id},
        ).first()
        path = Path(row.storage_path)
        original = path.read_bytes()
        path.write_bytes(b"tampered")
    assert client.post("/api/templates/reviewed-excel/save", data=data).status_code == 409
    path.write_bytes(original)
    path.unlink()
    assert client.post("/api/templates/reviewed-excel/save", data=data).status_code == 410
    path.write_bytes(original)
    with store.engine.begin() as conn:
        conn.execute(
            text("UPDATE etl_uploads SET expires_at=:expiry WHERE id=:id"),
            {"expiry": (datetime.now(UTC) - timedelta(days=1)).isoformat(), "id": row.id},
        )
    assert client.post("/api/templates/reviewed-excel/save", data=data).status_code == 410
    with store.engine.begin() as conn:
        conn.execute(
            text("UPDATE etl_uploads SET expires_at=:expiry WHERE id=:id"),
            {"expiry": (datetime.now(UTC) + timedelta(days=1)).isoformat(), "id": row.id},
        )
    existing_source = store.root / "templates" / "reviewed_excel" / "1" / "preexisting.source"
    existing_source.parent.mkdir(parents=True, exist_ok=True)
    existing_source.write_bytes(b"preexisting synthetic source")
    saved = client.post("/api/templates/reviewed-excel/save", data=data)
    assert saved.status_code == 200, saved.text
    assert list(existing_source.parent.glob("*.source")) == [existing_source]
    assert existing_source.read_bytes() == b"preexisting synthetic source"
    assert client.post("/api/templates/reviewed-excel/save", data=data).json()["replayed"] is True
    assert (
        client.post(
            "/api/templates/reviewed-excel/save", data={**data, "name": "不同名称"}
        ).status_code
        == 409
    )
    receipt = client.get(history_url).json()
    assert receipt["status"] == "template_review_saved" and receipt["confirmation"]["added"] == [
        {"sheet": "报价", "cell": "B3"}
    ]
    assert "forged" not in str(receipt)
    template_id = int(saved.json()["template_id"].split(":")[1])
    download = client.get(f"/api/templates/reviewed-excel/{template_id}/download")
    assert download.status_code == 200
    workbook = load_workbook(BytesIO(download.content))
    assert workbook.active["B3"].value is None and workbook.active["C3"].value == "=B3*2"
    store.actor["id"] = 8
    assert client.get(f"/api/templates/reviewed-excel/{template_id}/download").status_code == 200
    assert client.get(history_url).status_code == 404
    store.actor["read"] = False
    assert client.get(f"/api/templates/reviewed-excel/{template_id}/download").status_code == 403
    store.actor["read"] = True
    store.actor["id"] = 7
    store.actor["tenant"] = 2
    assert client.get(f"/api/templates/reviewed-excel/{template_id}/download").status_code == 404
    with store.engine.connect() as conn:
        assert conn.execute(text("SELECT count(*) FROM templates")).scalar_one() == 1


def test_quotation_proposals_identify_table_and_identity_without_touching_rules():
    from app.application.reviewed_excel_template import make_proposals

    wb = Workbook()
    ws = wb.active
    ws.title = "报价"
    for row in [
        ["客户", "样例公司"],
        ["付款条款", "到货付款"],
        ["产品名称", "数量", "单价", "小计"],
        ["样例A", 2, 3, "=B4*C4"],
        ["合计", None, None, "=D4"],
    ]:
        ws.append(row)
    out = BytesIO()
    wb.save(out)
    review = inspect_workbook(out.getvalue())
    proposals = make_proposals(review)
    assert len(proposals["candidates"]) == 3
    assert proposals["candidates"][0]["clear"] == []
    table = {(item["sheet"], item["cell"]) for item in proposals["candidates"][1]["clear"]}
    extended = {(item["sheet"], item["cell"]) for item in proposals["candidates"][2]["clear"]}
    assert table == {("报价", "A4"), ("报价", "B4"), ("报价", "C4")}
    assert extended == table | {("报价", "B1")}
    assert ("报价", "B2") not in extended and ("报价", "D4") not in extended
    assert proposals["version"] and proposals["model"]["status"] == "not_called"


def test_model_feedback_is_versioned_and_rejects_unknown_or_protected_cells(monkeypatch):
    from app.application import reviewed_excel_template as service
    from app.application.etl.llm_assist_runtime import LlmAssistResult

    body = sample()
    review = service.inspect_workbook(body)
    base = service.make_proposals(review)
    fake = LlmAssistResult(
        used_llm=True, data={"clear": [{"sheet": "报价", "cell": "B2"}], "reason": "测试建议"}
    )
    monkeypatch.setattr("app.application.etl.llm_assist._complete", lambda *args, **kwargs: fake)
    revised = service.refine_proposal(review, base["version"], "清除客户", owner_id=7)
    assert revised["parent_version"] == base["version"] and revised["version"] != base["version"]
    assert revised["candidates"][2]["clear"] == [{"sheet": "报价", "cell": "B2"}]
    assert revised["model"]["status"] == "returned"
    fake.data["clear"] = [{"sheet": "报价", "cell": "C3"}]
    with pytest.raises(ValueError):
        service.refine_proposal(review, base["version"], "清除公式", owner_id=7)
    fake.data["clear"] = [{"sheet": "报价", "cell": "A1"}]
    with pytest.raises(ValueError):
        service.refine_proposal(review, base["version"], "清除固定说明", owner_id=7)
    monkeypatch.setattr(
        "app.application.etl.llm_assist._complete", lambda *args, **kwargs: LlmAssistResult()
    )
    with pytest.raises(ValueError, match="模型未调用"):
        service.refine_proposal(review, base["version"], "清除客户", owner_id=7)


def test_feedback_endpoint_requires_consent_and_uses_stub_model(monkeypatch, review_store):
    from app.application.etl.llm_assist_runtime import LlmAssistResult
    from app.fastapi_routes.reviewed_excel import archive as archive_routes
    from app.fastapi_routes.template_api import router as product_template_router

    calls = []

    def stub(messages, **_kwargs):
        calls.append(messages)
        return LlmAssistResult(
            used_llm=True,
            data={
                "candidates": [
                    {"clear": [], "reason": "保留"},
                    {"clear": [{"sheet": "报价", "cell": "B2"}], "reason": "清理客户"},
                    {"clear": [{"sheet": "报价", "cell": "B3"}], "reason": "清理数量"},
                ]
            },
        )

    monkeypatch.setattr("app.application.etl.llm_assist._complete", stub)
    client = review_store.client
    source = sample()
    upload = {
        "file": (
            "sample.xlsx",
            source,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    preview = client.post(
        "/api/templates/reviewed-excel/preview", files=upload, data={"expected_actor": "1:7"}
    ).json()
    assert preview["proposal"]["model"]["status"] == "not_called" and calls == []
    form = {
        "review_id": preview["review_id"],
        "parent_version": preview["proposal"]["version"],
        "feedback": "清除客户",
    }
    assert (
        client.post("/api/templates/reviewed-excel/refine", files=upload, data=form).status_code
        == 400
    )
    assert calls == []
    form["consent"] = "true"
    form["sample_authorized"] = "true"
    response = client.post("/api/templates/reviewed-excel/refine", files=upload, data=form)
    assert response.status_code == 200
    assert response.json()["proposal"]["model"]["status"] == "returned" and len(calls) == 1
    assert len(response.json()["proposal"]["result_previews"]) == 3
    summary = response.json()["proposal"]["model"]["input_summary"]
    assert summary["input_sha256"] == input_digest(calls[0])
    assert "input_messages" not in response.json()["proposal"]["model"]
    assert response.json()["proposal"]["parent_version"] == preview["proposal"]["version"]
    history = client.get(f"/api/templates/reviewed-excel/reviews/{preview['review_id']}").json()
    assert (
        len(history["versions"]) == 2
        and history["versions"][1]["model"]["input_summary"]["input_sha256"]
        == input_digest(calls[0])
        and "input_messages" not in history["versions"][1]["model"]
    )
    assert client.post("/api/templates/reviewed-excel/refine", data=form).status_code == 409
    assert (
        client.post(
            "/api/templates/reviewed-excel/save",
            data={
                "review_id": preview["review_id"],
                "proposal_version": preview["proposal"]["version"],
                "candidate_id": "preserve",
                "publish_to_tenant": "true",
                "selections": "[]",
                "name": "过期版本",
            },
        ).status_code
        == 409
    )


def test_initial_model_only_for_authorized_samples_and_reports_coverage(monkeypatch, review_store):
    from app.application.etl.llm_assist_runtime import LlmAssistResult
    from app.fastapi_routes.reviewed_excel import archive as archive_routes
    from app.fastapi_routes.template_api import router as product_template_router

    calls = []

    def stub(messages, **kwargs):
        calls.append(messages)
        assert "candidates" in kwargs["schema"]["properties"]
        return LlmAssistResult(
            used_llm=True,
            data={
                "candidates": [
                    {"clear": [], "reason": "保留"},
                    {"clear": [{"sheet": "报价", "cell": "B2"}], "reason": "清理客户"},
                    {"clear": [{"sheet": "报价", "cell": "B3"}], "reason": "清理数量"},
                ]
            },
        )

    monkeypatch.setattr("app.application.etl.llm_assist._complete", stub)
    client = review_store.client
    source = sample()
    upload = {
        "file": (
            "sample.xlsx",
            source,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    assert (
        client.post(
            "/api/templates/reviewed-excel/preview",
            files=upload,
            data={"model_opt_in": "true", "expected_actor": "1:7"},
        ).status_code
        == 400
    )
    assert calls == []
    response = client.post(
        "/api/templates/reviewed-excel/preview",
        files=upload,
        data={"model_opt_in": "true", "sample_authorized": "true", "expected_actor": "1:7"},
    )
    assert response.status_code == 200
    proposal = response.json()["proposal"]
    assert proposal["model"]["status"] == "returned" and len(proposal["candidates"]) == 3
    assert proposal["model"]["sheet_overview"][0]["rows"] == 4
    assert proposal["model"]["input_summary"]["input_sha256"] == input_digest(calls[0])
    assert "input_messages" not in proposal["model"]


def test_owner_file_template_authorization_is_bound_to_mimo(monkeypatch, review_store):
    from app.application.etl.llm_assist_runtime import LlmAssistResult

    calls = []

    def fake(*args, **kwargs):
        calls.append(1)
        return LlmAssistResult(
            used_llm=True, data={"candidates": [{"clear": [], "reason": "合成"} for _ in range(3)]}
        )

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    monkeypatch.delenv("FHD_TEMPLATE_LLM_PROVIDER", raising=False)
    client = review_store.client
    data = {"model_opt_in": "true", "own_file_authorized": "true", "expected_actor": "1:7"}
    source = sample()
    upload = {"file": ("synthetic.xlsx", source)}
    rejected = client.post("/api/templates/reviewed-excel/preview", files=upload, data=data)
    assert rejected.status_code == 400 and "供应商不匹配" in rejected.json()["detail"] and not calls
    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    accepted = client.post("/api/templates/reviewed-excel/preview", files=upload, data=data)
    assert accepted.status_code == 200, accepted.text
    auth = accepted.json()["proposal"]["model"]["authorization"]
    assert (
        auth["kind"] == "owner_file" and auth["provider"] == "mimo" and auth["owner_user_id"] == 7
    )
    assert auth["source_sha256"] == hashlib.sha256(source).hexdigest() and calls == [1]
    refine_data = {
        "review_id": accepted.json()["review_id"],
        "parent_version": accepted.json()["proposal"]["version"],
        "feedback": "继续保留",
        "current": "[]",
        "consent": "true",
        "own_file_authorized": "true",
    }
    assert (
        client.post(
            "/api/templates/reviewed-excel/refine",
            data={**refine_data, "sample_authorized": "true"},
        ).status_code
        == 400
    )
    refined = client.post("/api/templates/reviewed-excel/refine", data=refine_data)
    assert refined.status_code == 200, refined.text
    assert (
        refined.json()["proposal"]["model"]["authorization"]["scope"]
        == "this_document_bounded_cells_and_feedback"
    )


def test_model_input_spreads_across_sheets_and_reports_uncovered(monkeypatch):
    import json

    from app.application import reviewed_excel_template as service
    from app.application.etl.llm_assist_runtime import LlmAssistResult

    wb = Workbook()
    ws = wb.active
    ws.title = "第一页"
    for i in range(120):
        ws.append([f"记录{i}"])
    wb.create_sheet("第二页")["A1"] = "后页证据"
    out = BytesIO()
    wb.save(out)
    review = service.inspect_workbook(out.getvalue())
    initial = service.make_proposals(review)
    seen = []

    def stub(messages, **_kwargs):
        seen.append(json.loads(messages[1]["content"]))
        return LlmAssistResult(used_llm=True, data={"clear": [], "reason": "保留"})

    monkeypatch.setattr("app.application.etl.llm_assist._complete", stub)
    revised = service.refine_proposal(review, initial["version"], "检查范围", owner_id=7)
    assert any(c["sheet"] == "第二页" for c in seen[0]["cells"])
    assert any(c["sheet"] == "第一页" and c["cell"] == "A120" for c in seen[0]["cells"])
    assert sum(item["uncovered"] for item in revised["model"]["sheet_overview"]) == 21
    assert revised["model"]["truncated"] is True


def test_concurrent_identical_confirmations_never_create_two_templates(review_store):
    from concurrent.futures import ThreadPoolExecutor

    client = review_store.client
    source = sample()
    upload = {
        "file": (
            "sample.xlsx",
            source,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    preview = client.post(
        "/api/templates/reviewed-excel/preview", files=upload, data={"expected_actor": "1:7"}
    ).json()
    data = {
        "review_id": preview["review_id"],
        "proposal_version": preview["proposal"]["version"],
        "candidate_id": "preserve",
        "publish_to_tenant": "true",
        "selections": "[]",
        "name": "并发模板",
    }
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(
            pool.map(
                lambda _: client.post("/api/templates/reviewed-excel/save", data=data), range(2)
            )
        )
    assert any(r.status_code == 200 for r in results)
    assert all(r.status_code in (200, 409) for r in results)
    with review_store.engine.connect() as conn:
        assert conn.execute(text("SELECT count(*) FROM templates")).scalar_one() == 1


def test_concurrent_feedback_keeps_first_committed_version(review_store, monkeypatch):
    import hashlib
    import threading
    from concurrent.futures import ThreadPoolExecutor

    from app.application.reviewed_excel_template import make_proposals

    client = review_store.client
    preview = client.post(
        "/api/templates/reviewed-excel/preview",
        files={"file": ("sample.xlsx", sample())},
        data={"expected_actor": "1:7"},
    ).json()
    barrier = threading.Barrier(2)

    def fake(review, parent_version, feedback, **kwargs):
        barrier.wait(timeout=5)
        proposal = make_proposals(review)
        proposal["version"] = hashlib.sha256(feedback.encode()).hexdigest()
        proposal["parent_version"] = parent_version
        return proposal

    monkeypatch.setattr(routes, "refine_proposal", fake)
    common = {
        "review_id": preview["review_id"],
        "parent_version": preview["proposal"]["version"],
        "consent": "true",
        "sample_authorized": "true",
    }
    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(
            pool.map(
                lambda feedback: client.post(
                    "/api/templates/reviewed-excel/refine", data={**common, "feedback": feedback}
                ),
                ["第一个", "第二个"],
            )
        )
    assert sorted(r.status_code for r in results) == [200, 409]
    winner = next(r.json()["proposal"]["version"] for r in results if r.status_code == 200)
    history = client.get(f"/api/templates/reviewed-excel/reviews/{preview['review_id']}").json()
    assert len(history["versions"]) == 2 and history["versions"][-1]["version"] == winner


def test_invalid_or_excessive_upload_is_rejected(review_store):
    client = review_store.client
    assert (
        client.post(
            "/api/templates/reviewed-excel/preview",
            data={"expected_actor": "1:7"},
            files={"file": ("bad.xlsx", b"not an xlsx")},
        ).status_code
        == 400
    )
    assert (
        client.post(
            "/api/templates/reviewed-excel/preview",
            data={"expected_actor": "1:7"},
            files={"file": ("bad.xlsm", sample())},
        ).status_code
        == 400
    )
    assert (
        client.post(
            "/api/templates/reviewed-excel/preview",
            files={"file": ("big.xlsx", b"x" * (20 * 1024 * 1024 + 1))},
            data={"expected_actor": "1:7"},
        ).status_code
        == 413
    )


def test_revocation_during_model_or_before_write_never_commits(review_store, monkeypatch):
    from fastapi import HTTPException

    from app.application.etl.llm_assist_runtime import LlmAssistResult

    store = review_store
    client = store.client
    upload = {
        "file": (
            "sample.xlsx",
            sample(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    preview = client.post(
        "/api/templates/reviewed-excel/preview", files=upload, data={"expected_actor": "1:7"}
    ).json()

    def revoke_model(*_args, **_kwargs):
        store.actor["manage"] = False
        return LlmAssistResult(
            used_llm=True, data={"candidates": [{"clear": [], "reason": "保留"} for _ in range(3)]}
        )

    monkeypatch.setattr("app.application.etl.llm_assist._complete", revoke_model)
    form = {
        "review_id": preview["review_id"],
        "parent_version": preview["proposal"]["version"],
        "feedback": "保留",
        "consent": "true",
        "sample_authorized": "true",
    }
    assert client.post("/api/templates/reviewed-excel/refine", data=form).status_code == 403
    history_url = f"/api/templates/reviewed-excel/reviews/{preview['review_id']}"
    assert len(client.get(history_url).json()["versions"]) == 1
    store.actor["manage"] = True

    def deny_before_write(_request, _expected):
        raise HTTPException(status_code=403, detail="权限已撤销")

    monkeypatch.setattr(routes, "_recheck", deny_before_write)
    save = {
        "review_id": preview["review_id"],
        "proposal_version": preview["proposal"]["version"],
        "candidate_id": "preserve",
        "publish_to_tenant": "true",
        "selections": "[]",
        "name": "撤权测试",
    }
    assert client.post("/api/templates/reviewed-excel/save", data=save).status_code == 403
    assert client.get(history_url).json()["status"] == "template_review_ready"
    with store.engine.connect() as conn:
        assert conn.execute(text("SELECT count(*) FROM templates")).scalar_one() == 0


@pytest.mark.parametrize("failure", ["mkdir", "open", "write"])
def test_local_save_failure_releases_only_own_locks_and_can_retry(
    review_store, monkeypatch, failure
):
    import os

    from app.legacy.routes import document_templates_compat

    client, body = review_store.client, sample()
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("sample.xlsx", body)},
    ).json()["files"][0]
    url = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"
    labels = [
        {"sheet": sheet["sheet"], "range": region["range"], "labels": region["labels"]}
        for sheet in item["sheets"]
        for region in sheet["regions"]
    ]
    confirmed = client.post(
        url + "/confirm", data={"version": item["version"], "labels_json": json.dumps(labels)}
    )
    assert confirmed.status_code == 200
    region = confirmed.json()["sheets"][0]["regions"][0]
    preview = client.post(
        "/api/templates/reviewed-excel/preview",
        data={
            "expected_actor": "1:7",
            "classification_id": item["classification_id"],
            "classification_version": confirmed.json()["version"],
            "region_sheet": "报价",
            "region_range": region["range"],
        },
        files={"file": ("sample.xlsx", body)},
    )
    assert preview.status_code == 200, preview.text
    review = preview.json()
    candidate = review["proposal"]["candidates"][0]
    payload = {
        "review_id": review["review_id"],
        "proposal_version": review["proposal"]["version"],
        "candidate_id": candidate["id"],
        "selections": json.dumps(candidate["clear"]),
        "name": "合成模板",
        "publish_to_tenant": "true",
    }
    real_mkdir, real_open, real_fdopen = Path.mkdir, os.open, os.fdopen
    calls = []
    with monkeypatch.context() as m:
        m.setattr(
            document_templates_compat,
            "run_archive_template_create",
            lambda *a, **kw: calls.append("called"),
        )
        if failure == "mkdir":

            def fail_mkdir(path, *args, **kwargs):
                if "templates/reviewed_excel" in str(path):
                    raise OSError("synthetic mkdir failure")
                return real_mkdir(path, *args, **kwargs)

            m.setattr(Path, "mkdir", fail_mkdir)
        elif failure == "open":

            def fail_open(path, *args, **kwargs):
                if "templates/reviewed_excel" in str(path):
                    raise OSError("synthetic open failure")
                return real_open(path, *args, **kwargs)

            m.setattr(os, "open", fail_open)
        else:

            class FailingWriter:
                def __init__(self, wrapped):
                    self.wrapped = wrapped

                def __enter__(self):
                    return self

                def __exit__(self, *args):
                    self.wrapped.close()

                def write(self, _data):
                    raise OSError("synthetic write failure")

            def fail_write(fd, *args, **kwargs):
                stream = real_fdopen(fd, *args, **kwargs)
                return FailingWriter(stream)

            m.setattr(os, "fdopen", fail_write)
        failed = client.post("/api/templates/reviewed-excel/save", data=payload)
    assert failed.status_code == 503 and "模板未提交" in failed.json()["detail"]
    assert calls == []
    with review_store.engine.connect() as conn:
        statuses = dict(
            conn.execute(
                text("SELECT id, status FROM etl_runs WHERE id IN (:review, :classification)"),
                {"review": review["review_id"], "classification": item["classification_id"]},
            ).all()
        )
    assert statuses[review["review_id"]] == "template_review_ready"
    assert statuses[item["classification_id"]] == "classification_ready"
    destination = review_store.root / "templates" / "reviewed_excel" / "1"
    assert not destination.exists() or not list(destination.glob("*.xlsx"))
    retried = client.post("/api/templates/reviewed-excel/save", data=payload)
    assert retried.status_code == 200, retried.text


def test_failed_save_does_not_unlock_another_review_owner(review_store):
    from fastapi import HTTPException

    from app.infrastructure.tenant_scope import tenant_scope

    client, body = review_store.client, sample()
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("sample.xlsx", body)},
    ).json()["files"][0]
    preview = client.post(
        "/api/templates/reviewed-excel/preview",
        data={"expected_actor": "1:7"},
        files={"file": ("sample.xlsx", body)},
    ).json()
    review_id, classification_id = preview["review_id"], item["classification_id"]
    with review_store.engine.begin() as conn:
        conn.execute(
            text(
                "UPDATE etl_runs SET status='template_review_saving', draft_json='new-owner' WHERE id=:id"
            ),
            {"id": review_id},
        )
        conn.execute(
            text("UPDATE etl_runs SET status='classification_template_saving' WHERE id=:id"),
            {"id": classification_id},
        )
    with tenant_scope(1), sessionmaker(bind=review_store.engine)() as db:
        with pytest.raises(HTTPException) as error:
            routes._release_failed_review(
                db, review_id, 1, 7, "stale-owner", "{}", classification_id
            )
        assert error.value.status_code == 409
    with review_store.engine.connect() as conn:
        statuses = dict(
            conn.execute(
                text("SELECT id, status FROM etl_runs WHERE id IN (:review, :classification)"),
                {"review": review_id, "classification": classification_id},
            ).all()
        )
    assert statuses == {
        review_id: "template_review_saving",
        classification_id: "classification_template_saving",
    }


def test_unknown_legacy_commit_outcome_is_not_replayed(review_store, monkeypatch):
    from app.legacy.routes import document_templates_compat

    client = TestClient(review_store.client.app, raise_server_exceptions=False)
    upload = {
        "file": (
            "sample.xlsx",
            sample(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    preview = client.post(
        "/api/templates/reviewed-excel/preview", files=upload, data={"expected_actor": "1:7"}
    ).json()
    data = {
        "review_id": preview["review_id"],
        "proposal_version": preview["proposal"]["version"],
        "candidate_id": "preserve",
        "publish_to_tenant": "true",
        "selections": "[]",
        "name": "未知回执",
    }
    actual = document_templates_compat.run_archive_template_create

    def lost_receipt(payload):
        result, code = actual(payload)
        assert result["success"] and code == 200
        raise RuntimeError("synthetic receipt loss")

    monkeypatch.setattr(document_templates_compat, "run_archive_template_create", lost_receipt)
    assert client.post("/api/templates/reviewed-excel/save", data=data).status_code == 500
    assert client.post("/api/templates/reviewed-excel/save", data=data).status_code == 409
    history = client.get(f"/api/templates/reviewed-excel/reviews/{preview['review_id']}").json()
    assert history["status"] == "template_review_saving"
    with review_store.engine.connect() as conn:
        assert conn.execute(text("SELECT count(*) FROM templates")).scalar_one() == 1


@pytest.mark.parametrize("count", [1, 20, 200, 201])
def test_classification_batch_file_limit(review_store, count):
    files = [("files", (f"{i}.xlsx", sample())) for i in range(count)]
    response = review_store.client.post(
        "/api/templates/reviewed-excel/classify", files=files, data={"expected_actor": "1:7"}
    )
    if count == 201:
        assert response.status_code == 400
        assert "200" in response.json()["detail"]
        with review_store.engine.connect() as conn:
            assert (
                conn.execute(
                    text(
                        "SELECT count(*) FROM etl_runs WHERE target_type='template_classification'"
                    )
                ).scalar_one()
                == 0
            )
    else:
        assert response.status_code == 200
        assert len(response.json()["files"]) == count


def test_classification_empty_batch_is_rejected(review_store):
    response = review_store.client.post(
        "/api/templates/reviewed-excel/classify", data={"expected_actor": "1:7"}
    )
    assert response.status_code in (400, 422)


def test_semantic_large_workbook_discloses_unread_regions_and_merge_omissions(monkeypatch):
    import json

    from app.application.etl.llm_assist_runtime import LlmAssistResult
    from app.application.reviewed_excel_template import classify_semantic_workbook

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    wb = Workbook()
    ws = wb.active
    ws.title = "混合"
    selected_indices = {round(index * 249 / 99) for index in range(100)} | {0, 249}
    unseen = [index + 1 for index in range(200, 249) if index not in selected_indices]
    business_row, legend_row = unseen[-2:]
    for row in range(1, 251):
        ws.cell(row, 1, f"合成记录{row}")
    ws["A1"] = "报价单"
    ws.cell(business_row, 1, "采购订单")
    ws.cell(legend_row, 1, "图例：单位为件")
    for row in range(1, 111):
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    out = BytesIO()
    wb.save(out)
    seen = []

    def fake(messages, schema, max_tokens):
        seen.append(json.loads(messages[1]["content"]))
        return LlmAssistResult(
            used_llm=True,
            data={
                "regions": [
                    {
                        "sheet": "混合",
                        "range": "A1:D1",
                        "labels": ["quotation"],
                        "evidence": [{"sheet": "混合", "cell": "A1"}],
                        "dependencies": [],
                        "reason": "标题",
                    }
                ]
            },
        )

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    result = classify_semantic_workbook(inspect_workbook(out.getvalue()))
    coverage = result["coverage"][0]
    assert (
        f"A{business_row}" not in coverage["unread"]
        and f"A{business_row}" in coverage["unassigned"]
    )
    assert f"A{legend_row}" not in coverage["unread"] and f"A{legend_row}" in coverage["unassigned"]
    assert coverage["merged_sent"] == 100 and coverage["merged_omitted"] == 10
    assert len(seen[0]["sheets"][0]["merged"]) == 100
    assert {f"A{business_row}", f"A{legend_row}"} <= {cell["cell"] for cell in seen[0]["cells"]}
    assert result["budget"]["unread"] == 0


def test_evidence_fairness_hidden_and_empty_sheets():
    from app.application.reviewed_excel_classification import (
        classify_review_workbook,
        plan_semantic_batches,
    )
    from app.application.reviewed_excel_evidence import build_structure_map

    wb = Workbook()
    wb.remove(wb.active)
    for i in range(70):
        ws = wb.create_sheet(f"S{i:02}")
        ws["A1"], ws["A2"] = "报价单", f"数据{i}"
    wb["S69"].sheet_state = "hidden"
    wb["S69"].row_dimensions[2].hidden = True
    wb["S69"].column_dimensions["A"].hidden = True
    wb.create_sheet("空白")
    output = BytesIO()
    wb.save(output)
    review = inspect_workbook(output.getvalue())
    rules = classify_review_workbook(review)
    structure = build_structure_map(review, rules)
    batches, _ = plan_semantic_batches(review, 1)
    assert len(batches) == 1
    chosen = batches[0]
    assert len(chosen) == 140
    assert {cell["sheet"] for cell in chosen} == {f"S{i:02}" for i in range(70)}
    hidden = next(sheet for sheet in structure if sheet["sheet"] == "S69")
    assert (
        hidden["sheet_state"] == "hidden"
        and hidden["hidden_rows"] == [2]
        and hidden["hidden_columns"] == ["A"]
    )
    assert next(sheet for sheet in structure if sheet["sheet"] == "空白")["nonempty_cells"] == 0


def test_semantic_sequential_batches_budget_and_conflict(monkeypatch):
    from app.application.etl.llm_assist_runtime import LlmAssistResult
    from app.application.reviewed_excel_template import classify_semantic_workbook

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    monkeypatch.setenv("FHD_TEMPLATE_CLASSIFICATION_MAX_CALLS", "4")
    review = inspect_workbook(large_quotation(rows=2001))
    calls = []
    region = {
        "sheet": "报价",
        "range": "A1:A1",
        "labels": ["quotation"],
        "evidence": [{"sheet": "报价", "cell": "A1"}],
        "dependencies": [],
        "reason": "标题",
    }

    def same_result(messages, schema, tokens):
        calls.append(messages)
        return LlmAssistResult(used_llm=True, data={"regions": [region]})

    monkeypatch.setattr(
        "app.application.reviewed_excel_template._template_mimo_complete", same_result
    )
    result = classify_semantic_workbook(review)
    assert len(calls) == result["budget"]["calls"] == 4
    assert result["budget"]["unread"] > 0
    assert result["coverage"][0]["sent"] + result["budget"]["unread"] == 2001
    assert len(result["regions"]) == 1
    assert len(result["input_summaries"]) == 4
    assert all("input_sha256" in summary for summary in result["input_summaries"])
    calls.clear()

    def conflict(messages, schema, tokens):
        calls.append(messages)
        changed = region if len(calls) == 1 else {**region, "range": "A1:A2"}
        return LlmAssistResult(used_llm=True, data={"regions": [changed]})

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", conflict)
    with pytest.raises(ValueError, match="分批模型区域冲突"):
        classify_semantic_workbook(review)
    assert len(calls) == 2
    monkeypatch.setenv("FHD_TEMPLATE_CLASSIFICATION_MAX_CALLS", "5")
    with pytest.raises(ValueError, match="调用预算"):
        classify_semantic_workbook(review)


def test_unprocessed_rule_regions_remain_separately_reviewable(review_store, monkeypatch):
    from app.application.etl.llm_assist_runtime import LlmAssistResult

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    wb = Workbook()
    ws = wb.active
    ws.title = "混合"
    ws["A1"], ws["A8"] = "报价单", "采购订单"
    wb.create_sheet("其他")["A1"] = "未知"
    out = BytesIO()
    wb.save(out)
    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("mixed.xlsx", out.getvalue())},
    ).json()["files"][0]

    def fake(messages, schema, tokens):
        return LlmAssistResult(
            used_llm=True,
            data={
                "regions": [
                    {
                        "sheet": "其他",
                        "range": "A1:A1",
                        "labels": ["unknown"],
                        "evidence": [],
                        "dependencies": [],
                        "reason": "待确认",
                    }
                ]
            },
        )

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    result = client.post(
        f"/api/templates/reviewed-excel/classifications/{item['classification_id']}/analyze",
        data={"version": item["version"], "sample_authorized": "true"},
    )
    assert result.status_code == 200, result.text
    regions = result.json()["sheets"][0]["regions"]
    assert [region["range"] for region in regions] == ["A1:A7", "A8:A8"]
    assert all(
        region["labels"] == ["unknown"] and region["needs_confirmation"] for region in regions
    )


def test_unresolved_fragments_never_cut_merged_cells():
    from app.application.reviewed_excel_evidence import unresolved_fragments

    candidate = {"range": "A1:D6", "evidence": [], "header_cells": [], "merged_ranges": []}
    fragments = unresolved_fragments(candidate, [{"range": "A1:D2"}], ["B3:C4"])
    assert [item["range"] for item in fragments] == ["A3:D6"]
    assert fragments[0]["labels"] == ["unknown"]
    # A partial overlap creates a residual fragment that would cut B3:C4; omit it.
    fragments = unresolved_fragments(candidate, [{"range": "A1:B3"}], ["B3:C4"])
    from openpyxl.utils.cell import range_boundaries

    for item in fragments:
        c1, r1, c2, r2 = range_boundaries(item["range"])
        assert not (c1 <= 3 and c2 >= 2 and r1 <= 4 and r2 >= 3) or (
            c1 <= 2 <= 3 <= c2 and r1 <= 3 <= 4 <= r2
        )


def test_multibatch_failure_does_not_commit_partial_history(review_store, monkeypatch):
    from app.application.etl.llm_assist_runtime import LlmAssistResult

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    monkeypatch.setenv("FHD_TEMPLATE_CLASSIFICATION_MAX_CALLS", "2")
    source = large_quotation()
    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("sample.xlsx", source)},
    ).json()["files"][0]
    calls = []

    def fake(messages, schema, tokens):
        calls.append(1)
        if len(calls) == 2:
            raise ValueError("合成失败")
        return LlmAssistResult(
            used_llm=True,
            data={
                "regions": [
                    {
                        "sheet": "报价",
                        "range": "A1:A1",
                        "labels": ["quotation"],
                        "evidence": [{"sheet": "报价", "cell": "A1"}],
                        "dependencies": [],
                        "reason": "标题",
                    }
                ]
            },
        )

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    url = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"
    response = client.post(
        url + "/analyze", data={"version": item["version"], "sample_authorized": "true"}
    )
    assert response.status_code == 200 and response.json()["phase"] == "pending" and len(calls) == 1
    response = client.post(
        url + "/analyze", data={"version": item["version"], "sample_authorized": "true"}
    )
    assert response.status_code == 400 and len(calls) == 2
    history = client.get(url).json()
    assert history["checkpoint"]["phase"] == "inflight" and history["checkpoint"]["completed"] == 1
    assert len(history["versions"]) == 1 and history["versions"][0]["version"] == item["version"]
    assert (
        client.post(
            url + "/analyze", data={"version": item["version"], "sample_authorized": "true"}
        ).status_code
        == 409
    )
    assert len(calls) == 2


def test_checkpoint_resume_after_stop_and_expiry(review_store, monkeypatch):
    from datetime import UTC, datetime, timedelta

    from app.application.etl.llm_assist_runtime import LlmAssistResult

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    monkeypatch.setenv("FHD_TEMPLATE_CLASSIFICATION_MAX_CALLS", "2")
    source = large_quotation()
    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("sample.xlsx", source)},
    ).json()["files"][0]
    calls = []

    def fake(messages, schema, tokens):
        calls.append(1)
        return LlmAssistResult(
            used_llm=True,
            data={
                "regions": [
                    {
                        "sheet": "报价",
                        "range": "A1:A1",
                        "labels": ["quotation"],
                        "evidence": [{"sheet": "报价", "cell": "A1"}],
                        "dependencies": [],
                        "reason": "标题",
                    }
                ]
            },
        )

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    url = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"
    data = {"version": item["version"], "sample_authorized": "true"}
    first = client.post(url + "/analyze", data=data)
    assert first.status_code == 200 and first.json()["checkpoint"]["completed"] == 1
    with review_store.engine.connect() as conn:
        persisted = conn.execute(
            text("SELECT draft_json FROM etl_runs WHERE id=:id"), {"id": item["classification_id"]}
        ).scalar_one()
    checkpoint_state = json.loads(persisted)["analysis_checkpoint"]
    assert '"input_messages"' not in persisted
    assert len(checkpoint_state["results"][0]["input_summary"]["input_sha256"]) == 64
    # Ending the request is a batch-boundary stop. The persisted checkpoint survives a new client/session.
    history = client.get(url).json()
    assert history["checkpoint"]["phase"] == "ready" and history["checkpoint"]["total"] == 2
    assert len(calls) == 1
    resumed = client.post(url + "/analyze", data=data)
    assert resumed.status_code == 200 and resumed.json()["budget"]["calls"] == 2
    assert len(calls) == 2 and client.get(url).json()["checkpoint"] is None
    # A separate checkpoint cannot resume after the original upload expires.
    item2 = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("sample2.xlsx", source)},
    ).json()["files"][0]
    url2 = f"/api/templates/reviewed-excel/classifications/{item2['classification_id']}"
    data2 = {"version": item2["version"], "sample_authorized": "true"}
    assert client.post(url2 + "/analyze", data=data2).json()["phase"] == "pending"
    with review_store.engine.begin() as conn:
        conn.execute(
            text(
                "UPDATE etl_uploads SET expires_at=:expiry WHERE id=(SELECT upload_id FROM etl_runs WHERE id=:id)"
            ),
            {
                "expiry": (datetime.now(UTC) - timedelta(days=1)).isoformat(),
                "id": item2["classification_id"],
            },
        )
    assert client.post(url2 + "/analyze", data=data2).status_code == 410
    assert len(calls) == 3


def test_concurrent_checkpoint_resume_never_sends_twice(review_store, monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Event

    from app.application.etl.llm_assist_runtime import LlmAssistResult

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    monkeypatch.setenv("FHD_TEMPLATE_CLASSIFICATION_MAX_CALLS", "2")
    source = large_quotation()
    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("sample.xlsx", source)},
    ).json()["files"][0]
    entered, release = Event(), Event()
    calls = []

    def fake(messages, schema, tokens):
        calls.append(1)
        if len(calls) == 2:
            entered.set()
            assert release.wait(5)
        return LlmAssistResult(
            used_llm=True,
            data={
                "regions": [
                    {
                        "sheet": "报价",
                        "range": "A1:A1",
                        "labels": ["quotation"],
                        "evidence": [{"sheet": "报价", "cell": "A1"}],
                        "dependencies": [],
                        "reason": "标题",
                    }
                ]
            },
        )

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    url = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"
    data = {"version": item["version"], "sample_authorized": "true"}
    assert client.post(url + "/analyze", data=data).json()["phase"] == "pending"
    with ThreadPoolExecutor(max_workers=2) as pool:
        first = pool.submit(client.post, url + "/analyze", data=data)
        assert entered.wait(5)
        second = pool.submit(client.post, url + "/analyze", data=data)
        assert second.result(timeout=5).status_code == 409
        release.set()
        assert first.result(timeout=5).status_code == 200
    assert len(calls) == 2


def test_dense_sheet_tail_block_is_sampled_without_keyword_classification(monkeypatch):
    import json

    from app.application.etl.llm_assist_runtime import LlmAssistResult
    from app.application.reviewed_excel_template import classify_semantic_workbook

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    monkeypatch.setenv("FHD_TEMPLATE_CLASSIFICATION_MAX_CALLS", "1")
    wb = Workbook()
    ws = wb.active
    ws.title = "往来资料"
    for row in range(3, 39):
        for col in range(1, 13):
            ws.cell(row, col, f"合成记录{row}-{col}")
    ws["A22"], ws["B22"] = "品名", "数量"
    tail = {
        "A41": "备料确认",
        "A42": "供方",
        "B42": "需方",
        "C42": "交期",
        "D42": "付款方式",
        "E42": "订单号",
        "A43": "远泽工贸",
        "B43": "星港项目",
    }
    for address, value in tail.items():
        ws[address] = value
    ws["A45"] = "合成备注"
    ws.merge_cells("H41:M43")
    ws["H41"] = "共享图例"
    out = BytesIO()
    wb.save(out)
    captured = []

    def fake(messages, schema, tokens):
        captured.append((json.loads(messages[1]["content"]), tokens))
        return LlmAssistResult(
            used_llm=True,
            data={
                "regions": [
                    {
                        "sheet": "往来资料",
                        "range": "A1:M45",
                        "labels": ["unknown"],
                        "evidence": [],
                        "dependencies": [],
                        "reason": "待确认",
                    }
                ]
            },
        )

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    result = classify_semantic_workbook(inspect_workbook(out.getvalue()))
    payload, output_tokens = captured[0]
    sent = {cell["cell"] for cell in payload["cells"]}
    assert set(tail) <= sent
    assert "structure_map" not in payload and "rule_hints" not in payload
    assert "实际发送" not in payload["cells"][0]["value"]
    assert not result["coverage"][0]["unread"] and len(sent) == 442
    assert (
        payload["batch_coverage"][0]["sent_nonempty"]
        == payload["batch_coverage"][0]["total_nonempty"]
        == 442
    )
    source_values = {
        (cell["sheet"], cell["cell"]): cell["value"]
        for cell in inspect_workbook(out.getvalue())["cells"]
    }
    assert {
        (cell["sheet"], cell["cell"]): cell["value"] for cell in payload["cells"]
    } == source_values
    assert 3000 <= output_tokens <= 4000


def test_semantic_finish_length_rejected_without_retry(monkeypatch):
    import json

    from app.application.reviewed_excel_template import classify_semantic_workbook
    from app.infrastructure.llm.providers.openai_compatible_provider import XiaomiMimoProvider

    for name, value in {
        "FHD_TEMPLATE_LLM_PROVIDER": "mimo",
        "FHD_TEMPLATE_CLASSIFICATION_MAX_CALLS": "1",
        "FHD_TEMPLATE_MIMO_BASE_URL": "https://token-plan-cn.xiaomimimo.com/v1",
        "FHD_TEMPLATE_MIMO_MODEL": "mimo-v2.6-flash",
        "MIMO_API_KEY": "synthetic-placeholder",
        "FHD_ETL_LLM": "on",
    }.items():
        monkeypatch.setenv(name, value)
    seen = []

    async def fake(self, messages, **kwargs):
        seen.append(kwargs["max_tokens"])
        return {
            "choices": [
                {"finish_reason": "length", "message": {"content": json.dumps({"regions": []})}}
            ]
        }

    monkeypatch.setattr(XiaomiMimoProvider, "chat_completion", fake)
    with pytest.raises(ValueError, match="不完整"):
        classify_semantic_workbook(inspect_workbook(sample()))
    assert len(seen) == 1 and 3000 <= seen[0] <= 4000


def test_real_failure_outside_region_evidence_is_rejected(monkeypatch):
    from app.application.etl.llm_assist_runtime import LlmAssistResult
    from app.application.reviewed_excel_template import classify_semantic_workbook

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    monkeypatch.setenv("FHD_TEMPLATE_CLASSIFICATION_MAX_CALLS", "1")
    wb = Workbook()
    ws = wb.active
    ws.title = "往来资料"
    for address in ["H20", "L20", "M21", "H21", "L19", "H22", "M40"]:
        ws[address] = "合成数据"
    out = BytesIO()
    wb.save(out)
    outside = [{"sheet": "往来资料", "cell": cell} for cell in ["H20", "L20", "M21", "H21", "L19"]]

    def fake(messages, schema, tokens):
        assert "every evidence coordinate MUST be inside" in messages[0]["content"]
        return LlmAssistResult(
            used_llm=True,
            data={
                "regions": [
                    {
                        "sheet": "往来资料",
                        "range": "H22:M40",
                        "labels": ["unknown"],
                        "evidence": outside,
                        "dependencies": [],
                        "reason": "合成复现",
                    }
                ]
            },
        )

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    with pytest.raises(ValueError, match="模型证据不在区域内"):
        classify_semantic_workbook(inspect_workbook(out.getvalue()))


def test_repeated_generic_header_does_not_establish_purchase(monkeypatch):
    from app.application.etl.llm_assist_runtime import LlmAssistResult
    from app.application.reviewed_excel_template import classify_semantic_workbook

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    monkeypatch.setenv("FHD_TEMPLATE_CLASSIFICATION_MAX_CALLS", "1")
    wb = Workbook()
    ws = wb.active
    ws.title = "往来资料"
    for address, value in {
        "A1": "报价单",
        "A3": "品名",
        "B3": "数量",
        "C3": "单价",
        "A22": "品名",
        "B22": "数量",
        "C22": "单价",
        "A41": "供方",
        "B41": "需方",
        "C41": "订单号",
        "A45": "数据",
        "H41": "图例",
    }.items():
        ws[address] = value
    ws.merge_cells("H41:M43")
    out = BytesIO()
    wb.save(out)

    def ref(cell):
        return {"sheet": "往来资料", "cell": cell}

    def fake(messages, schema, tokens):
        prompt = messages[0]["content"]
        assert "repeated table header" in prompt and "dependencies" in prompt
        return LlmAssistResult(
            used_llm=True,
            data={
                "regions": [
                    {
                        "sheet": "往来资料",
                        "range": "A1:F20",
                        "labels": ["quotation"],
                        "evidence": [ref("A1")],
                        "dependencies": [],
                        "reason": "报价标题",
                    },
                    {
                        "sheet": "往来资料",
                        "range": "A22:F40",
                        "labels": ["purchase"],
                        "evidence": [ref("A22"), ref("B22"), ref("C22")],
                        "dependencies": [ref("A1")],
                        "reason": "重复表头",
                    },
                    {
                        "sheet": "往来资料",
                        "range": "A41:F45",
                        "labels": ["purchase"],
                        "evidence": [ref("A41"), ref("B41"), ref("C41")],
                        "dependencies": [ref("H41")],
                        "reason": "双方",
                    },
                ]
            },
        )

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    regions = classify_semantic_workbook(inspect_workbook(out.getvalue()))["regions"]
    assert regions[0]["labels"] == ["quotation"]
    assert regions[1]["labels"] == ["unknown"] and regions[1]["suggested_labels"] == ["purchase"]
    assert regions[1]["continuation_context"]["parent_range"] == "A1:F20"
    assert regions[1]["continuation_context"]["parent_header_row"] == 3
    assert regions[1]["continuation_context"]["header_row"] == 22
    assert regions[2]["labels"] == ["purchase"] and regions[2]["dependencies"] == [ref("H41")]


def test_no_lexical_classification_gate_for_non_repeated_regions(monkeypatch):
    from app.application.etl.llm_assist_runtime import LlmAssistResult
    from app.application.reviewed_excel_template import classify_semantic_workbook

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    monkeypatch.setenv("FHD_TEMPLATE_CLASSIFICATION_MAX_CALLS", "1")
    examples = [
        ("purchase", "供方需方交期", "无订单号采购仍为候选"),
        ("purchase", "买方与订单号的报价", "模型错误仍仅为候选"),
        ("purchase", "本单不是采购订单", "否定不能由词面证明正确"),
        ("purchase", "采购部门员工考勤", "跨类别词面不代表正确"),
        ("delivery", "车辆交接", "非字面术语待审"),
        ("attendance", "到岗离岗", "非字面术语待审"),
        ("quotation", "价格建议", "非字面术语待审"),
        ("template", "共享说明", "非字面术语待审"),
    ]
    for label, value, reason in examples:
        wb = Workbook()
        ws = wb.active
        ws.title = "样本"
        ws["A1"] = value
        out = BytesIO()
        wb.save(out)

        def fake(messages, schema, tokens):
            return LlmAssistResult(
                used_llm=True,
                data={
                    "regions": [
                        {
                            "sheet": "样本",
                            "range": "A1:A1",
                            "labels": [label],
                            "evidence": [{"sheet": "样本", "cell": "A1"}],
                            "dependencies": [],
                            "reason": reason,
                        }
                    ]
                },
            )

        monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
        region = classify_semantic_workbook(inspect_workbook(out.getvalue()))["regions"][0]
        assert region["labels"] == [label] and region["needs_confirmation"] is True
        assert "suggested_labels" not in region
        if "不是" in value:
            assert "否定描述" in region["context_warning"]


@pytest.mark.parametrize("shifted_value", ["参数", "单位"])
def test_relationship_payload_shows_observed_header_positions_without_guessing(
    monkeypatch, shifted_value
):
    import json

    from app.application.etl.llm_assist_runtime import LlmAssistResult
    from app.application.reviewed_excel_template import classify_semantic_workbook

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    monkeypatch.setenv("FHD_TEMPLATE_CLASSIFICATION_MAX_CALLS", "1")
    wb = Workbook()
    ws = wb.active
    ws.title = "往来资料"
    for address, value in {
        "A1": "报价单",
        "B5": "参数",
        "C5": "数量",
        "C22": "数量",
        "D22": shifted_value,
        "A41": "采购订单",
        "H41": "共享图例",
        "A45": "尾部",
    }.items():
        ws[address] = value
    ws.merge_cells("H41:M43")
    out = BytesIO()
    wb.save(out)
    captured = []

    def fake(messages, schema, tokens):
        captured.append(json.loads(messages[1]["content"]))
        assert "observations, not proof" in messages[0]["content"]
        return LlmAssistResult(
            used_llm=True,
            data={
                "regions": [
                    {
                        "sheet": "往来资料",
                        "range": "A1:M45",
                        "labels": ["unknown"],
                        "evidence": [],
                        "dependencies": [],
                        "reason": "需确认",
                    }
                ]
            },
        )

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    result = classify_semantic_workbook(inspect_workbook(out.getvalue()))
    assert "relationship_context" not in captured[0]
    sent = {cell["cell"]: cell["value"] for cell in captured[0]["cells"]}
    assert sent["C5"] == sent["C22"] == "数量"
    assert sent["D22"] == shifted_value and sent["H41"] == "共享图例"
    assert "H41:M43" in captured[0]["sheets"][0]["merged"]
    assert result["regions"][0]["labels"] == ["unknown"]


def test_relationship_observations_do_not_invent_links_or_hidden_values():
    from app.application.reviewed_excel_evidence import build_relationship_context

    def region(area):
        return {"range": area, "evidence": [], "possible_shared_context": []}

    structure = [
        {
            "sheet": "无标题",
            "nonempty_cells": 5,
            "regions": [region("A1:C5"), region("A6:C10")],
            "merged_regions": [],
        }
    ]
    cells = [
        {"sheet": "无标题", "cell": address, "value": value, "kind": "review"}
        for address, value in [("A1", "甲"), ("B1", "乙"), ("A6", "丙"), ("B6", "丁")]
    ]
    context = build_relationship_context(structure, cells)
    sheet = context["sheets"][0]
    assert sheet["candidate_adjacencies"][0]["row_gap"] == 0
    assert sheet["sent_row_matches"] == [] and sheet["possible_context_refs"] == []
    assert sheet["unread_nonempty"] == 1
    assert "do not prove" in context["limitation"]


def test_same_headers_different_business_are_not_automatically_joined():
    from app.application.reviewed_excel_evidence import flag_repeated_header_conflicts

    def region(area, label):
        return {
            "sheet": "混合",
            "range": area,
            "labels": [label],
            "evidence": [],
            "needs_confirmation": True,
        }

    regions = [region("A1:B5", "quotation"), region("A6:B10", "attendance")]
    cells = [
        {"sheet": "混合", "cell": address, "value": value, "kind": "review"}
        for address, value in [("A1", "姓名"), ("B1", "日期"), ("A6", "姓名"), ("B6", "日期")]
    ]
    observed = flag_repeated_header_conflicts(regions, cells)
    assert observed[1]["labels"] == ["unknown"]
    assert observed[1]["suggested_labels"] == ["attendance"]
    assert "可能是续表" in observed[1]["context_warning"]
    assert regions[1]["labels"] == ["attendance"]


def test_unsent_shared_note_is_an_index_not_model_evidence():
    from app.application.reviewed_excel_evidence import build_relationship_context

    structure = [
        {
            "sheet": "表",
            "nonempty_cells": 2,
            "merged_regions": ["D9:F10"],
            "regions": [
                {
                    "range": "A1:F10",
                    "evidence": [{"cell": "A1", "matched": "标题"}],
                    "possible_shared_context": [{"cell": "D9", "kind": "可能的共享说明"}],
                }
            ],
        }
    ]
    cells = [{"sheet": "表", "cell": "A1", "value": "标题", "kind": "review"}]
    context = build_relationship_context(structure, cells)
    refs = context["sheets"][0]["possible_context_refs"]
    assert next(ref for ref in refs if ref["cell"] == "D9")["value_sent"] is False
    assert next(ref for ref in refs if ref["cell"] == "A1")["value_sent"] is True
    assert context["sheets"][0]["unread_nonempty"] == 1
    assert "D9:F10" in context["sheets"][0]["merged_ranges"]
    assert all("value" not in ref for ref in refs)


def test_adjacent_unrelated_suggestions_are_not_merged_or_relabelled():
    from app.application.reviewed_excel_evidence import flag_repeated_header_conflicts

    regions = [
        {
            "sheet": "表",
            "range": "A1:B5",
            "labels": ["quotation"],
            "evidence": [],
            "needs_confirmation": True,
        },
        {
            "sheet": "表",
            "range": "A6:B10",
            "labels": ["attendance"],
            "evidence": [],
            "needs_confirmation": True,
        },
    ]
    cells = [
        {"sheet": "表", "cell": address, "value": value, "kind": "review"}
        for address, value in [("A1", "项目"), ("B1", "价格"), ("A6", "姓名"), ("B6", "日期")]
    ]
    observed = flag_repeated_header_conflicts(regions, cells)
    assert [region["labels"] for region in observed] == [["quotation"], ["attendance"]]
    assert len(observed) == 2 and all(region["needs_confirmation"] for region in observed)


def test_overlong_cell_is_deferred_and_unread_region_stays_unknown(monkeypatch):
    from app.application.etl.llm_assist_runtime import LlmAssistResult
    from app.application.reviewed_excel_template import classify_semantic_workbook

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    monkeypatch.setenv("FHD_TEMPLATE_CLASSIFICATION_MAX_CALLS", "1")
    wb = Workbook()
    ws = wb.active
    ws.title = "表"
    ws["A1"], ws["A2"], ws["A3"] = "报价单", "长" * 4001, "末尾"
    out = BytesIO()
    wb.save(out)

    def fake(messages, schema, tokens):
        return LlmAssistResult(
            used_llm=True,
            data={
                "regions": [
                    {
                        "sheet": "表",
                        "range": "A1:A3",
                        "labels": ["quotation"],
                        "evidence": [{"sheet": "表", "cell": "A1"}],
                        "dependencies": [],
                        "reason": "标题",
                    }
                ]
            },
        )

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    result = classify_semantic_workbook(inspect_workbook(out.getvalue()))
    assert result["coverage"][0]["deferred_large_values"] == ["A2"]
    assert "A2" in result["coverage"][0]["unread"]
    assert result["regions"][0]["labels"] == ["unknown"]
    assert result["regions"][0]["suggested_labels"] == ["quotation"]
    assert "未发送" in result["regions"][0]["context_warning"]


def test_budget_pause_requires_explicit_resume_and_no_resend(review_store, monkeypatch):
    from app.application.etl.llm_assist_runtime import LlmAssistResult

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    monkeypatch.setenv("FHD_TEMPLATE_CLASSIFICATION_MAX_CALLS", "1")
    source = large_quotation()
    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("sample.xlsx", source)},
    ).json()["files"][0]
    calls = []

    def fake(messages, schema, tokens):
        calls.append(1)
        return LlmAssistResult(
            used_llm=True,
            data={
                "regions": [
                    {
                        "sheet": "报价",
                        "range": "A1:A1",
                        "labels": ["quotation"],
                        "evidence": [{"sheet": "报价", "cell": "A1"}],
                        "dependencies": [],
                        "reason": "标题",
                    }
                ]
            },
        )

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    url = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"
    data = {"version": item["version"], "sample_authorized": "true"}
    first = client.post(url + "/analyze", data=data)
    assert first.status_code == 200 and first.json()["phase"] == "paused"
    assert first.json()["checkpoint"]["remaining"] == 1 and len(calls) == 1
    assert client.post(url + "/analyze", data=data).status_code == 409 and len(calls) == 1
    second = client.post(url + "/analyze", data={**data, "continue_authorized": "true"})
    assert second.status_code == 200 and second.json()["budget"]["calls"] == 2
    assert second.json()["coverage"][0]["sent"] == 601 and len(calls) == 2


def test_oversized_single_row_refused_before_provider_and_checkpoint(review_store, monkeypatch):
    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "超" * 4001
    out = BytesIO()
    wb.save(out)
    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("sample.xlsx", out.getvalue())},
    ).json()["files"][0]
    calls = []
    monkeypatch.setattr(
        "app.application.reviewed_excel_template._template_mimo_complete",
        lambda *args: calls.append(1),
    )
    url = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"
    response = client.post(
        url + "/analyze", data={"version": item["version"], "sample_authorized": "true"}
    )
    assert response.status_code == 400 and "未调用模型" in response.json()["detail"]
    assert client.get(url).json()["checkpoint"] is None and not calls


def test_large_contiguous_batches_retain_title_and_shared_note(monkeypatch):
    import json

    from app.application.etl.llm_assist_runtime import LlmAssistResult
    from app.application.reviewed_excel_template import classify_semantic_workbook

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    monkeypatch.setenv("FHD_TEMPLATE_CLASSIFICATION_MAX_CALLS", "2")
    wb = Workbook()
    ws = wb.active
    ws.title = "大表"
    for row in range(1, 602):
        ws.cell(row, 1, f"合成记录{row}")
    ws["A1"], ws["H602"] = "报价单", "图例：单位为件"
    out = BytesIO()
    wb.save(out)
    captured = []

    def fake(messages, schema, tokens):
        captured.append(json.loads(messages[1]["content"]))
        return LlmAssistResult(
            used_llm=True,
            data={
                "regions": [
                    {
                        "sheet": "大表",
                        "range": "A1:A1",
                        "labels": ["quotation"],
                        "evidence": [{"sheet": "大表", "cell": "A1"}],
                        "dependencies": [{"sheet": "大表", "cell": "H602"}],
                        "reason": "标题和图例",
                    }
                ]
            },
        )

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    result = classify_semantic_workbook(inspect_workbook(out.getvalue()))
    assert len(captured) == 2 and result["budget"]["unread"] == 0
    assert all({"A1", "H602"} <= {cell["cell"] for cell in batch["cells"]} for batch in captured)
    assert next(cell["context_only"] for cell in captured[0]["cells"] if cell["cell"] == "H602")
    assert next(cell["context_only"] for cell in captured[1]["cells"] if cell["cell"] == "A1")
    assert all(len(batch["cells"]) <= 500 for batch in captured)
    assert result["coverage"][0]["sent"] == 602


def test_full_width_row_does_not_create_empty_batch_or_drop_cells():
    from app.application.reviewed_excel_classification import plan_semantic_batches

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "报价单"
    for col in range(1, 501):
        ws.cell(2, col, f"值{col}")
    ws["A3"] = "图例：合成"
    out = BytesIO()
    wb.save(out)
    batches, _ = plan_semantic_batches(inspect_workbook(out.getvalue()), 2)
    assert all(batches) and all(len(batch) <= 500 for batch in batches)
    assert any(sum(cell["cell"].endswith("2") for cell in batch) == 500 for batch in batches)


def test_fragment_links_are_candidates_not_label_based_merges():
    from app.application.reviewed_excel_evidence import (
        candidate_fragment_links,
        flag_repeated_header_conflicts,
    )

    def region(area, label, dependencies=None):
        return {
            "sheet": "表",
            "range": area,
            "labels": [label],
            "evidence": [],
            "dependencies": dependencies or [],
            "needs_confirmation": True,
        }

    first = region("A1:B5", "quotation")
    second = region("A6:B10", "quotation")
    assert candidate_fragment_links([first, second])["items"] == []
    cells = [
        {"sheet": "表", "cell": coord, "value": value, "kind": "review"}
        for coord, value in [("A1", "品名"), ("B1", "数量"), ("A6", "品名"), ("B6", "数量")]
    ]
    observed = flag_repeated_header_conflicts([first, second], cells)
    links = candidate_fragment_links(observed)["items"]
    assert len(observed) == 2 and [item["labels"] for item in observed] == [
        ["quotation"],
        ["quotation"],
    ]
    assert (
        len(links) == 1
        and links[0]["status"] == "candidate"
        and links[0]["repeated_header_observed"]
    )
    # A shared dependency is a candidate even for different labels, not proof of one document.
    ref = {"sheet": "表", "cell": "H20"}
    different = candidate_fragment_links(
        [region("A1:B5", "quotation", [ref]), region("D1:E5", "delivery", [ref])]
    )
    assert different["items"][0]["shared_dependencies"] == [ref]
    assert different["items"][0]["status"] == "candidate"


def test_fragment_links_distinguish_conflicting_document_identity():
    from app.application.reviewed_excel_evidence import candidate_fragment_links

    dependency = {"sheet": "表", "cell": "H1"}
    first = {
        "sheet": "表",
        "range": "A1:F10",
        "labels": ["quotation"],
        "dependencies": [dependency],
    }
    second = {
        "sheet": "表",
        "range": "A11:F20",
        "labels": ["quotation"],
        "dependencies": [dependency],
    }

    def cells(second_number):
        return [
            {"sheet": "表", "cell": ref, "value": value, "kind": "text"}
            for ref, value in [
                ("A2", "订单号"),
                ("B2", "PO-01"),
                ("A12", "订单号"),
                ("B12", second_number),
            ]
        ]

    conflict = candidate_fragment_links([first, second], cells("PO-02"))["items"][0]
    assert conflict["status"] == "identity_conflict"
    assert conflict["identity_conflicts"][0]["first_cell"] == "B2"
    assert conflict["identity_conflicts"][0]["second_cell"] == "B12"
    matching = candidate_fragment_links([first, second], cells("PO-01"))["items"][0]
    assert matching["status"] == "candidate" and matching["identity_matches"]
    assert candidate_fragment_links([first, second], [])["items"][0]["status"] == "candidate"
    for label, left, right in [("客户", "甲", "乙"), ("日期", "2026-01-01", "2026-01-02")]:
        evidence = [
            {"sheet": "表", "cell": ref, "value": value, "kind": "text"}
            for ref, value in [("A2", label), ("B2", left), ("A12", label), ("B12", right)]
        ]
        assert (
            candidate_fragment_links([first, second], evidence)["items"][0]["status"]
            == "identity_conflict"
        )
    # Ambiguous repeated identity keys are not evidence of either equality or conflict.
    ambiguous = cells("PO-02") + [
        {"sheet": "表", "cell": "A3", "value": "订单号", "kind": "text"},
        {"sheet": "表", "cell": "B3", "value": "PO-03", "kind": "text"},
    ]
    assert candidate_fragment_links([first, second], ambiguous)["items"][0]["status"] == "candidate"


def test_value_extent_is_display_only_and_keeps_original_region():
    from app.application.reviewed_excel_evidence import value_extent

    region = {"sheet": "表", "range": "A1:F40"}
    cells = [
        {"sheet": "表", "cell": "B3", "value": "表头"},
        {"sheet": "表", "cell": "D20", "value": "=1+1"},
        {"sheet": "表", "cell": "G35", "value": "outside"},
    ]
    assert value_extent(region, cells) == "B3:D20"
    assert region["range"] == "A1:F40"
    assert value_extent(region, []) is None


def test_fragment_identity_conflict_is_preserved_across_batches():
    from app.application.reviewed_excel_classification import combine_semantic_batches
    from app.application.reviewed_excel_template import inspect_workbook

    wb = Workbook()
    sheet = wb.active
    sheet.title = "表"
    for coord, value in [
        ("A2", "订单号"),
        ("B2", "PO-1"),
        ("A12", "订单号"),
        ("B12", "PO-2"),
        ("H1", "共享标题"),
    ]:
        sheet[coord] = value
    out = BytesIO()
    wb.save(out)
    review = inspect_workbook(out.getvalue())
    first = [cell for cell in review["cells"] if cell["cell"] in {"A2", "B2", "H1"}]
    second = [cell for cell in review["cells"] if cell["cell"] in {"A12", "B12", "H1"}]

    def result(area):
        return {
            "regions": [
                {
                    "sheet": "表",
                    "range": area,
                    "labels": ["quotation"],
                    "evidence": [],
                    "dependencies": [{"sheet": "表", "cell": "H1"}],
                }
            ],
            "coverage": [{"sheet": "表", "merged_sent": 0}],
            "input_summary": {},
        }

    combined = combine_semantic_batches(
        review, [first, second], [result("A1:F10"), result("A11:F12")], 2
    )
    assert combined["fragment_links"]["items"][0]["status"] == "identity_conflict"
    assert [region["range"] for region in combined["regions"]] == ["A1:F10", "A11:F12"]


def test_cross_batch_document_ids_are_namespaced_and_links_need_evidence():
    from app.application.reviewed_excel_classification import combine_semantic_batches

    wb = Workbook()
    ws = wb.active
    ws.title = "表"
    ws["A1"], ws["A5"], ws["H1"] = "第一片段", "续片", "共享标题"
    out = BytesIO()
    wb.save(out)
    review = inspect_workbook(out.getvalue())

    def result(area):
        region = {
            "sheet": "表",
            "range": area,
            "labels": ["quotation"],
            "evidence": [],
            "dependencies": [{"sheet": "表", "cell": "H1"}],
        }
        return {
            "regions": [region],
            "coverage": [{"sheet": "表", "merged_sent": 0}],
            "input_summary": {},
            "logical_documents": [
                {
                    "document_id": "d1",
                    "fragments": [{"sheet": "表", "range": area}],
                    "status": "candidate",
                }
            ],
        }

    batches = [
        [cell for cell in review["cells"] if cell["cell"] in ("A1", "H1")],
        [cell for cell in review["cells"] if cell["cell"] in ("A5", "H1")],
    ]
    combined = combine_semantic_batches(review, batches, [result("A1:A1"), result("A5:A5")], 2)
    assert [document["document_id"] for document in combined["suggested_documents"]] == [
        "batch-1:d1",
        "batch-2:d1",
    ]
    assert (
        combined["suggested_document_count"] is None
        and combined["document_count_status"] == "partial"
    )
    assert combined["fragment_links"]["items"][0]["shared_dependencies"] == [
        {"sheet": "表", "cell": "H1"}
    ]
    assert combined["fragment_links"]["items"][0]["status"] == "candidate"


def test_user_document_groups_are_versioned_split_and_stale_safe(review_store, monkeypatch):
    from app.application.etl.llm_assist_runtime import LlmAssistResult

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    wb = Workbook()
    ws = wb.active
    ws.title = "报价"
    for cell, value in [
        ("H1", "总标题"),
        ("A2", "订单号"),
        ("B2", "P1"),
        ("A3", "品名"),
        ("B3", "数量"),
        ("A8", "订单号"),
        ("B8", "P1"),
        ("A9", "品名"),
        ("B9", "数量"),
    ]:
        ws[cell] = value
    out = BytesIO()
    wb.save(out)
    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("sample.xlsx", out.getvalue())},
    ).json()["files"][0]

    def fake(*args):
        return LlmAssistResult(
            used_llm=True,
            data={
                "regions": [
                    {
                        "sheet": "报价",
                        "range": area,
                        "labels": ["quotation"],
                        "evidence": [{"sheet": "报价", "cell": evidence}],
                        "dependencies": [{"sheet": "报价", "cell": "H1"}],
                        "reason": "合成",
                    }
                    for area, evidence in [("A2:B3", "A2"), ("A8:B9", "A8")]
                ]
            },
        )

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    url = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"
    analyzed = client.post(
        url + "/analyze", data={"version": item["version"], "sample_authorized": "true"}
    ).json()
    labels = [
        {"sheet": sheet["sheet"], "range": region["range"], "labels": region["labels"]}
        for sheet in analyzed["sheets"]
        for region in sheet["regions"]
    ]
    edge = {"sheet": "报价", "first": "A2:B3", "second": "A8:B9"}
    data = {
        "version": analyzed["version"],
        "labels_json": json.dumps(labels),
        "links_json": json.dumps([edge]),
    }
    confirmed = client.post(url + "/confirm", data=data)
    assert confirmed.status_code == 200, confirmed.text
    version = confirmed.json()["version"]
    group = confirmed.json()["logical_documents"][0]
    assert (
        group["source_sha256"] == item["sha256"]
        and group["classification_version"] == analyzed["version"]
    )
    assert group["fragments"] == [
        {"sheet": "报价", "range": "A2:B3"},
        {"sheet": "报价", "range": "A8:B9"},
    ]
    assert client.get(url).json()["versions"][-1]["logical_documents"] == [group]
    assert client.post(url + "/confirm", data=data).status_code == 409
    split = client.post(url + "/confirm", data={**data, "version": version, "links_json": "[]"})
    assert split.status_code == 200 and split.json()["logical_documents"] == []


def test_document_group_rejects_transitive_identity_conflicts_and_unknown():
    from app.application.reviewed_excel_classification_store import _confirmed_documents

    ranges = ["A1:B3", "A4:B6", "A7:B9"]
    sheets = [
        {"sheet": "表", "regions": [{"range": area, "labels": ["quotation"]} for area in ranges]}
    ]
    links = {
        "items": [
            {"sheet": "表", "first": ranges[i], "second": ranges[i + 1], "status": "candidate"}
            for i in range(2)
        ]
    }
    selected = [{key: link[key] for key in ("sheet", "first", "second")} for link in links["items"]]
    cells = [
        {"sheet": "表", "cell": coordinate, "value": value, "kind": "review"}
        for coordinate, value in [("A1", "订单号"), ("B1", "A"), ("A7", "订单号"), ("B7", "B")]
    ]
    with pytest.raises(ValueError, match="身份冲突"):
        _confirmed_documents("hash", "version", sheets, links, selected, cells)
    sheets[0]["regions"][1]["labels"] = ["unknown"]
    with pytest.raises(ValueError, match="未知"):
        _confirmed_documents("hash", "version", sheets, links, selected, cells)
    with pytest.raises(ValueError, match="失效"):
        _confirmed_documents(
            "hash",
            "version",
            sheets,
            links,
            [{"sheet": "表", "first": "A1:B3", "second": "A7:B9"}],
            cells,
        )


def _workbook_with_print_area(area, *, index="3", extra_file=None, second_document=False):
    from xml.etree import ElementTree as ET
    from zipfile import ZIP_DEFLATED, ZipFile

    wb = Workbook()
    for name in ["表二", "表三", "报价"]:
        wb.create_sheet(name)
    wb["报价"]["A1"] = "报价单"
    if second_document:
        wb["报价"]["A25"] = "第二片段"
    out = BytesIO()
    wb.save(out)
    source = out.getvalue()
    patched = BytesIO()
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with ZipFile(BytesIO(source)) as old, ZipFile(patched, "w", ZIP_DEFLATED) as new:
        for item in old.infolist():
            payload = old.read(item.filename)
            if item.filename == "xl/workbook.xml":
                root = ET.fromstring(payload)
                names = root.find(f"{{{ns}}}definedNames")
                if names is None:
                    names = ET.SubElement(root, f"{{{ns}}}definedNames")
                entry = ET.SubElement(
                    names,
                    f"{{{ns}}}definedName",
                    {"name": "_xlnm.Print_Area", "localSheetId": index},
                )
                entry.text = area
                payload = ET.tostring(root, encoding="utf-8")
            new.writestr(item, payload)
        if extra_file:
            new.writestr(extra_file, b"synthetic")
    return patched.getvalue()


@pytest.mark.parametrize("area", ["报价!$1:$19", "报价!$A:$C", "报价!$1:$19,报价!$25:$30"])
def test_whole_row_print_area_readable_and_preserved_in_whole_workbook(review_store, area):
    body = _workbook_with_print_area(area)
    info = inspect_workbook(body)
    assert any(
        (cell["sheet"], cell["cell"], cell["value"]) == ("报价", "A1", "报价单")
        for cell in info["cells"]
    )
    assert info["read_only_print_areas"] == [{"sheet": "报价", "value": area}]
    cleaned = clean_workbook(body, info["sha256"], [])
    assert inspect_workbook(cleaned)["read_only_print_areas"] == [{"sheet": "报价", "value": area}]
    response = review_store.client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("print-area.xlsx", body)},
    )
    assert response.status_code == 200 and response.json()["files"][0]["status"] == "classified"
    assert response.json()["files"][0]["read_only_print_areas"]


@pytest.mark.parametrize("area", ["报价!$1:$19", "报价!$A:$C", "报价!$1:$19,报价!$25:$30"])
def test_document_print_area_requires_explicit_mapping(area):
    from app.application.reviewed_excel_template import derive_document_workbook

    body = _workbook_with_print_area(area)
    fragments = [{"sheet": "报价", "range": "A1:A1"}]
    with pytest.raises(ValueError, match="明确选择"):
        derive_document_workbook(body, fragments)
    derived = derive_document_workbook(body, fragments, print_area_mode="selected")
    wb = load_workbook(BytesIO(derived))
    assert wb.sheetnames == ["报价"] and str(wb["报价"].print_area) == "'报价'!$A$1"
    assert inspect_workbook(body)["read_only_print_areas"][0]["value"] == area


def test_multi_area_maps_each_selected_fragment_without_expansion():
    from openpyxl.utils.cell import range_boundaries

    from app.application.reviewed_excel_template import derive_document_workbook

    body = _workbook_with_print_area("报价!$1:$19,报价!$25:$30", second_document=True)
    fragments = [{"sheet": "报价", "range": "A1:A1"}, {"sheet": "报价", "range": "A25:A25"}]
    derived = derive_document_workbook(body, fragments, print_area_mode="selected")
    ws = load_workbook(BytesIO(derived))["报价"]
    assert sorted(range_boundaries(str(area)) for area in ws._print_area.ranges) == [
        (1, 1, 1, 1),
        (1, 25, 1, 25),
    ]
    expanded = derive_document_workbook(
        body, [{"sheet": "报价", "range": "A1:A25"}], print_area_mode="selected"
    )
    assert str(load_workbook(BytesIO(expanded))["报价"].print_area) == "'报价'!$A$1:$A$25"


def test_explicit_mapping_includes_document_outside_old_print_area_without_other_content():
    from openpyxl.utils.cell import range_boundaries

    from app.application.reviewed_excel_template import derive_document_workbook

    body = _workbook_with_print_area("报价!$1:$19", second_document=True)
    fragments = [{"sheet": "报价", "range": "A25:A25"}]
    with pytest.raises(ValueError, match="明确选择"):
        derive_document_workbook(body, fragments)
    selected = derive_document_workbook(body, fragments, print_area_mode="selected")
    sheet = load_workbook(BytesIO(selected))["报价"]
    assert sheet["A1"].value is None and sheet["A25"].value == "第二片段"
    assert [range_boundaries(str(area)) for area in sheet._print_area.ranges] == [(1, 25, 1, 25)]
    assert inspect_workbook(body)["read_only_print_areas"][0]["value"] == "报价!$1:$19"
    both = derive_document_workbook(
        body, [{"sheet": "报价", "range": "A1:A1"}, *fragments], print_area_mode="selected"
    )
    ranges = load_workbook(BytesIO(both))["报价"]._print_area.ranges
    assert sorted(range_boundaries(str(area)) for area in ranges) == [(1, 1, 1, 1), (1, 25, 1, 25)]


def test_signature_outside_old_rectangular_print_area_is_only_added_with_consent():
    from app.application.reviewed_excel_template import derive_document_workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "送货"
    ws["A1"] = "送货单"
    ws["A24"] = "签字栏"
    ws.print_area = "A1:A23"
    out = BytesIO()
    wb.save(out)
    body = out.getvalue()
    fragments = [{"sheet": "送货", "range": "A1:A24"}]
    default = load_workbook(BytesIO(derive_document_workbook(body, fragments)))["送货"]
    assert str(default.print_area) == "'送货'!$A$1:$A$23" and default["A24"].value == "签字栏"
    selected = load_workbook(
        BytesIO(derive_document_workbook(body, fragments, print_area_mode="selected"))
    )["送货"]
    assert str(selected.print_area) == "'送货'!$A$1:$A$24" and selected["A24"].value == "签字栏"
    assert str(load_workbook(BytesIO(body))["送货"].print_area) == "'送货'!$A$1:$A$23"


def test_product_saves_document_outside_original_print_with_explicit_choice(
    review_store, monkeypatch
):
    from app.application.etl.llm_assist_runtime import LlmAssistResult

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    original = _workbook_with_print_area("报价!$1:$19", second_document=True)
    physical = {
        "sheet": "报价",
        "range": "A25:A25",
        "labels": ["quotation"],
        "evidence": [{"sheet": "报价", "cell": "A25"}],
        "dependencies": [],
        "reason": "合成",
    }
    model = {
        "physical_regions": [physical],
        "logical_documents": [
            {"document_id": "d1", "fragments": [{"sheet": "报价", "range": "A25:A25"}]}
        ],
        "layout_references": [],
    }
    monkeypatch.setattr(
        "app.application.reviewed_excel_template._template_mimo_complete",
        lambda *args: LlmAssistResult(used_llm=True, data=model),
    )
    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("sample.xlsx", original)},
    ).json()["files"][0]
    url = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"
    analysis = client.post(
        url + "/analyze", data={"version": item["version"], "sample_authorized": "true"}
    )
    assert analysis.status_code == 200, analysis.text
    result = analysis.json()
    labels = [
        {"sheet": sheet["sheet"], "range": region["range"], "labels": region["labels"]}
        for sheet in result["sheets"]
        for region in sheet["regions"]
    ]
    confirmed = client.post(
        url + "/confirm", data={"version": result["version"], "labels_json": json.dumps(labels)}
    )
    assert confirmed.status_code == 200, confirmed.text
    params = {
        "expected_actor": "1:7",
        "classification_id": item["classification_id"],
        "classification_version": confirmed.json()["version"],
        "region_sheet": "报价",
        "region_range": "A25:A25",
    }
    blocked = client.post(
        "/api/templates/reviewed-excel/preview",
        data=params,
        files={"file": ("sample.xlsx", original)},
    )
    assert blocked.status_code == 400 and "明确选择" in blocked.json()["detail"]
    preview = client.post(
        "/api/templates/reviewed-excel/preview",
        data={**params, "print_area_mode": "selected"},
        files={"file": ("sample.xlsx", original)},
    )
    assert preview.status_code == 200, preview.text
    review = preview.json()
    assert review["selected_region"]["source_print_areas"][0]["value"] == "报价!$1:$19"
    assert review["selected_region"]["derived_print_area"] == "'报价'!$A$25"
    assert review["selected_region"]["print_area_mode"] == "selected"
    candidate = review["proposal"]["candidates"][0]
    saved = client.post(
        "/api/templates/reviewed-excel/save",
        data={
            "review_id": review["review_id"],
            "proposal_version": review["proposal"]["version"],
            "candidate_id": candidate["id"],
            "selections": json.dumps(candidate["clear"]),
            "name": "原打印区外的文档",
            "publish_to_tenant": "true",
        },
    )
    assert saved.status_code == 200, saved.text
    template_id = saved.json()["template_id"].split(":")[1]
    download = client.get(f"/api/templates/reviewed-excel/{template_id}/download")
    assert download.status_code == 200
    sheet = load_workbook(BytesIO(download.content))["报价"]
    assert sheet["A1"].value is None and sheet["A25"].value == "第二片段"
    assert str(sheet.print_area) == "'报价'!$A$25"
    assert inspect_workbook(original)["read_only_print_areas"][0]["value"] == "报价!$1:$19"


def _nine_batch_workbook():
    wb = Workbook()
    for index, count in enumerate((1900, 600, 600, 265)):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = f"表{index}"
        if index == 0:
            ws.sheet_state = "hidden"
        for row in range(1, count + 1):
            ws.cell(row, 1, f"合成记录{row}")
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def test_nine_batch_plan_is_fair_across_sheets_and_preserves_each_sheet_order():
    from app.application.reviewed_excel_classification import plan_semantic_batches

    batches, budget = plan_semantic_batches(inspect_workbook(_nine_batch_workbook()), 4)
    assert len(batches) == 9 and budget == 4
    assert [batch[0]["sheet"] for batch in batches[:4]] == ["表0", "表1", "表2", "表3"]
    for name in ("表0", "表1", "表2", "表3"):
        rows = [
            int(cell["cell"][1:])
            for batch in batches
            for cell in batch
            if cell["sheet"] == name and not cell.get("context_only")
        ]
        assert rows == sorted(rows)


@pytest.mark.parametrize("per_round", [3, 4])
def test_nine_batch_real_route_pauses_resumes_and_namespaces_model_ids(
    review_store, monkeypatch, per_round
):
    from app.application.etl.llm_assist_runtime import LlmAssistResult

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    monkeypatch.setenv("FHD_TEMPLATE_CLASSIFICATION_MAX_CALLS", str(per_round))
    calls = []

    def fake(messages, schema, tokens):
        payload = json.loads(messages[1]["content"])
        primary = [cell for cell in payload["cells"] if not cell["context_only"]]
        sheet = primary[0]["sheet"]
        rows = [int(cell["cell"][1:]) for cell in primary]
        area = f"A{min(rows)}:A{max(rows)}"
        evidence = primary[0]["cell"]
        calls.append((sheet, area))
        return LlmAssistResult(
            used_llm=True,
            data={
                "physical_regions": [
                    {
                        "sheet": sheet,
                        "range": area,
                        "labels": ["quotation"],
                        "evidence": [{"sheet": sheet, "cell": evidence}],
                        "dependencies": [],
                        "reason": "合成",
                    }
                ],
                "logical_documents": [
                    {"document_id": "same-id", "fragments": [{"sheet": sheet, "range": area}]}
                ],
                "layout_references": [],
            },
        )

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("synthetic.xlsx", _nine_batch_workbook())},
    ).json()["files"][0]
    url = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"
    form = {"version": item["version"], "own_file_authorized": "true"}
    for index in range(9):
        if index and index % per_round == 0:
            assert client.post(url + "/analyze", data=form).status_code == 409
            assert (
                client.post(
                    url + "/analyze",
                    data={
                        "version": item["version"],
                        "sample_authorized": "true",
                        "continue_authorized": "true",
                    },
                ).status_code
                == 409
            )
        response = client.post(
            url + "/analyze",
            data={
                **form,
                **({"continue_authorized": "true"} if index and index % per_round == 0 else {}),
            },
        )
        assert response.status_code == 200, response.text
        if index < 8 and (index + 1) % per_round == 0:
            assert (
                response.json()["phase"] == "paused"
                and response.json()["checkpoint"]["remaining"] == 8 - index
            )
        elif index < 8:
            assert response.json()["phase"] == "pending"
        if index == 1:
            assert client.get(url).json()["checkpoint"]["completed"] == 2
    result = response.json()
    assert len(calls) == 9 and [sheet for sheet, _ in calls[:4]] == ["表0", "表1", "表2", "表3"]
    assert (
        result["suggested_document_count"] is None and result["document_count_status"] == "partial"
    )
    assert len({doc["document_id"] for doc in result["suggested_documents"]}) == 9
    assert all(not item["unread"] for item in result["coverage"])
    authorization = client.get(url).json()["versions"][-1]["model_authorization"]
    assert (
        authorization["kind"] == "owner_file" and authorization["source_sha256"] == item["sha256"]
    )
    labels = [
        {"sheet": sheet["sheet"], "range": region["range"], "labels": region["labels"]}
        for sheet in result["sheets"]
        for region in sheet["regions"]
    ]
    confirmed = client.post(
        url + "/confirm",
        data={
            "version": result["version"],
            "labels_json": json.dumps(labels),
            "independent_confirmed": "true",
        },
    )
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["confirmed_document_count"] == 9


def test_nine_batch_concurrent_resume_does_not_resend(review_store, monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Event

    from app.application.etl.llm_assist_runtime import LlmAssistResult

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    monkeypatch.setenv("FHD_TEMPLATE_CLASSIFICATION_MAX_CALLS", "4")
    entered, release, calls = Event(), Event(), []

    def fake(messages, schema, tokens):
        primary = [
            cell for cell in json.loads(messages[1]["content"])["cells"] if not cell["context_only"]
        ]
        calls.append(1)
        if len(calls) == 5:
            entered.set()
            assert release.wait(10)
        rows = [int(cell["cell"][1:]) for cell in primary]
        sheet = primary[0]["sheet"]
        return LlmAssistResult(
            used_llm=True,
            data={
                "physical_regions": [
                    {
                        "sheet": sheet,
                        "range": f"A{min(rows)}:A{max(rows)}",
                        "labels": ["quotation"],
                        "evidence": [{"sheet": sheet, "cell": primary[0]["cell"]}],
                        "dependencies": [],
                        "reason": "合成",
                    }
                ],
                "logical_documents": [],
                "layout_references": [],
            },
        )

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("synthetic.xlsx", _nine_batch_workbook())},
    ).json()["files"][0]
    url = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"
    data = {"version": item["version"], "own_file_authorized": "true"}
    for index in range(4):
        result = client.post(url + "/analyze", data=data)
        assert result.status_code == 200
    assert result.json()["phase"] == "paused"
    with ThreadPoolExecutor(max_workers=2) as pool:
        first = pool.submit(
            client.post, url + "/analyze", data={**data, "continue_authorized": "true"}
        )
        assert entered.wait(10)
        second = pool.submit(
            client.post, url + "/analyze", data={**data, "continue_authorized": "true"}
        )
        assert second.result(timeout=10).status_code == 409
        release.set()
        assert first.result(timeout=10).status_code == 200
    assert len(calls) == 5


def test_two_documents_print_mapping_product_save_and_readback(review_store):
    from xml.etree import ElementTree as ET
    from zipfile import ZIP_DEFLATED, ZipFile

    wb = Workbook()
    ws = wb.active
    ws.title = "报价"
    for cell, value in [
        ("A1", "报价单"),
        ("A2", "数量"),
        ("B2", 4),
        ("C2", "=B2*2"),
        ("A8", "采购订单"),
        ("B9", "不能进入报价模板"),
    ]:
        ws[cell] = value
    ws.column_dimensions["B"].width = 26
    ws.row_dimensions[2].height = 28
    ws.oddHeader.center.text = "固定页眉"
    ws.page_setup.orientation = "landscape"
    ws.page_margins.left = 0.45
    source = BytesIO()
    wb.save(source)
    body = BytesIO()
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with ZipFile(BytesIO(source.getvalue())) as old, ZipFile(body, "w", ZIP_DEFLATED) as new:
        for entry in old.infolist():
            data = old.read(entry.filename)
            if entry.filename == "xl/workbook.xml":
                root = ET.fromstring(data)
                names = root.find(f"{{{ns}}}definedNames")
                if names is None:
                    names = ET.SubElement(root, f"{{{ns}}}definedNames")
                node = ET.SubElement(
                    names, f"{{{ns}}}definedName", {"name": "_xlnm.Print_Area", "localSheetId": "0"}
                )
                node.text = "报价!$1:$19"
                data = ET.tostring(root, encoding="utf-8")
            new.writestr(entry, data)
    original = body.getvalue()
    source_hash = hashlib.sha256(original).hexdigest()
    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("sample.xlsx", original)},
    ).json()["files"][0]
    labels = [
        {"sheet": sheet["sheet"], "range": region["range"], "labels": region["labels"]}
        for sheet in item["sheets"]
        for region in sheet["regions"]
    ]
    url = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"
    confirmed = client.post(
        url + "/confirm", data={"version": item["version"], "labels_json": json.dumps(labels)}
    )
    assert confirmed.status_code == 200, confirmed.text
    params = {
        "expected_actor": "1:7",
        "classification_id": item["classification_id"],
        "classification_version": confirmed.json()["version"],
        "region_sheet": "报价",
        "region_range": "A1:C7",
    }
    blocked = client.post(
        "/api/templates/reviewed-excel/preview",
        data=params,
        files={"file": ("sample.xlsx", original)},
    )
    assert blocked.status_code == 400 and "明确选择" in blocked.json()["detail"]
    response = client.post(
        "/api/templates/reviewed-excel/preview",
        data={**params, "print_area_mode": "selected"},
        files={"file": ("sample.xlsx", original)},
    )
    assert response.status_code == 200, response.text
    review = response.json()
    assert review["selected_region"]["source_print_areas"] == [
        {"sheet": "报价", "value": "报价!$1:$19"}
    ]
    assert review["selected_region"]["derived_print_area"] == "'报价'!$A$1:$C$7"
    candidate = review["proposal"]["candidates"][0]
    saved = client.post(
        "/api/templates/reviewed-excel/save",
        data={
            "review_id": review["review_id"],
            "proposal_version": review["proposal"]["version"],
            "candidate_id": candidate["id"],
            "selections": json.dumps(candidate["clear"]),
            "name": "报价打印测试",
            "publish_to_tenant": "true",
        },
    )
    assert saved.status_code == 200, saved.text
    template_id = saved.json()["template_id"].split(":")[1]
    download = client.get(f"/api/templates/reviewed-excel/{template_id}/download")
    assert download.status_code == 200
    sheet = load_workbook(BytesIO(download.content), data_only=False)["报价"]
    assert sheet["B9"].value is None and sheet["C2"].value == "=B2*2"
    assert sheet.oddHeader.center.text == "固定页眉" and sheet.page_setup.orientation == "landscape"
    assert sheet.page_margins.left == 0.45 and sheet.column_dimensions["B"].width == 26
    assert sheet.row_dimensions[2].height == 28 and str(sheet.print_area) == "'报价'!$A$1:$C$7"
    assert hashlib.sha256(original).hexdigest() == source_hash


@pytest.mark.parametrize(
    "area,index",
    [
        ("报价!$0:$3", "3"),
        ("别表!$1:$19", "3"),
        ("报价!$19:$1", "3"),
        ("报价!$1:$19,", "3"),
        ("OFFSET(报价!$A$1,0,0,1,1)", "3"),
        ("报价!$1:$19", "100"),
    ],
)
def test_invalid_print_area_is_never_silently_allowed(area, index):
    with pytest.raises((ValueError, TypeError)):
        inspect_workbook(_workbook_with_print_area(area, index=index))


@pytest.mark.parametrize(
    "extra_file",
    ["xl/vbaProject.bin", "xl/externalLinks/externalLink1.xml", "xl/embeddings/oleObject1.bin"],
)
def test_print_area_does_not_bypass_unsupported_embedded_content(extra_file):
    with pytest.raises(ValueError):
        inspect_workbook(_workbook_with_print_area("报价!$1:$19", extra_file=extra_file))


def test_known_print_warning_does_not_hide_unrelated_warning(monkeypatch):
    import warnings

    from app.application import reviewed_excel_template as service

    original = service.load_workbook

    def loader(*args, **kwargs):
        result = original(*args, **kwargs)
        warnings.warn("unsupported custom extension", UserWarning, stacklevel=1)
        return result

    monkeypatch.setattr(service, "load_workbook", loader)
    with pytest.raises(ValueError, match="无法保证保真"):
        inspect_workbook(_workbook_with_print_area("报价!$1:$19"))


def test_semantic_contract_distinguishes_fragments_documents_and_layout(monkeypatch):
    from app.application.etl.llm_assist_runtime import LlmAssistResult
    from app.application.reviewed_excel_classification import _classify_semantic_batch

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    wb = Workbook()
    ws = wb.active
    ws.title = "流水"
    for row in (1, 4, 7, 10):
        ws.cell(row, 1, "业务记录")
    ws.merge_cells("A13:C13")
    out = BytesIO()
    wb.save(out)
    review = inspect_workbook(out.getvalue())
    areas = ["A1:C3", "A4:C6", "A7:C9", "A10:C12"]
    physical = [
        {
            "sheet": "流水",
            "range": area,
            "labels": ["quotation"],
            "evidence": [{"sheet": "流水", "cell": f"A{row}"}],
            "dependencies": [],
            "reason": "内容",
        }
        for area, row in zip(areas, (1, 4, 7, 10), strict=True)
    ]
    result = {
        "physical_regions": physical,
        "logical_documents": [
            {"document_id": "d1", "fragments": [{"sheet": "流水", "range": area} for area in areas]}
        ],
        "layout_references": [{"sheet": "流水", "range": "A13:C13"}],
    }
    monkeypatch.setattr(
        "app.application.reviewed_excel_template._template_mimo_complete",
        lambda *args: LlmAssistResult(used_llm=True, data=result),
    )
    parsed = _classify_semantic_batch(review)
    assert len(parsed["regions"]) == 4 and parsed["logical_document_count"] == 1
    assert parsed["layout_references"] == [{"sheet": "流水", "range": "A13:C13"}]
    physical[0]["evidence"] = [{"sheet": "流水", "cell": "A1 业务记录"}]
    with pytest.raises(ValueError, match="证据"):
        _classify_semantic_batch(review)
    physical[0]["evidence"] = [{"sheet": "流水", "cell": "A13"}]
    with pytest.raises(ValueError, match="证据"):
        _classify_semantic_batch(review)
    physical[0]["evidence"] = [{"sheet": "流水", "cell": "A1"}]
    result["logical_documents"].append(
        {"document_id": "d2", "fragments": [{"sheet": "流水", "range": areas[0]}]}
    )
    with pytest.raises(ValueError, match="重复"):
        _classify_semantic_batch(review)
    result["logical_documents"].pop()
    result["layout_references"] = [{"sheet": "流水", "range": "A13"}]
    with pytest.raises(ValueError, match="布局引用"):
        _classify_semantic_batch(review)
    result["layout_references"] = []
    result["sheetcount"] = 1
    with pytest.raises(ValueError, match="结构无效"):
        _classify_semantic_batch(review)


def test_provider_accepts_explicit_empty_document_and_layout_arrays(monkeypatch):
    from app.application.reviewed_excel_template import _template_mimo_complete
    from app.infrastructure.llm.providers.openai_compatible_provider import XiaomiMimoProvider

    for name, value in {
        "FHD_TEMPLATE_MIMO_BASE_URL": "https://token-plan-cn.xiaomimimo.com/v1",
        "FHD_TEMPLATE_MIMO_MODEL": "mimo-v2.6-flash",
        "MIMO_API_KEY": "synthetic-placeholder",
        "FHD_ETL_LLM": "on",
    }.items():
        monkeypatch.setenv(name, value)

    async def fake(self, messages, **kwargs):
        payload = {
            "physical_regions": [{"sheet": "表"}],
            "logical_documents": [],
            "layout_references": [],
        }
        return {"choices": [{"finish_reason": "stop", "message": {"content": json.dumps(payload)}}]}

    monkeypatch.setattr(XiaomiMimoProvider, "chat_completion", fake)
    schema = {
        "type": "object",
        "required": ["physical_regions", "logical_documents", "layout_references"],
        "additionalProperties": False,
        "properties": {
            "physical_regions": {"type": "array", "items": {"type": "object"}},
            "logical_documents": {"type": "array"},
            "layout_references": {"type": "array"},
        },
    }
    assert _template_mimo_complete([], schema, 100).data["logical_documents"] == []


def test_template_schema_validates_required_empty_arrays_and_nested_constraints():
    from app.application.reviewed_excel_template import _validate_template_schema

    ref = {
        "type": "object",
        "required": ["sheet", "cell"],
        "additionalProperties": False,
        "properties": {"sheet": {"type": "string"}, "cell": {"type": "string"}},
    }
    schema = {
        "type": "object",
        "required": ["regions", "layout_references", "clear"],
        "additionalProperties": False,
        "properties": {
            "regions": {"type": "array", "minItems": 1, "items": ref},
            "layout_references": {"type": "array", "items": ref},
            "clear": {"type": "array", "items": ref},
        },
    }
    good = {"regions": [{"sheet": "表", "cell": "A1"}], "layout_references": [], "clear": []}
    assert _validate_template_schema(schema, good)
    assert not _validate_template_schema(
        schema, {key: value for key, value in good.items() if key != "layout_references"}
    )
    assert not _validate_template_schema(schema, {**good, "regions": []})
    assert not _validate_template_schema(schema, {**good, "clear": "not-array"})
    assert not _validate_template_schema(schema, {**good, "regions": [{"sheet": "表"}]})
    assert not _validate_template_schema(
        schema, {**good, "regions": [{"sheet": "表", "cell": "A1", "extra": 1}]}
    )
    assert not _validate_template_schema({**schema, "unexpectedConstraint": True}, good)


def test_classification_provider_schema_with_empty_layout_and_strict_region(monkeypatch):
    from app.application.reviewed_excel_classification import _classify_semantic_batch
    from app.infrastructure.llm.providers.openai_compatible_provider import XiaomiMimoProvider

    for name, value in {
        "FHD_TEMPLATE_LLM_PROVIDER": "mimo",
        "FHD_TEMPLATE_MIMO_BASE_URL": "https://token-plan-cn.xiaomimimo.com/v1",
        "FHD_TEMPLATE_MIMO_MODEL": "mimo-v2.6-flash",
        "MIMO_API_KEY": "synthetic-placeholder",
        "FHD_ETL_LLM": "on",
    }.items():
        monkeypatch.setenv(name, value)
    review = inspect_workbook(sample())
    region = {
        "sheet": "报价",
        "range": "A1:C4",
        "labels": ["quotation"],
        "evidence": [{"sheet": "报价", "cell": "A1"}],
        "dependencies": [],
        "reason": "合成",
    }
    good = {"physical_regions": [region], "logical_documents": [], "layout_references": []}
    responses = [
        good,
        {key: value for key, value in good.items() if key != "layout_references"},
        {**good, "physical_regions": []},
        {**good, "physical_regions": [{**region, "extra": True}]},
        {
            **good,
            "physical_regions": [{**region, "evidence": [{"sheet": "报价", "cell": "A1 值说明"}]}],
        },
    ]

    async def fake(self, messages, **kwargs):
        payload = responses.pop(0)
        return {
            "choices": [
                {
                    "finish_reason": "stop",
                    "message": {"content": json.dumps(payload, ensure_ascii=False)},
                }
            ]
        }

    monkeypatch.setattr(XiaomiMimoProvider, "chat_completion", fake)
    assert _classify_semantic_batch(review)["layout_references"] == []
    for _ in range(4):
        with pytest.raises(ValueError):
            _classify_semantic_batch(review)


def test_document_derivation_preserves_coordinates_layout_and_rejects_external_reference():
    from app.application.reviewed_excel_template import derive_document_workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "混合"
    ws["A1"] = "报价单"
    ws["A2"] = "实例"
    ws["B2"] = 2
    ws["C2"] = "=B2*2"
    ws.merge_cells("A1:C1")
    ws.column_dimensions["B"].width = 24
    ws.row_dimensions[2].height = 30
    ws["A8"] = "另一份采购"
    ws["B8"] = "不应出现"
    ws.print_area = "A1:C3"
    wb.create_sheet("其他")["A1"] = "不应出现"
    source = BytesIO()
    wb.save(source)
    body = source.getvalue()
    before = hashlib.sha256(body).hexdigest()
    output = derive_document_workbook(body, [{"sheet": "混合", "range": "A1:C3"}])
    result = load_workbook(BytesIO(output), data_only=False)
    assert result.sheetnames == ["混合"]
    sheet = result["混合"]
    assert sheet["A1"].value == "报价单" and sheet["C2"].value == "=B2*2"
    assert sheet["A8"].value is None and sheet["B8"].value is None
    assert sheet.column_dimensions["B"].width == 24 and sheet.row_dimensions[2].height == 30
    assert str(sheet.merged_cells) == "A1:C1" and str(sheet.print_area) == "'混合'!$A$1:$C$3"
    assert hashlib.sha256(body).hexdigest() == before
    ws["C2"] = "=B8*2"
    changed = BytesIO()
    wb.save(changed)
    supported = load_workbook(
        BytesIO(
            derive_document_workbook(changed.getvalue(), [{"sheet": "混合", "range": "A1:C3"}])
        ),
        data_only=False,
    )
    assert supported["混合"]["C2"].value == "='依赖支撑'!B8*2"
    assert supported["依赖支撑"]["B8"].value == "不应出现"


def test_document_derivation_blocks_uncertain_print_merge_and_dynamic_formula():
    from app.application.reviewed_excel_template import derive_document_workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "表"
    ws["A1"] = "标题"
    ws["C1"] = "说明"
    ws["A2"] = '=INDIRECT("A1")'
    ws["A6"] = "其他"
    ws.print_area = "A1:C6"
    out = BytesIO()
    wb.save(out)
    part = [{"sheet": "表", "range": "A1:C3"}]
    with pytest.raises(ValueError, match="打印区域"):
        derive_document_workbook(out.getvalue(), part)
    ws.print_area = "A1:C3"
    out = BytesIO()
    wb.save(out)
    with pytest.raises(ValueError, match="公式函数"):
        derive_document_workbook(out.getvalue(), part)
    ws["A2"] = "内容"
    ws.merge_cells("B3:B5")
    out = BytesIO()
    wb.save(out)
    with pytest.raises(ValueError, match="合并"):
        derive_document_workbook(out.getvalue(), part)


def test_model_document_to_template_library_roundtrip(review_store, monkeypatch):
    from app.application.etl.llm_assist_runtime import LlmAssistResult

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    wb = Workbook()
    ws = wb.active
    ws.title = "混合"
    for cell, value in [
        ("A1", "报价单"),
        ("A2", "客户"),
        ("B2", "合成客户"),
        ("C2", "=1+1"),
        ("A8", "采购订单"),
        ("B9", "不应进入报价模板"),
    ]:
        ws[cell] = value
    ws.merge_cells("A1:C1")
    ws.column_dimensions["B"].width = 24
    ws.print_area = "A1:C3"
    out = BytesIO()
    wb.save(out)
    body = out.getvalue()
    source_hash = hashlib.sha256(body).hexdigest()
    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("sample.xlsx", body)},
    ).json()["files"][0]

    def fake(*args):
        return LlmAssistResult(
            used_llm=True,
            data={
                "physical_regions": [
                    {
                        "sheet": "混合",
                        "range": area,
                        "labels": [label],
                        "evidence": [{"sheet": "混合", "cell": cell}],
                        "dependencies": [],
                        "reason": "合成",
                    }
                    for area, label, cell in [
                        ("A1:C3", "quotation", "A1"),
                        ("A8:C9", "purchase", "A8"),
                    ]
                ],
                "logical_documents": [
                    {"document_id": "quote", "fragments": [{"sheet": "混合", "range": "A1:C3"}]},
                    {"document_id": "order", "fragments": [{"sheet": "混合", "range": "A8:C9"}]},
                ],
                "layout_references": [{"sheet": "混合", "range": "A1:C1"}],
            },
        )

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    url = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"
    analyzed = client.post(
        url + "/analyze", data={"version": item["version"], "sample_authorized": "true"}
    )
    assert analyzed.status_code == 200, analyzed.text
    result = analyzed.json()
    labels = [
        {"sheet": sheet["sheet"], "range": region["range"], "labels": region["labels"]}
        for sheet in result["sheets"]
        for region in sheet["regions"]
    ]
    confirmed = client.post(
        url + "/confirm", data={"version": result["version"], "labels_json": json.dumps(labels)}
    )
    assert confirmed.status_code == 200, confirmed.text
    params = {
        "expected_actor": "1:7",
        "classification_id": item["classification_id"],
        "classification_version": confirmed.json()["version"],
        "region_sheet": "混合",
        "region_range": "A1:C3",
    }
    preview = client.post(
        "/api/templates/reviewed-excel/preview", data=params, files={"file": ("sample.xlsx", body)}
    )
    assert preview.status_code == 200, preview.text
    review = preview.json()
    assert review["selected_region"]["source_sha256"] == source_hash
    assert len(review["proposal"]["result_previews"]) == 3
    candidate = review["proposal"]["candidates"][0]
    candidate_url = f"/api/templates/reviewed-excel/reviews/{review['review_id']}/versions/{review['proposal']['version']}/candidates/{candidate['id']}/download"
    download = client.get(candidate_url)
    assert download.status_code == 200
    document = load_workbook(BytesIO(download.content), data_only=False)
    assert document["混合"]["A8"].value is None and document["混合"]["B9"].value is None
    assert (
        document["混合"]["C2"].value == "=1+1"
        and document["混合"].column_dimensions["B"].width == 24
    )
    save_data = {
        "review_id": review["review_id"],
        "proposal_version": review["proposal"]["version"],
        "candidate_id": candidate["id"],
        "selections": json.dumps(candidate["clear"]),
        "name": "合成报价",
        "publish_to_tenant": "true",
    }
    # Simulate classification changing after the review lease commits but before the library write.
    with review_store.engine.begin() as conn:
        original_draft = conn.execute(
            text("SELECT draft_json FROM etl_runs WHERE id=:id"), {"id": item["classification_id"]}
        ).scalar_one()
    recheck = routes._recheck

    def change_version(request, expected):
        recheck(request, expected)
        changed = json.loads(original_draft)
        changed["versions"][-1]["version"] = "changed-after-review-lease"
        with review_store.engine.begin() as conn:
            conn.execute(
                text("UPDATE etl_runs SET draft_json=:draft WHERE id=:id"),
                {"id": item["classification_id"], "draft": json.dumps(changed, ensure_ascii=False)},
            )

    monkeypatch.setattr(routes, "_recheck", change_version)
    rejected = client.post("/api/templates/reviewed-excel/save", data=save_data)
    assert rejected.status_code == 409
    with review_store.engine.begin() as conn:
        assert conn.execute(text("SELECT count(*) FROM templates")).scalar_one() == 0
        conn.execute(
            text("UPDATE etl_runs SET draft_json=:draft WHERE id=:id"),
            {"id": item["classification_id"], "draft": original_draft},
        )
    monkeypatch.setattr(routes, "_recheck", recheck)
    from fastapi import HTTPException

    from app.application.etl.service_support import load_json
    from app.application.reviewed_excel_classification_store import (
        _classification_cas,
        _classification_run,
    )
    from app.infrastructure.tenant_scope import tenant_scope
    from app.legacy.routes import document_templates_compat

    actual_create = document_templates_compat.run_archive_template_create

    def attempt_concurrent_classification(payload):
        with tenant_scope(1), sessionmaker(bind=review_store.engine)() as db:
            classification = _classification_run(db, item["classification_id"], 1, 7)
            old = classification.draft_json
            changed = load_json(old, {})
            changed["versions"][-1]["version"] = "concurrent-change"
            with pytest.raises(HTTPException) as conflict:
                _classification_cas(db, classification, old, changed)
            assert conflict.value.status_code == 409
        return actual_create(payload)

    monkeypatch.setattr(
        document_templates_compat, "run_archive_template_create", attempt_concurrent_classification
    )
    saved = client.post("/api/templates/reviewed-excel/save", data=save_data)
    assert saved.status_code == 200, saved.text
    template_id = saved.json()["template_id"].split(":")[1]
    reopened = client.get(f"/api/templates/reviewed-excel/{template_id}/download")
    assert (
        reopened.status_code == 200
        and load_workbook(BytesIO(reopened.content))["混合"]["B9"].value is None
    )
    assert hashlib.sha256(body).hexdigest() == source_hash
    reconfirmed = client.post(
        url + "/confirm",
        data={"version": confirmed.json()["version"], "labels_json": json.dumps(labels)},
    )
    assert reconfirmed.status_code == 200
    assert client.get(candidate_url).status_code == 409
    assert client.get(f"/api/templates/reviewed-excel/{template_id}/download").status_code == 200
    second = client.post(
        "/api/templates/reviewed-excel/preview",
        data={**params, "classification_version": reconfirmed.json()["version"]},
        files={"file": ("sample.xlsx", body)},
    ).json()
    second_candidate = second["proposal"]["candidates"][0]

    def lost_receipt(payload):
        result, code = actual_create(payload)
        assert result["success"] and code == 200
        raise RuntimeError("synthetic receipt loss")

    monkeypatch.setattr(document_templates_compat, "run_archive_template_create", lost_receipt)
    retry_data = {
        "review_id": second["review_id"],
        "proposal_version": second["proposal"]["version"],
        "candidate_id": second_candidate["id"],
        "selections": json.dumps(second_candidate["clear"]),
        "name": "回执结果未知",
        "publish_to_tenant": "true",
    }
    no_raise = TestClient(client.app, raise_server_exceptions=False)
    assert no_raise.post("/api/templates/reviewed-excel/save", data=retry_data).status_code == 500
    assert client.post("/api/templates/reviewed-excel/save", data=retry_data).status_code == 409
    assert (
        client.post(
            url + "/confirm",
            data={"version": reconfirmed.json()["version"], "labels_json": json.dumps(labels)},
        ).status_code
        == 409
    )
    assert client.get(url).json()["template_write_state"] == "locked"
    with review_store.engine.connect() as conn:
        assert (
            conn.execute(
                text("SELECT status FROM etl_runs WHERE id=:id"), {"id": item["classification_id"]}
            ).scalar_one()
            == "classification_template_saving"
        )


def test_model_multifragment_document_is_a_real_candidate(review_store, monkeypatch):
    from app.application.etl.llm_assist_runtime import LlmAssistResult

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    wb = Workbook()
    ws = wb.active
    ws.title = "表"
    for cell, value in [("A1", "报价单"), ("A3", "部分一"), ("B4", "其他内容"), ("A5", "部分二")]:
        ws[cell] = value
    out = BytesIO()
    wb.save(out)
    body = out.getvalue()
    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("sample.xlsx", body)},
    ).json()["files"][0]
    physical = [
        {
            "sheet": "表",
            "range": area,
            "labels": ["quotation"],
            "evidence": [{"sheet": "表", "cell": ref}],
            "dependencies": [],
            "reason": "合成",
        }
        for area, ref in [("A1:A3", "A1"), ("A5:A5", "A5")]
    ]
    model = {
        "physical_regions": physical,
        "logical_documents": [
            {
                "document_id": "doc1",
                "fragments": [{"sheet": "表", "range": "A1:A3"}, {"sheet": "表", "range": "A5:A5"}],
            }
        ],
        "layout_references": [],
    }
    monkeypatch.setattr(
        "app.application.reviewed_excel_template._template_mimo_complete",
        lambda *args: LlmAssistResult(used_llm=True, data=model),
    )
    url = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"
    analyzed = client.post(
        url + "/analyze", data={"version": item["version"], "sample_authorized": "true"}
    )
    assert analyzed.status_code == 200, analyzed.text
    result = analyzed.json()
    assert (
        result["suggested_document_count"] is None and result["document_count_status"] == "partial"
    )
    link = result["fragment_links"]["items"][0]
    assert link["model_group_suggested"] and link["shared_dependencies"] == []
    labels = [
        {"sheet": sheet["sheet"], "range": region["range"], "labels": region["labels"]}
        for sheet in result["sheets"]
        for region in sheet["regions"]
    ]
    edge = {key: link[key] for key in ("sheet", "first", "second")}
    confirm = client.post(
        url + "/confirm",
        data={
            "version": result["version"],
            "labels_json": json.dumps(labels),
            "links_json": json.dumps([edge]),
        },
    )
    assert confirm.status_code == 200, confirm.text
    params = {
        "expected_actor": "1:7",
        "classification_id": item["classification_id"],
        "classification_version": confirm.json()["version"],
        "region_sheet": "表",
        "region_range": "A1:A3",
    }
    preview = client.post(
        "/api/templates/reviewed-excel/preview", data=params, files={"file": ("sample.xlsx", body)}
    )
    assert preview.status_code == 200, preview.text
    review = preview.json()
    assert len(review["selected_region"]["fragments"]) == 2
    candidate = review["proposal"]["candidates"][0]
    url = f"/api/templates/reviewed-excel/reviews/{review['review_id']}/versions/{review['proposal']['version']}/candidates/{candidate['id']}/download"
    download = client.get(url)
    document = load_workbook(BytesIO(download.content))["表"]
    assert (
        document["A3"].value == "部分一"
        and document["A5"].value == "部分二"
        and document["B4"].value is None
    )


def test_legacy_prompt_history_remains_but_is_not_copied_to_new_confirmation(review_store):
    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("sample.xlsx", sample())},
    ).json()["files"][0]
    with review_store.engine.begin() as conn:
        raw = conn.execute(
            text("SELECT draft_json FROM etl_runs WHERE id=:id"), {"id": item["classification_id"]}
        ).scalar_one()
        state = json.loads(raw)
        state["versions"][0]["input_messages"] = [{"role": "user", "content": "旧历史合成内容"}]
        conn.execute(
            text("UPDATE etl_runs SET draft_json=:draft WHERE id=:id"),
            {"draft": json.dumps(state, ensure_ascii=False), "id": item["classification_id"]},
        )
    labels = [
        {"sheet": sheet["sheet"], "range": region["range"], "labels": region["labels"]}
        for sheet in item["sheets"]
        for region in sheet["regions"]
    ]
    url = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"
    result = client.post(
        url + "/confirm", data={"version": item["version"], "labels_json": json.dumps(labels)}
    )
    assert result.status_code == 200
    versions = client.get(url).json()["versions"]
    assert versions[0]["input_messages"][0]["content"] == "旧历史合成内容"
    assert "input_messages" not in versions[1]


@pytest.mark.parametrize("commit_unknown", [False, True])
def test_archive_classification_cleans_only_uncommitted_upload_copy(
    review_store, monkeypatch, commit_unknown
):
    from sqlalchemy.orm import Session

    from app.db.models.etl import EtlRun
    from app.fastapi_routes.reviewed_excel import archive as archive_routes

    body = sample()
    client = review_store.client
    assert (
        client.post(
            "/api/templates/reviewed-excel/archive-quota",
            data={"expected_actor": "1:7", "limit_bytes": 1000000},
        ).status_code
        == 200
    )
    if commit_unknown:
        original_commit = Session.commit

        def commit_then_lose_receipt(db):
            classification = any(
                isinstance(item, EtlRun) and item.target_type == "template_classification"
                for item in db.new
            )
            original_commit(db)
            if classification:
                raise RuntimeError("synthetic lost commit receipt")

        monkeypatch.setattr(Session, "commit", commit_then_lose_receipt)
    else:

        def fail_after_upload(*args, **kwargs):
            raise ValueError("synthetic failure before commit")

        monkeypatch.setattr(archive_routes, "_classification_version", fail_after_upload)
    url = "/api/templates/reviewed-excel/archive-classify"
    data = {"expected_actor": "1:7", "request_key": "copy-cleanup"}
    if commit_unknown:
        with pytest.raises(RuntimeError, match="lost commit receipt"):
            client.post(url, data=data, files={"file": ("sample.xlsx", body)})
    else:
        response = client.post(url, data=data, files={"file": ("sample.xlsx", body)})
        assert response.status_code == 400
    copies = list((review_store.root / "etl" / "uploads").rglob("*.xlsx"))
    assert len(copies) == int(commit_unknown)
    originals = list((review_store.root / "etl" / "archive" / "1" / "7").iterdir())
    assert len(originals) == 1 and originals[0].read_bytes() == body
    with review_store.engine.connect() as connection:
        assert connection.execute(text("SELECT count(*) FROM etl_uploads")).scalar_one() == int(
            commit_unknown
        )
        classification_id = connection.execute(
            text("SELECT classification_id FROM etl_archive_versions")
        ).scalar_one()
        assert bool(classification_id) is commit_unknown
    if commit_unknown:
        assert copies[0].read_bytes() == body
        replay = client.post(url, data=data, files={"file": ("sample.xlsx", body)})
        assert replay.status_code == 200
        assert replay.json()["classification_id"] == classification_id


@pytest.mark.parametrize("entry", ["classify", "preview", "model"])
def test_reviewed_excel_never_returns_underlying_exception_payload(
    entry, review_store, monkeypatch
):
    sensitive = "synthetic-secret-at-/private/customer-workbook.xlsx"

    def fail(*args, **kwargs):
        raise ValueError(sensitive)

    files = {"file": ("sample.xlsx", sample())}
    data = {"expected_actor": "1:7"}
    if entry == "model":
        monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
        monkeypatch.setattr(routes, "refine_proposal", fail)
        data.update({"model_opt_in": "true", "sample_authorized": "true"})
    else:
        monkeypatch.setattr(routes, "inspect_workbook", fail)
    if entry == "classify":
        response = review_store.client.post(
            "/api/templates/reviewed-excel/classify",
            data=data,
            files={"files": files["file"]},
        )
        assert response.status_code == 200
        assert response.json()["files"][0]["status"] == "unsupported"
    else:
        response = review_store.client.post(
            "/api/templates/reviewed-excel/preview", data=data, files=files
        )
        assert response.status_code == (200 if entry == "model" else 400)
        if entry == "model":
            assert response.json()["proposal"]["model"]["status"] == "unavailable"
    assert sensitive not in response.text
    assert "文件或请求无法处理" in response.text


def test_auto_slice_api_without_confirmation_and_owner_isolation(review_store, monkeypatch):
    from app.application.etl.llm_assist_runtime import LlmAssistResult

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    source = sample()
    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("sample.xlsx", source)},
    ).json()["files"][0]
    path = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"

    def fake(messages, schema, max_tokens):
        data = {
            "physical_regions": [
                {
                    "sheet": "报价",
                    "range": "A1:C4",
                    "labels": ["unknown"],
                    "evidence": [],
                    "dependencies": [],
                    "reason": "不确定类型",
                }
            ],
            "logical_documents": [
                {"document_id": "fixture", "fragments": [{"sheet": "报价", "range": "A1:C4"}]}
            ],
            "layout_references": [],
        }
        return LlmAssistResult(used_llm=True, data=data)

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    analyzed = client.post(
        path + "/analyze", data={"version": item["version"], "sample_authorized": "true"}
    )
    assert analyzed.status_code == 200, analyzed.text
    value = analyzed.json()
    assert not value["confirmed"]
    route = path + "/auto-slices/" + value["version"]
    listing = client.get(route)
    assert listing.status_code == 200, listing.text
    result = listing.json()["slices"]
    assert len(result) == 1 and result[0]["basis"] == "model_group" and not result[0]["confirmed"]
    download = client.get(route + "/0")
    assert download.status_code == 200
    assert result[0]["filename"].endswith(".xlsx")
    assert "filename*=UTF-8''" in download.headers["content-disposition"]
    assert hashlib.sha256(download.content).hexdigest() == result[0]["sha256"]
    assert load_workbook(BytesIO(download.content)).sheetnames == ["报价"]
    assert client.get(path + "/auto-slices/" + item["version"]).status_code == 409
    review_store.actor["tenant"] = 2
    assert client.get(route).status_code == 404
    review_store.actor["tenant"] = 1
    review_store.actor["read"] = False
    assert client.get(route + "/0").status_code == 403


def test_model_cannot_create_region_or_evidence_from_context_only(monkeypatch):
    from app.application.etl.errors import ReviewedExcelError
    from app.application.etl.llm_assist_runtime import LlmAssistResult
    from app.application.reviewed_excel_classification import _classify_semantic_batch

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    review = inspect_workbook(sample())
    first = next(cell for cell in review["cells"] if cell["cell"] == "A1")
    second = next(cell for cell in review["cells"] if cell["cell"] == "B2")
    third = next(cell for cell in review["cells"] if cell["cell"] == "A3")
    review["cells"] = [first, {**second, "context_only": True}]
    region = {
        "sheet": "报价",
        "range": "B2:B2",
        "labels": ["unknown"],
        "evidence": [],
        "dependencies": [],
        "reason": "仅重复上下文",
    }

    def fake(messages, schema, tokens):
        return LlmAssistResult(
            used_llm=True,
            data={"physical_regions": [region], "logical_documents": [], "layout_references": []},
        )

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    with pytest.raises(ReviewedExcelError, match="仅包含上下文或未读内容"):
        _classify_semantic_batch(review)
    region.update(range="A1:C4", evidence=[{"sheet": "报价", "cell": "B2"}])
    review["full_cells"] = [first]
    with pytest.raises(ReviewedExcelError, match="证据仅为重复上下文"):
        _classify_semantic_batch(review)
    region["evidence"] = []
    review["full_cells"] = [first, second]
    clipped = _classify_semantic_batch(review)["regions"][0]
    assert clipped["range"] == "A1:C1" and clipped["model_original_range"] == "A1:C4"
    assert clipped["open_fragment"]
    review["cells"].append(third)
    review["full_cells"].append(third)
    with pytest.raises(ReviewedExcelError, match="包含本批未读内容"):
        _classify_semantic_batch(review)


@pytest.mark.parametrize("group_not_sent", [False, True])
def test_auto_group_resume_budget_and_download_multibatch(
    review_store, monkeypatch, group_not_sent
):
    from app.application.etl.llm_assist_runtime import LlmAssistResult

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    monkeypatch.setenv("FHD_TEMPLATE_CLASSIFICATION_MAX_CALLS", "2")
    calls = []

    def fake(messages, schema, tokens):
        payload = json.loads(messages[1]["content"])
        if "groups" in schema["properties"]:
            if group_not_sent and "not_sent" not in calls:
                from app.application.reviewed_excel_audit import ModelRequestNotSent

                calls.append("not_sent")
                raise ModelRequestNotSent("配置未完成；未调用模型")
            calls.append("group")
            fragments = [part["region_id"] for part in payload["fragments"]]
            evidence = [cell["evidence_id"] for cell in payload["cells"]]
            return LlmAssistResult(
                used_llm=True,
                data={
                    "groups": [
                        {
                            "document_id": "d1",
                            "fragment_ids": fragments,
                            "evidence_ids": evidence,
                            "reason": "已读片段有相同文档标识",
                        }
                    ]
                },
            )
        calls.append("batch")
        primary = [c for c in payload["cells"] if not c["context_only"]]
        rows = [int(c["cell"][1:]) for c in primary]
        area = f"A{min(rows)}:A{max(rows)}"
        return LlmAssistResult(
            used_llm=True,
            data={
                "physical_regions": [
                    {
                        "sheet": "报价",
                        "range": area,
                        "evidence_ids": [],
                        "dependency_ids": [],
                        "primary_id": primary[0]["evidence_id"],
                        "reason": "合成",
                    }
                ],
                "logical_documents": [],
                "layout_references": [],
            },
        )

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("batch.xlsx", large_quotation())},
    ).json()["files"][0]
    path = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"
    form = {"version": item["version"], "sample_authorized": "true", "auto_group": "true"}
    first = client.post(path + "/analyze", data=form)
    assert first.status_code == 200 and first.json()["phase"] == "pending"
    second = client.post(path + "/analyze", data=form)
    assert second.status_code == 200 and second.json()["phase"] == "paused"
    assert calls == ["batch", "batch"]
    assert client.post(path + "/analyze", data=form).status_code == 409
    done = client.post(path + "/analyze", data={**form, "continue_authorized": "true"})
    if group_not_sent:
        assert done.status_code == 400, done.text
        checkpoint = client.get(path).json()["checkpoint"]
        assert checkpoint["phase"] == "ready" and checkpoint["completed"] == 2
        assert checkpoint["calls_in_round"] == 0
        done = client.post(path + "/analyze", data=form)
        calls.remove("not_sent")
    assert done.status_code == 200, done.text
    data = done.json()
    assert calls == ["batch", "batch", "group"]
    assert data["budget"]["grouping_calls"] == 1 and not data["confirmed"]
    assert all(
        summary["reference_response"] and summary["reference_index"]
        for summary in data["input_summaries"]
    )
    assert len(data["suggested_documents"]) == 1
    url = path + "/auto-slices/" + data["version"]
    listing = client.get(url)
    assert listing.status_code == 200, listing.text
    assert len(listing.json()["slices"]) == 1
    download = client.get(url + "/0")
    assert download.status_code == 200
    book = load_workbook(BytesIO(download.content))
    assert book["报价"]["A601"].value is not None


def test_audit_response_is_committed_before_parsing_and_is_owner_scoped(review_store, monkeypatch):
    import base64

    from app.application.etl.llm_assist_runtime import LlmAssistResult
    from app.application.reviewed_excel_audit import response_capture

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("合成.xlsx", sample())},
    ).json()["files"][0]
    raw = b'{"usage":{"total_tokens":12},"choices":[]}'

    def fake(messages, schema, tokens):
        response_capture()(raw, 200, "synthetic-req")
        return LlmAssistResult(
            used_llm=True,
            provider="openai",
            model="synthetic-model",
            data={
                "regions": [
                    {
                        "sheet": "报价",
                        "range": "A1:C4",
                        "labels": ["unknown"],
                        "evidence": [],
                        "dependencies": [],
                        "reason": "fixture",
                    }
                ]
            },
        )

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    url = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"
    result = client.post(
        url + "/analyze", data={"version": item["version"], "sample_authorized": "true"}
    )
    assert result.status_code == 200, result.text
    assert result.json()["actual_model_providers"] == ["openai"]
    assert result.json()["input_summaries"][0]["model"] == "synthetic-model"
    assert result.json()["structure_map_source"] == "program_observation_not_sent_to_model"
    with review_store.engine.connect() as conn:
        persisted = conn.execute(
            text("SELECT draft_json FROM etl_runs WHERE id=:id"), {"id": item["classification_id"]}
        ).scalar_one()
    record = json.loads(persisted)["model_audit"][0]
    assert record["state"] == "response_received" and record["request_id"] == "synthetic-req"
    assert base64.b64decode(record["response_base64"]) == raw
    assert record["usage"] == {"total_tokens": 12}
    assert record["classification_version"] == item["version"]
    assert "Authorization" not in persisted
    review_store.actor["tenant"] = 2
    assert client.get(url).status_code == 404


def test_audit_write_failure_blocks_replay(review_store, monkeypatch):
    from fastapi import HTTPException

    from app.application.reviewed_excel_audit import response_capture

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("合成.xlsx", sample())},
    ).json()["files"][0]
    real_cas = routes._classification_cas

    def fail_response(db, run, old, state):
        if (
            state.get("model_audit", [])
            and state["model_audit"][-1]["state"] == "response_received"
        ):
            raise HTTPException(status_code=409, detail="synthetic write failure")
        return real_cas(db, run, old, state)

    monkeypatch.setattr(routes, "_classification_cas", fail_response)
    calls = []

    def fake(messages, schema, tokens):
        calls.append(1)
        response_capture()(b'{"choices":[]}', 200, "synthetic-req")
        raise AssertionError("must stop on failed audit")

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    url = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}/analyze"
    data = {"version": item["version"], "sample_authorized": "true"}
    assert client.post(url, data=data).status_code == 409
    assert client.post(url, data=data).status_code == 409
    assert len(calls) == 1
    with review_store.engine.connect() as conn:
        persisted = conn.execute(
            text("SELECT draft_json FROM etl_runs WHERE id=:id"), {"id": item["classification_id"]}
        ).scalar_one()
    assert json.loads(persisted)["model_audit"][-1]["state"] == "inflight"


def test_completed_invalid_response_is_counted_and_requires_explicit_retry(
    review_store, monkeypatch
):
    from app.application.etl.errors import ReviewedExcelError
    from app.application.etl.llm_assist_runtime import LlmAssistResult
    from app.application.reviewed_excel_audit import response_capture

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    monkeypatch.setenv("FHD_TEMPLATE_CLASSIFICATION_MAX_CALLS", "2")
    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("合成.xlsx", large_quotation())},
    ).json()["files"][0]
    calls = []

    def fake(messages, schema, tokens):
        calls.append(1)
        response_capture()(b'{"choices":[],"usage":{"total_tokens":1}}', 200, "synthetic")
        if len(calls) > 1:
            raise ReviewedExcelError("synthetic incomplete")
        return LlmAssistResult(
            used_llm=True,
            data={
                "regions": [
                    {
                        "sheet": "报价",
                        "range": "A1:A1",
                        "labels": ["quotation"],
                        "evidence": [{"sheet": "报价", "cell": "A1"}],
                        "dependencies": [],
                        "reason": "fixture",
                    }
                ]
            },
        )

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    base = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"
    form = {"version": item["version"], "sample_authorized": "true"}
    assert client.post(base + "/analyze", data=form).status_code == 200
    assert client.post(base + "/analyze", data=form).status_code == 400
    checkpoint = client.get(base).json()["checkpoint"]
    assert checkpoint["phase"] == "failed" and checkpoint["completed"] == 1
    assert checkpoint["calls_in_round"] == 2
    assert client.post(base + "/analyze", data=form).status_code == 409
    assert (
        client.post(base + "/analyze", data={**form, "retry_authorized": "true"}).status_code == 409
    )
    assert len(calls) == 2
    retry = client.post(
        base + "/analyze", data={**form, "retry_authorized": "true", "continue_authorized": "true"}
    )
    assert retry.status_code == 400 and len(calls) == 3
    resumed = client.get(base).json()["checkpoint"]
    assert resumed["phase"] == "failed" and resumed["completed"] == 1
    assert resumed["calls_in_round"] == 1


def test_archived_model_slice_can_be_selected_for_template_without_coordinate_confirmation(
    review_store, monkeypatch
):
    from app.application.etl.llm_assist_runtime import LlmAssistResult
    from app.application.reviewed_excel_audit import response_capture

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    with review_store.engine.begin() as conn:
        conn.execute(
            text(
                "INSERT INTO etl_archive_quotas (tenant_id, limit_bytes, reserved_bytes) VALUES (1, 1000000, 0)"
            )
        )
    body = sample()
    client = review_store.client
    archived = client.post(
        "/api/templates/reviewed-excel/archive-classify",
        data={"expected_actor": "1:7", "request_key": "synthetic-auto-flow"},
        files={"file": ("合成.xlsx", body)},
    )
    assert archived.status_code == 200, archived.text
    item = archived.json()

    def fake(messages, schema, tokens):
        primary = schema["properties"]["physical_regions"]["items"]["properties"]["primary_id"][
            "enum"
        ][0]
        data = {
            "physical_regions": [
                {
                    "sheet": "报价",
                    "range": "A1:C4",
                    "evidence_ids": [primary],
                    "dependency_ids": [],
                    "primary_id": primary,
                    "reason": "synthetic",
                }
            ],
            "logical_documents": [
                {"document_id": "fixture", "fragments": [{"sheet": "报价", "range": "A1:C4"}]}
            ],
            "layout_references": [],
        }
        response_capture()(
            json.dumps({"choices": [], "usage": {}}, ensure_ascii=False).encode(), 200, "synthetic"
        )
        return LlmAssistResult(used_llm=True, model="synthetic", provider="synthetic", data=data)

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", fake)
    path = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"
    analyzed = client.post(
        path + "/analyze",
        data={"version": item["version"], "sample_authorized": "true", "auto_group": "true"},
    )
    assert analyzed.status_code == 200, analyzed.text
    latest = analyzed.json()
    assert not latest["confirmed"]
    slices = client.get(path + "/auto-slices/" + latest["version"])
    assert slices.status_code == 200 and len(slices.json()["slices"]) == 1
    params = {
        "expected_actor": "1:7",
        "classification_id": item["classification_id"],
        "classification_version": latest["version"],
        "auto_slice_index": "0",
    }
    preview = client.post(
        "/api/templates/reviewed-excel/preview", data=params, files={"file": ("合成.xlsx", body)}
    )
    assert preview.status_code == 200, preview.text
    review = preview.json()
    assert review["selected_region"]["selection_basis"] == "model_group"
    assert review["selected_region"]["archive_version_id"]
    candidate = review["proposal"]["candidates"][0]
    saved = client.post(
        "/api/templates/reviewed-excel/save",
        data={
            "review_id": review["review_id"],
            "proposal_version": review["proposal"]["version"],
            "candidate_id": candidate["id"],
            "selections": json.dumps(candidate["clear"]),
            "name": "合成模板",
            "publish_to_tenant": "true",
        },
    )
    assert saved.status_code == 200, saved.text
    template_id = saved.json()["template_id"].split(":")[-1]
    output = client.get("/api/templates/reviewed-excel/" + template_id + "/download")
    assert output.status_code == 200
    assert load_workbook(BytesIO(output.content)).sheetnames == ["报价"]
    assert client.get(path).json()["versions"][-1]["confirmed"] is False
    review_store.actor["tenant"] = 2
    assert (
        client.post(
            "/api/templates/reviewed-excel/preview",
            data={**params, "expected_actor": "2:7"},
            files={"file": ("合成.xlsx", body)},
        ).status_code
        == 404
    )


def test_auto_slice_blocks_unassigned_cells(review_store, monkeypatch):
    """An apparently valid model group must not silently omit source cells."""
    from app.application.etl.llm_assist_runtime import LlmAssistResult

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("合成.xlsx", sample())},
    ).json()["files"][0]
    base = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"

    def incomplete(messages, schema, tokens):
        return LlmAssistResult(
            used_llm=True,
            model="synthetic",
            data={
                "physical_regions": [
                    {
                        "sheet": "报价",
                        "range": "A1:C1",
                        "labels": ["unknown"],
                        "evidence": [],
                        "dependencies": [],
                        "reason": "仅分类首格，其余原件不可静默遗漏",
                    }
                ],
                "logical_documents": [
                    {
                        "document_id": "only-first",
                        "fragments": [{"sheet": "报价", "range": "A1:C1"}],
                    }
                ],
                "layout_references": [],
            },
        )

    monkeypatch.setattr(
        "app.application.reviewed_excel_template._template_mimo_complete", incomplete
    )
    response = client.post(
        base + "/analyze", data={"version": item["version"], "sample_authorized": "true"}
    )
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["coverage"][0]["unassigned"]
    endpoint = base + "/auto-slices/" + result["version"]
    listing = client.get(endpoint)
    assert listing.status_code == 400 and "完整覆盖" in listing.text
    download = client.get(endpoint + "/0")
    assert download.status_code == 400 and "完整覆盖" in download.text


@pytest.mark.parametrize("fault", ["off", "missing_key", "invalid_base", "invalid_model"])
def test_unsent_model_configuration_does_not_lock_classification(review_store, monkeypatch, fault):
    from app.application.etl.llm_assist_runtime import LlmAssistResult
    from app.application.reviewed_excel_audit import response_capture
    from app.domain.shared_services.conversation.llm_adapter import OpenAICompatibleAdapter

    monkeypatch.setenv("FHD_TEMPLATE_LLM_PROVIDER", "mimo")
    monkeypatch.setenv("FHD_ETL_LLM", "on")
    monkeypatch.setenv("FHD_TEMPLATE_MIMO_BASE_URL", "https://token-plan-cn.xiaomimimo.com/v1")
    monkeypatch.setenv("FHD_TEMPLATE_MIMO_MODEL", "mimo-v2.6-flash")
    monkeypatch.setenv("MIMO_API_KEY", "synthetic-placeholder")
    if fault == "off":
        monkeypatch.setenv("FHD_ETL_LLM", "off")
    elif fault == "invalid_base":
        monkeypatch.setenv("FHD_TEMPLATE_MIMO_BASE_URL", "https://unsupported.invalid/v1")
    elif fault == "invalid_model":
        monkeypatch.setenv("FHD_TEMPLATE_MIMO_MODEL", "unsupported")
    else:
        monkeypatch.setattr(OpenAICompatibleAdapter, "is_configured", property(lambda self: False))
    client = review_store.client
    item = client.post(
        "/api/templates/reviewed-excel/classify",
        data={"expected_actor": "1:7"},
        files={"files": ("synthetic.xlsx", sample())},
    ).json()["files"][0]
    base = f"/api/templates/reviewed-excel/classifications/{item['classification_id']}"
    form = {"version": item["version"], "sample_authorized": "true"}
    failed = client.post(base + "/analyze", data=form)
    assert failed.status_code == 400 and "未调用模型" in failed.json()["detail"]
    checkpoint = client.get(base).json()["checkpoint"]
    assert checkpoint["phase"] == "ready" and checkpoint["calls_in_round"] == 0
    with review_store.engine.connect() as conn:
        state = json.loads(
            conn.execute(
                text("SELECT draft_json FROM etl_runs WHERE id=:id"),
                {"id": item["classification_id"]},
            ).scalar_one()
        )
    assert state["model_audit"][-1]["state"] == "not_sent"
    assert "未调用模型" in state["model_audit"][-1]["reason"]

    def repaired(messages, schema, tokens):
        response_capture()(b'{"choices":[]}', 200, "synthetic")
        return LlmAssistResult(
            used_llm=True,
            data={
                "regions": [
                    {
                        "sheet": "报价",
                        "range": "A1:C4",
                        "labels": ["unknown"],
                        "evidence": [],
                        "dependencies": [],
                        "reason": "synthetic",
                    }
                ]
            },
        )

    monkeypatch.setattr("app.application.reviewed_excel_template._template_mimo_complete", repaired)
    retry = client.post(base + "/analyze", data=form)
    assert retry.status_code == 200, retry.text
