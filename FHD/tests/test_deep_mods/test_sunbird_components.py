"""Synthetic attendance conversions exercise resource, output and policy ownership."""

import importlib
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook


@pytest.fixture
def converter(monkeypatch):
    backend = Path(__file__).resolve().parents[2] / "mods/sunbird-attendance-custom/backend"
    previous = {
        name: module
        for name, module in sys.modules.items()
        if name == "sunbird_attendance" or name.startswith("sunbird_attendance.")
    }
    for name in previous:
        sys.modules.pop(name)
    monkeypatch.syspath_prepend(str(backend))
    module = importlib.import_module("sunbird_attendance.convert")
    owner_config = importlib.import_module("sunbird_attendance.owner_config")
    monkeypatch.setattr(owner_config, "read_policy", lambda: {"conversion_owner": "inner"})
    try:
        yield module
    finally:
        for name in list(sys.modules):
            if name == "sunbird_attendance" or name.startswith("sunbird_attendance."):
                sys.modules.pop(name)
        sys.modules.update(previous)


def _sample(directory):
    from sunbird_attendance.verification import write_conversion_sample

    return write_conversion_sample(directory)


def _invalid_template(path):
    workbook = Workbook()
    workbook.active.title = "明细"
    workbook.save(path)
    workbook.close()


def _convert(converter, source, output, template):
    return converter.convert_attendance_file(
        str(source), str(output), template_path=str(template), month="2026-09", use_llm=False
    )


def test_invalid_template_preserves_previous_output(converter, tmp_path):
    source, template = _sample(tmp_path)
    _invalid_template(template)
    output = tmp_path / "previous.xlsx"
    previous = b"previous successful delivery"
    output.write_bytes(previous)

    result = _convert(converter, source, output, template)

    assert result["success"] is False
    assert "员工块" in result["error"]
    assert output.read_bytes() == previous


def test_invalid_template_does_not_publish_output(converter, tmp_path):
    source, template = _sample(tmp_path)
    _invalid_template(template)
    output = tmp_path / "missing" / "converted.xlsx"

    result = _convert(converter, source, output, template)

    assert result["success"] is False
    assert not output.exists()


def test_workbook_is_closed_when_template_validation_returns(converter, tmp_path, monkeypatch):
    source, template = _sample(tmp_path)
    _invalid_template(template)
    opened = []
    closed = []
    original_open = converter.open_output_workbook
    original_close = Workbook.close

    def open_workbook(output, selected_template):
        workbook = original_open(output, selected_template)
        opened.append(workbook)
        return workbook

    def close_workbook(workbook):
        closed.append(workbook)
        original_close(workbook)

    monkeypatch.setattr(converter, "open_output_workbook", open_workbook)
    monkeypatch.setattr(Workbook, "close", close_workbook)
    result = _convert(converter, source, tmp_path / "output.xlsx", template)

    assert result["success"] is False
    assert len(opened) == 1
    assert opened[0] in closed


def test_failed_save_preserves_output_closes_workbook_and_removes_own_temp(
    converter, tmp_path, monkeypatch
):
    source, template = _sample(tmp_path)
    output = tmp_path / "previous.xlsx"
    output.write_bytes(b"previous successful delivery")
    unrelated = tmp_path / ".previous.xlsx-someone-else.tmp"
    unrelated.write_bytes(b"another conversion")
    before = {path.name: path.read_bytes() for path in tmp_path.iterdir()}
    closed = []
    saved = []
    original_close = Workbook.close

    def failed_save(workbook, destination):
        saved.append(workbook)
        Path(destination).write_bytes(b"partial workbook")
        raise OSError("synthetic storage failure")

    def close_workbook(workbook):
        closed.append(workbook)
        original_close(workbook)

    monkeypatch.setattr(Workbook, "save", failed_save)
    monkeypatch.setattr(Workbook, "close", close_workbook)
    result = _convert(converter, source, output, template)

    assert result == {"success": False, "error": "synthetic storage failure"}
    assert len(saved) == 1
    assert saved[0] in closed
    assert {path.name: path.read_bytes() for path in tmp_path.iterdir()} == before


@pytest.mark.parametrize("valid_template", [True, False])
def test_conversion_restores_outer_policy(converter, tmp_path, valid_template):
    from sunbird_attendance.rules import ACTIVE_POLICY, set_attendance_policy

    source, template = _sample(tmp_path)
    if not valid_template:
        _invalid_template(template)
    set_attendance_policy({"conversion_owner": "outer"})
    try:
        result = _convert(converter, source, tmp_path / "output.xlsx", template)
        assert result["success"] is valid_template
        assert ACTIVE_POLICY.get("conversion_owner") == "outer"
    finally:
        set_attendance_policy(None)


def test_success_preserves_template_and_writes_real_linked_workbook(converter, tmp_path):
    source, template = _sample(tmp_path)
    original = template.read_bytes()
    output = tmp_path / "new" / "output.xlsx"

    result = _convert(converter, source, output, template)

    assert result["success"] is True
    assert result["rows_in"] == result["rows_used_for_template"] == result["employees_matched"] == 1
    assert result["used_llm"] is False
    assert template.read_bytes() == original
    workbook = load_workbook(output)
    try:
        assert workbook.sheetnames == ["明细", "月度统计"]
        assert workbook["明细"].cell(4, 3).value == "交付验证样例"
        assert any(
            isinstance(cell.value, str) and cell.value.startswith("=") and "明细" in cell.value
            for row in workbook["月度统计"]
            for cell in row
        )
    finally:
        workbook.close()


def test_record_conversion_restores_caller_policy_and_keeps_database_contract(converter, tmp_path):
    from sunbird_attendance.rules import ACTIVE_POLICY, set_attendance_policy

    source, template = _sample(tmp_path)
    records = converter.parse_attendance_workbook(source, month="2026-09", use_llm=False).records
    set_attendance_policy({"conversion_owner": "outer"})
    try:
        result = converter.convert_attendance_records(
            records, tmp_path / "output.xlsx", template_path=template, month_label="2026-09"
        )
        assert result["success"] is True
        assert result["input"] == "database"
        assert result["rows_in"] == result["employees_matched"] == 1
        assert result["header_info"] is None
        assert result["used_llm"] is False
        assert ACTIVE_POLICY.get("conversion_owner") == "outer"
    finally:
        set_attendance_policy(None)


def test_workbook_writer_runs_without_public_converter_dependency(converter, tmp_path, monkeypatch):
    from sunbird_attendance.workbook_writer import AttendanceWorkbookWriter

    source, template = _sample(tmp_path)
    records = converter.parse_attendance_workbook(source, month="2026-09", use_llm=False).records

    def facade_is_not_an_implementation_dependency(*args, **kwargs):
        raise AssertionError("the workbook component looked back through its public facade")

    monkeypatch.setattr(
        converter, "open_output_workbook", facade_is_not_an_implementation_dependency
    )
    result = AttendanceWorkbookWriter().write(
        records, tmp_path / "output.xlsx", template, "2026-09", None
    )
    assert result["success"] is True
    assert result["employees_matched"] == 1


def test_nested_policy_scope_restores_its_parent_after_exception(converter):
    from sunbird_attendance.rules import ACTIVE_POLICY, attendance_policy_scope

    assert ACTIVE_POLICY.get("conversion_owner") is None
    with attendance_policy_scope({"conversion_owner": "outer"}):
        with pytest.raises(ValueError, match="synthetic nested failure"):
            with attendance_policy_scope({"conversion_owner": "inner"}):
                assert ACTIVE_POLICY.get("conversion_owner") == "inner"
                raise ValueError("synthetic nested failure")
        assert ACTIVE_POLICY.get("conversion_owner") == "outer"
    assert ACTIVE_POLICY.get("conversion_owner") is None


def test_same_name_records_keep_department_identity_and_source_unmodified(converter, tmp_path):
    from dataclasses import replace

    from sunbird_attendance.roster_identity import identify_roster_records

    source, _ = _sample(tmp_path)
    sample = converter.parse_attendance_workbook(source, month="2026-09", use_llm=False).records[0]
    records = [replace(sample, department="one"), replace(sample, department="two")]
    roster = [("one", "day", sample.employee_name), ("two", "day", sample.employee_name)]
    matched, keyed, display = identify_roster_records(records, roster)
    assert matched[0].employee_name != matched[1].employee_name
    assert [record.employee_name for record in matched] == [entry[2] for entry in keyed]
    assert [display[record.employee_name] for record in matched] == [sample.employee_name] * 2
    assert records[0].employee_name == records[1].employee_name == sample.employee_name
    assert roster[0][2] == roster[1][2] == sample.employee_name
    with pytest.raises(ValueError, match="缺少唯一身份"):
        identify_roster_records([replace(sample, raw_times=list(sample.daily_times))], roster)
