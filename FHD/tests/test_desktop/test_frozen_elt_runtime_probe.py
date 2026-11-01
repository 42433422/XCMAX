"""The packaged dependency gate executes the real synthetic XLSX derivation path."""

from __future__ import annotations

import builtins
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.application import reviewed_excel_template
from XCAGI import run_fastapi


def test_elt_probe_derives_and_reads_back_without_starting_server(monkeypatch, capsys):
    calls = []
    derive = reviewed_excel_template.derive_document_workbook

    def record(source, fragments, **kwargs):
        result = derive(source, fragments, **kwargs)
        calls.append((source, result, fragments))
        return result

    monkeypatch.setattr(reviewed_excel_template, "derive_document_workbook", record)
    run_fastapi._verify_frozen_elt_runtime()

    assert len(calls) == 1
    source, result, fragments = calls[0]
    assert fragments == [{"sheet": "Synthetic", "range": "A1:B3"}]
    for body in (source, result):
        workbook = load_workbook(BytesIO(body), data_only=False)
        try:
            assert workbook.sheetnames == ["Synthetic"]
            assert workbook.active["B3"].value == "=SUM(B2:B2)"
            assert workbook.active.auto_filter.ref == "A1:B3"
        finally:
            workbook.close()
    assert "XLSX inspect/derive/readback" in capsys.readouterr().out


def test_elt_probe_fails_closed_when_xml_runtime_is_missing(monkeypatch):
    original_import = builtins.__import__

    def without_safe_xml(name, *args, **kwargs):
        if name == "defusedxml" or name.startswith("defusedxml."):
            raise ModuleNotFoundError("No module named 'defusedxml'", name="defusedxml")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", without_safe_xml)
    with pytest.raises(ModuleNotFoundError, match="defusedxml"):
        run_fastapi._verify_frozen_elt_runtime()


def test_elt_probe_rejects_corrupt_derived_output(monkeypatch):
    derive = reviewed_excel_template.derive_document_workbook

    def corrupt(source, fragments, **kwargs):
        workbook = load_workbook(BytesIO(derive(source, fragments, **kwargs)))
        try:
            workbook.active["B3"] = 999
            output = BytesIO()
            workbook.save(output)
            return output.getvalue()
        finally:
            workbook.close()

    monkeypatch.setattr(reviewed_excel_template, "derive_document_workbook", corrupt)
    with pytest.raises(RuntimeError, match="did not preserve"):
        run_fastapi._verify_frozen_elt_runtime()
