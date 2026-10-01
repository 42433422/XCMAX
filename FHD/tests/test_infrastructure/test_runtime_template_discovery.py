from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.infrastructure.templates.template_store_impl import FileSystemTemplateStore


def test_attendance_template_is_discovered_from_desktop_data_dir(tmp_path, monkeypatch) -> None:
    runtime_root = tmp_path / "runtime"
    tenant_templates = runtime_root / "tenants" / "424" / "templates"
    tenant_templates.mkdir(parents=True)
    template = tenant_templates / "考勤-2026-3月份考勤统计表.xlsx"
    template.write_bytes(b"PK test workbook")
    code_root = tmp_path / "code"
    code_root.mkdir()

    monkeypatch.setattr(
        "app.infrastructure.templates.template_store_impl.get_app_data_dir",
        lambda: str(runtime_root),
    )
    monkeypatch.setattr(
        "app.infrastructure.tenant_scope.current_tenant_id",
        lambda: 424,
    )
    store = FileSystemTemplateStore(str(code_root))

    items = store._discover_excel_templates()
    found = next(item for item in items if item["filename"] == template.name)
    assert found["template_type"] == "考勤记录"
    assert found["business_scope"] == "shipmentRecords"
    assert store.resolve_template_file(found["id"]) == str(template)


def test_bundled_delivery_template_is_empty_resolvable_order_form():
    root = Path(__file__).resolve().parents[2]
    store = FileSystemTemplateStore(str(root))
    item = next(t for t in store._discover_excel_templates() if t["filename"] == "标准发货单.xlsx")
    assert item["business_scope"] == "orders"
    assert item["category"] == "excel" and item["source"] == "fs_scan"
    assert store.resolve_template_file(item["id"]) == item["path"]
    wb = load_workbook(item["path"], data_only=True)
    assert wb.sheetnames == ["发货单"]
    ws = wb.active
    assert ws["A1"].value == "标准发货单"
    assert [v for v in next(ws.iter_rows(min_row=3, max_row=3, values_only=True)) if v] == (
        ["产品型号", "产品名称", "数量/件", "规格/KG", "数量/KG", "单价/元", "金额/元"]
    )
    assert all(cell.value is None for row in ws.iter_rows(min_row=4) for cell in row)
    from app.legacy.documents.legacy_shipment_document import (
        load_legacy_shipment_document_generator,
    )

    loaded = load_legacy_shipment_document_generator(caller_file=__file__)
    generator = object.__new__(loaded.ShipmentDocumentGenerator)
    generator._clear_template_data(ws)
    generator._fill_from_template(
        ws,
        "WIN-ORDER",
        loaded.PurchaseUnitInfo(name="验收客户"),
        [
            {
                "model_number": "MODEL",
                "name": "验收产品",
                "quantity_tins": 2,
                "tin_spec": 10,
                "quantity_kg": 20,
                "unit_price": 12.5,
                "amount": 250,
            }
        ],
        20,
        2,
        250,
        "",
        "2026-10-01",
    )
    assert "验收客户" in ws["A2"].value and "WIN-ORDER" in ws["A2"].value
    assert [ws.cell(4, c).value for c in (1, 4, 5, 6, 7, 8, 9)] == [
        "MODEL",
        "验收产品",
        2,
        10,
        20,
        12.5,
        250,
    ]
    wb.close()
