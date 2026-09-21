"""回归：客户自带标准送货单模板（含合并单元格）生成发货单不得中断。

真实客户模板与产品内置送货单版式一致：第 2 行购货单位为 A2:J2 合并、第 3 行表头、
产品型号列 A3:C3 与数据行 A4:C4…A14:C14 合并、第 15 行合计 A15:D15 合并、
第 16 行人民币大写 B16:G16 合并（金额列为 SUM 公式）。

旧实现把金额合计写死到 D16，而 D16 落在 B16:G16 合并区内，openpyxl 对 MergedCell
赋值抛 "attribute 'value' is read-only"，导致 `/api/shipment/generate` 返回 500
（界面显示「生成失败: 服务器内部错误」）。本用例锁定合并模板可正常生成。
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.legacy.documents.legacy_shipment_document import (
    load_legacy_shipment_document_generator,
)


def _build_merged_customer_template(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "送货单"
    ws["A1"] = "送货单"
    ws.merge_cells("A1:J1")
    ws["A2"] = "购货单位：              联系人：              日期：              订单编号："
    ws.merge_cells("A2:J2")
    for column, header in enumerate(
        [
            "产品型号",
            None,
            None,
            "产品名称",
            "数量/件",
            "规格/KG",
            "数量/KG",
            "单价/元",
            "金额/元",
            "备注",
        ],
        1,
    ):
        if header:
            ws.cell(row=3, column=column, value=header)
    for row in range(3, 15):
        ws.merge_cells(f"A{row}:C{row}")
    ws["A15"] = "合计"
    ws.merge_cells("A15:D15")
    ws["E15"] = "=SUM(E4:E14)"
    ws["G15"] = "=SUM(G4:G14)"
    ws["A16"] = "人民币大写："
    ws.merge_cells("B16:G16")
    ws["H16"] = "（小写）¥"
    ws["I16"] = "=SUM(I4:I14)"
    wb.save(path)
    return path


def test_generate_with_merged_template_preserves_layout_and_totals(tmp_path):
    template = _build_merged_customer_template(tmp_path / "客户送货单模板.xlsx")
    legacy = load_legacy_shipment_document_generator(caller_file=__file__)
    generator = legacy.ShipmentDocumentGenerator(
        db_path=str(tmp_path / "products.db"), output_dir=str(tmp_path / "out")
    )

    doc = generator.generate_document(
        "",
        {
            "date": "2026-09-21",
            "products": [
                {
                    "name": "UI验收商品-0921",
                    "model_number": "UI-0921-A",
                    "quantity_tins": 24,
                    "tin_spec": 25,
                    "quantity_kg": 600,
                    "unit_price": 3.5,
                    "amount": 2100,
                }
            ],
        },
        purchase_unit=legacy.PurchaseUnitInfo(name="验收客户-PHASE-C"),
        template_name=str(template),
        custom_order_number="26-09-00003A",
    )

    ws = load_workbook(Path(doc.filepath)).active
    # 购货单位行写在合并区锚点，日期与订单号可见
    assert "验收客户-PHASE-C" in ws["A2"].value
    assert "26-09-00003A" in ws["A2"].value
    # 数据行落在表头下一行（第 4 行），各列与表头对齐
    assert [ws.cell(row=4, column=column).value for column in (1, 4, 5, 6, 7, 8, 9)] == [
        "UI-0921-A",
        "UI验收商品-0921",
        24,
        25,
        600,
        3.5,
        2100,
    ]
    # 合计行：数量/件、数量/KG 与「合计」标签都在
    assert (ws["A15"].value, ws["E15"].value, ws["G15"].value) == ("合计", 24, 600)
    # 金额合计不写进 B16:G16 合并区，保留模板 SUM 公式
    assert ws["B16"].value in (None, "")
    assert ws["I16"].value == "=SUM(I4:I14)"
