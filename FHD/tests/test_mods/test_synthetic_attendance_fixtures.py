"""Published attendance fixtures are synthetic, complete and self-contained."""

from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
FIXTURES = ROOT / "tests/fixtures/synthetic_attendance"


def test_template_contains_only_explicit_synthetic_employee_keys():
    path = FIXTURES / "synthetic-attendance-template.xlsx"
    workbook = load_workbook(path, read_only=True)
    try:
        sheet = workbook["明细"]
        keys = [sheet.cell(row, 3).value for row in range(4, 910, 6)]
        assert keys[:80] == [f"SYNTHETIC-{i:03d}" for i in range(1, 81)]
        assert keys[80:] == [None] * 71
        assert sheet["A1"].value == "SYNTHETIC TEST DATA ONLY"
    finally:
        workbook.close()
    runtime = ROOT.parent / (
        "成都修茈科技有限公司/MODstore_deploy/examples/"
        "taiyangniao-attendance-employee/backend/templates/424/考勤-2026-3月份考勤统计表.xlsx"
    )
    assert runtime.read_bytes() == path.read_bytes()


def test_daily_and_raw_records_are_complete_and_use_no_customer_identifiers():
    workbook = load_workbook(
        FIXTURES / "synthetic-dingtalk-source.xlsx", read_only=True
    )
    expected_names = {f"SYNTHETIC-{i:03d}" for i in range(1, 81)}
    try:
        daily = list(workbook["每日统计"].iter_rows(min_row=3, values_only=True))
        raw = list(workbook["原始记录"].iter_rows(min_row=3, values_only=True))
        assert len(daily) == 2480
        assert len(raw) == 14880
        assert {row[0] for row in daily} == expected_names
        assert {row[0] for row in raw} == expected_names
        assert all(row[2] == "合成测试部门" for row in daily)
        assert all(str(row[3]).startswith("TEST-") for row in daily)
        assert all(str(row[4]).startswith("synthetic-user-") for row in daily)
        assert len({(row[0], row[5]) for row in daily}) == 2480
    finally:
        workbook.close()


def test_fixtures_have_no_embedded_customer_media_macros_or_external_links():
    for path in FIXTURES.glob("*.xlsx"):
        with ZipFile(path) as package:
            names = package.namelist()
            assert not any(
                token in name.lower()
                for name in names
                for token in (
                    "externallinks",
                    "vbaproject",
                    "embeddings",
                    "media/",
                    "comments",
                )
            )
            for name in names:
                if name.endswith(".rels"):
                    assert b'TargetMode="External"' not in package.read(name)
