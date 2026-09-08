from __future__ import annotations

from unittest.mock import MagicMock, patch

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from starlette.requests import Request

from app.application.agent_orchestrator import (
    AgentOrchestrator,
    InMemoryAgentRunRepository,
    get_agent_run_repository,
    set_agent_run_repository_for_tests,
)
from app.application.agent_orchestrator.approval_grant import (
    clear_consumed_approval_grants_for_tests,
)
from app.application.agent_orchestrator.run_control import clear_run_controls_for_tests
from app.application.agent_orchestrator.run_models import AgentRun, AgentStep, ToolCall
from app.application.agent_orchestrator.task_dispatcher import (
    set_agent_task_dispatcher_for_tests,
    stop_agent_task_dispatcher,
)
from app.application.agent_orchestrator.task_execution_repository import (
    InMemoryTaskExecutionRepository,
    get_task_execution_repository,
    set_task_execution_repository_for_tests,
)
from app.fastapi_routes.domains.agent.routes import router
from app.infrastructure.auth.agent_principal import AgentPrincipal, require_agent_principal


def test_agent_principal_keeps_actor_but_uses_verified_tutorial_tenant() -> None:
    request = Request(
        {
            "type": "http",
            "method": "GET",
            "path": "/api/agent/tasks",
            "headers": [],
            "query_string": b"",
            "server": ("test", 80),
            "client": ("test", 1),
            "scheme": "http",
        }
    )
    request.state.tutorial_active = True
    request.state.tenant_id = 29
    user = MagicMock(id=17, username="learner", tenant_id=7, tier="", role="user")
    with patch("app.infrastructure.auth.agent_principal.resolve_session_user", return_value=user):
        principal = require_agent_principal(request)

    assert principal.user_id == "17"
    assert principal.tenant_id == "29"


@pytest.fixture(autouse=True)
def _isolated_agent_task_repositories(tmp_path, monkeypatch):
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker

    from app.application.agent_orchestrator.approval_consumption_repository import (
        SQLAlchemyApprovalConsumptionRepository,
    )
    from app.db.models.agent_approval import AgentApprovalConsumption

    engine = create_engine(f"sqlite:///{tmp_path / 'approval-consumption.db'}")
    AgentApprovalConsumption.__table__.create(engine)
    repository = SQLAlchemyApprovalConsumptionRepository(sessionmaker(bind=engine))
    monkeypatch.setattr(
        "app.application.agent_orchestrator.approval_grant._consumption_repository",
        lambda: repository,
    )
    set_agent_run_repository_for_tests(InMemoryAgentRunRepository())
    set_task_execution_repository_for_tests(InMemoryTaskExecutionRepository())
    set_agent_task_dispatcher_for_tests(None)
    yield
    stop_agent_task_dispatcher(timeout=0)
    set_agent_task_dispatcher_for_tests(None)
    set_task_execution_repository_for_tests(None)
    set_agent_run_repository_for_tests(None)
    engine.dispose()


@pytest.mark.parametrize(
    "method,path",
    [
        ("GET", "/api/agent/tasks/private-task"),
        ("POST", "/api/agent/tasks/private-task/read"),
        ("POST", "/api/agent/tasks/private-task/archive"),
    ],
)
def test_empty_tenant_cannot_access_scoped_task(method, path):
    from app.application.agent_orchestrator.task_models import AgentTask

    repo = get_agent_run_repository()
    task = AgentTask(
        task_id="private-task",
        user_id="u1",
        tenant_id="other",
        title="private tenant result",
        status="completed",
        attention_state="result_unread",
    )
    repo.save_task(task)
    before = repo.get_task(user_id="u1", task_id=task.task_id, tenant_id="other").to_dict()
    response = _client().request(method, path)
    assert response.status_code == 404
    assert "private tenant result" not in response.text
    assert repo.get_task(user_id="u1", task_id=task.task_id, tenant_id="other").to_dict() == before


@pytest.mark.parametrize("path", ["/api/agent/tasks", "/api/agent/tasks/events/stream?once=true"])
def test_empty_tenant_task_feeds_exclude_scoped_tasks(path):
    from app.application.agent_orchestrator.task_models import AgentTask

    repo = get_agent_run_repository()
    repo.save_task(
        AgentTask(
            task_id="private-task", user_id="u1", tenant_id="other", title="private tenant result"
        )
    )
    repo.save_task(AgentTask(task_id="public-task", user_id="u1", title="own unscoped task"))
    response = _client().get(path)
    assert response.status_code == 200
    assert "private tenant result" not in response.text
    assert "own unscoped task" in response.text


def test_empty_tenant_does_not_deduplicate_another_tenants_task():
    body = {
        "task_id": "same-id",
        "title": "查询产品",
        "tool_id": "products",
        "action": "query",
        "params": {"keyword": "5003"},
    }
    scoped = _client(tenant_id="other").post("/api/agent/tasks", json=body)
    unscoped = _client().post("/api/agent/tasks", json=body)
    assert scoped.status_code == unscoped.status_code == 202
    assert unscoped.json()["deduplicated"] is False
    assert scoped.json()["data"]["run_id"] != unscoped.json()["data"]["run_id"]
    repeated = _client().post("/api/agent/tasks", json=body)
    assert repeated.json()["deduplicated"] is True
    assert repeated.json()["data"]["run_id"] == unscoped.json()["data"]["run_id"]


@pytest.mark.parametrize("second_mod", ["mod-b", ""])
def test_unified_task_deduplication_rejects_different_mod_scope(second_mod):
    body = {
        "task_id": "same-mod-task-id",
        "title": "查询产品",
        "tool_id": "products",
        "action": "query",
        "params": {"keyword": "5003"},
    }
    first_client = _client(
        mod_authorization={"mod_id": "mod-a", "session_row_id": 1, "user_id": "u1"}
    )
    first = first_client.post("/api/agent/tasks", json=body)
    assert first.status_code == 202
    run_id = first.json()["data"]["run_id"]
    before = get_agent_run_repository().get(run_id).to_dict()
    second = _client(mod_authorization={"mod_id": second_mod} if second_mod else None)
    rejected = second.post("/api/agent/tasks", json=body)
    assert rejected.status_code == 409
    assert run_id not in rejected.text
    assert get_agent_run_repository().get(run_id).to_dict() == before
    same = first_client.post("/api/agent/tasks", json=body)
    assert same.json()["deduplicated"] is True
    assert same.json()["data"]["run_id"] == run_id


@pytest.mark.parametrize(
    "method,suffix", [("GET", ""), ("GET", "/events"), ("POST", "/cancel"), ("POST", "/pause")]
)
@pytest.mark.parametrize("other_mod", ["mod-b", ""])
def test_run_access_rejects_cross_mod_scope(method, suffix, other_mod):
    owner = _client(mod_authorization={"mod_id": "mod-a", "session_row_id": 1, "user_id": "u1"})
    created = owner.post(
        "/api/agent/tasks",
        json={
            "task_id": "private-mod-task",
            "title": "查询产品",
            "tool_id": "products",
            "action": "query",
            "params": {"keyword": "5003"},
        },
    )
    assert created.status_code == 202
    run_id = created.json()["data"]["run_id"]
    before = get_agent_run_repository().get(run_id).to_dict()
    other = _client(mod_authorization={"mod_id": other_mod} if other_mod else None)
    rejected = other.request(method, f"/api/agent/runs/{run_id}{suffix}")
    assert rejected.status_code == 403
    assert run_id not in rejected.text
    assert get_agent_run_repository().get(run_id).to_dict() == before
    assert owner.get(f"/api/agent/runs/{run_id}").status_code == 200


@pytest.mark.parametrize("method,suffix", [("GET", ""), ("POST", "/read"), ("POST", "/archive")])
@pytest.mark.parametrize("other_mod", ["mod-b", ""])
def test_task_access_rejects_cross_mod_scope(method, suffix, other_mod):
    owner = _client(mod_authorization={"mod_id": "mod-a", "session_row_id": 1, "user_id": "u1"})
    created = owner.post(
        "/api/agent/tasks",
        json={
            "task_id": "private-mod-task",
            "title": "查询产品",
            "tool_id": "products",
            "action": "query",
            "params": {"keyword": "5003"},
        },
    )
    assert created.status_code == 202
    orchestrator = AgentOrchestrator()
    before = orchestrator.get_task(user_id="u1", task_id="private-mod-task", tenant_id="").to_dict()
    other = _client(mod_authorization={"mod_id": other_mod} if other_mod else None)
    response = other.request(method, f"/api/agent/tasks/private-mod-task{suffix}")
    assert response.status_code == 404
    assert "private-mod-task" not in response.text
    assert (
        orchestrator.get_task(user_id="u1", task_id="private-mod-task", tenant_id="").to_dict()
        == before
    )
    assert owner.get("/api/agent/tasks/private-mod-task").status_code == 200


def test_task_collection_filters_mod_before_limit_and_stream_snapshot():
    owner = _client(mod_authorization={"mod_id": "mod-a", "session_row_id": 1, "user_id": "u1"})
    other = _client(mod_authorization={"mod_id": "mod-b", "session_row_id": 2, "user_id": "u1"})
    body = {
        "title": "查询产品",
        "tool_id": "products",
        "action": "query",
        "params": {"keyword": "5003"},
    }
    assert (
        owner.post("/api/agent/tasks", json={**body, "task_id": "visible-task"}).status_code == 202
    )
    for index in range(101):
        assert (
            other.post(
                "/api/agent/tasks", json={**body, "task_id": f"hidden-task-{index}"}
            ).status_code
            == 202
        )
    listing = owner.get("/api/agent/tasks?limit=1").json()
    assert [task["task_id"] for task in listing["data"]] == ["visible-task"]
    assert listing["count"] == 1
    runs = owner.get("/api/agent/runs?limit=1").json()
    assert runs["count"] == 1
    assert runs["data"][0]["metadata"]["task_context"]["task_id"] == "visible-task"
    assert _client().get("/api/agent/runs").json()["data"] == []
    stream = owner.get("/api/agent/tasks/events/stream?once=true")
    assert stream.status_code == 200
    assert "visible-task" in stream.text
    assert "hidden-task" not in stream.text
    runtime = owner.get("/api/agent/task-runtime").json()
    assert runtime["data"]["progress"]["task_count"] == 1
    assert _client().get("/api/agent/tasks").json()["data"] == []


def test_runtime_active_count_excludes_other_accounts_and_mods():
    repository = get_agent_run_repository()
    active = []
    for user, mod, tenant in [
        ("u1", "a", ""),
        ("u1", "b", ""),
        ("u2", "a", ""),
        ("u1", "a", "other"),
    ]:
        run = AgentRun(user_id=user, message="active")
        run.metadata["runtime_context"] = {
            "tenant_id": tenant,
            "_mod_authorization": {"mod_id": mod},
        }
        repository.save(run)
        active.append({"run_id": run.run_id})
    dispatcher = MagicMock()
    dispatcher.snapshot.return_value = {
        "running": True,
        "max_workers": 4,
        "active_count": 4,
        "active": active,
    }
    with patch(
        "app.fastapi_routes.domains.agent.task_routes.get_agent_task_dispatcher",
        return_value=dispatcher,
    ):
        response = _client(mod_authorization={"mod_id": "a"}).get("/api/agent/task-runtime")
    assert response.status_code == 200
    assert response.json()["data"]["active_count"] == 1


def _drain_background_run(run_id: str) -> AgentRun:
    queue = get_task_execution_repository()
    claimed = queue.claim("route-test-worker", lease_seconds=30)
    assert claimed is not None and claimed.run_id == run_id
    run = AgentOrchestrator().execute_dispatched_run(
        run_id,
        recovered=claimed.recovery_count > 0,
    )
    assert run is not None
    state = {
        "completed": "completed",
        "failed": "failed",
        "cancelled": "cancelled",
        "paused": "paused",
    }.get(run.status, "blocked")
    queue.finish(run_id, "route-test-worker", state)
    return run


def _client(
    user_id: str | None = "u1", *, tenant_id: str = "", mod_authorization=None
) -> TestClient:
    app = FastAPI()
    app.include_router(router)
    if user_id is not None:
        app.dependency_overrides[require_agent_principal] = lambda: AgentPrincipal(
            user_id=user_id,
            tenant_id=tenant_id,
            mod_authorization=mod_authorization,
        )
    return TestClient(app, raise_server_exceptions=False)


def _planner_fallback_patches():
    return (
        patch("app.application.workflow.planner.get_ai_conversation_service"),
        patch(
            "app.application.workflow.planner.LLMWorkflowPlanner._plan_with_react_multiagent",
            return_value=None,
        ),
        patch(
            "app.application.normal_chat_dispatch.resolve_tool_execution_profile",
            return_value="pro_default",
        ),
        patch("app.application.get_user_memory_rag_app_service", side_effect=ImportError),
    )


def test_create_get_and_list_agent_run() -> None:
    get_agent_run_repository().clear()
    clear_consumed_approval_grants_for_tests()
    client = _client()
    patches = _planner_fallback_patches()

    with (
        patches[0],
        patches[1],
        patches[2],
        patches[3],
        patch(
            "app.application.facades.tools_facade.execute_registered_workflow_tool",
            return_value={"success": True, "data": [{"model_number": "XG-5003"}]},
        ),
    ):
        response = client.post(
            "/api/agent/runs",
            json={
                "message": "查数据库产品 XG-5003",
                "user_id": "forged-user",
                "runtime_context": {"source": "route-test"},
            },
        )
        completed_run = _drain_background_run(response.json()["data"]["run_id"])

    assert response.status_code == 202
    payload = response.json()
    assert payload["success"] is True
    assert payload["data"]["status"] == "queued"
    assert payload["execution"]["state"] == "queued"
    run = completed_run.to_dict()
    assert run["status"] == "completed"
    assert run["user_id"] == "u1"
    assert run["intent"] == "business_db_read"
    assert run["metadata"]["runtime_context"]["source"] == "route-test"
    assert "run.completed" in [event["event_type"] for event in run["events"]]

    get_response = client.get(f"/api/agent/runs/{run['run_id']}")
    assert get_response.status_code == 200
    assert get_response.json()["data"]["run_id"] == run["run_id"]

    events_response = client.get(f"/api/agent/runs/{run['run_id']}/events")
    assert events_response.status_code == 200
    events_payload = events_response.json()
    assert events_payload["count"] == len(run["events"])
    assert events_payload["data"][-1]["event_type"] == "run.completed"

    first_event_id = events_payload["data"][0]["event_id"]
    tail_response = client.get(
        f"/api/agent/runs/{run['run_id']}/events",
        params={"after_event_id": first_event_id},
    )
    assert tail_response.status_code == 200
    assert tail_response.json()["count"] == len(run["events"]) - 1

    stream_response = client.get(f"/api/agent/runs/{run['run_id']}/events/stream")
    assert stream_response.status_code == 200
    assert stream_response.headers["content-type"].startswith("text/event-stream")
    assert "event: run.completed" in stream_response.text
    assert "event: stream.closed" in stream_response.text

    list_response = client.get("/api/agent/runs", params={"user_id": "u1"})
    assert list_response.status_code == 200
    list_payload = list_response.json()
    assert list_payload["count"] == 1
    assert list_payload["data"][0]["run_id"] == run["run_id"]


def test_record_observed_tool_run_is_owned_and_task_scoped() -> None:
    get_agent_run_repository().clear()
    response = _client("u1", tenant_id="shadow-tenant").post(
        "/api/agent/runs/observed-tool",
        json={
            "message": "查 5003 产品",
            "tool_id": "products",
            "action": "query",
            "params": {"keyword": "5003"},
            "output": {"success": True, "data": []},
            "response": "未找到",
            "runtime_context": {
                "task_id": "task-product-5003",
                "conversation_id": "conversation-product-5003",
            },
        },
    )
    assert response.status_code == 200
    reference = response.json()["data"]
    assert reference["status"] == "completed"
    assert reference["task_id"] == "task-product-5003"
    stored = get_agent_run_repository().get(reference["run_id"])
    assert stored is not None
    run = stored.to_dict()
    assert run["user_id"] == "u1"
    assert run["metadata"]["task_context"]["conversation_id"] == "conversation-product-5003"
    assert run["metadata"]["runtime_context"]["tenant_id"] == "shadow-tenant"
    assert run["tool_calls"][0]["tool_id"] == "products"
    assert run["metadata"]["runtime_context"]["observation_trust"] == ("authenticated_client")
    assert run["metadata"]["non_retryable"] is True


def test_record_observed_tool_rejects_write_or_unknown_tools() -> None:
    response = _client("u1").post(
        "/api/agent/runs/observed-tool",
        json={
            "message": "伪造写入",
            "tool_id": "business_db_write",
            "action": "create",
            "params": {},
            "output": {"success": True},
        },
    )
    assert response.status_code == 400


def test_failed_observation_cannot_be_retried_as_an_execution_plan() -> None:
    get_agent_run_repository().clear()
    client = _client("u1")
    observed = client.post(
        "/api/agent/runs/observed-tool",
        json={
            "message": "查产品失败",
            "tool_id": "products",
            "action": "query",
            "params": {"keyword": "5003"},
            "output": {"success": False, "message": "network unavailable"},
        },
    )
    assert observed.status_code == 200
    reference = observed.json()["data"]
    assert reference["status"] == "failed"

    retry = client.post(f"/api/agent/runs/{reference['run_id']}/retry", json={})
    assert retry.status_code == 409
    assert "观察记录" in retry.json()["message"]


def test_continue_waiting_agent_run() -> None:
    get_agent_run_repository().clear()
    clear_consumed_approval_grants_for_tests()
    client = _client()
    patches = _planner_fallback_patches()

    with (
        patches[0],
        patches[1],
        patches[2],
        patches[3],
        patch(
            "app.application.facades.tools_facade.execute_registered_workflow_tool"
        ) as mock_execute,
    ):
        create_response = client.post(
            "/api/agent/runs",
            json={
                "message": "请把客户 星光贸易 写入数据库",
                "user_id": "u1",
                "runtime_context": {"source": "route-continue-test"},
            },
        )
        assert create_response.status_code == 202
        queued = create_response.json()["data"]
        assert queued["status"] == "queued"
        waiting = _drain_background_run(queued["run_id"])
        approval_response = client.get(f"/api/agent/runs/{waiting.run_id}").json()
        approval_grant = approval_response["approval"]["grant"]
        assert waiting.status == "waiting_user"
        mock_execute.assert_not_called()

        mock_execute.return_value = {"success": True, "message": "客户已写入"}
        continue_response = client.post(
            f"/api/agent/runs/{waiting.run_id}/continue",
            json={"approval_grant": approval_grant, "approved_by": "forged-approver"},
        )
        assert continue_response.json()["data"]["status"] == "queued"
        completed_run = _drain_background_run(waiting.run_id)

    assert continue_response.status_code == 202
    payload = continue_response.json()
    assert payload["success"] is True
    completed = completed_run.to_dict()
    assert completed["status"] == "completed"
    assert completed["steps"][0]["status"] == "completed"
    event_types = [event["event_type"] for event in completed["events"]]
    assert "step.approved" in event_types
    assert "tool.completed" in event_types
    mock_execute.assert_called_once()


def test_agent_routes_require_authentication_and_enforce_ownership() -> None:
    get_agent_run_repository().clear()
    anonymous = _client(user_id=None)
    assert anonymous.get("/api/agent/runs").status_code == 401

    owner = _client("owner")
    patches = _planner_fallback_patches()
    with (
        patches[0],
        patches[1],
        patches[2],
        patches[3],
        patch(
            "app.application.facades.tools_facade.execute_registered_workflow_tool",
            return_value={"success": True, "data": []},
        ),
    ):
        created = owner.post("/api/agent/runs", json={"message": "查产品"}).json()["data"]

    stranger = _client("stranger")
    assert stranger.get(f"/api/agent/runs/{created['run_id']}").status_code == 403
    assert stranger.get(f"/api/agent/runs/{created['run_id']}/events").status_code == 403
    assert stranger.get("/api/agent/runs").json()["count"] == 0


def test_continue_rejects_missing_mismatched_and_replayed_grants() -> None:
    get_agent_run_repository().clear()
    clear_consumed_approval_grants_for_tests()
    owner = _client("owner")
    patches = _planner_fallback_patches()
    with (
        patches[0],
        patches[1],
        patches[2],
        patches[3],
        patch(
            "app.application.facades.tools_facade.execute_registered_workflow_tool",
            return_value={"success": True},
        ) as execute,
    ):
        response = owner.post("/api/agent/runs", json={"message": "请把客户 星光贸易 写入数据库"})
        run_id = response.json()["data"]["run_id"]
        _drain_background_run(run_id)
        grant = owner.get(f"/api/agent/runs/{run_id}").json()["approval"]["grant"]
        second_grant = owner.get(f"/api/agent/runs/{run_id}").json()["approval"]["grant"]

        missing_grant = owner.post(f"/api/agent/runs/{run_id}/continue", json={})
        assert missing_grant.status_code == 403
        assert missing_grant.json()["message"] == ("approval_grant 无效、过期或与当前步骤不匹配")
        assert (
            _client("stranger")
            .post(f"/api/agent/runs/{run_id}/continue", json={"approval_grant": grant})
            .status_code
            == 403
        )
        approved = owner.post(
            f"/api/agent/runs/{run_id}/continue",
            json={"approval_grant": grant},
        )
        assert approved.status_code == 202
        assert approved.json()["execution"]["state"] == "queued"
        assert (
            owner.post(
                f"/api/agent/runs/{run_id}/continue",
                json={"approval_grant": second_grant},
            ).status_code
            == 403
        )
        completed = _drain_background_run(run_id)
        assert completed.status == "completed"
        assert len(completed.tool_calls) == 1
        assert (
            owner.post(
                f"/api/agent/runs/{run_id}/continue",
                json={"approval_grant": grant},
            ).status_code
            == 403
        )
        execution = get_task_execution_repository().get(run_id)
        assert execution is not None
        assert execution.state == "completed"
        assert execution.execution_count == 2
        execute.assert_called_once()


def test_approval_storage_failure_preserves_waiting_run_and_allows_retry(tmp_path, monkeypatch):
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker

    from app.application.agent_orchestrator import approval_grant
    from app.application.agent_orchestrator.approval_consumption_repository import (
        SQLAlchemyApprovalConsumptionRepository,
    )

    owner = _client("owner")
    patches = _planner_fallback_patches()
    engine = create_engine(f"sqlite:///{tmp_path / 'unmigrated.db'}")
    durable_repository = approval_grant._consumption_repository()
    unavailable_repository = SQLAlchemyApprovalConsumptionRepository(sessionmaker(bind=engine))
    try:
        with (
            patches[0],
            patches[1],
            patches[2],
            patches[3],
            patch(
                "app.application.facades.tools_facade.execute_registered_workflow_tool",
                return_value={"success": True},
            ) as execute,
        ):
            response = owner.post(
                "/api/agent/runs", json={"message": "请把客户 星光贸易 写入数据库"}
            )
            run_id = response.json()["data"]["run_id"]
            _drain_background_run(run_id)
            grant = owner.get(f"/api/agent/runs/{run_id}").json()["approval"]["grant"]
            monkeypatch.setattr(
                approval_grant, "_consumption_repository", lambda: unavailable_repository
            )
            with patch("app.fastapi_routes.domains.agent.routes._enqueue_run") as enqueue:
                rejected = owner.post(
                    f"/api/agent/runs/{run_id}/continue", json={"approval_grant": grant}
                )
                assert rejected.status_code == 503
                assert rejected.json()["message"] == "审批存储暂时不可用，请稍后重试"
                enqueue.assert_not_called()
            assert get_agent_run_repository().get(run_id).status == "waiting_user"
            execute.assert_not_called()
            monkeypatch.setattr(
                approval_grant, "_consumption_repository", lambda: durable_repository
            )
            accepted = owner.post(
                f"/api/agent/runs/{run_id}/continue", json={"approval_grant": grant}
            )
            assert accepted.status_code == 202
            assert _drain_background_run(run_id).status == "completed"
            execute.assert_called_once()
    finally:
        engine.dispose()


@pytest.mark.parametrize("renew_session", [False, True])
def test_durable_http_approval_rolls_back_enqueue_failure_and_retries(
    tmp_path, monkeypatch, renew_session
):
    from sqlalchemy import create_engine
    from sqlalchemy.exc import OperationalError
    from sqlalchemy.orm import sessionmaker

    from app.application.agent_orchestrator.approval_grant import issue_approval_grant
    from app.application.agent_orchestrator.run_sql_repository import SQLAlchemyAgentRunRepository
    from app.application.agent_orchestrator.task_execution_sql_repository import (
        SQLAlchemyTaskExecutionRepository,
    )
    from app.db.models.agent_approval import AgentApprovalConsumption

    engine = create_engine(f"sqlite:///{tmp_path / 'http-atomic.db'}")
    factory = sessionmaker(bind=engine)
    runs = SQLAlchemyAgentRunRepository(session_factory=factory)
    queue = SQLAlchemyTaskExecutionRepository(session_factory=factory)
    set_agent_run_repository_for_tests(runs)
    set_task_execution_repository_for_tests(queue)
    AgentApprovalConsumption.__table__.create(engine)
    run = AgentRun(
        user_id="owner",
        message="approve",
        status="waiting_user",
        steps=[
            AgentStep(node_id="node", tool_id="sales", action="create_order", status="waiting_user")
        ],
    )
    previous_binding = {
        "session_row_id": 1,
        "user_id": "owner",
        "mod_id": "sales",
        "account_tenant_id": "",
        "account_role": "user",
    }
    if renew_session:
        run.metadata["runtime_context"] = {"_mod_authorization": previous_binding}
    runs.save(run)
    token = issue_approval_grant(run, principal_id="owner")["grant"]
    enqueue = queue.enqueue_in_session

    def interrupted(db, run, **kwargs):
        enqueue(db, run, **kwargs)
        db.flush()
        raise OperationalError("injected failure", {}, RuntimeError("interrupted"))

    try:
        replacement = {**previous_binding, "session_row_id": 2} if renew_session else None
        client = _client("owner", mod_authorization=replacement)
        with patch(
            "app.fastapi_routes.domains.agent.routes.notify_agent_task_dispatcher"
        ) as notify:
            monkeypatch.setattr(queue, "enqueue_in_session", interrupted)
            failed = client.post(
                f"/api/agent/runs/{run.run_id}/continue", json={"approval_grant": token}
            )
            assert failed.status_code == 503
            notify.assert_not_called()
            assert runs.get(run.run_id).status == "waiting_user"
            if renew_session:
                assert (
                    runs.get(run.run_id).metadata["runtime_context"]["_mod_authorization"]
                    == previous_binding
                )
            assert queue.get(run.run_id) is None
            with factory() as db:
                assert db.query(AgentApprovalConsumption).count() == 0
            monkeypatch.setattr(queue, "enqueue_in_session", enqueue)
            accepted = client.post(
                f"/api/agent/runs/{run.run_id}/continue", json={"approval_grant": token}
            )
            assert accepted.status_code == 202
            notify.assert_called_once()
            assert runs.get(run.run_id).status == "queued"
            if renew_session:
                assert (
                    runs.get(run.run_id).metadata["runtime_context"]["_mod_authorization"]
                    == replacement
                )
            assert queue.get(run.run_id).state == "queued"
            with factory() as db:
                assert db.query(AgentApprovalConsumption).count() == 1
            replay = client.post(
                f"/api/agent/runs/{run.run_id}/continue", json={"approval_grant": token}
            )
            assert replay.status_code == 403
            notify.assert_called_once()
    finally:
        engine.dispose()


def test_approval_cannot_replace_persisted_tenant_and_does_not_consume_grant():
    from app.application.agent_orchestrator.approval_grant import issue_approval_grant

    run = AgentRun(
        user_id="owner",
        message="inbound",
        status="waiting_user",
        steps=[
            AgentStep(
                node_id="in",
                tool_id="inventory",
                action="stock_in",
                status="waiting_user",
                params={"product_id": 1, "warehouse_id": 1, "quantity": 50},
            )
        ],
    )
    run.metadata["runtime_context"] = {"tenant_id": "1"}
    get_agent_run_repository().save(run)
    token = issue_approval_grant(run, principal_id="owner")["grant"]
    client = _client("owner", tenant_id="1")
    with patch("app.fastapi_routes.domains.agent.routes._enqueue_run") as enqueue:
        response = client.post(
            f"/api/agent/runs/{run.run_id}/continue",
            json={"approval_grant": token, "runtime_context": {"tenant_id": "2"}},
        )
        assert response.status_code == 400
        enqueue.assert_not_called()
        stored = get_agent_run_repository().get(run.run_id)
        assert stored.status == "waiting_user"
        assert stored.metadata["runtime_context"]["tenant_id"] == "1"
        retry = client.post(
            f"/api/agent/runs/{run.run_id}/continue", json={"approval_grant": token}
        )
        assert retry.status_code == 202
        enqueue.assert_called_once()


def test_resume_rejects_tenant_change_without_control_command():
    run = AgentRun(user_id="owner", message="paused", status="paused")
    run.metadata["runtime_context"] = {"tenant_id": "1"}
    repository = get_agent_run_repository()
    repository.save(run)
    before = repository.get(run.run_id).to_dict()
    with patch("app.fastapi_routes.domains.agent.routes._enqueue_run") as enqueue:
        response = _client("owner", tenant_id="1").post(
            f"/api/agent/runs/{run.run_id}/resume",
            json={"runtime_context": {"tenant_id": "2"}},
        )
        assert response.status_code == 400
        assert response.json()["message"] == "不能更改任务的租户范围"
        enqueue.assert_not_called()
    assert repository.get(run.run_id).to_dict() == before
    assert repository.latest_task_control(run.run_id) is None


@pytest.mark.parametrize(
    "method,path",
    [
        ("get", ""),
        ("get", "/events"),
        ("post", "/pause"),
        ("post", "/cancel"),
        ("post", "/resume"),
        ("post", "/retry"),
        ("post", "/continue"),
    ],
)
def test_same_user_cannot_access_run_from_another_tenant(method, path):
    run = AgentRun(user_id="owner", message="private task", status="paused")
    run.metadata["runtime_context"] = {"tenant_id": "tenant-a"}
    repository = get_agent_run_repository()
    repository.save(run)
    before = repository.get(run.run_id).to_dict()
    response = getattr(_client("owner", tenant_id="tenant-b"), method)(
        f"/api/agent/runs/{run.run_id}{path}"
    )
    assert response.status_code == 403
    assert "private task" not in response.text
    assert repository.get(run.run_id).to_dict() == before
    assert repository.latest_task_control(run.run_id) is None


def test_run_list_filters_tenant_before_limit():
    repository = get_agent_run_repository()
    for tenant in ("tenant-a", "tenant-b"):
        run = AgentRun(user_id="owner", message=tenant)
        run.metadata["runtime_context"] = {"tenant_id": tenant}
        repository.save(run)
    response = _client("owner", tenant_id="tenant-a").get("/api/agent/runs?limit=1")
    assert response.status_code == 200
    assert [item["message"] for item in response.json()["data"]] == ["tenant-a"]


@pytest.mark.parametrize("tenant", ["", "verified"])
def test_create_task_uses_authenticated_tenant_only(tenant):
    response = _client("owner", tenant_id=tenant).post(
        "/api/agent/tasks",
        json={
            "task_id": "tenant-binding",
            "title": "客户查询",
            "tool_id": "customers",
            "action": "query",
            "params": {"keyword": "示例"},
            "runtime_context": {"tenant_id": "forged", "source": "test"},
        },
    )
    assert response.status_code == 202
    run = get_agent_run_repository().get(response.json()["data"]["run_id"])
    context = run.metadata["runtime_context"]
    assert context.get("tenant_id", "") == tenant
    assert context["source"] == "test"


def test_create_agent_run_validates_request_body() -> None:
    get_agent_run_repository().clear()
    client = _client()

    missing_message = client.post("/api/agent/runs", json={"user_id": "u1"})
    assert missing_message.status_code == 400
    assert missing_message.json()["success"] is False

    bad_context = client.post(
        "/api/agent/runs",
        json={"message": "查库存", "runtime_context": ["bad"]},
    )
    assert bad_context.status_code == 400
    assert bad_context.json()["success"] is False


def test_agent_routes_do_not_expose_internal_exception_details() -> None:
    secret = "Traceback: database password=do-not-expose"
    repository = get_agent_run_repository()
    repository.clear()
    run = AgentRun(user_id="u1", message="执行任务", status="failed", error=secret)
    run.steps.append(
        AgentStep(
            node_id="n1",
            tool_id="products",
            action="query",
            status="failed",
            error=secret,
            output={"success": False, "error": secret},
        )
    )
    run.tool_calls.append(
        ToolCall(
            step_id=run.steps[0].step_id,
            node_id="n1",
            tool_id="products",
            action="query",
            status="failed",
            error=secret,
        )
    )
    run.add_event("run.failed", "Agent run 失败", {"error": secret})

    client = _client()
    with patch(
        "app.fastapi_routes.domains.agent.routes.AgentOrchestrator.start_run",
        return_value=run,
    ):
        response = client.post("/api/agent/runs", json={"message": "执行任务"})

    assert response.status_code == 200
    assert secret not in response.text
    payload = response.json()["data"]
    assert payload["error"] == "Agent 执行失败，详细信息已记录"
    assert payload["steps"][0]["output"]["error"] == "Agent 执行失败，详细信息已记录"
    assert payload["events"][0]["data"]["error"] == "Agent 执行失败，详细信息已记录"

    repository.save(run)
    list_response = client.get("/api/agent/runs")
    events_response = client.get(f"/api/agent/runs/{run.run_id}/events")
    assert secret not in list_response.text
    assert secret not in events_response.text


def test_agent_route_internal_failure_returns_stable_public_message() -> None:
    client = _client()
    with patch(
        "app.fastapi_routes.domains.agent.routes.AgentOrchestrator.start_run",
        side_effect=RuntimeError("Traceback: database password=do-not-expose"),
    ):
        response = client.post("/api/agent/runs", json={"message": "执行任务"})

    assert response.status_code == 500
    assert response.json() == {
        "success": False,
        "message": "Agent 服务暂时不可用，请稍后重试",
    }


def test_pause_resume_and_cancel_agent_run() -> None:
    get_agent_run_repository().clear()
    clear_run_controls_for_tests()
    client = _client("owner")
    patches = _planner_fallback_patches()
    with (
        patches[0],
        patches[1],
        patches[2],
        patches[3],
        patch(
            "app.application.facades.tools_facade.execute_registered_workflow_tool",
            return_value={"success": True, "data": []},
        ) as execute,
    ):
        created = client.post(
            "/api/agent/runs",
            json={"message": "查产品", "auto_execute": False},
        ).json()["data"]
        paused = client.post(f"/api/agent/runs/{created['run_id']}/pause").json()["data"]
        assert paused["status"] == "paused"
        execute.assert_not_called()

        resumed_response = client.post(
            f"/api/agent/runs/{created['run_id']}/resume",
            json={"runtime_context": {"source": "resume-test"}},
        )
        assert resumed_response.status_code == 202
        assert resumed_response.json()["data"]["status"] == "queued"
        resumed = _drain_background_run(created["run_id"]).to_dict()
        assert resumed["status"] == "completed"
        execute.assert_called_once()

        second = client.post(
            "/api/agent/runs",
            json={"message": "查产品", "auto_execute": False},
        ).json()["data"]
        cancelled = client.post(f"/api/agent/runs/{second['run_id']}/cancel").json()["data"]
        assert cancelled["status"] == "cancelled"
        assert cancelled["steps"][0]["status"] == "skipped"
        assert cancelled["final_output"]["cancelled"] is True


def test_retry_agent_run_preserves_task_identity_and_creates_new_attempt() -> None:
    repository = get_agent_run_repository()
    repository.clear()
    clear_run_controls_for_tests()
    client = _client("owner")
    patches = _planner_fallback_patches()
    with (
        patches[0],
        patches[1],
        patches[2],
        patches[3],
        patch(
            "app.application.facades.tools_facade.execute_registered_workflow_tool",
            return_value={"success": True, "data": []},
        ),
    ):
        created = client.post(
            "/api/agent/runs",
            json={
                "message": "核对库存",
                "auto_execute": False,
                "runtime_context": {
                    "conversation_id": "conversation-42",
                    "task_id": "conversation-42",
                    "task_title": "库存核对任务",
                },
            },
        ).json()["data"]
        cancelled = client.post(f"/api/agent/runs/{created['run_id']}/cancel").json()["data"]
        assert cancelled["status"] == "cancelled"

        response = client.post(
            f"/api/agent/runs/{created['run_id']}/retry",
            json={"runtime_context": {"task_id": "other-task", "conversation_id": "other-chat"}},
        )
        replay = client.post(f"/api/agent/runs/{created['run_id']}/retry", json={})

    assert response.status_code == 200
    retried_reference = response.json()["data"]
    assert replay.status_code == 200
    assert replay.json()["data"]["run_id"] == retried_reference["run_id"]
    assert retried_reference["run_id"] != created["run_id"]
    retried_stored = repository.get(retried_reference["run_id"])
    assert retried_stored is not None
    retried = retried_stored.to_dict()
    assert retried["metadata"]["task_context"] == {
        "task_id": "conversation-42",
        "title": "库存核对任务",
        "conversation_id": "conversation-42",
        "root_run_id": created["run_id"],
        "parent_run_id": created["run_id"],
        "attempt": 2,
        "workspace_id": "",
        "workspace_path": "",
        "isolation": "business_workspace",
    }
    previous = repository.get(created["run_id"])
    assert previous is not None
    assert previous.events[-1].event_type == "run.retry_created"


def test_retry_agent_run_rejects_non_terminal_task() -> None:
    repository = get_agent_run_repository()
    repository.clear()
    run = AgentRun(user_id="owner", message="仍在运行", status="running")
    repository.save(run)

    response = _client("owner").post(f"/api/agent/runs/{run.run_id}/retry", json={})

    assert response.status_code == 409
    assert "可以重试" in response.json()["message"]


def test_get_agent_run_returns_404_for_missing_run() -> None:
    get_agent_run_repository().clear()
    client = _client()

    response = client.get("/api/agent/runs/run_missing")

    assert response.status_code == 404
    assert response.json()["success"] is False

    events_response = client.get("/api/agent/runs/run_missing/events")
    assert events_response.status_code == 404
    assert events_response.json()["success"] is False


def test_unified_task_requires_approval_and_deduplicates_execution() -> None:
    repository = get_agent_run_repository()
    repository.clear()
    clear_consumed_approval_grants_for_tests()
    client = _client("task-owner")
    body = {
        "task_id": "chat-shipment-001",
        "title": "生成客户甲发货单",
        "message": "给客户甲生成一张发货单",
        "tool_id": "shipment_orders",
        "action": "generate",
        "params": {
            "unit_name": "客户甲",
            "products": [{"model_number": "A-01", "quantity": 2}],
        },
        "runtime_context": {"conversation_id": "chat-001"},
    }

    with patch(
        "app.application.facades.tools_facade.execute_registered_workflow_tool",
        return_value={
            "success": True,
            "message": "发货单已生成",
            "order_id": 41,
            "file_path": "/tmp/shipment-41.docx",
        },
    ) as execute:
        created = client.post("/api/agent/tasks", json=body)
        assert created.status_code == 202
        created_payload = created.json()
        run = created_payload["data"]
        assert run["status"] == "waiting_user"
        assert run["steps"][0]["status"] == "waiting_user"
        assert created_payload["capabilities"]["approve"] is True
        assert created_payload["capabilities"]["pause"] is True
        assert created_payload["deduplicated"] is False
        execute.assert_not_called()

        paused = client.post(f"/api/agent/runs/{run['run_id']}/pause").json()
        assert paused["data"]["status"] == "paused"
        assert paused["capabilities"]["resume"] is True

        resumed = client.post(f"/api/agent/runs/{run['run_id']}/resume", json={}).json()
        assert resumed["data"]["status"] == "waiting_user"
        assert resumed["capabilities"]["approve"] is True
        execute.assert_not_called()

        approval_grant = resumed["approval"]["grant"]
        approved = client.post(
            f"/api/agent/runs/{run['run_id']}/continue",
            json={"approval_grant": approval_grant},
        )
        assert approved.status_code == 202
        assert approved.json()["data"]["status"] == "queued"
        completed = _drain_background_run(run["run_id"]).to_dict()
        assert completed["status"] == "completed"
        assert completed["tool_calls"][0]["status"] == "completed"
        assert completed["final_output"]["node_outputs"]
        assert "step.approved" in [event["event_type"] for event in completed["events"]]
        execute.assert_called_once()
        assert execute.call_args.args[:2] == ("shipment_orders", "generate")
        assert execute.call_args.args[2]["unit_name"] == "客户甲"
        assert execute.call_args.args[2]["products"] == [{"model_number": "A-01", "quantity": 2}]
        assert isinstance(execute.call_args.args[2]["_runtime_context"], dict)

        duplicate = client.post("/api/agent/tasks", json=body)
        assert duplicate.status_code == 200
        assert duplicate.json()["deduplicated"] is True
        assert duplicate.json()["data"]["run_id"] == run["run_id"]
        execute.assert_called_once()

    mismatched = client.post(
        "/api/agent/tasks",
        json={**body, "params": {"unit_name": "客户乙", "products": [{"quantity": 1}]}},
    )
    assert mismatched.status_code == 409


def test_unified_low_risk_task_cancel_and_retry_preserve_exact_approval() -> None:
    repository = get_agent_run_repository()
    repository.clear()
    client = _client("task-owner")
    body = {
        "task_id": "chat-shipment-retry",
        "title": "可恢复客户查询任务",
        "tool_id": "customers",
        "action": "query",
        "params": {"keyword": "客户乙"},
    }
    created = client.post("/api/agent/tasks", json=body).json()["data"]
    cancelled = client.post(f"/api/agent/runs/{created['run_id']}/cancel").json()["data"]
    assert cancelled["status"] == "cancelled"

    with patch("app.application.facades.tools_facade.execute_registered_workflow_tool") as execute:
        retried = client.post(f"/api/agent/runs/{created['run_id']}/retry", json={})

    assert retried.status_code == 200
    retried_run = repository.get(retried.json()["data"]["run_id"])
    assert retried_run is not None
    assert retried_run.status == "waiting_user"
    assert retried_run.steps[0].tool_id == "customers"
    assert retried_run.steps[0].action == "query"
    assert retried_run.steps[0].params == body["params"]
    assert retried_run.metadata["task_context"]["task_id"] == body["task_id"]
    assert retried_run.metadata["task_context"]["attempt"] == 2
    execute.assert_not_called()


def test_task_ssot_lists_details_and_archives_per_tenant() -> None:
    repository = get_agent_run_repository()
    repository.clear()
    clear_consumed_approval_grants_for_tests()
    client = _client("task-owner", tenant_id="tenant-a")
    body = {
        "task_id": "task-monthly-report",
        "title": "生成月度经营报告",
        "tool_id": "customers",
        "action": "query",
        "params": {"keyword": "客户甲"},
    }

    created = client.post("/api/agent/tasks", json=body)
    assert created.status_code == 202
    run = created.json()["data"]

    listing = client.get("/api/agent/tasks")
    assert listing.status_code == 200
    assert listing.json()["count"] == 1
    task = listing.json()["data"][0]
    assert task["task_id"] == body["task_id"]
    assert task["tenant_id"] == "tenant-a"
    assert task["attention_state"] == "approval_required"
    assert task["approval_required"] is True
    assert task["unread_count"] == 0
    assert task["conversation_id"] == body["task_id"]
    assert task["workspace_id"] == body["task_id"]
    assert task["active_run_id"] == run["run_id"]
    assert task["active_run"]["run_id"] == run["run_id"]
    assert task["capabilities"]["approve"] is True
    assert task["progress"]["percent"] == 0
    assert task["progress"]["stage"] == "等待审批或用户确认"
    assert task["progress"]["total_units"] == 1

    detail = client.get(f"/api/agent/tasks/{body['task_id']}")
    assert detail.status_code == 200
    assert detail.json()["data"]["run_count"] == 1
    assert len(detail.json()["data"]["runs"]) == 1

    assert client.post(f"/api/agent/tasks/{body['task_id']}/archive").status_code == 409

    cancelled = client.post(f"/api/agent/runs/{run['run_id']}/cancel")
    assert cancelled.status_code == 200
    archived = client.post(f"/api/agent/tasks/{body['task_id']}/archive")
    assert archived.status_code == 200
    assert archived.json()["data"]["archived_at"]
    assert client.get("/api/agent/tasks").json()["count"] == 0
    archived_listing = client.get(
        "/api/agent/tasks",
        params={"include_archived": True},
    ).json()
    assert archived_listing["count"] == 1

    tenant_b_client = _client("task-owner", tenant_id="tenant-b")
    assert tenant_b_client.get(f"/api/agent/tasks/{body['task_id']}").status_code == 404
    tenant_b_created = tenant_b_client.post("/api/agent/tasks", json=body)
    assert tenant_b_created.status_code == 202
    assert tenant_b_client.get("/api/agent/tasks").json()["count"] == 1
    assert client.get("/api/agent/tasks").json()["count"] == 0


def test_completed_workspace_result_is_durably_marked_read_per_tenant() -> None:
    repository = get_agent_run_repository()
    repository.clear()
    client = _client("workspace-owner", tenant_id="tenant-a")
    body = {
        "task_id": "workspace-result-001",
        "title": "客户B销售结果",
        "tool_id": "customers",
        "action": "query",
        "params": {"keyword": "客户B"},
    }
    created = client.post("/api/agent/tasks", json=body).json()["data"]
    run = repository.get(created["run_id"])
    assert run is not None
    run.status = "completed"
    run.steps[0].status = "completed"
    repository.save(run)

    unread = client.get("/api/agent/tasks").json()["data"][0]
    assert unread["attention_state"] == "result_unread"
    assert unread["unread_count"] == 1

    marked = client.post(f"/api/agent/tasks/{body['task_id']}/read")
    assert marked.status_code == 200
    assert marked.json()["data"]["attention_state"] == ""
    assert marked.json()["data"]["unread_count"] == 0
    assert marked.json()["data"]["metadata"]["read_at"]
    assert client.post(f"/api/agent/tasks/{body['task_id']}/read").status_code == 200
    persisted = client.get("/api/agent/tasks").json()["data"][0]
    assert persisted["attention_state"] == ""
    assert persisted["unread_count"] == 0
    assert (
        _client("workspace-owner", tenant_id="tenant-b")
        .post(f"/api/agent/tasks/{body['task_id']}/read")
        .status_code
        == 404
    )


def test_task_center_stream_emits_tenant_scoped_snapshot() -> None:
    client = _client("task-owner", tenant_id="tenant-a")
    created = client.post(
        "/api/agent/tasks",
        json={
            "task_id": "stream-task",
            "title": "流式任务",
            "tool_id": "customers",
            "action": "query",
            "params": {"keyword": "客户甲"},
        },
    )
    assert created.status_code == 202
    assert (
        _client("task-owner", tenant_id="tenant-b")
        .post(
            "/api/agent/tasks",
            json={
                "task_id": "other-tenant-task",
                "title": "其它租户任务",
                "tool_id": "customers",
                "action": "query",
                "params": {"keyword": "客户乙"},
            },
        )
        .status_code
        == 202
    )

    streamed = client.get("/api/agent/tasks/events/stream", params={"once": True})

    assert streamed.status_code == 200
    assert streamed.headers["content-type"].startswith("text/event-stream")
    assert "event: task.snapshot" in streamed.text
    assert '"task_id": "stream-task"' in streamed.text
    assert '"progress": {' in streamed.text
    assert '"stage": "等待审批或用户确认"' in streamed.text
    assert "other-tenant-task" not in streamed.text
    assert "event: stream.closed" in streamed.text
    runtime = client.get("/api/agent/task-runtime")
    assert runtime.status_code == 200
    assert runtime.json()["data"]["max_workers"] == 4
    assert runtime.json()["data"]["active_count"] == 0
    assert runtime.json()["data"]["progress"] == {
        "task_count": 1,
        "active_count": 0,
        "attention_count": 1,
        "completed_count": 0,
        "overall_percent": 0,
    }


def test_task_control_command_survives_process_local_control_reset() -> None:
    repository = get_agent_run_repository()
    repository.clear()
    clear_run_controls_for_tests()
    run = AgentRun(user_id="task-owner", message="长任务", status="running")
    run.metadata["runtime_context"] = {"task_id": "durable-control-task"}
    run.metadata["execution"] = {"state": "active"}
    run.metadata["task_context"] = {
        "task_id": "durable-control-task",
        "title": "长任务",
        "root_run_id": run.run_id,
        "attempt": 1,
    }
    repository.save(run)

    orchestrator = AgentOrchestrator(repository=repository)
    paused = orchestrator.pause_run(run.run_id, requested_by="task-owner")
    assert paused is not None
    assert paused.status == "running"
    assert paused.metadata["control_request"]["status"] == "requested"
    command = repository.latest_task_control(run.run_id)
    assert command is not None
    assert command.action == "pause"
    assert command.status == "requested"

    clear_run_controls_for_tests()
    worker_snapshot = repository.get(run.run_id)
    assert worker_snapshot is not None
    assert worker_snapshot.status == "running"
    assert "control_request" not in worker_snapshot.metadata
    assert (
        AgentOrchestrator(repository=repository)._apply_requested_control(worker_snapshot) is True
    )
    applied = repository.latest_task_control(run.run_id)
    assert applied is not None
    assert applied.command_id == command.command_id
    assert applied.status == "applied"
    assert applied.applied_at
    persisted = repository.get(run.run_id)
    assert persisted is not None
    assert persisted.status == "paused"


def test_clarification_answer_is_bound_validated_and_does_not_approve_write():
    from app.application.workflow.clarification_node import build_clarify_node
    from app.application.workflow.types import PlanGraph, WorkflowNode

    plan = PlanGraph(
        plan_id="input",
        intent="create",
        nodes=[
            build_clarify_node("请提供客户", ambient={"target_node_id": "create"}),
            WorkflowNode(
                node_id="create",
                tool_id="customers",
                action="create",
                params={},
            ),
        ],
    )
    run = AgentOrchestrator().start_run_from_plan(user_id="u1", message="新增产品", plan=plan)
    step_id = run.steps[0].step_id
    url = f"/api/agent/runs/{run.run_id}/clarification"
    original = get_agent_run_repository().get(run.run_id).to_dict()
    for parameters in (
        {},
        {"unit_name": ""},
        {"name_or_model": "changed"},
        {"_runtime_context": {"tenant_id": 99}},
    ):
        response = _client().post(url, json={"step_id": step_id, "parameters": parameters})
        assert response.status_code == 400
        assert get_agent_run_repository().get(run.run_id).to_dict() == original
    denied = _client("someone-else").post(
        url, json={"step_id": step_id, "parameters": {"unit_name": "客户甲"}}
    )
    assert denied.status_code in {403, 404}
    response = _client().post(url, json={"step_id": step_id, "parameters": {"unit_name": "客户甲"}})
    assert response.status_code == 202, response.text
    with patch("app.application.facades.tools_facade.execute_registered_workflow_tool") as execute:
        resumed = _drain_background_run(run.run_id)
        execute.assert_not_called()
    assert resumed.steps[0].status == "completed"
    assert resumed.steps[1].params == {"unit_name": "客户甲"}
    assert resumed.steps[1].status == "waiting_user"
    assert resumed.status == "waiting_user"
    repeated = _client().post(url, json={"step_id": step_id, "parameters": {"unit_name": "客户乙"}})
    assert repeated.status_code == 400


def test_exported_spreadsheet_download_checks_owner_and_file_content(tmp_path, monkeypatch):
    import json
    from io import BytesIO

    from openpyxl import load_workbook

    from app.application.agent_orchestrator.artifact_files import artifact_path
    from app.application.agent_orchestrator.run_models import artifact_from_dict
    from app.services.tools_workflow_registered_part01_part02 import _registered_router_reports

    monkeypatch.setenv("XCAGI_DATA_DIR", str(tmp_path))
    run = AgentRun(user_id="u1", message="导出销售报表", status="completed")
    exported = _registered_router_reports(
        "export",
        {"report_type": "sales", "data": [{"product_name": "产品甲", "amount": 125.5}]},
        {"run_id": run.run_id},
        "pro_default",
        run.message,
    )
    assert exported["success"] and exported["row_count"] == 1
    json.dumps(exported)
    run.artifacts = [artifact_from_dict(item) for item in exported["artifacts"]]
    get_agent_run_repository().save(run)
    artifact = run.artifacts[0]
    response = _client().get(artifact.uri)
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
    assert list(workbook.active.values) == [("product_name", "amount"), ("产品甲", 125.5)]
    workbook.close()
    assert _client("someone-else").get(artifact.uri).status_code in {403, 404}
    artifact_path(run.run_id, artifact.artifact_id).write_bytes(b"corrupt")
    assert _client().get(artifact.uri).status_code == 410


def test_generated_shipment_download_contains_real_business_cells(tmp_path, monkeypatch):
    from io import BytesIO

    from openpyxl import load_workbook
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker

    from app.application.agent_orchestrator.run_models import artifact_from_dict
    from app.db.base import Base
    from app.db.models import Product, PurchaseUnit
    from app.infrastructure.tenant_scope import tenant_scope
    from app.services.tools_workflow_registered_part02_part01 import (
        _registered_router_shipment_orders,
    )

    monkeypatch.setenv("XCAGI_DATA_DIR", str(tmp_path / "data"))
    engine = create_engine(f"sqlite:///{tmp_path / 'shipment.db'}")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)
    monkeypatch.setattr("app.db.session.SessionLocal", factory)
    monkeypatch.setattr("app.db.SessionLocal", factory)
    try:
        with tenant_scope(1):
            with factory() as db:
                db.add(PurchaseUnit(unit_name="下载测试客户"))
                db.add(
                    Product(name="下载测试产品", model_number="9803", price=25, specification="12")
                )
                db.commit()
            run = AgentRun(user_id="u1", message="生成发货单", status="completed")
            result = _registered_router_shipment_orders(
                "generate",
                {
                    "unit_name": "下载测试客户",
                    "products": [
                        {"model_number": "9803", "quantity_tins": 3, "tin_spec": 12},
                    ],
                },
                {"run_id": run.run_id},
                "pro_default",
                run.message,
            )
            assert result["success"]
            run.artifacts = [artifact_from_dict(item) for item in result["artifacts"]]
            get_agent_run_repository().save(run)
            response = _client().get(run.artifacts[0].uri)
            assert response.status_code == 200
            workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
            try:
                sheet = workbook["发货单"]
                assert "下载测试客户" in sheet["A2"].value
                assert [sheet[cell].value for cell in ["A5", "D5", "E5", "F5", "G5", "D7"]] == [
                    "9803",
                    "下载测试产品",
                    3,
                    12,
                    36,
                    900,
                ]
            finally:
                workbook.close()
            assert _client("another-user").get(run.artifacts[0].uri).status_code in {403, 404}
    finally:
        engine.dispose()
