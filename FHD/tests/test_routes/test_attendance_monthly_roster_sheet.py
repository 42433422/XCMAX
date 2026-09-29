"""月度统计必须按当前人员名单改姓名列，并清掉模板里多出来的旧人。"""

from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

from openpyxl import Workbook


def _load_mapper():
    path = (
        Path(__file__).resolve().parents[2]
        / "XCAGI"
        / "mods"
        / "attendance-industry"
        / "backend"
        / "attendance_engine"
        / "mapper.py"
    )
    spec = importlib.util.spec_from_file_location("attendance_engine_mapper", path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def test_monthly_sheet_replaces_name_column_and_drops_removed_people():
    mapper = _load_mapper()
    wb = Workbook()
    ws = wb.active
    ws.title = "月度统计"
    ws["A1"] = "序号"
    ws["B1"] = "部门"
    ws["C2"] = "姓名"
    ws["D2"] = "正常上班"
    ws["C3"] = "老王"
    ws["C4"] = "老李"
    ws["C5"] = "老赵"
    ws["B3"] = "旧部门"
    ws["D3"] = 8

    mapper.write_monthly_sheet(
        wb,
        [
            {"姓名": "新人甲", "部门": "生产部", "正常上班": 1},
            {"姓名": "留任乙", "部门": "质检", "正常上班": 2},
        ],
        link_detail_side_totals=False,
    )

    assert ws["C3"].value == "新人甲"
    assert ws["B3"].value == "生产部"
    assert ws["C4"].value == "留任乙"
    assert ws["B4"].value == "质检"
    assert ws["C5"].value is None
    assert ws["B5"].value is None
    assert ws["D5"].value is None
    assert ws["A3"].value == "=ROW()-2"
    assert ws["D3"].value == 1


def test_monthly_sheet_header_on_first_row_still_follows_roster():
    mapper = _load_mapper()
    wb = Workbook()
    ws = wb.active
    ws.title = "月度统计"
    ws["A1"] = "序号"
    ws["B1"] = "姓名"
    ws["C1"] = "部门"
    ws["D1"] = "正常上班"
    ws["B2"] = "老王"
    ws["B3"] = "老李"
    ws["C2"] = "旧部门"

    mapper.write_monthly_sheet(
        wb,
        [{"姓名": "新人甲", "部门": "生产部", "正常上班": 3}],
        link_detail_side_totals=False,
    )

    assert ws["B2"].value == "新人甲"
    assert ws["C2"].value == "生产部"
    assert ws["B3"].value is None
    assert ws["C3"].value is None
    assert ws["D2"].value == 3
